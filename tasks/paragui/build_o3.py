"""o3 - one Writer document (an NHRF grant proposal) + its Calc budget: A applies the funder's revised cost rules
(scanned letter) to the budget workbook and the proposal's budget section / cover table; B reformats the proposal to
the funder's formatting annex (scan: page, fonts, heading styles, header reference, page numbers, references) and
exports the submission PDF."""
from __future__ import annotations

import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_oscan import Page, save_scan  # noqa: E402

REF = "NHRF-2026-0417"
TITLE = ("TeleRehab-Home: a randomised feasibility study of tablet-supported home physiotherapy after stroke in rural "
         "communities")

RATES_OLD = {"salary": 34500, "oncost": 0.28, "months": 24, "visits": 120, "miles": 46, "mileage": 0.45, "fee": 85,
             "conference": 1200, "tablets": 30, "tablet_cost": 320, "mocap": 7800, "consumables": 2300, "ppi_n": 12,
             "ppi_fee": 150, "indirect": 0.20}
RATES_NEW = dict(RATES_OLD, salary=32000, mileage=0.35, indirect=0.15, mocap=0)


def budget(r):
    staff = r["salary"] * (1 + r["oncost"]) * r["months"] / 12
    sessions = r["visits"] * r["fee"]
    travel = r["visits"] * r["miles"] * r["mileage"] + r["conference"]
    equipment = r["tablets"] * r["tablet_cost"] + r["mocap"]
    cons = r["consumables"]
    ppi = r["ppi_n"] * r["ppi_fee"]
    direct = staff + sessions + travel + equipment + cons + ppi
    ind = direct * r["indirect"]
    return {k: round(v, 2) for k, v in dict(staff=staff, sessions=sessions, travel=travel, equipment=equipment,
                                             consumables=cons, ppi=ppi, direct=direct, indirect=ind,
                                             total=direct + ind).items()}


def gbp(x):
    return f"£{x:,.2f}"


ROWS = [("staff", "Staff costs"), ("sessions", "Clinical sessions (home physiotherapy visits)"), ("travel", "Travel"),
        ("equipment", "Equipment"), ("consumables", "Consumables"), ("ppi", "Patient and public involvement"),
        ("direct", "Direct costs subtotal"), ("indirect", "Indirect costs ({pct}% of direct costs)"),
        ("total", "Total requested")]

REFS = [
    "Osei A, Brennan K, Hughes T. Feasibility of tablet-based home exercise after stroke: a mixed-methods pilot. "
    "Clin Rehabil. 2025;39(2):145-156.",
    "Langhorne P, Bernhardt J, Kwakkel G. Stroke rehabilitation. Lancet. 2011;377(9778):1693-1702.",
    "Chen L, Martins R. Telerehabilitation for upper-limb recovery after stroke: a systematic review. J Telemed "
    "Telecare. 2023;29(4):233-247.",
    "Wade DT. Measurement in neurological rehabilitation. Oxford: Oxford University Press; 1992.",
    "Fugl-Meyer AR, Jaasko L, Leyman I, Olsson S, Steglind S. The post-stroke hemiplegic patient. Scand J Rehabil Med. "
    "1975;7(1):13-31.",
    "Brennan K, Osei A. Rural access to community stroke services in northern England. BMC Health Serv Res. "
    "2024;24:812.",
    "Eldridge SM, Chan CL, Campbell MJ, et al. CONSORT 2010 statement: extension to randomised pilot and feasibility "
    "trials. BMJ. 2016;355:i5239.",
    "Hughes T, Patel N, Osei A. The cost of community physiotherapy visits in rural areas. Physiotherapy. "
    "2022;114:e88-e89.",
]


def sections(b, pct):
    """[(level, title, [paragraph texts or ('TABLE', None)])]"""
    r = RATES_OLD
    return [
        (1, "1. Lay Summary", [
            "Every year around 100,000 people in the UK have a stroke, and most of them leave hospital with problems "
            "moving an arm or a leg. Regular physiotherapy in the first months after a stroke makes a large difference "
            "to recovery, but people living in rural areas often receive far fewer sessions than people in towns, "
            "because therapists spend much of their working day driving between patients.",
            "TeleRehab-Home will test whether a tablet computer with a simple exercise app, combined with fewer but "
            "longer home visits and weekly video calls with a physiotherapist, can give rural stroke survivors as much "
            "therapy as their urban neighbours. We will recruit 60 people in Cumbria and North Yorkshire, give half of "
            "them the TeleRehab-Home programme for twelve weeks and compare them with half receiving usual care. This "
            "feasibility study will tell us whether a full trial is possible and worthwhile.",
        ]),
        (1, "2. Background and Rationale", [
            "Intensive, task-specific practice is the cornerstone of motor recovery after stroke (Langhorne et al., "
            "2011). National guidance recommends at least 45 minutes of each relevant therapy on five days a week, yet "
            "audit data show that community stroke teams in rural districts deliver less than half of that dose "
            "(Brennan and Osei, 2024). Travel time is the main reason: in our previous survey, therapists in rural "
            "teams spent 38% of their working hours driving (Hughes et al., 2022).",
            "Telerehabilitation can increase the amount of practice without increasing travel. A recent systematic "
            "review found moderate-quality evidence that remotely supervised exercise produces upper-limb gains "
            "similar to face-to-face therapy (Chen and Martins, 2023). Our own pilot with 14 participants showed that "
            "older adults with no previous tablet experience could use a simplified exercise app safely, with 81% of "
            "prescribed sessions completed (Osei et al., 2025). What is not known is whether such a programme can be "
            "delivered at scale by rural NHS teams, whether patients will stay with it for three months, and which "
            "outcome measure should be used in a definitive trial.",
        ]),
        (1, "3. Aims and Objectives", [
            "The aim of this study is to establish whether a definitive randomised controlled trial of TeleRehab-Home "
            "is feasible. The specific objectives are:",
            "(a) to estimate recruitment, retention and adherence rates over twelve weeks;",
            "(b) to assess the acceptability of the programme to patients, carers and therapists;",
            "(c) to estimate the variability of candidate outcome measures to inform a sample size calculation;",
            "(d) to estimate the cost of delivering the programme in rural community stroke teams.",
        ]),
        (1, "4. Methods", []),
        (2, "4.1 Study Design", [
            "A two-arm, parallel-group randomised feasibility study with an embedded qualitative process evaluation, "
            "reported according to the CONSORT extension for pilot and feasibility trials (Eldridge et al., 2016). "
            "Participants will be randomised 1:1 using a web-based system with minimisation by site and baseline arm "
            "function.",
        ]),
        (2, "4.2 Participants and Recruitment", [
            "Adults within six months of a first or recurrent stroke, living in a rural postcode (ONS classification) "
            "and discharged home from one of three participating hospitals, will be eligible. People with severe "
            "cognitive impairment or no home internet connection after the provision of a mobile data plan will be "
            "excluded. Community stroke teams will identify eligible patients at discharge; the research assistant "
            "will obtain written consent at a home visit.",
        ]),
        (2, "4.3 Outcome Measures", [
            "Feasibility outcomes are recruitment rate, retention at twelve weeks, and adherence (proportion of "
            "prescribed sessions completed, recorded by the app). Candidate clinical outcomes are the Fugl-Meyer "
            "Assessment of motor recovery (Fugl-Meyer et al., 1975), the Action Research Arm Test and the Barthel "
            "Index (Wade, 1992), assessed at baseline, six and twelve weeks by an assessor blind to allocation.",
        ]),
        (2, "4.4 Analysis", [
            "Feasibility outcomes will be reported descriptively with 95% confidence intervals. Clinical outcomes will "
            "be summarised by arm without formal hypothesis testing; their standard deviations will inform the sample "
            "size of a definitive trial. Interviews with 15 participants and 8 therapists will be analysed using "
            "framework analysis.",
        ]),
        (1, "5. Budget and Justification", [
            "The table below summarises the costs requested from the Fund for the 24-month project.",
            ("TABLE", None),
            f"Staff costs: one full-time research assistant for {r['months']} months at {gbp(r['salary'])} per annum "
            f"plus {int(r['oncost'] * 100)}% employer on-costs ({gbp(b['staff'])}). The research assistant will "
            "recruit participants, deliver the home visits jointly with NHS physiotherapists and collect all outcome "
            "data.",
            f"Clinical sessions: {r['visits']} home physiotherapy visits at the NHS community rate of {gbp(r['fee'])} "
            f"per visit ({gbp(b['sessions'])}).",
            f"Travel: {r['visits']} home visits averaging {r['miles']} miles at {gbp(r['mileage'])} per mile, plus "
            f"one national conference ({gbp(r['conference'])}), in total {gbp(b['travel'])}.",
            f"Equipment: {r['tablets']} tablets at {gbp(r['tablet_cost'])} each ({gbp(r['tablets'] * r['tablet_cost'])}) "
            f"and one portable motion capture system ({gbp(r['mocap'])}) for objective gait assessment.",
            f"Consumables ({gbp(b['consumables'])}) cover mobile data plans, exercise bands and printed materials; "
            f"patient and public involvement ({gbp(b['ppi'])}) covers {r['ppi_n']} advisory group meetings at "
            f"{gbp(r['ppi_fee'])} per member per meeting.",
            f"Indirect costs are charged at {pct}% of direct costs ({gbp(b['indirect'])}). The total requested from the "
            f"Fund is {gbp(b['total'])}.",
        ]),
        (1, "6. Impact and Dissemination", [
            "If TeleRehab-Home proves feasible, the results will support an application for a definitive multicentre "
            "trial. The programme has been co-designed with the Cumbria Stroke Survivors Group, who will co-author a "
            "plain-language summary of the findings. Results will be presented to the regional stroke network, "
            "published open access and shared with the three participating NHS trusts through their community "
            "therapy leads.",
        ]),
        (1, "7. References", [("REFS", None)]),
    ]


def _letter(path):
    p = Page()
    p.text(90, 80, "NORTHERN HEALTH RESEARCH FUND", 38, "bold")
    p.text(90, 130, "Grants Office - 12 Rose Street, York YO1 7HB - grants@nhrf.org.uk", 22)
    p.hline(90, 1150, 170, 3)
    y = 200
    for ln in ("2 October 2026", "", "To: all applicants, Community Health Innovation call 2026",
               "Ref: NHRF/CHI/2026 - revised cost rules"):
        y = p.text(90, y, ln, 25, "bold" if ln.startswith("Ref") else "sans")
    y += 20
    y = p.para(90, y, "Dear applicant,", 25)
    y += 10
    y = p.para(90, y, "Following the decision of the Board of Trustees on 30 September 2026, the cost rules below "
                      "replace section 6 of the call guidance with immediate effect. They apply to every application "
                      "to this call, including applications already in preparation.", 25, width=1060)
    y += 15
    rules = ["(a) Indirect costs may be requested at no more than 15% of total direct costs (previously 20%).",
             "(b) Research staff salaries are eligible up to a maximum of £32,000 per annum (full-time equivalent, "
             "before on-costs). Employer on-costs remain eligible at the institution's standard rate.",
             "(c) Travel by private car is reimbursed at £0.35 per mile (previously £0.45 per mile).",
             "(d) Single items of equipment costing more than £5,000 (excluding VAT) are no longer eligible and must "
             "be removed from the budget.",
             "(e) Payments for patient and public involvement remain at £150 per member per meeting."]
    for r in rules:
        y = p.para(120, y, r, 25, width=1020) + 12
    y += 10
    y = p.para(90, y, "Applications that do not comply with these rules will be returned without review. The closing "
                      "date of 13 November 2026 is unchanged.", 25, width=1060)
    y += 30
    p.text(90, y, "Yours faithfully,", 25)
    p.hand(100, y + 45, "R. Whitfield", 44, (20, 30, 110), angle=3, shear=0.35)
    p.text(90, y + 130, "Dr Ruth Whitfield, Head of Grants", 24)
    save_scan([p], path, seed=7301)


ANNEX = [
    ("1.", "Page: A4 portrait (210 x 297 mm) with margins of 2.5 cm on all four sides."),
    ("2.", "Body text, including tables: Liberation Sans or Arial, 11 pt, line spacing 1.15."),
    ("3.", "Section headings (1., 2., ...) must use the paragraph style \"Heading 1\" set to 14 pt bold; sub-section "
           "headings (4.1, 4.2, ...) must use \"Heading 2\" set to 12 pt bold."),
    ("4.", "Header on every page: the application reference only, right-aligned. No other header text (remove "
           "draft markings)."),
    ("5.", "Footer on every page: page numbers centred, in the form \"Page X of Y\"."),
    ("6.", "References: in alphabetical order of the first author's surname, each with a hanging indent of 1 cm."),
    ("7.", "Length: no more than 6 pages in total."),
    ("8.", "Submit one PDF named <application reference>_<PI surname>.pdf, e.g. NHRF-2026-0001_Smith.pdf."),
]


def _annex(path):
    p = Page()
    p.text(90, 80, "NORTHERN HEALTH RESEARCH FUND", 38, "bold")
    p.text(90, 130, "Community Health Innovation call 2026", 24)
    p.hline(90, 1150, 170, 3)
    p.rect((90, 195, 1150, 345), outline=(60, 60, 60), w=2)
    p.text(115, 215, "Application registered", 25, "bold")
    p.para(115, 255, f"Your outline was accepted and registered under reference {REF} (PI: Dr A. Osei, University of "
                     "Eastfield). Quote this reference in all correspondence.", 24, width=1010)
    y = 390
    p.text(90, y, "ANNEX C - FORMATTING REQUIREMENTS FOR FULL PROPOSALS", 30, "bold")
    y += 60
    y = p.para(90, y, "Full proposals that do not follow these requirements will be returned without review.", 24)
    y += 15
    for n, rule in ANNEX:
        p.text(95, y, n, 25, "bold")
        y = p.para(140, y, rule, 25, width=1000) + 14
    p.stamp(820, 1450, ["NHRF", "GRANTS OFFICE"], color=(30, 60, 140), size=28, angle=-7)
    save_scan([p], path, seed=7302)


def build(dst):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from openpyxl import Workbook
    from openpyxl.styles import Font
    os.makedirs(f"{dst}/Grant", exist_ok=True)
    old, new = budget(RATES_OLD), budget(RATES_NEW)
    pct_old = int(RATES_OLD["indirect"] * 100)

    # ---------------------------------------------------------------- proposal (deliberately off-template)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # US Letter (template came from a US co-applicant)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Liberation Serif", Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.0
    hp = sec.header.paragraphs[0]
    hp.text = "DRAFT - internal - do not circulate"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def para(text, size=12, bold=False, italic=False, align=None, indent=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name, run.font.size, run.font.bold, run.font.italic = "Liberation Serif", Pt(size), bold, italic
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.left_indent = indent
        return p

    para(TITLE, 18, bold=True)
    para("Full proposal to the Northern Health Research Fund - Community Health Innovation call 2026", 12, italic=True)
    cover = doc.add_table(rows=0, cols=2)
    cover.style = "Table Grid"
    for k, v in (("Principal investigator", "Dr Amara Osei, School of Health Sciences, University of Eastfield"),
                 ("Co-investigators", "Dr Kieran Brennan (University of Eastfield); Tom Hughes MCSP (Cumbria Community "
                                      "Health NHS Trust)"),
                 ("Duration", "24 months (1 April 2027 - 31 March 2029)"),
                 ("Amount requested", gbp(old["total"]))):
        row = cover.add_row().cells
        for cell, text, bold in ((row[0], k, True), (row[1], v, False)):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.font.name, run.font.size, run.font.bold = "Liberation Serif", Pt(12), bold
    doc.add_paragraph()
    for level, title, paras in sections(old, pct_old):
        para(title, 14 if level == 1 else 12, bold=True)
        for t in paras:
            if t == ("TABLE", None):
                tbl = doc.add_table(rows=1, cols=2)
                tbl.style = "Table Grid"
                for cell, text in zip(tbl.rows[0].cells, ("Cost category", "Amount (GBP)")):
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(text)
                    run.font.name, run.font.size, run.font.bold = "Liberation Serif", Pt(12), True
                for key, label in ROWS:
                    cells = tbl.add_row().cells
                    for cell, text, al in ((cells[0], label.format(pct=pct_old), None),
                                           (cells[1], gbp(old[key]), WD_ALIGN_PARAGRAPH.RIGHT)):
                        cell.text = ""
                        run = cell.paragraphs[0].add_run(text)
                        run.font.name, run.font.size = "Liberation Serif", Pt(12)
                        run.font.bold = key in ("direct", "total")
                        if al is not None:
                            cell.paragraphs[0].alignment = al
                doc.add_paragraph()
            elif t == ("REFS", None):
                for ref in REFS:
                    para(ref, 12)
            else:
                para(t, 12)
    doc.core_properties.author = "Amara Osei"
    doc.core_properties.title = "NHRF full proposal - TeleRehab-Home"
    doc.save(f"{dst}/Grant/NHRF_proposal_Osei.docx")

    # ---------------------------------------------------------------- budget workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"], ws["A1"].font = "Budget assumptions - TeleRehab-Home (NHRF CHI 2026)", Font(bold=True, size=13)
    rates = [("salary", "Research assistant salary (GBP per annum, FTE)"), ("oncost", "Employer on-costs (fraction)"),
             ("months", "Project duration (months)"), ("visits", "Home physiotherapy visits"),
             ("miles", "Average round-trip miles per visit"), ("mileage", "Mileage rate (GBP per mile)"),
             ("fee", "NHS community physiotherapy rate (GBP per visit)"), ("conference", "Conference (GBP)"),
             ("tablets", "Tablets (number)"), ("tablet_cost", "Tablet unit cost (GBP)"),
             ("consumables", "Consumables (GBP)"), ("ppi_n", "PPI advisory group meetings"),
             ("ppi_fee", "PPI payment per meeting (GBP)"), ("indirect", "Indirect cost rate (fraction of direct costs)")]
    ws["A3"], ws["B3"] = "Parameter", "Value"
    ws["A3"].font = ws["B3"].font = Font(bold=True)
    cell = {}
    for i, (k, label) in enumerate(rates, 4):
        ws[f"A{i}"], ws[f"B{i}"] = label, RATES_OLD[k]
        cell[k] = f"Rates!$B${i}"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 14
    wb2 = wb.create_sheet("Budget")
    wb2["A1"], wb2["A1"].font = "Budget summary (GBP)", Font(bold=True, size=13)
    wb2["A3"], wb2["B3"] = "Cost line", "Amount (GBP)"
    wb2["A3"].font = wb2["B3"].font = Font(bold=True)
    c = cell
    lines = [("Staff costs - research assistant", f"={c['salary']}*(1+{c['oncost']})*{c['months']}/12"),
             ("Clinical sessions", f"={c['visits']}*{c['fee']}"),
             ("Travel - home visits (mileage)", f"={c['visits']}*{c['miles']}*{c['mileage']}"),
             ("Travel - conference", f"={c['conference']}"),
             ("Equipment - tablets", f"={c['tablets']}*{c['tablet_cost']}"),
             ("Equipment - portable motion capture system", RATES_OLD["mocap"]),
             ("Consumables", f"={c['consumables']}"),
             ("Patient and public involvement", f"={c['ppi_n']}*{c['ppi_fee']}")]
    for i, (label, f) in enumerate(lines, 4):
        wb2[f"A{i}"], wb2[f"B{i}"] = label, f
        wb2[f"B{i}"].number_format = "#,##0.00"
    end = 3 + len(lines)
    wb2[f"A{end + 1}"], wb2[f"B{end + 1}"] = "Direct costs subtotal", f"=SUM(B4:B{end})"
    wb2[f"A{end + 2}"], wb2[f"B{end + 2}"] = "Indirect costs", f"=B{end + 1}*{c['indirect']}"
    wb2[f"A{end + 3}"], wb2[f"B{end + 3}"] = "Total requested", f"=B{end + 1}+B{end + 2}"
    for r in range(end + 1, end + 4):
        wb2[f"B{r}"].number_format = "#,##0.00"
        wb2[f"A{r}"].font = wb2[f"B{r}"].font = Font(bold=True)
    wb2.column_dimensions["A"].width, wb2.column_dimensions["B"].width = 46, 16
    wb.save(f"{dst}/Grant/NHRF_budget_Osei.xlsx")

    _letter(f"{dst}/Grant/NHRF_cost_rules_letter.pdf")
    _annex(f"{dst}/Grant/NHRF_Annex_C_formatting.pdf")
    heads1 = [t for lvl, t, _ in sections(old, pct_old) if lvl == 1]
    heads2 = [t for lvl, t, _ in sections(old, pct_old) if lvl == 2]
    json.dump({"ref": REF, "pdf_name": f"{REF}_Osei.pdf", "rates_old": RATES_OLD, "rates_new": RATES_NEW,
               "old": old, "new": new, "rows": ROWS, "pct_new": int(RATES_NEW["indirect"] * 100),
               "heading1": heads1, "heading2": heads2, "refs": REFS,
               "refs_sorted": sorted(REFS, key=lambda s: s.split()[0].lower()),
               "stale_numbers": [old["total"], RATES_OLD["salary"], old["staff"], old["travel"], old["equipment"],
                                 old["direct"], old["indirect"], RATES_OLD["mocap"], RATES_OLD["mileage"]],
               "fresh_numbers": [new["total"], RATES_NEW["salary"], new["staff"], new["travel"], new["equipment"],
                                 new["direct"], new["indirect"], RATES_NEW["mileage"]]},
              open(f"{dst}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/o3")
