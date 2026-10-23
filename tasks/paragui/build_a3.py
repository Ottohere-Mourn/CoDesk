"""a3 - one open Calc stock workbook at Copper Lane Cycles (bike shop workshop):
A (pga3goodsin) books three scanned supplier delivery notes into the stock sheet (add received qty, Last received,
  Goods-in log, back-order list);
B (pga3stocktake) applies last Friday's scanned, handwritten stock-take count sheets (adjust On hand by the variance,
  Last counted, delete discontinued rows, insert new SKUs in order, Variance sheet).
Both edit the same On hand cells (read-modify-write) and B's row deletes/inserts shift the rows A edits.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import random
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_x import Page, scanned_pdf  # noqa: E402

SHOP = "Copper Lane Cycles"
ADDR = "Unit 4, Copper Lane, Kendal LA9 4QT"
CATS = {"BAR": "Bar tape & grips", "BRG": "Bearings & BBs", "BRK": "Brakes", "CBL": "Cables", "CHN": "Chains",
        "CST": "Cassettes", "LGT": "Lights", "TUB": "Tubes & tubeless", "TYR": "Tyres"}
BINS = {"BAR": "G", "BRG": "D", "BRK": "B", "CBL": "B", "CHN": "T", "CST": "D", "LGT": "L", "TUB": "T", "TYR": "T"}

CATALOGUE = [  # SKU, description, reorder point
    ("BAR-201", "Lizard Skins DSP 2.5 bar tape black", 4), ("BAR-204", "Fizik Vento Microtex tape black", 3),
    ("BAR-207", "ODI Elite Pro lock-on grips", 3), ("BAR-210", "Ergon GA3 grips large", 2),
    ("BAR-213", "Supacaz Super Sticky Kush tape", 2), ("BAR-216", "Brooks Cambium rubber tape", 1),
    ("BAR-219", "Velox bar end plugs (pair)", 10), ("BAR-222", "Ritchey Comp True Grip", 2),
    ("BRG-301", "6901-2RS cartridge bearing", 10), ("BRG-303", "6902-2RS cartridge bearing", 10),
    ("BRG-305", "6802-2RS cartridge bearing", 8), ("BRG-307", "Steel balls 3/16in (bag 100)", 2),
    ("BRG-309", "Shimano BB-RS500 BSA bottom bracket", 3), ("BRG-311", "SRAM DUB BSA bottom bracket", 2),
    ("BRG-313", "Headset bearing 1-1/8in 45x45", 4), ("BRG-315", "Hope PF41 bottom bracket", 1),
    ("BRK-101", "Shimano B01S resin disc pads", 10), ("BRK-104", "Shimano E01S resin disc pads", 2),
    ("BRK-107", "SRAM Level organic disc pads", 6), ("BRK-110", "Clarks VRX820 disc pads", 4),
    ("BRK-113", "SwissStop Flash Pro rim pads", 4), ("BRK-116", "Kool-Stop Salmon V-brake pads", 6),
    ("BRK-119", "Shimano RT-EM600 rotor 160mm", 3), ("BRK-122", "Shimano SM-RT64 rotor 180mm", 3),
    ("BRK-125", "SRAM Centerline rotor 160mm", 2), ("BRK-128", "Shimano mineral oil 100ml", 3),
    ("BRK-131", "SRAM DOT 5.1 fluid 120ml", 2),
    ("CBL-401", "Jagwire brake inner 1.6x1700 slick", 20), ("CBL-404", "Jagwire gear inner 1.1x2300", 25),
    ("CBL-407", "Shimano SP41 gear outer (per m)", 10), ("CBL-410", "Jagwire brake outer 5mm (per m)", 10),
    ("CBL-413", "Ferrules 4mm (bag 50)", 2), ("CBL-416", "Shimano BH59 hydraulic hose 1000mm", 3),
    ("CBL-419", "Cable end crimps (bag 100)", 1), ("CBL-422", "Jagwire Pro gear cable kit", 3),
    ("CHN-501", "KMC X8 chain 116L", 6), ("CHN-504", "KMC X10 chain 116L", 6),
    ("CHN-507", "Shimano CN-HG601 11-speed chain", 5), ("CHN-510", "SRAM PC-1110 11-speed chain", 3),
    ("CHN-513", "KMC X12 chain 126L", 3), ("CHN-516", "KMC MissingLink 11-speed (pair)", 12),
    ("CHN-519", "Shimano SM-CN910 quick link 12sp", 8), ("CHN-522", "KMC Z1 single-speed chain", 2),
    ("CST-601", "Shimano CS-HG50 8sp 11-32", 3), ("CST-604", "Shimano CS-HG500 10sp 11-34", 2),
    ("CST-607", "Shimano CS-R7000 11sp 11-30", 2), ("CST-610", "SRAM PG-1130 11-42", 1),
    ("CST-613", "Sunrace CSMS8 11-46", 1), ("CST-616", "Shimano CS-M7100 10-51", 1),
    ("LGT-310", "Lezyne Femto Drive front", 3), ("LGT-315", "Knog Blinder rear", 2),
    ("LGT-320", "Lezyne Strip Drive rear", 3), ("LGT-330", "Cateye Volt 800 front", 2),
    ("LGT-335", "Moon Comet front", 2), ("LGT-340", "Busch+Muller IQ-X dynamo light", 1),
    ("TUB-701", "Schwalbe SV17 700x18-28 presta", 20), ("TUB-704", "Schwalbe SV20 700x28-45 presta", 15),
    ("TUB-707", "Schwalbe SV13 26x1.5-2.4 presta", 5), ("TUB-710", "Schwalbe SV19 27.5x1.5-2.4", 6),
    ("TUB-716", "Continental Race 28 700x20-25 60mm", 6), ("TUB-719", "Schwalbe AV15 20x1.5", 3),
    ("TUB-722", "Muc-Off tubeless sealant 140ml", 4), ("TUB-725", "Stan's tubeless valves 44mm (pair)", 4),
    ("TYR-803", "Schwalbe Marathon Plus 700x35c", 4), ("TYR-806", "Schwalbe Marathon 700x38c", 3),
    ("TYR-809", "Continental GP5000 700x25", 4), ("TYR-812", "Continental GP5000 700x28", 4),
    ("TYR-815", "Schwalbe Big Apple 26x2.15", 2), ("TYR-818", "Maxxis Minion DHF 29x2.5", 2),
    ("TYR-821", "Schwalbe G-One Allround 700x40", 2), ("TYR-827", "Vittoria Rubino Pro 700x25", 4),
    ("TYR-830", "Michelin Power Cup 700x28", 2),
]

# On hand now (Thu 5 Nov 2026, morning) for the SKUs the scenario cares about; the rest are random.
SEED_ONHAND = {"BRK-101": 12, "BRK-104": 3, "BRK-107": 7, "BRK-113": 9, "BRK-119": 5, "BRK-122": 1,
               "CBL-404": 44, "CBL-410": 12, "CBL-413": 3, "CHN-507": 6, "CHN-510": 4, "CHN-516": 16,
               "TUB-701": 27, "TUB-704": 22, "TUB-707": 6, "TYR-803": 6, "TYR-812": 3, "TYR-821": 2}

# delivery notes: supplier, note no, date, lines (supplier code, our ref printed, description, ordered, delivered, b/o)
DELIVERIES = [
    {"supplier": "Madison Cycle Parts", "no": "MCP-558120", "date": "2026-11-02", "addr": "Abbot's Hill, Milton Keynes MK10 0AB",
     "lines": [("SHBP-B01S", "BRK-101", "Shimano B01S resin disc pads", 10, 10, 0),
               ("SRLV-ORG", "BRK-107", "SRAM Level organic disc pads", 6, 6, 0),
               ("SHRT-RT64-180", "BRK-122", "Shimano SM-RT64 rotor 180mm", 4, 0, 4),
               ("SHRT-EM600-160", "BRK-119", "Shimano RT-EM600 rotor 160mm", 4, 4, 0),
               ("JWGI-1123", "CBL-404", "Jagwire gear inner 1.1x2300", 20, 20, 0)]},
    {"supplier": "Bikeline Wholesale", "no": "BW-77341", "date": "2026-11-03", "addr": "Riverside Park, Leeds LS11 5JY",
     "lines": [("HG601-116", "CHN-507", "Shimano CN-HG601 11-speed chain", 6, 6, 0),
               ("KMCML11-P", "CHN-516", "KMC MissingLink 11-speed (pair)", 10, 10, 0),
               ("SW-SV20", "TUB-704", "Schwalbe SV20 700x28-45 presta", 12, 12, 0),
               ("SW-MP35", "TYR-803", "Schwalbe Marathon Plus 700x35c", 4, 4, 0)]},
    {"supplier": "Velotech Distribution", "no": "VT-20931", "date": "2026-11-04", "addr": "Unit 9, Hamilton Way, Coventry CV3 2SX",
     "lines": [("CO-GP5K-28", "TYR-813", "Continental GP5000 700x28", 6, 6, 0),
               ("SW-SV17", "TUB-701", "Schwalbe SV17 700x18-28 presta", 20, 12, 8),
               ("FER4-50", "CBL-413", "Ferrules 4mm (bag 50)", 2, 2, 0),
               ("SW-GONE-40", "TYR-821", "Schwalbe G-One Allround 700x40", 4, 4, 0)]},
]
DAMAGED = {"TYR-803": 1}          # one tyre damaged, sent back with the driver (handwritten on BW-77341)
REF_TYPO = {"TYR-813": "TYR-812"}  # Velotech printed a wrong ref; the description identifies our SKU

# stock-take Friday 30 Oct 2026: SKU, printed system qty, counted, remark
COUNT_DATE = "2026-10-30"
COUNT_SHEETS = [
    {"bay": "B", "title": "Bay B - brakes & cables", "lines": [
        ("BRK-101", 14, 13, ""), ("BRK-104", 3, 3, "DISC - remove"), ("BRK-107", 8, 10, ""), ("BRK-113", 9, 9, ""),
        ("BRK-119", 5, 4, ""), ("CBL-404", 46, 41, ""), ("CBL-410", 12, 11, "DISC - remove"), ("CBL-413", 3, 3, "")],
     "new": [("CBL-411", "Jagwire 2X brake cable kit", "B-07", 5)]},
    {"bay": "T", "title": "Bay T - chains, tubes & tyres", "lines": [
        ("CHN-507", 7, 6, ""), ("CHN-510", 5, 5, ""), ("CHN-516", 18, 22, ""), ("TUB-701", 31, 28, ""),
        ("TUB-704", 24, 24, ""), ("TUB-707", 6, 6, "DISC - remove"), ("TYR-803", 6, 4, ""), ("TYR-812", 4, 6, "")],
     "new": [("TYR-824", "Panaracer GravelKing SK 700x38", "T-14", 4)]},
]


def _catalogue_rows(rng):
    rows = []
    for k, (sku, desc, rp) in enumerate(CATALOGUE):
        pre = sku[:3]
        onhand = SEED_ONHAND.get(sku, rng.randint(0, 3 * rp + 4))
        bin_ = f"{BINS[pre]}-{(k % 14) + 1:02d}"
        last_rec = dt.date(2026, 10, 1) + dt.timedelta(days=rng.randrange(0, 26))
        last_cnt = dt.date(2026, 7, 31) if BINS[pre] in ("B", "T") else dt.date(2026, 4, 30)
        notes = ""
        if sku == "BRK-104":
            notes = "Superseded by E01S/B01S - run down"
        if sku == "TUB-707":
            notes = "26in - low demand"
        rows.append({"sku": sku, "desc": desc, "cat": CATS[pre], "bin": bin_, "onhand": onhand, "reorder": rp,
                     "last_rec": last_rec.isoformat(), "last_cnt": last_cnt.isoformat(), "notes": notes})
    return rows


def _workbook(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    head = ["SKU", "Description", "Category", "Bin", "On hand", "Reorder point", "Last received", "Last counted", "Notes"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
    for r in rows:
        ws.append([r["sku"], r["desc"], r["cat"], r["bin"], r["onhand"], r["reorder"],
                   dt.date.fromisoformat(r["last_rec"]), dt.date.fromisoformat(r["last_cnt"]), r["notes"]])
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = "DD/MM/YYYY"
        row[7].number_format = "DD/MM/YYYY"
        row[4].alignment = Alignment(horizontal="right")
    for col, w in zip("ABCDEFGHI", (10, 38, 17, 7, 9, 13, 13, 13, 32)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    gi = wb.create_sheet("Goods in")
    gi.append(["Date", "Supplier", "Delivery note", "SKU", "Qty", "Booked by"])
    for c in gi[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
    old = [("2026-10-06", "Madison Cycle Parts", "MCP-554871", "BRK-128", 4), ("2026-10-06", "Madison Cycle Parts", "MCP-554871", "CBL-401", 25),
           ("2026-10-13", "Bikeline Wholesale", "BW-76902", "CHN-501", 6), ("2026-10-13", "Bikeline Wholesale", "BW-76902", "TUB-716", 10),
           ("2026-10-20", "Velotech Distribution", "VT-20514", "TYR-809", 6), ("2026-10-20", "Velotech Distribution", "VT-20514", "LGT-320", 4)]
    for d, s, n, sku, q in old:
        gi.append([dt.date.fromisoformat(d), s, n, sku, q, "DK"])
    for row in gi.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
    for col, w in zip("ABCDEF", (12, 24, 15, 10, 7, 10)):
        gi.column_dimensions[col].width = w
    wb.save(path)


def _delivery_page(dn, seed):
    p = Page()
    p.text(80, 70, dn["supplier"].upper(), size=46, bold=True)
    p.text(80, 135, dn["addr"] + "  -  trade sales", size=22, fill=(70, 70, 70))
    p.text(1160, 80, "DELIVERY NOTE", size=40, bold=True, anchor="ra")
    p.hline(80, 1160, 185)
    y = p.lines(80, 215, ["Deliver to:", SHOP + " - Workshop", ADDR, "Account: CLC-0042"], size=24)
    d = dt.date.fromisoformat(dn["date"])
    p.lines(760, 215, [f"Note no:  {dn['no']}", f"Date:  {d.strftime('%d/%m/%Y')}", "Carrier: own van",
                       f"Your order: PO-26-{310 + seed}"], size=24)
    cols = [("Item code", 230, "l"), ("Your ref", 140, "l"), ("Description", 420, "l"), ("Ordered", 110, "r"),
            ("Delivered", 120, "r"), ("B/O", 60, "r")]
    rows = [(c, ref, desc, str(o), str(dv), str(bo) if bo else "-") for c, ref, desc, o, dv, bo in dn["lines"]]
    y = p.table(70, 420, cols, rows, size=21, row_h=46)
    p.text(80, y + 30, "B/O = back-ordered, will follow on a separate delivery note.", size=20, fill=(80, 80, 80))
    p.text(80, y + 62, "Goods remain the property of the supplier until paid in full. Report discrepancies within 48h.",
           size=18, fill=(100, 100, 100))
    p.text(80, 1500, "Received by (print / sign): ____________________", size=24)
    if dn["no"] == "BW-77341":
        # handwritten correction by the mechanic who signed for the delivery
        row_y = 420 + 46 * 4 + 23  # 4th data row = Marathon Plus; "Delivered" column spans x 970..1090
        p.d.line((1030, row_y + 12, 1088, row_y - 12), fill=(20, 40, 120), width=3)
        p.handwriting(1160, row_y - 22, "3", size=40, seed=10)
        p.handwriting(80, y + 110, "1 x Marathon Plus 700x35 - sidewall cut,", size=34, seed=11)
        p.handwriting(80, y + 155, "returned with driver. Kept 3.  - DK", size=34, seed=12)
        p.handwriting(470, 1490, "D. Kirby", size=40, seed=13)
    else:
        p.handwriting(470, 1490, "D. Kirby", size=40, seed=14 + seed)
        p.stamp(860, 1380, "RECEIVED", angle=-8)
    return p


def _count_page(sheet, seed):
    p = Page()
    p.text(80, 70, f"{SHOP} - stock count", size=40, bold=True)
    p.text(80, 130, f"{sheet['title']}   |   count date Friday 30/10/2026", size=26)
    p.text(80, 172, "System qty = quantity on the stock sheet when this list was printed (30/10/2026 17:05).", size=20,
           fill=(70, 70, 70))
    p.text(80, 200, "Stock can move after printing - write the COUNTED quantity only. Mark discontinued lines.", size=20,
           fill=(70, 70, 70))
    cols = [("SKU", 140, "l"), ("Description", 440, "l"), ("Bin", 90, "l"), ("System qty", 140, "r"),
            ("Counted", 140, "r"), ("Remarks", 150, "l")]
    desc = {s: d for s, d, _ in CATALOGUE}
    rows = [(s, desc[s], "", str(sys_q), "", "") for s, sys_q, _, _ in sheet["lines"]]
    top, rh = 260, 62
    y = p.table(60, top, cols, rows, size=21, row_h=rh)
    r = random.Random(seed)
    for i, (s, sys_q, cnt, rem) in enumerate(sheet["lines"]):
        cy = top + rh * (i + 1) + 12
        p.handwriting(60 + 140 + 440 + 90 + 140 + 45 + r.randint(-6, 6), cy, str(cnt), size=38, seed=seed * 10 + i)
        if rem:
            p.handwriting(60 + 140 + 440 + 90 + 140 + 140 + 8, cy + 4, "DISC", size=30, fill=(150, 20, 20), seed=seed + i)
            p.handwriting(980, cy + 34, "remove", size=22, fill=(150, 20, 20), seed=seed + 50 + i)
    y += 40
    p.handwriting(80, y, "Found on shelf - not on list:", size=34, seed=seed + 7)
    for k, (s, d, b, cnt) in enumerate(sheet["new"]):
        p.handwriting(110, y + 60 + 55 * k, f"NEW  {s}  {d}  bin {b}  - counted {cnt}", size=32, seed=seed + 20 + k)
    p.text(80, 1560, "Counted by: ______________   Checked by: ______________", size=24)
    p.handwriting(300, 1545, "Mo", size=38, seed=seed + 31)
    p.handwriting(860, 1545, "DK", size=38, seed=seed + 32)
    return p


def build(out_dir):
    os.makedirs(f"{out_dir}/Deliveries", exist_ok=True)
    os.makedirs(f"{out_dir}/Stocktake", exist_ok=True)
    rng = random.Random(20261105)
    rows = _catalogue_rows(rng)
    _workbook(f"{out_dir}/parts_stock.xlsx", rows)
    files = {}
    for k, dn in enumerate(DELIVERIES):
        name = f"Deliveries/{dn['no']}_{dn['date']}.pdf"
        scanned_pdf(f"{out_dir}/{name}", [_delivery_page(dn, k)], seed=100 + k)
        files[name] = f"Documents/Stock/{name}"
    for k, sh in enumerate(COUNT_SHEETS):
        name = f"Stocktake/count_bay{sh['bay']}_2026-10-30.pdf"
        scanned_pdf(f"{out_dir}/{name}", [_count_page(sh, 300 + k)], seed=200 + k)
        files[name] = f"Documents/Stock/{name}"

    seed = {r["sku"]: r for r in rows}
    receipts, last_rec, booked_lines, backorders = {}, {}, [], []
    for dn in DELIVERIES:
        for code, ref, desc, ordered, delivered, bo in dn["lines"]:
            sku = REF_TYPO.get(ref, ref)
            qty = delivered - DAMAGED.get(sku, 0)
            if bo:
                backorders.append({"supplier": dn["supplier"], "note": dn["no"], "sku": sku, "desc": desc, "qty": bo})
            if qty > 0:
                receipts[sku] = receipts.get(sku, 0) + qty
                last_rec[sku] = dn["date"]
                booked_lines.append({"date": dn["date"], "supplier": dn["supplier"], "note": dn["no"], "sku": sku, "qty": qty})
    variances, counted, discs, new = {}, [], [], []
    for sh in COUNT_SHEETS:
        for sku, sys_q, cnt, rem in sh["lines"]:
            if rem:
                discs.append(sku)
            else:
                counted.append(sku)
                variances[sku] = cnt - sys_q
        for sku, d, b, cnt in sh["new"]:
            new.append({"sku": sku, "desc": d, "bin": b, "counted": cnt, "cat": CATS[sku[:3]]})
    assert all(s in seed for s in receipts) and all(s in seed for s in variances)
    exp = {
        "workbook": "Documents/Stock/parts_stock.xlsx", "files": files,
        "seed": {s: {"onhand": r["onhand"], "last_rec": r["last_rec"], "last_cnt": r["last_cnt"], "desc": r["desc"],
                     "cat": r["cat"]} for s, r in seed.items()},
        "order": [r["sku"] for r in rows],
        "receipts": receipts, "last_received": last_rec, "booked_lines": booked_lines, "backorders": backorders,
        "damaged": DAMAGED, "ref_typo": REF_TYPO, "printed_system": {s: q for sh in COUNT_SHEETS for s, q, _, _ in sh["lines"]},
        "counted_qty": {s: c for sh in COUNT_SHEETS for s, _, c, _ in sh["lines"]},
        "count_date": COUNT_DATE, "variances": variances, "counted": counted, "discontinued": discs, "new": new,
    }
    json.dump(exp, open(f"{out_dir}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/a3_assets")
