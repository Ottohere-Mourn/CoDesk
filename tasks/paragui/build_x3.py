"""X3 - Brightwater Studio's 15th-anniversary client evening: RSVPs + door list vs the catering order + budget.

build(out_dir) writes:
  app_x3.py                    Brightwater Hub web app (copied)
  hub_seed.db                  seeded sqlite db (2 events, 36 invitees + 6 hosts, suppliers, catalogues, 4 purchase orders)
  Anniversary_Budget_2026.xlsx the event budget sheet (uploaded to ~/Documents/Events)
  thunderbird-profile.tar.gz   Rosa Lindqvist's offline mailbox (16 RSVP mails incl. a photographed reply card)
  expected.json                ground truth (host only)
"""
from __future__ import annotations

import datetime as dt
import json
import os
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from email.utils import parsedate_to_datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lib_x as X  # noqa: E402
from lib_thunderbird import make_profile  # noqa: E402

ME = ("Rosa Lindqvist", "rosa@brightwater-studio.co.uk")
MARTIN = ("Martin Achebe", "martin.achebe@brightwater-studio.co.uk")
CLARA = ("Clara Benn", "events@theglasshouse-bristol.co.uk")
DEV = ("Dev Sharma", "orders@fennelandfig.co.uk")
TZ = dt.timezone(dt.timedelta(hours=0))
SUBJ = "Invitation: Brightwater Studio at 15 - Thursday 19 November"
BEEF, BASS, RISOTTO, VEGAN = "FF-M01", "FF-M02", "FF-M03", "FF-M04"
MAIN_NAME = {BEEF: "Fillet of beef", BASS: "Pan-roasted sea bass", RISOTTO: "Wild mushroom risotto (v)",
             VEGAN: "Squash & lentil wellington (vegan)"}

# invitee name -> (company, email); the 13 who reply are first
INVITEES = [
    ("Helen Park", "Park & Rowe Developments", "helen.park@parkrowe.co.uk"),
    ("James Okoro", "Okoro Engineering", "james.okoro@okoro-eng.co.uk"),
    ("Priya Natarajan", "Bristol City Council - Planning", "priya.natarajan@bristol.gov.uk"),
    ("Tom Reilly", "Reilly Construction", "tom@reillyconstruction.co.uk"),
    ("Anna Kowalski", "Harbourside Trust", "a.kowalski@harbourside-trust.org.uk"),
    ("Ben Carter", "Carter Property Group", "ben.carter@carterpg.co.uk"),
    ("Grace Liu", "Liu Design Partners", "grace@liudesign.co.uk"),
    ("Daniel Fischer", "Avon Housing Association", "d.fischer@avonhousing.org.uk"),
    ("Olivia Grant", "Grant & Sons Joinery", "olivia@grantjoinery.co.uk"),
    ("Fiona MacLeod", "Bristol Architecture Forum", "fiona.macleod@bristolarchforum.org"),
    ("Raj Malhotra", "Malhotra Hotels", "raj@malhotrahotels.co.uk"),
    ("Elena Rossi", "Rossi Structural", "elena@rossistructural.co.uk"),
    ("Hugo Brandt", "Brandt Glass GmbH", "hugo.brandt@brandtglass.de"),
    # no reply yet
    ("Oliver Grant", "Westbury Homes", "oliver.grant@westburyhomes.co.uk"),
    ("Samuel Adeyemi", "Clifton Estates", "s.adeyemi@cliftonestates.co.uk"),
    ("Chloe Bennett", "Bennett Landscape", "chloe@bennettlandscape.co.uk"),
    ("Isabel Moreno", "Moreno Lighting", "isabel@morenolighting.com"),
    ("Kwame Asante", "University of Bristol Estates", "k.asante@bristol.ac.uk"),
    ("Rachel Dunn", "Dunn & Partners Solicitors", "rdunn@dunnpartners.co.uk"),
    ("Stefan Novak", "Novak Timber Frames", "stefan@novaktimber.co.uk"),
    ("Lucy Hartley", "Hartley Interiors", "lucy@hartleyinteriors.co.uk"),
    ("Mohammed Rahman", "Rahman Property Holdings", "m.rahman@rahmanholdings.co.uk"),
    ("Sophie Turner", "Bristol Civic Society", "sophie.turner@bristolcivic.org.uk"),
    ("George Whitfield", "Whitfield Quantity Surveyors", "george@whitfieldqs.co.uk"),
    ("Hannah Cole", "Cole Heritage Consulting", "hannah@coleheritage.co.uk"),
    ("Patrick O'Neill", "O'Neill M&E Services", "patrick@oneill-me.co.uk"),
    ("Amelia Shaw", "Shaw Hotels Group", "amelia.shaw@shawhotels.co.uk"),
    ("Victor Lindgren", "Nordic Timber Imports", "victor@nordictimber.se"),
    ("Zara Ahmed", "Ahmed Architects Collective", "zara@aac-studio.co.uk"),
    ("Neil Foster", "Foster Steelworks", "neil@fostersteel.co.uk"),
    ("Charlotte Price", "Avon Arts Trust", "charlotte@avonarts.org.uk"),
    ("Ian McKenzie", "McKenzie Planning", "ian@mckenzieplanning.co.uk"),
    ("Laura Bianchi", "Bianchi Stone", "laura@bianchistone.it"),
    ("Edward Hughes", "Hughes Developments", "ed@hughesdev.co.uk"),
    ("Mei Chen", "Chen Acoustic Design", "mei@chenacoustics.co.uk"),
    ("Robert Lang", "Lang Glazing", "rob@langglazing.co.uk"),
]
COMPANY = {n: c for n, c, _ in INVITEES}
# expected outcome of the replies: invitee -> (status, plus-one or None)
OUTCOME = {
    "Helen Park": ("Attending", "David Park"), "James Okoro": ("Attending", None), "Priya Natarajan": ("Declined", None),
    "Tom Reilly": ("Attending", "Sarah Reilly"), "Anna Kowalski": ("Attending", "Mark Lewis"), "Ben Carter": ("Declined", None),
    "Grace Liu": ("Attending", None), "Daniel Fischer": ("Declined", None), "Olivia Grant": ("Attending", "Peter Grant"),
    "Fiona MacLeod": ("Attending", None), "Raj Malhotra": ("Attending", "Anita Malhotra"), "Elena Rossi": ("Attending", "Marco Rossi"),
    "Hugo Brandt": ("Attending", None),
}
SUBSTITUTE = ("Laura Mensah", "Avon Housing Association", "laura.mensah@avonhousing.org.uk")
FORWARDED = "Marcus Webb"
# every person eating: (name, main, vegan, soft drinks only, dietary note for the caterer)
DINERS = [
    ("Helen Park", BASS, False, False, ""), ("David Park", BEEF, False, False, "nut allergy (severe)"),
    ("James Okoro", BEEF, False, False, ""), ("Tom Reilly", BEEF, False, False, ""), ("Sarah Reilly", RISOTTO, False, False, ""),
    ("Anna Kowalski", BASS, False, False, ""), ("Mark Lewis", BASS, False, False, ""), ("Grace Liu", VEGAN, True, True, "vegan"),
    ("Laura Mensah", RISOTTO, False, False, "gluten-free"), ("Olivia Grant", BEEF, False, False, ""),
    ("Peter Grant", BASS, False, False, "coeliac - gluten-free"), ("Fiona MacLeod", RISOTTO, False, False, ""),
    ("Raj Malhotra", BASS, False, False, ""), ("Anita Malhotra", VEGAN, True, False, "vegan"), ("Elena Rossi", BEEF, False, False, ""),
    ("Marco Rossi", BEEF, False, False, ""), ("Hugo Brandt", BASS, False, False, "shellfish allergy"),
]
# Brightwater hosts (visible only on the Hub event page): name, role, main, vegan, soft drinks, dietary column text
HOSTS = [
    ("Rosa Lindqvist", "Office manager", RISOTTO, False, False, ""), ("Martin Achebe", "Director", BEEF, False, False, ""),
    ("Nadia Petrova", "Associate director", BASS, False, False, ""), ("Callum Price", "Project architect", BEEF, False, False, ""),
    ("Yuki Tanaka", "Architect", VEGAN, True, False, "Vegan"), ("Owen Hughes", "Architectural technologist", BASS, False, True,
                                                                  "No alcohol - soft drinks"),
]
DIETARY_NAMES = ["David Park", "Peter Grant", "Hugo Brandt", "Laura Mensah"]
# Fennel & Fig catalogue: code, name, unit, price, description
CATALOGUE_FF = [
    ("FF-C01", "Canapé reception", "head", "14.50", "Six canapés per person incl. vegetarian and vegan pieces"),
    ("FF-M01", "Fillet of beef, dauphinoise, red wine jus", "portion", "32.00", "Main course"),
    ("FF-M02", "Pan-roasted sea bass, samphire, beurre blanc", "portion", "29.00", "Main course"),
    ("FF-M03", "Wild mushroom risotto (v)", "portion", "22.00", "Main course, vegetarian"),
    ("FF-M04", "Squash & lentil wellington (vegan)", "portion", "22.00", "Main course, vegan"),
    ("FF-M05", "Children's main (chicken goujons)", "portion", "12.00", "Main course for under-12s"),
    ("FF-D01", "Dark chocolate torte, crème fraîche", "portion", "9.00", "Dessert"),
    ("FF-D02", "Poached pear, blackberry sorbet (vegan)", "portion", "9.00", "Dessert, vegan"),
    ("FF-B01", "Drinks package - 3 hours", "head", "28.00", "Fizz on arrival, house wine, beer, soft drinks"),
    ("FF-B02", "Soft drinks package - 3 hours", "head", "12.00", "Alcohol-free cocktails, juices, soft drinks"),
    ("FF-B03", "Premium wine upgrade", "head", "9.00", "Upgrade to the reserve list"),
    ("FF-L01", "Late bites (bacon / halloumi rolls)", "head", "7.50", "Served at 21:30"),
    ("FF-S01", "Additional waiting staff (per hour)", "hour", "22.00", "Beyond the included team of four"),
]


def D(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"), ROUND_HALF_UP)


def expected_lines():
    q = {}
    everyone = [(n, m, v, s) for n, m, v, s, _ in DINERS] + [(n, m, v, s) for n, _, m, v, s, _ in HOSTS]
    for _, main, vegan, soft in everyone:
        for code in ("FF-C01", main, "FF-D02" if vegan else "FF-D01", "FF-B02" if soft else "FF-B01"):
            q[code] = q.get(code, 0) + 1
    return q


def po_total(lines: dict, service_pct="12.5") -> Decimal:
    price = {c: Decimal(p) for c, _, _, p, _ in CATALOGUE_FF}
    sub = sum((price[c] * n for c, n in lines.items()), Decimal("0"))
    svc = D(sub * Decimal(service_pct) / 100)
    return sub + svc + D((sub + svc) * Decimal("0.2"))


# ---------------------------------------------------------------- documents
def reply_card(path):
    p = X.Page(1500, 1060, bg=(252, 250, 244))
    p.rect((30, 30, 1470, 1030), outline=(18, 52, 59), width=4)
    p.text(750, 70, "BRIGHTWATER STUDIO AT 15", size=44, bold=True, serif=True, fill=(18, 52, 59), anchor="ma")
    p.text(750, 135, "Thursday 19 November 2026 · 6.30 pm · The Glasshouse, Harbourside, Bristol", size=26, serif=True, anchor="ma")
    p.text(750, 180, "Reply card - kindly return by 4 November", size=24, serif=True, fill=(90, 90, 90), anchor="ma")
    y = 260
    for label in ("Name", "Guest (one per invitation)"):
        p.text(90, y, label, size=26, serif=True)
        p.hline(470, 1400, y + 34, width=2, fill=(120, 120, 120))
        y += 90
    p.handwriting(500, 248, "Helen Park", size=40, seed=1)
    p.handwriting(500, 338, "David Park (my husband)", size=40, seed=2)
    p.text(90, 450, "[  ] will attend          [  ] regretfully declines", size=28, serif=True)
    p.handwriting(103, 440, "X", size=40, seed=3)
    p.text(90, 540, "Main course - please write the initial of each person next to their choice:", size=24, serif=True)
    choices = ["Fillet of beef, dauphinoise", "Pan-roasted sea bass, samphire", "Wild mushroom risotto (v)",
               "Squash & lentil wellington (vegan)"]
    for i, c in enumerate(choices):
        p.text(160, 590 + i * 52, "o  " + c, size=26, serif=True)
    p.handwriting(720, 582, "D", size=40, seed=4)
    p.handwriting(720, 634, "H", size=40, seed=5)
    p.text(90, 820, "Dietary requirements / allergies", size=26, serif=True)
    p.hline(560, 1400, 856, width=2, fill=(120, 120, 120))
    p.hline(90, 1400, 930, width=2, fill=(120, 120, 120))
    p.handwriting(580, 812, "David - nut allergy (severe!)", size=40, seed=6)
    X.photo_jpg(path, p, seed=7, size=(1400, 1050))


def budget_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Brightwater Studio at 15 - client evening, Thu 19 Nov 2026 - event budget"
    ws["A1"].font = Font(bold=True, size=13)
    heads = ["Item", "Supplier", "PO / ref", "Amount (£ incl. VAT)", "Status"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(3, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCEBED")
    rows = [("Venue hire", "The Glasshouse", "PO-2026-0141", 3240.00, "Confirmed"),
            ("AV & lighting", "Bristol AV Hire", "PO-2026-0142", 1068.00, "Confirmed"),
            ("Flowers & styling", "Wild Bloom Studio", "Quote Q-882", 720.00, "Quoted"),
            ("Photographer", "Lena Ortiz Photography", "", 540.00, "Booked"),
            ("Catering", "Fennel & Fig Catering", "", None, "To order"),
            ("Guest gifts", "Bristol Tea Co.", "", 380.00, "Ordered")]
    for i, r in enumerate(rows, 4):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            if j == 4:
                c.number_format = "#,##0.00"
    others = sum(r[3] for r in rows if r[3])
    for i, (lbl, v) in enumerate([("Total committed", others), ("Approved budget", 9000.00), ("Remaining", 9000.00 - others)], 11):
        ws.cell(i, 1, lbl).font = Font(bold=True)
        ws.cell(i, 4, v).number_format = "#,##0.00"
    ws["A15"] = "Updated by Rosa, 14 Oct 2026"
    ws["A15"].font = Font(italic=True, color="777777")
    for col, w in zip("ABCDE", (22, 26, 16, 22, 14)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return others, 9000.00


def mailbox(d, card_jpg):
    when = lambda s: dt.datetime.fromisoformat(s).replace(tzinfo=TZ)  # noqa: E731
    ids = {}
    host = "brightwater-studio.co.uk"

    def mid(key, dom):
        ids[key] = f"<x3.{key}.{len(ids) + 1:02d}@{dom}>"
        return ids[key]

    inv_body = ("Dear friends,\n\nBrightwater Studio is fifteen this year, and we would love to celebrate with the clients and "
                "partners who made it possible.\n\nThursday 19 November 2026, 6.30 pm until late\nThe Glasshouse, Harbourside, Bristol\n"
                "Drinks, canapés and a seated dinner\n\nPlease reply to this email by Wednesday 4 November and let us know:\n"
                "- whether you can come, and the name of your guest if you are bringing one (one guest per invitation);\n"
                "- a main course for each of you: Fillet of beef / Pan-roasted sea bass / Wild mushroom risotto (vegetarian) / "
                "Squash & lentil wellington (vegan);\n- any dietary requirements or allergies.\n\n"
                "We very much hope to see you there.\n\nRosa Lindqvist\nOffice Manager, Brightwater Studio\n14 Queen Square, Bristol BS1 4NT")
    inv = X.mail(ME, [ME], SUBJ, inv_body, when("2026-10-12T10:00:00"), mid("invitation", host))
    re_s = "Re: " + SUBJ

    def rsvp(key, frm, body, at, unread=True, atts=(), subject=None, irt=True):
        dom = frm[1].split("@")[1]
        m = X.mail(frm, [ME], subject or re_s, body, when(at), mid(key, dom), irt=ids["invitation"] if irt else None, atts=atts)
        return m, unread, key

    card = open(card_jpg, "rb").read()
    inbox = [
        (X.mail(MARTIN, [ME], "Anniversary evening - arrangements",
                "Hi Rosa,\n\nA few things for the anniversary evening on the 19th now that replies are coming in:\n\n"
                "- Everyone on our invitation list may bring one guest.\n"
                "- If someone asks to come on a forwarded invitation, please don't confirm them or add them to the list - tell me "
                "who it is and I'll decide. If an invitee can't come and sends a colleague in their place, that's fine: mark the "
                "invitee as declined and put the colleague on the list.\n"
                "- The Glasshouse needs the total number of people - our guests plus the six of us hosting - and a door list of "
                "the guests (name and company, not us) by 6 November. Please copy me when you send it.\n"
                "- Catering is with Fennel & Fig. I've started a draft order in Hub (PO-2026-0144) with rough numbers - please "
                "finish it with the real numbers. Everyone, guests and hosts, gets the canapé reception, one main course of their "
                "choice, a dessert and the drinks package. Vegans get the vegan main and the vegan dessert; anyone who doesn't "
                "drink gets the soft-drinks package instead of the drinks package. I added late bites as an idea, but we can't "
                "afford them - take them off. Charge it to the anniversary evening and submit it. Our own choices are on the "
                "event page in Hub.\n"
                "- Keep the budget sheet (Documents/Events/Anniversary_Budget_2026.xlsx) up to date: PO number and amount "
                "including service and VAT, the total and what's left.\n\nThanks,\nMartin",
                when("2026-10-14T08:40:00"), mid("martin", host)), True, "martin"),
        rsvp("james", ("James Okoro", "james.okoro@okoro-eng.co.uk"),
             "Hi Rosa,\n\nThanks for the invitation - I'd be delighted to come. Just me, and the beef for me please. "
             "No dietary requirements.\n\nBest,\nJames", "2026-10-13T09:15:00", False),
        rsvp("priya", ("Priya Natarajan", "priya.natarajan@bristol.gov.uk"),
             "Dear Rosa,\n\nThank you for the kind invitation. Unfortunately I am on leave that week, so I must decline. "
             "Congratulations on fifteen years!\n\nPriya", "2026-10-13T14:02:00", False),
        rsvp("tom", ("Tom Reilly", "tom@reillyconstruction.co.uk"),
             "Rosa - count me in, and I'll bring my wife, Sarah Reilly. Beef for me, Sarah will have the risotto.\n\nCheers,\nTom",
             "2026-10-14T17:31:00", False),
        rsvp("anna", ("Anna Kowalski", "a.kowalski@harbourside-trust.org.uk"),
             "Hello Rosa,\n\nYes please! I'll come with my partner, Mark Lewis. The sea bass sounds lovely - same for both of us.\n\n"
             "Anna", "2026-10-15T11:48:00"),
        rsvp("ben1", ("Ben Carter", "ben.carter@carterpg.co.uk"),
             "Yes, I'll be there - the beef please.\n\nBen", "2026-10-16T08:05:00", False),
        rsvp("grace", ("Grace Liu", "grace@liudesign.co.uk"),
             "Hi Rosa,\n\nI'd love to come - on my own. I'm vegan, so please put me down for the vegan option. I also don't "
             "drink alcohol, so soft drinks are perfect for me.\n\nThank you!\nGrace", "2026-10-19T13:20:00"),
        rsvp("daniel", ("Daniel Fischer", "d.fischer@avonhousing.org.uk"),
             "Dear Rosa,\n\nI'm afraid I am at a conference in Manchester that week. My colleague Laura Mensah "
             "(laura.mensah@avonhousing.org.uk), our Head of Development, would very much like to come in my place if that's "
             "all right. She would like the mushroom risotto and she is gluten-free.\n\nKind regards,\nDaniel Fischer\n"
             "Avon Housing Association", "2026-10-20T10:11:00"),
        rsvp("olivia1", ("Olivia Grant", "olivia@grantjoinery.co.uk"),
             "Hi Rosa,\n\nYes from me, plus my husband Peter Grant. I'll have the beef; Peter is still deciding - I'll let you know.\n\n"
             "Olivia", "2026-10-21T19:44:00"),
        (X.mail(("Callum Price", "callum.price@brightwater-studio.co.uk"), [ME], "Winter party - count me in (+1)",
                "Hi Rosa, count me and Jess in for the winter party on 11 Dec. Both veggie please! Callum",
                when("2026-10-22T12:30:00"), mid("winter", host)), True, "winter"),
        rsvp("fiona", ("Fiona MacLeod", "fiona.macleod@bristolarchforum.org"),
             "Dear Rosa,\n\nDelighted to accept. Just me - the wild mushroom risotto, please.\n\nWarm wishes,\nFiona MacLeod",
             "2026-10-23T09:00:00"),
        rsvp("helen", ("Jo Whitmore", "jo.whitmore@parkrowe.co.uk"),
             "Dear Rosa,\n\nHelen has filled in the reply card that came with your invitation - a photo of it is attached.\n\n"
             "Kind regards,\nJo Whitmore\nPA to Helen Park, Park & Rowe Developments", "2026-10-26T15:27:00",
             atts=[("Reply_card_Helen_Park.jpg", card, "image/jpeg")], subject="RSVP - Helen Park (Park & Rowe Developments)", irt=False),
        rsvp("raj", ("Raj Malhotra", "raj@malhotrahotels.co.uk"),
             "Hello Rosa,\n\nMy wife, Anita Malhotra, and I would be pleased to join you. I'll have the sea bass. Anita is vegan, "
             "so the vegan main for her please.\n\nBest regards,\nRaj Malhotra\nMalhotra Hotels", "2026-10-27T08:16:00"),
        (X.mail(DEV, [ME], "Your anniversary dinner on 19 Nov - menu and order",
                "Hi Rosa,\n\nGood to talk yesterday. As discussed, the menu is in your Hub catalogue (canapé reception, four mains "
                "incl. vegetarian and vegan, two desserts incl. vegan, drinks or soft-drinks package; 12.5% service). Martin "
                "mentioned he'd started the order.\n\nOnce the order is submitted in Hub, please email me the PO number and a list "
                "of any guests with allergies or dietary requirements, by name, so the kitchen can label their plates.\n\n"
                "Best,\nDev Sharma\nFennel & Fig Catering", when("2026-10-28T16:05:00"), mid("dev", "fennelandfig.co.uk")), True, "dev"),
        rsvp("elena", ("Elena Rossi", "elena@rossistructural.co.uk"),
             "Hi Rosa - yes for two, myself and Marco Rossi. Beef for both of us, please.\n\nElena", "2026-10-29T12:55:00"),
        (X.mail(("Marcus Webb", "marcus.webb@reillyconstruction.co.uk"), [ME], "Fwd: " + SUBJ,
                "Hi Rosa,\n\nTom Reilly forwarded me your invitation below - I'd love to come along if there's space. Beef for me, "
                "please.\n\nMarcus Webb\nSite Director, Reilly Construction\n\n-------- Forwarded Message --------\n"
                f"Subject: {SUBJ}\n\n" + inv_body, when("2026-10-30T09:37:00"), mid("marcus", "reillyconstruction.co.uk")), True, "marcus"),
        rsvp("hugo", ("Hugo Brandt", "hugo.brandt@brandtglass.de"),
             "Dear Ms Lindqvist,\n\nThank you for your kind invitation, which I accept with pleasure. I will come alone. I would "
             "like the sea bass; please note that I have a shellfish allergy.\n\nWith best regards,\nHugo Brandt\nBrandt Glass GmbH",
             "2026-11-02T07:58:00"),
        rsvp("ben2", ("Ben Carter", "ben.carter@carterpg.co.uk"),
             "Hi Rosa,\n\nApologies - I've just been asked to be in Leeds on the 19th, so I have to drop out after all. I hope it's "
             "a great evening.\n\nBen", "2026-11-02T18:20:00"),
        rsvp("olivia2", ("Olivia Grant", "olivia@grantjoinery.co.uk"),
             "Hi Rosa,\n\nPeter would like the sea bass, please. He's also coeliac, so everything gluten-free for him.\n\nThanks!\n"
             "Olivia", "2026-11-03T08:49:00"),
        (X.mail(CLARA, [ME], "Brightwater Studio, 19 November - final numbers",
                "Dear Rosa,\n\nWe're looking forward to hosting your anniversary evening. Could you send me the final number of "
                "people (your guests and your own team) and a door list with your guests' names and companies by Friday "
                "6 November? Our front-of-house team will use it at the entrance.\n\nKind regards,\nClara Benn\nEvents Coordinator, "
                "The Glasshouse", when("2026-11-03T11:10:00"), mid("clara", "theglasshouse-bristol.co.uk")), True, "clara"),
        (X.mail(("RIBA Journal", "newsletter@ribaj.com"), [ME], "This week: timber towers and the new fire rules",
                "Read the latest from RIBA Journal...\n\nUnsubscribe", when("2026-11-04T07:00:00"), mid("riba", "ribaj.com")), True, "riba"),
        (X.mail(("Wild Bloom Studio", "hello@wildbloomstudio.co.uk"), [ME], "Flowers for 19 Nov - quote Q-882",
                "Hi Rosa, attached is our quote for the table flowers and entrance styling (£720.00 incl. VAT). Let us know! Wild Bloom",
                when("2026-11-04T10:30:00"), mid("flowers", "wildbloomstudio.co.uk")), True, "flowers"),
    ]
    inbox.sort(key=lambda t: parsedate_to_datetime(t[0]["Date"]))
    folders = {"Inbox": [(m, parsedate_to_datetime(m["Date"]), u) for m, u, _ in inbox],
               "Sent": [(inv, parsedate_to_datetime(inv["Date"]), False)]}
    make_profile(f"{d}/thunderbird-profile.tar.gz", ME, folders)
    rsvp_keys = ["james", "priya", "tom", "anna", "ben1", "grace", "daniel", "olivia1", "fiona", "helen", "raj", "elena",
                 "marcus", "hugo", "ben2", "olivia2"]
    for k in ("clara", "martin"):  # the headcount must not be guessable from quoted text
        body = next(m for m, _, key in inbox if key == k).get_body().get_content()
        assert "23" not in body
    return ids, rsvp_keys


SCHEMA = """
CREATE TABLE settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE users(id INTEGER PRIMARY KEY, username TEXT UNIQUE, name TEXT, role TEXT, salt TEXT, pw TEXT);
CREATE TABLE events(id INTEGER PRIMARY KEY, name TEXT, date TEXT, time TEXT, venue TEXT, description TEXT, budget REAL);
CREATE TABLE guests(id INTEGER PRIMARY KEY, event_id INTEGER, kind TEXT, name TEXT, company TEXT, email TEXT, status TEXT,
    plus_one TEXT, main TEXT, dietary TEXT, notes TEXT, updated_at TEXT);
CREATE TABLE suppliers(id INTEGER PRIMARY KEY, name TEXT, category TEXT, contact TEXT, email TEXT, terms TEXT, service_pct REAL);
CREATE TABLE catalogue(id INTEGER PRIMARY KEY, supplier_id INTEGER, code TEXT, name TEXT, unit TEXT, price REAL, description TEXT);
CREATE TABLE pos(id INTEGER PRIMARY KEY, number TEXT UNIQUE, supplier_id INTEGER, event_id INTEGER, status TEXT, notes TEXT,
    created_by TEXT, created_at TEXT, submitted_at TEXT);
CREATE TABLE po_lines(id INTEGER PRIMARY KEY, po_id INTEGER, item_id INTEGER, qty INTEGER);
CREATE TABLE sessions(sid TEXT PRIMARY KEY, user_id INTEGER, created TEXT, state TEXT);
CREATE TABLE audit(id INTEGER PRIMARY KEY, ts TEXT, user TEXT, action TEXT, detail TEXT);
"""


def build(d):
    os.makedirs(d, exist_ok=True)
    X.copy_app("app_x3.py", d)
    card = f"{d}/Reply_card_Helen_Park.jpg"
    reply_card(card)
    ids, rsvp_keys = mailbox(d, card)
    os.remove(card)
    others, approved = budget_xlsx(f"{d}/Anniversary_Budget_2026.xlsx")
    lines = expected_lines()
    total = po_total(lines)
    guests_attending = [n for n, (s, _) in OUTCOME.items() if s == "Attending"]
    door = [n for n in guests_attending] + [p for s, p in OUTCOME.values() if s == "Attending" and p] + [SUBSTITUTE[0]]
    assert sorted(door) == sorted(n for n, *_ in DINERS), (sorted(door), sorted(n for n, *_ in DINERS))
    headcount = len(door) + len(HOSTS)
    # ---- database
    rows = {"settings": [{"key": "default_event", "value": "2"}, {"key": "company", "value": "Brightwater Studio"}],
            "users": [{"id": 1, "username": "rosa", "name": "Rosa Lindqvist", "role": "office", "salt": "bw15",
                       "pw": X.pw_hash("Glasshouse15", "bw15")},
                      {"id": 2, "username": "martin", "name": "Martin Achebe", "role": "director", "salt": "mA77",
                       "pw": X.pw_hash("Queen-Square-14", "mA77")}],
            "events": [{"id": 1, "name": "Anniversary Evening 2026", "date": "2026-11-19", "time": "18:30-23:00",
                        "venue": "The Glasshouse, Harbourside, Bristol",
                        "description": "Brightwater Studio at 15 - client and partner evening: drinks, canapés, seated dinner.",
                        "budget": approved},
                       {"id": 2, "name": "Winter Party 2026", "date": "2026-12-11", "time": "19:00-late",
                        "venue": "The Old Granary, Welsh Back", "description": "Staff winter party (staff + partners).", "budget": 4500.0}],
            "guests": [], "suppliers": [], "catalogue": [], "pos": [], "po_lines": [], "audit": []}
    g = rows["guests"]
    for n, c, e in INVITEES:
        g.append({"event_id": 1, "kind": "invitee", "name": n, "company": c, "email": e, "status": "Invited", "plus_one": "",
                  "main": "", "dietary": "", "notes": "", "updated_at": "2026-10-12 10:02:11"})
    for n, role, main, _, _, diet in HOSTS:
        g.append({"event_id": 1, "kind": "staff", "name": n, "company": role, "email": n.split()[0].lower() + "@brightwater-studio.co.uk",
                  "status": "Attending", "plus_one": "", "main": MAIN_NAME[main], "dietary": diet, "notes": "Host",
                  "updated_at": "2026-10-16 12:00:00"})
    winter = [("Rosa Lindqvist", "Attending", "Tom Lindqvist"), ("Martin Achebe", "Attending", "Ngozi Achebe"),
              ("Nadia Petrova", "Invited", ""), ("Callum Price", "Invited", ""), ("Yuki Tanaka", "Attending", ""),
              ("Owen Hughes", "Declined", ""), ("Ellie Marsh", "Attending", "Sam Marsh"), ("Joe Barker", "Invited", ""),
              ("Priti Shah", "Attending", ""), ("Dan Okafor", "Invited", ""), ("Lily Evans", "Attending", "Ben Evans"),
              ("Marta Kowal", "Invited", "")]
    for n, st, p in winter:
        g.append({"event_id": 2, "kind": "invitee", "name": n, "company": "Brightwater Studio",
                  "email": n.split()[0].lower() + "@brightwater-studio.co.uk", "status": st, "plus_one": p, "main": "", "dietary": "",
                  "notes": "", "updated_at": "2026-10-20 09:00:00"})
    for i, r in enumerate(g, 1):
        r["id"] = i
    sups = [(1, "Fennel & Fig Catering", "Catering", "Dev Sharma", DEV[1], "Payment 30 days after the event.", 12.5),
            (2, "The Glasshouse", "Venue", "Clara Benn", CLARA[1], "50% deposit on booking.", 0),
            (3, "Bristol AV Hire", "AV & lighting", "Mike Dunne", "bookings@bristolavhire.co.uk", "Payment on delivery.", 0),
            (4, "Harbourside Canapés Co.", "Catering", "Jenny Lowe", "hello@harboursidecanapes.co.uk", "Payment 14 days.", 10),
            (5, "Wild Bloom Studio", "Flowers & styling", "Ana Ruiz", "hello@wildbloomstudio.co.uk", "Payment on invoice.", 0)]
    rows["suppliers"] = [{"id": a, "name": b, "category": c, "contact": dd, "email": e, "terms": f, "service_pct": s}
                         for a, b, c, dd, e, f, s in sups]
    cat = [(1, *x) for x in CATALOGUE_FF] + [
        (2, "GH-01", "Evening venue hire (Main Hall, 18:00-23:30)", "event", "2700.00", "Includes furniture and cloakroom"),
        (3, "AV-01", "Evening AV package (PA, 2 mics, uplighting)", "event", "890.00", "Delivery, set-up and technician"),
        (4, "HC-01", "Canapé selection (8 pieces)", "head", "16.00", "Seasonal canapés"),
        (4, "HC-02", "Bowl food (3 bowls)", "head", "24.00", "Standing dinner alternative"),
        (5, "WB-01", "Table centrepieces", "table", "45.00", "Seasonal flowers")]
    rows["catalogue"] = [{"id": i, "supplier_id": s, "code": c, "name": n, "unit": u, "price": float(p), "description": ds}
                         for i, (s, c, n, u, p, ds) in enumerate(cat, 1)]
    code_id = {r["code"]: r["id"] for r in rows["catalogue"]}
    pos = [(1, "PO-2026-0141", 2, 1, "submitted", "Venue hire 19 Nov, Main Hall.", "rosa", "2026-09-02 10:12:00", "2026-09-02 10:15:31",
            {"GH-01": 1}),
           (2, "PO-2026-0142", 3, 1, "submitted", "AV for the anniversary evening.", "rosa", "2026-09-30 14:40:00", "2026-09-30 14:41:09",
            {"AV-01": 1}),
           (3, "PO-2026-0143", 1, 2, "submitted", "Winter party buffet dinner, The Old Granary.", "rosa", "2026-10-08 09:10:00",
            "2026-10-08 09:22:45", {"FF-C01": 24, "FF-M01": 8, "FF-M02": 6, "FF-M03": 7, "FF-M04": 3, "FF-D01": 21, "FF-D02": 3,
                                    "FF-B01": 22, "FF-B02": 2}),
           (4, "PO-2026-0144", 1, 1, "draft", "Anniversary dinner 19 Nov - numbers to be confirmed (MA).", "martin",
            "2026-10-14 08:31:00", None, {"FF-C01": 30, "FF-M01": 10, "FF-M02": 10, "FF-M03": 10, "FF-D01": 30, "FF-B01": 30, "FF-L01": 30})]
    for pid, num, sup, ev, st, notes, by, at, sub, ls in pos:
        rows["pos"].append({"id": pid, "number": num, "supplier_id": sup, "event_id": ev, "status": st, "notes": notes,
                            "created_by": by, "created_at": at, "submitted_at": sub})
        for code, q in ls.items():
            rows["po_lines"].append({"po_id": pid, "item_id": code_id[code], "qty": q})
    rows["audit"] = [{"ts": "2026-09-02 10:15:31", "user": "rosa", "action": "PO submitted", "detail": "PO-2026-0141 £3,240.00"},
                     {"ts": "2026-09-30 14:41:09", "user": "rosa", "action": "PO submitted", "detail": "PO-2026-0142 £1,068.00"},
                     {"ts": "2026-10-08 09:22:45", "user": "rosa", "action": "PO submitted", "detail": "PO-2026-0143"},
                     {"ts": "2026-10-14 08:31:00", "user": "martin", "action": "PO created", "detail": "PO-2026-0144 (draft)"}]
    X.seed_db(f"{d}/hub_seed.db", SCHEMA, rows)
    exp = {
        "me": ME[1], "martin": MARTIN[1], "clara": CLARA[1], "dev": DEV[1], "msg": ids,
        "rsvp_ids": [ids[k] for k in rsvp_keys], "winter_id": ids["winter"],
        "outcome": {n: {"status": s, "plus_one": p} for n, (s, p) in OUTCOME.items()},
        "substitute": SUBSTITUTE[0], "forwarded": FORWARDED,
        "invitees": [n for n, _, _ in INVITEES], "seed_guests": [{"id": r["id"], "name": r["name"], "event_id": r["event_id"],
                                                                  "status": r["status"], "plus_one": r["plus_one"]} for r in g],
        "seed_guest_max_id": len(g), "door_list": sorted(door), "headcount": headcount,
        "not_on_door_list": ["Ben Carter", "Daniel Fischer", "Priya Natarajan", FORWARDED],
        "lines": lines, "po_total": float(total), "seed_po_max_id": 4, "draft_po": "PO-2026-0144",
        "budget_others": others, "approved": approved, "budget_total": float(D(Decimal(str(others)) + total)),
        "remaining": float(D(Decimal(str(approved)) - Decimal(str(others)) - total)), "dietary_names": DIETARY_NAMES,
    }
    json.dump(exp, open(f"{d}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp())
