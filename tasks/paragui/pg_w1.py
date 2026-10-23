"""W1 - one help-desk web app (Larkspur Support Desk on :8101) in one Chrome session: triaging the untriaged queue
per the SLA policy vs answering six customer tickets with knowledge-base articles. See spec_w1.md."""
from __future__ import annotations

import re

from evaluation_examples.task_class.pg_common import (HOME, expected, fetch, open_chrome, result, sh, sql,
                                                      start_webapp, upload, webapp_db)

NAME, PORT = "w1", 8101
URL = f"http://127.0.0.1:{PORT}/"
DOCS = ["Support_Triage_Policy_v3.pdf", "support_log_2026-10.xlsx"]
LOG = f"{HOME}/Documents/support_log_2026-10.xlsx"
WORKING = {"Open", "Pending", "On-hold", "Solved"}


def setup(sc):
    start_webapp(sc, f"{NAME}/app_w1.py", NAME, PORT, f"{NAME}/seed.db")
    # documents: never overwrite (setup runs again before the second task of a sequential run)
    upload(sc, [(f"{NAME}/{d}", f"{HOME}/.cache/pg{NAME}/{d}") for d in DOCS])
    sh(sc, f"mkdir -p ~/Documents && for f in {' '.join(DOCS)}; do cp -n ~/.cache/pg{NAME}/$f ~/Documents/$f; done")
    open_chrome(sc, [URL])


def _state(env, exp):
    m = exp["seed_max"]
    return sql(env, webapp_db(NAME), {
        "tickets": "SELECT id, status, priority, team, assignee_id, tags FROM tickets",
        "comments": f"SELECT id, ticket_id, user_id, public, body FROM comments WHERE id > {m['comments']} ORDER BY id",
        "events": f"SELECT id, ticket_id, rid, field, old, new FROM events WHERE id > {m['events']} ORDER BY id",
        "posts": f"SELECT id, user_id, title, body FROM posts WHERE id > {m['posts']} ORDER BY id"})


def _mentions_count(text, label, n):
    """True if the text pairs the label with the number n ('Billing: 3', 'Billing 3 tickets', '3 Billing')."""
    for m in re.finditer(rf"{label}\b[^0-9\n]{{0,25}}?(\d+)", text, re.I):
        if int(m.group(1)) == n:
            return True
    return any(int(m.group(1)) == n for m in re.finditer(rf"\b(\d+)\b[^0-9\n]{{0,15}}?\b{label}", text, re.I))


def evaluate_triage(env, dest: str) -> dict:
    exp = expected(NAME)
    st = _state(env, exp)
    if not st:
        return result(0.0, {"app_db": "unreadable"})
    tickets = {str(t["id"]): t for t in st["tickets"]}
    staff = {int(k): v for k, v in exp["staff"].items()}
    checks, wrong = {}, []
    # 1) the regular tickets: team, priority, assignee, status
    regular = {k: v for k, v in exp["triage"].items() if isinstance(v, dict)}
    pts, teams_ok, status_ok = 0.0, 0, 0
    for tid, want in regular.items():
        t = tickets.get(tid, {})
        got = {"team": t.get("team"), "priority": t.get("priority"), "assignee": staff.get(t.get("assignee_id"))}
        ok = [got[k] == want[k] for k in ("team", "priority", "assignee")]
        pts += sum(ok) / 3 if got["team"] else 0  # an untouched ticket earns nothing (not even "unassigned")
        teams_ok += ok[0]
        status_ok += t.get("status") in WORKING
        if not all(ok) or t.get("status") not in WORKING:
            wrong.append(f"#{tid}: got {got['team']}/{got['priority']}/{got['assignee']}/{t.get('status')} "
                         f"want {want['team']}/{want['priority']}/{want['assignee']}")
    n = len(regular)
    checks["tickets_fully_correct"] = f"{n - len(wrong)}/{n}"
    checks["triage_fields"] = round(pts / n, 3)
    checks["triaged_status_not_new"] = f"{status_ok}/{n}"
    checks["wrong"] = wrong
    # 2) spam and the duplicate
    spam_ok = 0
    for tid in [k for k, v in exp["triage"].items() if v == "spam"]:
        t = tickets.get(tid, {})
        good = t.get("status") == "Closed" and "spam" in (t.get("tags") or "") and not t.get("team")
        spam_ok += good
        checks[f"spam_#{tid}"] = good
    dup_ok = 0
    for tid, orig in exp["duplicate_of"].items():
        t = tickets.get(tid, {})
        note = any(c["ticket_id"] == int(tid) and not c["public"] and c["user_id"]
                   and re.search(rf"(?<!\d){orig}(?!\d)", c["body"]) for c in st["comments"])
        good = t.get("status") == "Closed" and "duplicate" in (t.get("tags") or "") and bool(note)
        dup_ok += good
        checks[f"duplicate_#{tid}"] = {"closed": t.get("status") == "Closed", "tag": "duplicate" in (t.get("tags") or ""),
                                       "note_names_original": bool(note)}
    # 3) the policy's "keep Pending/On-hold/Solved" rule: one submission that triaged a ticket (team change) and
    #    also moved it out of a colleague's working status
    by_req = {}
    for e in st["events"]:
        by_req.setdefault((e["rid"], e["ticket_id"]), []).append(e)
    reopened = sorted({tk for (rid, tk), evs in by_req.items()
                       if any(e["field"] == "team" for e in evs)
                       and any(e["field"] == "status" and e["old"] in ("Pending", "On-hold", "Solved") for e in evs)})
    checks["reopened_colleague_tickets"] = reopened
    # 4) nothing else re-triaged (task B's own tickets excluded: that task is told not to touch those fields)
    collateral = []
    for tid, s in exp["seed"].items():
        if tid in exp["triage"] or tid in exp["b_tickets"]:
            continue
        t = tickets.get(tid, {})
        if (t.get("team"), t.get("priority"), t.get("assignee_id")) != (s["team"], s["priority"], s["assignee_id"]):
            collateral.append(int(tid))
    checks["changed_other_tickets"] = collateral
    # 5) the handover post
    posts = [p for p in st["posts"] if p["user_id"] == exp["sam_id"]]
    text = "\n".join(f"{p['title']}\n{p['body']}" for p in posts)
    urgent_ok = bool(posts) and all(re.search(rf"(?<!\d){u}(?!\d)", text) for u in exp["urgent"])
    counts_ok = bool(posts) and all(_mentions_count(text, t, c) for t, c in exp["team_counts"].items())
    closed_ok = bool(posts) and _mentions_count(text, "spam", exp["spam_count"]) and _mentions_count(text, "duplicates?", exp["dup_count"])
    checks.update({"handover_posted": bool(posts), "handover_urgent_ids": urgent_ok, "handover_team_counts": counts_ok,
                   "handover_spam_dup_counts": closed_ok})
    done = teams_ok / n  # the "did no harm" checks only count in proportion to the work actually done
    score = (0.55 * pts / n + 0.05 * status_ok / n + 0.05 * spam_ok + 0.05 * dup_ok + 0.05 * done * (not reopened)
             + 0.05 * done * (not collateral) + 0.07 * urgent_ok + 0.07 * counts_ok + 0.01 * closed_ok)
    if teams_ok < n / 2:  # the queue was essentially not triaged
        score = min(score, 0.4)
        checks["cap"] = "fewer than half of the tickets have the right team -> score <= 0.4"
    return result(score, checks)


def _row_status(cells):
    for c in cells:
        s = re.sub(r"[\s_-]+", "", str(c or "")).lower()
        if s in ("open", "pending", "onhold", "solved", "closed", "new"):
            return {"onhold": "On-hold"}.get(s, s.capitalize())
    return None


def evaluate_reply(env, dest: str) -> dict:
    exp = expected(NAME)
    st = _state(env, exp)
    if not st:
        return result(0.0, {"app_db": "unreadable"})
    tickets = {str(t["id"]): t for t in st["tickets"]}
    checks, part, notes_ok, n_notes, any_reply = {}, 0.0, 0, 0, False
    for tid, want in exp["b_tickets"].items():
        mine = [c for c in st["comments"] if str(c["ticket_id"]) == tid and c["user_id"] == exp["sam_id"]]
        pub = "\n".join(c["body"] for c in mine if c["public"])
        internal = "\n".join(c["body"] for c in mine if not c["public"])
        any_reply |= bool(pub.strip())
        kb = bool(re.search(rf"\bKB[\s-]?{want['kb'][3:]}\b", pub, re.I))
        facts = {k: bool(re.search(rx, pub, re.I)) for k, rx in want["facts"].items()}
        status = tickets.get(tid, {}).get("status")
        part += (kb + sum(facts.values()) / len(facts) + (status == want["status"])) / 3
        c = {"public_reply": bool(pub.strip()), "cites_" + want["kb"]: kb, "facts": facts,
             "status": f"{status} (want {want['status']})"}
        for k, rx in want.get("note", {}).items():
            n_notes += 1
            ok = bool(re.search(rx, internal, re.I))
            notes_ok += ok
            c["internal_note_" + k] = ok
        checks[f"#{tid}"] = c
    nt = len(exp["b_tickets"])
    # the support log: one row per ticket with the KB article and the new status; the old rows must survive
    path = fetch(env, LOG, dest + ".xlsx")
    rows_ok, kept = 0, False
    if path:
        from openpyxl import load_workbook
        try:
            ws = load_workbook(path, data_only=True).worksheets[0]
            rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        except Exception as e:  # noqa: BLE001
            rows, checks["log_error"] = [], str(e)[:200]
        kept = len([r for r in rows[1:] if any(c not in (None, "") for c in r)]) >= exp["log_rows"]
        found = {}
        for tid, want in exp["b_tickets"].items():
            for r in rows:
                if any(re.fullmatch(rf"#?\s*{tid}(\.0)?", str(c or "").strip()) for c in r):
                    kb = any(re.fullmatch(rf"KB[\s-]?{want['kb'][3:]}", str(c or "").strip(), re.I) for c in r)
                    stat = _row_status(r)
                    good = kb and stat == want["status"] and stat == tickets.get(tid, {}).get("status")
                    found[tid] = good or found.get(tid, False)
            rows_ok += found.get(tid, False)
        checks["log_rows_ok"] = f"{rows_ok}/{nt}"
        checks["log_old_rows_kept"] = kept
    else:
        checks["log"] = "missing"
    score = 0.6 * part / nt + 0.1 * (notes_ok / n_notes if n_notes else 1) + 0.25 * rows_ok / nt + 0.05 * (kept and rows_ok > 0)
    if not any_reply:
        score = min(score, 0.3)
        checks["cap"] = "no public reply on any ticket -> score <= 0.3"
    return result(score, checks)
