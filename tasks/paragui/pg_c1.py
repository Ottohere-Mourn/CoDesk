"""C1 - one workbook: cleaning the CRM export (Raw sheet) vs building the regional summary in the same file."""
from __future__ import annotations

import datetime as dt
import re
import zipfile

from evaluation_examples.task_class.pg_common import HOME, expected, fetch, result, upload

XLSX = f"{HOME}/Documents/sales_2026Q3.xlsx"
REGIONS = ["North America", "EMEA", "APAC", "LATAM"]
MONTHS = {7: "July", 8: "August", 9: "September"}


def setup(sc):
    upload(sc, [("c1/sales_2026Q3.xlsx", XLSX)])
    sc._open_setup(XLSX)


def _canon(region) -> str | None:
    key = re.sub(r"[^a-z]", "", str(region or "").lower())
    table = {"northamerica": "North America", "namerica": "North America", "emea": "EMEA", "apac": "APAC",
             "asiapacific": "APAC", "latam": "LATAM", "latinamerica": "LATAM"}
    return table.get(key)


def _date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y/%m/%d", "%m-%d-%Y", "%d.%m.%Y", "%b %d, %Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _raw(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "raw"), None)
    if ws is None:
        return wb, None, []
    head = [str(c.value or "").strip().lower() for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2):
        vals = {h: c for h, c in zip(head, r)}
        if vals.get("order id") is None or vals["order id"].value in (None, ""):
            continue
        rows.append(vals)
    return wb, ws, rows


def evaluate_summary(env, dest: str) -> dict:
    path = fetch(env, XLSX, dest + ".xlsx")
    if not path:
        return result(0.0, {"file": "missing"})
    wb, ws, rows = _raw(path)
    want = {(r, m): 0.0 for r in REGIONS for m in MONTHS}
    for row in rows:
        reg, day = _canon(row.get("region").value if row.get("region") else None), _date(row["date"].value if "date" in row else None)
        try:
            rev = float(row["revenue"].value)
        except (KeyError, TypeError, ValueError):
            continue
        if reg and day and day.month in MONTHS:
            want[(reg, day.month)] += rev
    summ = next((wb[n] for n in wb.sheetnames if "summary" in n.lower()), None)
    nums = [float(c.value) for line in (summ.iter_rows() if summ else []) for c in line
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    hit = sum(1 for v in want.values() if any(abs(x - v) <= max(1.0, 0.001 * v) for x in nums))
    with zipfile.ZipFile(path) as z:
        chart = any(n.startswith("xl/charts/chart") for n in z.namelist())
    checks = {"summary_sheet": summ is not None, "totals_matched": f"{hit}/12", "chart": chart}
    return result((hit / 12) * 0.8 + 0.2 * chart if summ is not None else 0.0, checks)


def evaluate_clean(env, dest: str) -> dict:
    path = fetch(env, XLSX, dest + ".xlsx")
    if not path:
        return result(0.0, {"file": "missing"})
    exp = expected("c1")["unique"]
    _, ws, rows = _raw(path)
    if ws is None:
        return result(0.0, {"raw_sheet": "missing"})
    ids = [str(r["order id"].value).strip() for r in rows]
    dedup = len(ids) == len(set(ids))
    complete = set(ids) == set(exp)
    ok_date = ok_region = ok_rev = 0
    for r, oid in zip(rows, ids):
        e = exp.get(oid)
        if not e:
            continue
        c = r.get("date")
        fmt = re.sub(r'[\\"]|;@', "", str(c.number_format if c is not None else "")).lower()
        if c is not None and isinstance(c.value, (dt.date, dt.datetime)) and _date(c.value).isoformat() == e["date"] \
                and "yyyy-mm-dd" in fmt:
            ok_date += 1
        if r.get("region") is not None and str(r["region"].value) in REGIONS and r["region"].value == e["region"]:
            ok_region += 1
        try:
            ok_rev += abs(float(r["revenue"].value) - e["revenue"]) < 0.01
        except (KeyError, TypeError, ValueError):
            pass
    n = len(exp)
    parts = [dedup, complete, ok_date / n, ok_region / n, ok_rev / n]
    checks = {"no_duplicates": dedup, "all_orders_kept": complete, "dates_ok": f"{ok_date}/{n}",
              "regions_ok": f"{ok_region}/{n}", "revenue_unchanged": f"{ok_rev}/{n}"}
    return result(sum(float(p) for p in parts) / len(parts), checks)
