"""X2 - recruiting at Fernleaf Analytics (Thunderbird + HireDesk ATS in Chrome + LibreOffice Calc).

A (pgx2screen): CV screening for REQ-031 Data Analyst - application mails with CVs (three scanned, one updated CV, one
   applicant already on file) -> candidates, scorecards and stages in HireDesk -> hiring manager's tracker in Calc ->
   reply to Grace with the tracker attached (draft).
B (pgx2sched): panel interviews for REQ-027 Product Designer - availability replies + Ben's constraints -> slots booked
   in HireDesk (video links only there), withdrawn candidate updated -> Ben's panel sheet in Calc -> invitation reply
   drafts to each candidate.
Shared live state: one Thunderbird, one Chrome with one HireDesk login session whose requisition selector (scopes the
pipeline and the interview list, and pre-selects the requisition in the new-candidate form) and filters are per session.
"""
from __future__ import annotations

import re

from evaluation_examples.task_class import pg_common as pc
from evaluation_examples.task_class import pg_xcommon as xc

APP, PORT = "hiredesk", 8122
TRACKER = f"{pc.HOME}/Documents/Recruiting/REQ-031_screening.xlsx"
PANEL = f"{pc.HOME}/Documents/Recruiting/ProductDesigner_interviews.xlsx"
SCORE_KEYS = ["score_sql", "score_code", "score_bi", "score_degree"]


def setup(sc):
    pc.sh(sc, "mkdir -p ~/Documents/Recruiting")
    pc.start_webapp(sc, "x2/app_x2.py", APP, PORT, "x2/hiredesk_seed.db")
    pc.start_thunderbird(sc, "x2/thunderbird-profile.tar.gz", "pgx2")
    pc.open_chrome(sc, [f"http://127.0.0.1:{PORT}/login"])


def _db(env):
    return pc.sql(env, pc.webapp_db(APP), {
        "cands": "SELECT id, name, lower(trim(email)) AS email FROM candidates",
        "apps": "SELECT a.id, a.candidate_id, r.code, a.stage, a.score_sql, a.score_code, a.score_bi, a.score_degree, a.rtw, "
                "a.reason FROM applications a JOIN requisitions r ON r.id=a.req_id",
        "slots": "SELECT id, application_id FROM slots"})


def _int(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else None
    n = xc.numbers(str(v))
    return int(n[0]) if len(n) == 1 and float(n[0]).is_integer() else None


def _sheet_rows(path, must=("candidate",)):
    """Rows below the first header row containing all `must` words, as {header-flat: value} dicts (all sheets)."""
    from openpyxl import load_workbook
    out = []
    for ws in load_workbook(path, data_only=True).worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for i, row in enumerate(grid):
            heads = [xc.flat(v) for v in row]
            if all(any(m in h for h in heads) for m in must):
                for below in grid[i + 1:]:
                    if any(v is not None and str(v).strip() for v in below):
                        out.append({h: v for h, v in zip(heads, below) if h})
                break
    return out


def _col(row: dict, *words):
    for h, v in row.items():
        if any(w in h for w in words):
            return v
    return None


# ---------------------------------------------------------------- A: screening
def evaluate_screen(env, dest: str) -> dict:
    exp = pc.expected("x2")
    q = _db(env)
    if q is None:
        return pc.result(0.0, {"app_db": "unreadable"})
    checks = {}
    seed = exp["seed_max"]
    by_email = {}
    for c in q["cands"]:
        by_email.setdefault(c["email"], []).append(c["id"])
    apps_of = {}
    for a in q["apps"]:
        apps_of.setdefault(a["candidate_id"], []).append(a)
    exists = scores = stages = 0.0
    problems, dupes, wrong_req = [], [], []
    for e in exp["applicants"]:
        cids = by_email.get(e["email"].lower(), [])
        if len(cids) > 1:
            dupes.append(e["email"])
        apps = [a for cid in cids for a in apps_of.get(cid, [])]
        wrong_req += [f"{e['name']} on {a['code']}" for a in apps if a["id"] > seed["applications"] and a["code"] != "REQ-031"]
        mine = [a for a in apps if a["code"] == "REQ-031"]
        if not mine:
            problems.append(f"{e['name']}: no REQ-031 application")
            continue
        a = max(mine, key=lambda x: x["id"])
        exists += 1
        got = [a[k] for k in SCORE_KEYS] + [a["rtw"]]
        want = e["scores"] + [e["rtw"]]
        ok = sum(g == w for g, w in zip(got, want))
        scores += ok / 5
        stage_ok = a["stage"] == e["stage"]
        if stage_ok and e["rtw"] == "No":
            stage_ok = bool(re.search(r"right to work|visa|sponsor|rtw", a["reason"] or "", re.I))
        stages += stage_ok
        if ok < 5 or not stage_ok:
            problems.append(f"{e['name']}: scores {got} (want {want}), stage {a['stage']!r} reason {a['reason']!r}")
    # brand-new candidates must not have been put on another requisition either (requisition selector contention)
    mine_ids = {cid for e in exp["applicants"] for cid in by_email.get(e["email"].lower(), [])}
    new_ids = {c["id"] for c in q["cands"] if c["id"] > seed["candidates"]} - mine_ids
    wrong_req += [f"candidate #{a['candidate_id']} on {a['code']}" for a in q["apps"]
                  if a["candidate_id"] in new_ids and a["code"] != "REQ-031"]
    n = len(exp["applicants"])
    clean = not dupes and not wrong_req and exists > 0
    checks.update({"applications_on_REQ-031": f"{exists:g}/{n}", "scorecards": f"{scores:.2f}/{n}", "stages": f"{stages:g}/{n}",
                   "duplicate_candidates": dupes, "wrong_requisition": wrong_req, "db_problems": problems})

    # tracker
    path = pc.fetch(env, TRACKER, dest + "_tracker.xlsx")
    rows_ok, tracker_ok = 0.0, False
    if path:
        try:
            rows = _sheet_rows(path, ("candidate", "decision"))
            tracker_ok = True
        except Exception as e:
            rows = []
            checks["tracker_error"] = repr(e)[:200]
        detail = []
        for e in exp["applicants"]:
            last = xc.flat(e["name"].split()[-1])
            row = next((r for r in rows if e["email"].lower() in str(_col(r, "email") or "").lower()
                        or last in xc.flat(_col(r, "candidate", "name"))), None)
            if not row:
                detail.append(f"{e['name']}: no row")
                continue
            r = row
            rtw = str(_col(r, "right", "rtw") or "").strip().lower()[:1]
            dec = xc.flat(_col(r, "decision"))
            dec_ok = ("phone" in dec or "shortlist" in dec) if e["stage"] == "Phone screen" else ("reject" in dec)
            fields = [_int(_col(r, "sql")) == e["scores"][0], _int(_col(r, "python")) == e["scores"][1],
                      _int(_col(r, "bi")) == e["scores"][2], _int(_col(r, "degree")) == e["scores"][3],
                      rtw == e["rtw"][0].lower(), _int(_col(r, "total")) == e["total"], dec_ok]
            rows_ok += sum(fields) / len(fields)
            if not all(fields):
                detail.append(f"{e['name']}: {fields}")
        checks["tracker_rows"] = f"{rows_ok:.2f}/{n}"
        checks["tracker_problems"] = detail
    checks["tracker_file"] = tracker_ok

    # reply to Grace
    boxes = xc.mailboxes(env) or {}
    mine = [m for m in xc.outgoing(boxes, exp["msg"].values()) if xc.addressed(m, exp["grace"])]
    best, att_ok = None, False
    for k, m in enumerate(mine):
        for i, a in enumerate(m["attachments"]):
            if a["name"].lower().endswith(".xlsx"):
                p = xc.fetch_attachment(env, a, f"{dest}_att{k}_{i}")
                try:
                    text = " ".join(str(c[3]) for c in xc.xlsx_cells(p)).lower() if p else ""
                except Exception:
                    text = ""
                if all(e["email"].lower() in text or e["name"].lower() in text for e in exp["applicants"]):
                    best, att_ok = m, True
        best = best or m
    body = xc.flat(best["text"]) if best else ""
    named = all(xc.flat(n) in body for n in exp["shortlist"])
    is_reply = bool(best) and exp["msg"]["grace"] in best["irt"]
    checks.update({"reply_draft": bool(best), "reply_has_tracker": att_ok, "reply_names_shortlist": named, "is_reply_to_grace": is_reply})

    score = (0.15 * exists / n + 0.225 * scores / n + 0.15 * stages / n + 0.10 * clean + 0.025 * (stages == n)
             + 0.175 * rows_ok / n + 0.10 * att_ok + 0.05 * named + 0.025 * is_reply)
    # a wrong shortlist/reject decision is the costly error: cap by the number of wrong decisions
    score = xc.capped(score, [(exists == 0, 0.3), (not tracker_ok, 0.6), (0 < exists and stages < n, 0.85 - 0.1 * (n - stages))])
    return pc.result(score, checks)


# ---------------------------------------------------------------- B: interview scheduling
def _date_rx(start: str) -> str:
    d = int(start[8:10])
    return (rf"(?<!\d){d}(?:st|nd|rd|th)?\s*(?:of\s+)?nov|nov(?:ember)?\.?\s+{d}(?!\d)|{start[:10]}|(?<!\d)0?{d}[/.]11(?!\d)")


def _time_rx(start: str) -> str:
    h, m = int(start[11:13]), start[14:16]
    h12 = h - 12 if h > 12 else h
    ampm = "p" if h >= 12 else "a"
    alts = [rf"(?<!\d)0?{h}[:.]{m}"]
    alts.append(rf"(?<!\d){h12}(?:[:.]{m})?\s*{ampm}\.?m" if m == "00" else rf"(?<!\d){h12}[:.]{m}\s*{ampm}\.?m")
    return "|".join(alts)


def evaluate_sched(env, dest: str) -> dict:
    exp = pc.expected("x2")
    q = _db(env)
    if q is None:
        return pc.result(0.0, {"app_db": "unreadable"})
    checks = {}
    slot_app = {s["id"]: s["application_id"] for s in q["slots"]}
    stage = {a["id"]: a["stage"] for a in q["apps"]}
    booked = [b for b in exp["bookings"] if slot_app.get(b["slot"]) == b["app_id"]]
    wd = exp["withdrawn"]
    wrong = []
    if wd["app_id"] in slot_app.values():
        wrong.append("withdrawn candidate booked")
    wrong += [f"blocked slot {s} booked" for s in exp["blocked_slots"] if slot_app.get(s)]
    if slot_app.get(exp["dan"]["slot"]) != exp["dan"]["app_id"]:
        wrong.append("Dan Price's booking changed")
    for b in exp["bookings"]:
        others = [s for s, a in slot_app.items() if a == b["app_id"] and s != b["slot"]]
        if others:
            wrong.append(f"{b['name']} booked in slot {others}")
    withdrawn_ok = stage.get(wd["app_id"]) == "Withdrawn"
    checks.update({"bookings_ok": f"{len(booked)}/3", "booked": [b["name"] for b in booked], "wrong_bookings": wrong,
                   "withdrawn_stage": stage.get(wd["app_id"])})

    # panel sheet
    path = pc.fetch(env, PANEL, dest + "_panel.xlsx")
    rows_ok, panel_ok = 0, False
    if path:
        try:
            rows = _sheet_rows(path, ("date", "candidate"))
            panel_ok = True
        except Exception as e:
            rows = []
            checks["panel_error"] = repr(e)[:200]
        names = {b["slot"]: b["name"] for b in exp["bookings"]}
        detail = []
        for sid, (dlabel, tm) in exp["slot_labels"].items():
            row = next((r for r in rows if xc.flat(dlabel) in xc.flat(" ".join(str(v) for v in r.values() if v is not None))
                        and xc.flat(tm) in xc.flat(" ".join(str(v) for v in r.values() if v is not None))), None)
            cand = xc.flat(_col(row, "candidate")) if row else ""
            want = names.get(int(sid))
            if want:
                ok = xc.flat(want.split()[-1]) in cand
                rows_ok += ok
                if not ok:
                    detail.append(f"{dlabel} {tm}: '{cand}' (want {want})")
            elif int(sid) != exp["dan"]["slot"]:
                stray = [b["name"] for b in exp["bookings"]] + [wd["name"]]
                if any(xc.flat(n.split()[-1]) in cand for n in stray):
                    rows_ok -= 0.5
                    detail.append(f"{dlabel} {tm}: unexpected '{cand}'")
        rows_ok = max(0.0, rows_ok)
        checks["panel_rows"] = f"{rows_ok:g}/3"
        checks["panel_problems"] = detail
    checks["panel_file"] = panel_ok

    # invitation drafts
    boxes = xc.mailboxes(env) or {}
    out = xc.outgoing(boxes, exp["msg"].values())
    have = content = replies = 0
    detail = []
    for b in exp["bookings"]:
        ms = [m for m in out if xc.addressed(m, b["email"])]
        if not ms:
            detail.append(f"{b['name']}: no draft")
            continue
        have += 1
        good = [m for m in ms if re.search(_date_rx(b["start"]), m["text"], re.I) and re.search(_time_rx(b["start"]), m["text"], re.I)
                and b["link"] in m["text"].lower()]
        content += bool(good)
        replies += any(b["latest"] in m["irt"] for m in (good or ms))
        if not good:
            detail.append(f"{b['name']}: date/time/link missing in {len(ms)} draft(s)")
    checks.update({"invites": f"{have}/3", "invites_with_date_time_link": f"{content}/3", "invites_reply_to_latest": f"{replies}/3",
                   "invite_problems": detail})

    score = (0.30 * len(booked) / 3 + 0.10 * (not wrong and len(booked) > 0) + 0.10 * withdrawn_ok + 0.15 * rows_ok / 3
             + 0.15 * have / 3 + 0.15 * content / 3 + 0.05 * replies / 3)
    score = xc.capped(score, [(not booked, 0.3), (not panel_ok, 0.75)])
    return pc.result(score, checks)
