"""m4 - one Thunderbird: composing an HTML customer newsletter (template + draft) vs working the support queue.

Persona: Hana Kim, owner of Fern & Fennel, a small online plant & pottery shop in Portland (hello@fernandfennel.shop).
Task A (pgm4newsletter): HTML newsletter with three inline product photos, prices/links/coupon from an image-only promo
                         sheet, Bcc = opted-in customers minus unsubscribe requests, with one changed address;
                         saved as a template and as a draft.
Task B (pgm4support): ten support e-mails -> reply drafts with the matching KB macro (order numbers partly only in
                      photos), originals to Support/Done, two scams to Junk, a support log workbook.
The unsubscribe / address-change requests A needs are support e-mails that B answers and moves away; both compose in
the same Thunderbird (two composer windows, one Drafts folder).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import random

from PIL import Image, ImageDraw, ImageFilter

from lib_mmail import DocPage, build_profile, dump_expected, font, md5, msg, pdf_scan, scan_look

ME = ("Hana Kim", "hello@fernandfennel.shop")
HOST = "pop.fernandfennel.shop"
HANA = f"Fern & Fennel <{ME[1]}>"
PT = dt.timezone(dt.timedelta(hours=-7))


def at(m, d, hh, mm=0):
    return dt.datetime(2026, m, d, hh, mm, tzinfo=PT)


def mid(tag, dom):
    return f"<pgm4.{tag}@{dom}>"


PRODUCTS = [
    dict(key="monstera", name="Monstera Deliciosa (6-inch nursery pot)", regular="$42.00", price="$34.00",
         url="https://fernandfennel.shop/p/monstera-6in", photo="monstera_dec.jpg", color=(46, 125, 50)),
    dict(key="planter", name="Speckled Stoneware Planter - Large", regular="$36.00", price="$28.50",
         url="https://fernandfennel.shop/p/speckled-planter-l", photo="planter_speckled_L.jpg", color=(176, 150, 120)),
    dict(key="pothos", name="Pothos 'Marble Queen' Hanging Basket", regular="$26.00", price="$22.00",
         url="https://fernandfennel.shop/p/marble-queen-basket", photo="marble_queen_basket.jpg", color=(120, 160, 90)),
]
DECOYS = [("monstera_2025.jpg", (60, 110, 60), "Monstera (2025 shoot)"),
          ("planter_speckled_S.jpg", (190, 170, 140), "Speckled planter S"),
          ("pothos_golden.jpg", (150, 170, 60), "Golden Pothos")]
COUPON = "COZYDEC15"
UNSUB = "https://fernandfennel.shop/unsubscribe"
SUBJECT = "Fern & Fennel - December picks"


def photo(color, label, seed):
    r = random.Random(seed)
    im = Image.new("RGB", (800, 600), (240, 236, 228))
    d = ImageDraw.Draw(im)
    for y in range(600):  # soft studio backdrop
        g = 240 - y // 12
        d.line((0, y, 800, y), fill=(g, g - 4, g - 10))
    d.rounded_rectangle((300, 330, 500, 540), radius=26, fill=tuple(max(0, c - 40) for c in color))
    for _ in range(14):
        x, y = r.randrange(260, 540), r.randrange(80, 340)
        d.ellipse((x - 60, y - 34, x + 60, y + 34), fill=tuple(min(255, c + r.randrange(-20, 40)) for c in color))
    d.rectangle((0, 548, 800, 600), fill=(255, 255, 255))
    d.text((24, 558), label, fill=(60, 60, 60), font=font(26))
    b = io.BytesIO()
    im.filter(ImageFilter.GaussianBlur(0.8)).save(b, "JPEG", quality=82)
    return b.getvalue()


def slip_photo(order, item, seed):
    """Phone photo of a broken pot with the packing slip next to it - the order number exists only as pixels."""
    r = random.Random(seed)
    im = Image.new("RGB", (1200, 900), (150, 120, 95))
    d = ImageDraw.Draw(im)
    for _ in range(400):
        x, y = r.randrange(1200), r.randrange(900)
        d.point((x, y), fill=(140 + r.randrange(20), 112 + r.randrange(16), 88 + r.randrange(14)))
    d.pieslice((80, 300, 520, 760), 200, 340, fill=(200, 190, 175))
    d.polygon([(160, 520), (330, 470), (300, 640), (190, 700)], fill=(185, 175, 160))
    d.line((250, 480, 330, 700), fill=(80, 70, 60), width=5)
    slip = Image.new("RGB", (560, 420), (252, 252, 250))
    s = ImageDraw.Draw(slip)
    s.text((24, 20), "Fern & Fennel", fill=(30, 90, 50), font=font(34, True))
    s.text((24, 70), "PACKING SLIP", fill=(40, 40, 40), font=font(24, True))
    s.text((24, 130), f"Order no.  {order}", fill=(20, 20, 20), font=font(30, True))
    s.text((24, 190), f"1 x {item}", fill=(40, 40, 40), font=font(22))
    s.text((24, 240), "Thank you for shopping small!", fill=(90, 90, 90), font=font(20))
    slip = slip.rotate(r.uniform(-9, -5), expand=True, fillcolor=(150, 120, 95))
    im.paste(slip, (590, 240))
    b = io.BytesIO()
    im.filter(ImageFilter.GaussianBlur(1.0)).save(b, "JPEG", quality=80)
    return b.getvalue()


MACROS = [
    ("KB-01", "Item arrived damaged",
     "Hi {name}, we're so sorry your order {order} arrived damaged! A free replacement is on its way and you don't "
     "need to send anything back. Thank you for the photo - it helps us improve our packaging.",
     "free replacement is on its way"),
    ("KB-02", "Where is my order?",
     "Hi {name}, thanks for checking in. Your order {order} has been handed to the carrier - you can follow it with "
     "the tracking link in your shipping confirmation. Orders usually arrive within 3-5 business days.",
     "follow it with the tracking link"),
    ("KB-03", "Yellow leaves / plant care",
     "Hi {name}, yellowing lower leaves are usually a sign of overwatering. Let the top 2 inches of soil dry out "
     "before you water again, make sure the pot drains, and give the plant bright, indirect light.",
     "sign of overwatering"),
    ("KB-04", "Return outside the 30-day window",
     "Hi {name}, our return window is 30 days from delivery, so unfortunately we can't accept a return for order "
     "{order}. As a thank-you for shopping with us, here is 10% off your next order: THANKYOU10.",
     "return window is 30 days"),
    ("KB-05", "Unsubscribe request",
     "Hi {name}, you have been removed from our newsletter list and won't receive any more newsletters from us. "
     "Order and shipping emails are not affected.",
     "removed from our newsletter list"),
    ("KB-06", "Email address change",
     "Hi {name}, we've updated your email address in our records - from now on newsletters and order updates will "
     "go to your new address.",
     "updated your email address in our records"),
    ("KB-07", "Wholesale inquiry",
     "Hi {name}, thank you for your interest in carrying Fern & Fennel products! Our wholesale partner manager, "
     "Marco Rossi (marco@fernandfennel.shop), will get in touch within two business days.",
     "wholesale partner manager"),
]


def macros_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Fern & Fennel - Customer Support Macros", 0)
    d.add_paragraph("Copy the macro that matches the customer's request into your reply. Replace {name} with the "
                    "customer's first name and {order} with the order number (FF-xxxxx). One macro per email.")
    for code, title, text, _ in MACROS:
        d.add_heading(f"{code}  {title}", 2)
        d.add_paragraph(text)
    d.add_heading("Spam / scams", 2)
    d.add_paragraph("Never reply to payout 'verification' requests or SEO offers - mark them as junk.")
    d.save(path)


FIRST = ["Emma", "Noah", "Ava", "Liam", "Sophia", "Ethan", "Isabella", "Mason", "Mia", "Lucas", "Harper", "Elijah",
         "Amelia", "James", "Evelyn", "Benjamin", "Abigail", "Henry", "Ella", "Jack", "Chloe", "Owen", "Grace", "Leo"]
LAST = ["Walker", "Nguyen", "Patel", "Brooks", "Rivera", "Foster", "Kim", "Hughes", "Price", "Bennett", "Coleman",
        "Reyes", "Sanders", "Ward", "Ortiz", "Jenkins", "Perry", "Long", "Powell", "Barnes", "Ross", "Fisher", "Gray", "Hayes"]
DOMAINS = ["gmail.com", "outlook.com", "yahoo.com", "icloud.com", "proton.me", "hotmail.com"]


def customers():
    r = random.Random(404)
    rows = []
    for i, (f, l) in enumerate(zip(FIRST, LAST)):
        email = f"{f.lower()}.{l.lower()}@{DOMAINS[i % len(DOMAINS)]}" if i % 3 else f"{f[0].lower()}{l.lower()}{r.randrange(10, 99)}@{DOMAINS[i % len(DOMAINS)]}"
        opt = "yes" if i % 3 != 2 else "no"
        rows.append({"Name": f"{f} {l}", "Email": email, "Newsletter": opt,
                     "City": r.choice(["Portland", "Seattle", "Eugene", "Bend", "Tacoma", "Salem", "Boise"])})
    return rows


def build(d):
    os.makedirs(f"{d}/Pictures/newsletter", exist_ok=True)
    os.makedirs(f"{d}/Documents/KB", exist_ok=True)
    os.makedirs(f"{d}/attachments", exist_ok=True)
    for i, p in enumerate(PRODUCTS):
        open(f"{d}/Pictures/newsletter/{p['photo']}", "wb").write(photo(p["color"], p["name"], 10 + i))
    for i, (fn, col, label) in enumerate(DECOYS):
        open(f"{d}/Pictures/newsletter/{fn}", "wb").write(photo(col, label, 20 + i))
    macros_docx(f"{d}/Documents/KB/Support_Macros.docx")
    cust = customers()
    with open(f"{d}/Documents/customers.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Name", "Email", "Newsletter", "City"])
        w.writeheader()
        w.writerows(cust)
    optin = [c for c in cust if c["Newsletter"] == "yes"]
    unsub = [optin[1], optin[6], optin[11]]
    changer = optin[4]
    new_addr = f"{changer['Name'].split()[0].lower()}.{changer['Name'].split()[1].lower()}.pdx@fastmail.com"
    recipients = sorted({c["Email"] for c in optin if c not in unsub and c is not changer} | {new_addr})

    sheet = DocPage("FERN & FENNEL", "Plants & pottery - Portland, OR", "December newsletter - promo sheet",
                    meta=[("Prepared by:", "Marco Rossi"), ("Promo period:", "1 - 24 December 2026")],
                    table=(["Product", "Regular", "December", "Photo file"],
                           [(p["name"], p["regular"], p["price"], p["photo"]) for p in PRODUCTS]),
                    widths=[250, 70, 70, 122],
                    notes=["**Product pages", *[f"{p['name'][:28]}...  {p['url']}" for p in PRODUCTS], "",
                           f"**Coupon code: {COUPON}  (15% off everything, valid until 24 December)", "",
                           "Photos are in the shared Pictures/newsletter folder - use the December shots listed above,",
                           "not last year's ones."])
    pdf_scan(f"{d}/attachments/December_Promo_Sheet.pdf", [scan_look(sheet.image(), seed=404)])

    mails, sup = [], []

    def add(key, frm, subject, when, body, unread=True, atts=(), kind=None, **extra):
        dom = frm.split("@")[1].rstrip(">")
        m = msg(frm, HANA, subject, when, mid(key, dom), body=body, attachments=atts)
        mails.append((m, when, unread))
        rec = {"key": key, "id": m["Message-ID"], "from": frm.split("<")[1].rstrip(">").lower() if "<" in frm else frm,
               "subject": subject, "kind": kind, "date": when.date().isoformat(), **extra}
        sup.append(rec)
        return m["Message-ID"]

    promo = add("promo", "Marco Rossi <marco@fernandfennel.shop>", "December newsletter - featured products",
                at(10, 19, 16, 30),
                "Hi Hana,\n\nhere is the promo sheet for the December newsletter: the three featured products, their "
                "December prices, the product links, which photos to use and the coupon code. Can you put the "
                "newsletter together in Thunderbird? I'd like to send it on 1 December.\n\nMarco",
                unread=False, kind="other",
                atts=[("December_Promo_Sheet.pdf", open(f"{d}/attachments/December_Promo_Sheet.pdf", "rb").read(),
                       "application/pdf")])

    def cname(c):
        return c["Name"]

    add("s1", "Olivia Bennett <olivia.bennett88@gmail.com>", "Cracked pot!!", at(10, 20, 8, 12),
        "Hi, my speckled planter arrived with a big crack in it :( Photo of the pot and the packing slip attached.\n\n"
        "Olivia", kind="support", macro="KB-01", order="FF-10517",
        atts=[("IMG_4471.jpg", slip_photo("FF-10517", "Speckled Stoneware Planter (L)", 1), "image/jpeg")])
    add("s2", "Daniel Cho <dcho@outlook.com>", "Where is my order FF-10482?", at(10, 20, 11, 40),
        "Hello, I ordered a fiddle leaf fig 9 days ago (order FF-10482) and it still hasn't arrived. Can you check?\n\n"
        "Thanks, Daniel", kind="support", macro="KB-02", order="FF-10482")
    add("s3", "Priya Nair <priya.nair@yahoo.com>", "Monstera leaves turning yellow", at(10, 20, 19, 5),
        "Hi! The lower leaves of the monstera I bought from you are turning yellow. I water it every two days. "
        "What am I doing wrong?\n\nPriya", kind="support", macro="KB-03", order=None)
    add("s4", "Greg Holt <greg.holt@gmail.com>", "Return request - order FF-10233", at(10, 21, 9, 15),
        "Hi, I'd like to return the ceramic pot set from order FF-10233. It was delivered on 2 August but I only "
        "unpacked it now and the colour doesn't match my living room.\n\nGreg", kind="support", macro="KB-04",
        order="FF-10233")
    add("s5", f"{cname(unsub[0])} <{unsub[0]['Email']}>", "Unsubscribe", at(10, 21, 12, 0),
        "Please unsubscribe me from your newsletter. Thanks.", kind="support", macro="KB-05", order=None)
    add("s6", f"{cname(unsub[1])} <{unsub[1]['Email']}>", "Please stop the newsletters", at(10, 21, 14, 22),
        "Hi, I love your shop but I get too many emails - please take me off the newsletter list.\n\n"
        f"{unsub[1]['Name'].split()[0]}", kind="support", macro="KB-05", order=None)
    add("s7", f"{cname(changer)} <{changer['Email']}>", "New email address", at(10, 21, 16, 48),
        f"Hi Fern & Fennel team,\n\nI'm closing this mailbox at the end of the month. Please use my new address "
        f"{new_addr} for everything from now on (newsletter included).\n\nBest,\n{changer['Name']}",
        kind="support", macro="KB-06", order=None)
    add("s8", "Maya Lopez <maya.lopez@icloud.com>", "Plant arrived broken", at(10, 22, 7, 30),
        "Hello, the pot of my pothos basket broke in transit and soil is everywhere. Picture attached (with the slip).\n\n"
        "Maya", kind="support", macro="KB-01", order="FF-10561",
        atts=[("photo.jpg", slip_photo("FF-10561", "Pothos 'Marble Queen' Basket", 2), "image/jpeg")])
    add("s9", "Sam Rivera <sam@bloomcafe.co>", "Wholesale inquiry - Bloom Cafe", at(10, 22, 9, 5),
        "Hi there, we run a small cafe in Eugene and would love to sell your planters. Do you do wholesale?\n\n"
        "Sam Rivera, Bloom Cafe", kind="support", macro="KB-07", order=None)
    add("s10", f"{cname(unsub[2])} <{unsub[2]['Email']}>", "remove me from your mailing list", at(10, 22, 10, 10),
        "remove me from your mailing list pls", kind="support", macro="KB-05", order=None)
    add("j1", "Shopify Payments <payouts@shopify-verify-payouts.com>", "Your payout is on hold - verify your account",
        at(10, 21, 3, 14), "We could not process your payout. Verify your bank details within 24 hours at "
        "http://shopify-verify-payouts.com/login or your store will be suspended.", kind="junk")
    add("j2", "SEO Expert Team <rank.first.page@gmail.com>", "Get your shop on Google page 1 !!!", at(10, 20, 22, 1),
        "Dear owner, we can bring fernandfennel.shop to the first page of Google in 7 days. Only $99.", kind="junk")
    add("inv", "Greenleaf Nursery Supply <billing@greenleafnursery.com>", "Invoice #4471 - October delivery",
        at(10, 19, 10, 0), "Invoice #4471 for 120 nursery pots and 40 bags of potting mix: $1,284.00, due 18 Nov.",
        unread=False, kind="other")
    add("shopify", "Shopify Community <community@shopify.com>", "Your weekly community digest", at(10, 20, 6, 0),
        "Top discussions this week: holiday shipping cut-off dates, product photography tips...", kind="other")
    add("ups", "UPS <pkginfo@ups.com>", "UPS pickup confirmed for Oct 23", at(10, 21, 18, 0),
        "Your pickup request PRN 7G4K2 is confirmed for Friday, October 23, 2026.", kind="other")
    mails.sort(key=lambda x: x[1])
    build_profile(f"{d}/thunderbird-profile.tar.gz", ME, {"Inbox": mails}, tz="America/Los_Angeles",
                  extra_prefs={"mail.spam.manualMark": True, "mail.spam.manualMarkMode": 0,
                               "mail.identity.id1.compose_html": True})
    dump_expected(d, {
        "host": HOST, "me": ME[1], "subject": SUBJECT, "coupon": COUPON, "unsub_url": UNSUB,
        "products": [{k: p[k] for k in ("key", "name", "price", "url", "photo")} for p in PRODUCTS],
        "photo_md5": {p["photo"]: md5(f"{d}/Pictures/newsletter/{p['photo']}") for p in PRODUCTS},
        "decoy_md5": {fn: md5(f"{d}/Pictures/newsletter/{fn}") for fn, _, _ in DECOYS},
        "recipients": recipients, "excluded": sorted({c["Email"] for c in unsub} | {changer["Email"]}),
        "all_customers": [c["Email"] for c in cust], "mails": sup, "promo_id": promo,
        "macros": {code: phrase for code, _, _, phrase in MACROS},
    })
