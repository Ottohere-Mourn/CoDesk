"""Evaluation helpers shared by the F6 cross-application pairs (pg_x1 .. pg_x4); deployed next to them in task_class/.

Everything Guest-facing goes through the pg_common module object (pc.guest, pc.fetch, ...) so a self-test can
monkeypatch pg_common alone.
"""
from __future__ import annotations

import os
import re
import zipfile

from evaluation_examples.task_class import pg_common as pc

OUT_FOLDERS = ("drafts", "sent", "unsent messages", "outbox")

# Like pg_common.MBOX, plus: html bodies unescaped, and every attachment of an outgoing message (Drafts, Sent,
# Unsent Messages, Outbox) written to a temp dir in the Guest so the evaluator can fetch and open it.
MAILX = r"""python3 - <<'PY'
import email, email.policy, glob, html, json, os, re, tempfile
OUTF = %r
tmp = tempfile.mkdtemp(prefix='pgx_mail_')
out, n = {}, 0
for prof in glob.glob(os.path.expanduser('~/.thunderbird/*/Mail')):
    for d, _, fs in os.walk(prof):
        for f in fs:
            p = os.path.join(d, f)
            if '.' in f or os.path.getsize(p) == 0:
                continue
            data = open(p, 'rb').read()
            if not data.startswith(b'From '):
                continue
            folder = os.path.relpath(p, prof).replace('.sbd/', '/')
            leaf = folder.rsplit('/', 1)[-1].lower()
            for chunk in re.split(rb'(?m)^From .*\r?\n', data)[1:]:
                try:
                    m = email.message_from_bytes(chunk, policy=email.policy.default)
                    status = int((m.get('X-Mozilla-Status') or '0').strip() or '0', 16)
                    if status & 0x0008:
                        continue
                    body = m.get_body(('plain', 'html'))
                    text = body.get_content() if body is not None else ''
                    if body is not None and body.get_content_type() == 'text/html':
                        text = re.sub(r'(?is)<(script|style).*?</\1>', ' ', text)
                        text = re.sub(r'(?i)<br\s*/?>|</p>|</div>|</li>|</tr>', '\n', text)
                        text = html.unescape(re.sub(r'<[^>]+>', ' ', text))
                    atts = []
                    for a in m.iter_attachments():
                        name = a.get_filename() or ''
                        info = {'name': name, 'type': a.get_content_type()}
                        if leaf in OUTF:
                            n += 1
                            payload = a.get_payload(decode=True) or b''
                            dst = os.path.join(tmp, str(n))
                            os.makedirs(dst)
                            fp = os.path.join(dst, re.sub(r'[^A-Za-z0-9._ -]', '_', name) or 'attachment')
                            open(fp, 'wb').write(payload)
                            info.update(path=fp, size=len(payload))
                        atts.append(info)
                    out.setdefault(folder, []).append({
                        'id': (m.get('Message-ID') or '').strip(), 'subject': str(m.get('Subject') or ''),
                        'from': str(m.get('From') or ''), 'to': str(m.get('To') or ''), 'cc': str(m.get('Cc') or ''),
                        'irt': str(m.get('In-Reply-To') or '') + ' ' + str(m.get('References') or ''),
                        'read': bool(status & 0x0001), 'flagged': bool(status & 0x0004),
                        'keys': (m.get('X-Mozilla-Keys') or '').split(), 'attachments': atts,
                        'text': text[:8000]})
                except Exception as e:
                    out.setdefault(folder, []).append({'id': '', 'subject': '', 'from': '', 'to': '', 'cc': '',
                        'irt': '', 'read': False, 'flagged': False, 'keys': [], 'attachments': [],
                        'text': '', 'error': repr(e)})
print('MAILX' + json.dumps(out))
PY"""


def mailboxes(env) -> dict | None:
    """{folder relative to Mail/ ('pop.x.com/Inbox', 'pop.x.com/Gala 2026/RSVPs'): [live messages]}."""
    return pc.guest_json(env, MAILX % (OUT_FOLDERS,), tag="MAILX")


def leaf(folder: str) -> str:
    return folder.rsplit("/", 1)[-1].lower()


def outgoing(boxes: dict, seeded=()) -> list[dict]:
    """Drafts plus anything the agent tried to send (Sent / Outbox / Unsent Messages), minus the seeded messages."""
    seeded = set(seeded)
    return [dict(m, folder=f) for f, msgs in (boxes or {}).items() if leaf(f) in OUT_FOLDERS for m in msgs
            if m["id"] not in seeded]


def where(boxes: dict) -> dict:
    """{Message-ID: set of folders (full relative paths)} for every live message."""
    out = {}
    for f, msgs in (boxes or {}).items():
        for m in msgs:
            out.setdefault(m["id"], set()).add(f)
    return out


def addressed(m: dict, addr: str) -> bool:
    return addr.lower() in (m.get("to", "") + " " + m.get("cc", "")).lower()


def fetch_attachment(env, att: dict, dest: str) -> str | None:
    """Fetch an outgoing message's attachment (dict from mailboxes()) to dest + its extension."""
    if not att.get("path"):
        return None
    ext = os.path.splitext(att.get("name") or "")[1].lower() or ".bin"
    return pc.fetch(env, att["path"], dest + ext)


# ---------------------------------------------------------------- office files
def xlsx_cells(path: str) -> list[tuple]:
    """[(sheet, row, col, value)] for every non-empty cell (cached values of formulas)."""
    from openpyxl import load_workbook
    out = []
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip() != "":
                    out.append((ws.title, c.row, c.column, c.value))
    return out


def xlsx_rows(path: str) -> list[tuple[str, list]]:
    """[(sheet, [values...])] per non-empty row."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in row):
                out.append((ws.title, list(row)))
    return out


def xlsx_charts(path: str) -> int:
    try:
        return sum(1 for n in zipfile.ZipFile(path).namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
    except Exception:
        return 0


def docx_text(path: str) -> str:
    """All text of word/document.xml (body, tables, text frames), one line per paragraph."""
    try:
        xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf8", "replace")
    except Exception:
        return ""
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    parts = re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>|(\n|\t)", xml)
    import html
    return html.unescape("".join(a or b for a, b in parts))


def docx_visuals(path: str) -> dict:
    try:
        names = zipfile.ZipFile(path).namelist()
    except Exception:
        return {"charts": 0, "media": 0, "objects": 0}
    return {"charts": sum(n.startswith("word/charts/chart") for n in names),
            "media": sum(n.startswith("word/media/") for n in names),
            "objects": sum(n.startswith("word/embeddings/") for n in names)}


def pdf_text(path: str) -> str:
    try:
        import fitz
        with fitz.open(path) as doc:
            return "\n".join(p.get_text() for p in doc)
    except Exception:
        return ""


def pdf_visuals(path: str) -> dict:
    """Raster images and vector drawing paths (a chart pasted from Calc exports as vector paths)."""
    try:
        import fitz
        with fitz.open(path) as doc:
            return {"pages": len(doc), "images": sum(len(p.get_images()) for p in doc),
                    "drawings": sum(len(p.get_drawings()) for p in doc)}
    except Exception:
        return {"pages": 0, "images": 0, "drawings": 0}


def is_pdf(path: str | None) -> bool:
    try:
        return bool(path) and open(path, "rb").read(5) == b"%PDF-"
    except OSError:
        return False


# ---------------------------------------------------------------- text matching
NUM = re.compile(r"(?<![\d.])[-\u2212]?\d{1,3}(?:[,\u00a0\u202f]\d{3})+(?:\.\d+)?|(?<![\d.,])[-\u2212]?\d+(?:\.\d+)?")


def numbers(text: str) -> list[float]:
    out = []
    for tok in NUM.findall(str(text or "")):
        try:
            out.append(float(re.sub(r"[,\u00a0\u202f]", "", tok).replace("\u2212", "-")))
        except ValueError:
            pass
    return out


def has_amount(text_or_values, value: float, tol: float = 0.006) -> bool:
    vals = text_or_values if isinstance(text_or_values, list) else numbers(text_or_values)
    return any(abs(abs(v) - abs(value)) <= tol for v in vals)


def cell_numbers(cells: list[tuple]) -> list[float]:
    out = []
    for *_, v in cells:
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)):
            out.append(float(v))
        else:
            out += numbers(str(v))
    return out


def percents(text: str) -> list[float]:
    return [float(m.replace("−", "-").replace(",", ".")) for m in
            re.findall(r"([-+−]?\d+(?:[.,]\d+)?)\s*%", str(text or ""))]


def flat(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def in_order(text: str, names: list[str]) -> bool:
    """All names occur in text (case/punctuation-insensitive) and their first occurrences are in this order."""
    t = flat(text)
    pos = [t.find(flat(n)) for n in names]
    return all(p >= 0 for p in pos) and pos == sorted(pos)


def capped(score: float, caps: list[tuple[bool, float]]) -> float:
    """Apply hard caps: [(condition_that_triggers_cap, max_score)]."""
    for cond, cap in caps:
        if cond:
            score = min(score, cap)
    return score
