"""W3 - one sales CRM (Bluepeak CRM on :8103) in one Chrome session: merging duplicate accounts per the data quality
guide vs logging the week's calls, opportunities and contact changes from meeting notes. See spec_w3.md."""
from __future__ import annotations

import csv
import io
import re

from evaluation_examples.task_class.pg_common import (HOME, expected, fetch, open_chrome, result, sh, sql,
                                                      start_webapp, upload, webapp_db)

NAME, PORT = "w3", 8103
URL = f"http://127.0.0.1:{PORT}/"
DOCS = {"call_notes_wk43.docx": "Documents", "pipeline_q4.xlsx": "Documents",
        "card_ben_alvarez.png": "Pictures", "card_grace_liu.png": "Pictures"}
MERGE_LOG = f"{HOME}/Documents/crm_merge_log.csv"
PIPELINE = f"{HOME}/Documents/pipeline_q4.xlsx"
FIELDS = ["name", "website", "tax_id", "erp_no", "phone", "industry", "street", "city", "postcode", "country", "owner_id"]


def setup(sc):
    start_webapp(sc, f"{NAME}/app_w3.py", NAME, PORT, f"{NAME}/seed.db")
    # documents: never overwrite (setup runs again before the second task of a sequential run)
    upload(sc, [(f"{NAME}/{f}", f"{HOME}/.cache/pg{NAME}/{f}") for f in DOCS])
    sh(sc, "mkdir -p ~/Documents ~/Pictures && " + " && ".join(
        f"cp -n ~/.cache/pg{NAME}/{f} ~/{d}/{f}" for f, d in DOCS.items()))
    open_chrome(sc, [URL])


def _state(env, exp):
    m = exp["seed_max"]
    return sql(env, webapp_db(NAME), {
        "accounts": "SELECT * FROM accounts",
        "kids": "SELECT account_id, COUNT(*) AS n FROM (SELECT account_id FROM contacts UNION ALL SELECT account_id FROM opps "
                "UNION ALL SELECT account_id FROM activities) GROUP BY account_id",
        "merges": "SELECT master_id, merged_id FROM merges",
        "acts": f"SELECT * FROM activities WHERE id > {m['activities']}",
        "opps": f"SELECT * FROM opps WHERE id > {m['opps']}",
        "contacts": "SELECT * FROM contacts"})


def _resolver(accounts):
    by = {a["id"]: a for a in accounts}

    def resolve(aid):
        for _ in range(20):
            a = by.get(aid)
            if not a or a["status"] != "merged" or not a["merged_into"]:
                return aid
            aid = a["merged_into"]
        return aid
    return by, resolve


def _norm(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def _digits(s):
    return re.sub(r"\D", "", s or "")[-10:]


def evaluate_dedupe(env, dest: str) -> dict:
    exp = expected(NAME)
    st = _state(env, exp)
    if not st:
        return result(0.0, {"app_db": "unreadable"})
    by, resolve = _resolver(st["accounts"])
    kids = {k["account_id"]: k["n"] for k in st["kids"]}
    checks, part, right = {}, 0.0, 0
    pair_ids = set()
    for k, p in exp["pairs"].items():
        m, d = by.get(p["master"], {}), by.get(p["dup"], {})
        pair_ids |= {p["master"], p["dup"]}
        merged_ok = d.get("status") == "merged" and resolve(p["dup"]) == p["master"] and m.get("status") == "active"
        wrong = [f for f in FIELDS if _norm(m.get(f)) != _norm(p["final"][f])] if merged_ok else FIELDS
        moved = kids.get(p["dup"], 0) == 0
        right += merged_ok
        part += (0.6 * merged_ok + 0.4 * (len(FIELDS) - len(wrong)) / len(FIELDS)) if merged_ok and moved else 0.0
        checks[k] = {"pair": " / ".join(p["names"]), "merged_into_right_master": merged_ok,
                     "dup_status": d.get("status"), "dup_into": d.get("merged_into"), "children_moved": moved,
                     "wrong_fields": wrong if merged_ok else "n/a"}
        if merged_ok and wrong:
            checks[k]["got"] = {f: m.get(f) for f in wrong}
            checks[k]["want"] = {f: p["final"][f] for f in wrong}
    n = len(exp["pairs"])
    look_ok, look_bad = 0, []
    look_ids = set()
    for k, (a, b) in exp["lookalikes"].items():
        look_ids |= {a, b}
        bad = any(by.get(x, {}).get("status") != "active" for x in (a, b)) or \
            any(r["master_id"] in (a, b) and r["merged_id"] in (a, b) for r in st["merges"])
        look_ok += not bad
        if bad:
            look_bad.append(k)
    checks["lookalikes_merged"] = look_bad
    collateral = sorted(a["id"] for a in st["accounts"] if a["status"] == "merged" and a["id"] not in pair_ids | look_ids)
    collateral += sorted(p["master"] for p in exp["pairs"].values() if by.get(p["master"], {}).get("status") == "merged")
    checks["other_accounts_merged"] = collateral
    # the merge log
    log_score, rows = 0.0, []
    path = fetch(env, MERGE_LOG, dest + ".csv")
    if path:
        try:
            text = open(path, encoding="utf-8-sig", errors="replace").read()
            for row in csv.reader(io.StringIO(text)):
                nums = [int(x) for x in re.findall(r"(?<!\d)(\d{4})(?!\d)", " ".join(row))]
                if len(nums) >= 2:
                    rows.append((nums[0], nums[1]))
        except Exception as e:  # noqa: BLE001
            checks["log_error"] = str(e)[:200]
        done = {(r["master_id"], r["merged_id"]) for r in st["merges"]}
        want = {(p["master"], p["dup"]) for p in exp["pairs"].values()}
        good = {r for r in rows if r in want and r in done}
        bad = [r for r in rows if r not in done]
        log_score = max(0, len(good) - len(bad)) / n
        checks["log_rows"] = rows
        checks["log_correct"] = f"{len(good)}/{n}"
        checks["log_rows_not_matching_a_merge"] = bad
    else:
        checks["log"] = "missing"
    done_frac = right / n
    score = (0.65 * part / n + 0.10 * done_frac * look_ok / len(exp["lookalikes"])
             + 0.05 * done_frac * (not collateral) + 0.20 * log_score)
    checks["pairs_merged_right"] = f"{right}/{n}"
    if look_bad:
        score = min(score, 0.5)
        checks["cap"] = "merged different companies (look-alikes) -> score <= 0.5"
    if right < n / 2:
        score = min(score, 0.4)
        checks["cap"] = checks.get("cap", "") + " fewer than half of the duplicate pairs merged correctly -> score <= 0.4"
    return result(score, checks)


def evaluate_calls(env, dest: str) -> dict:
    exp = expected(NAME)
    st = _state(env, exp)
    if not st:
        return result(0.0, {"app_db": "unreadable"})
    by, resolve = _resolver(st["accounts"])
    checks = {}
    # 1) calls
    calls_pts, any_call = 0.0, False
    for c in exp["calls"]:
        acc = resolve(c["account"])
        here = [a for a in st["acts"] if a["account_id"] == acc and a["date"] == c["date"]]
        any_call |= any(a["type"] == "Call" for a in st["acts"])
        full = [a for a in here if a["type"] == "Call" and a["contact_id"] == c["contact"]
                and re.search(c["kw"], f"{a['subject']} {a['notes']}", re.I)]
        pts = 1.0 if full else (0.5 if here else 0.0)
        calls_pts += pts
        checks[f"call {c['date']} {c['name']}"] = "ok" if full else ("logged but type/contact/summary wrong" if here else "missing")
    # 2) opportunities
    opp_pts, by_name = 0.0, {}
    for dl in exp["deals"]:
        acc = resolve(dl["account"])
        cands = [o for o in st["opps"] if o["account_id"] == acc and re.search(dl["rx"], o["name"] or "", re.I)]
        best, bs = None, 0.0
        for o in cands:
            s = 0.4 + 0.2 * (abs((o["amount"] or 0) - dl["amount"]) < 0.01) + 0.2 * (o["stage"] == dl["stage"]) \
                + 0.2 * (o["close_date"] == dl["close"])
            if s > bs:
                best, bs = o, s
        opp_pts += bs
        by_name[dl["name"]] = best
        checks[f"opp {dl['name']}"] = ({"id": f"OPP-{best['id']}", "amount": best["amount"], "stage": best["stage"],
                                        "close": best["close_date"], "score": round(bs, 2)} if best else "missing")
    # 3) contacts and the moved address
    contacts = {c["id"]: c for c in st["contacts"]}
    items = []
    for cid, want in exp["contact_updates"].items():
        c = contacts.get(int(cid), {})
        ok = [_digits(c.get("phone")) == want["phone"][-10:]]
        if "title" in want:
            ok.append(bool(re.search(want["title"], c.get("title") or "", re.I)))
        items.append(sum(ok) / len(ok))
        checks[f"contact update {want['name']}"] = {"title": c.get("title"), "phone": c.get("phone")}
    for nc in exp["new_contacts"]:
        acc = resolve(nc["account"])
        hit = [c for c in st["contacts"] if c["account_id"] == acc and c["id"] > exp["seed_max"]["contacts"]
               and (c["first"] or "").strip().lower() == nc["first"].lower() and (c["last"] or "").strip().lower() == nc["last"].lower()]
        s = max((0.4 + 0.3 * ((c["email"] or "").strip().lower() == nc["email"]) + 0.3 * (_digits(c["phone"]) == nc["phone"][-10:])
                 for c in hit), default=0.0)
        items.append(s)
        checks[f"new contact {nc['first']} {nc['last']}"] = round(s, 2)
    co = by.get(exp["coastal"]["account"], {})
    f_ok = [bool(re.search(rx, _norm(co.get(f)), re.I)) for f, rx in exp["coastal"]["fields"].items()]
    items.append(sum(f_ok) / len(f_ok))
    checks["coastal address"] = {f: co.get(f) for f in exp["coastal"]["fields"]}
    # 4) the pipeline sheet: CRM IDs of the new opportunities
    sheet_ok, kept = 0, False
    path = fetch(env, PIPELINE, dest + ".xlsx")
    if path:
        from openpyxl import load_workbook
        try:
            rows = [list(r) for r in load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True)]
        except Exception as e:  # noqa: BLE001
            rows, checks["sheet_error"] = [], str(e)[:200]
        flat = " ".join(str(c) for r in rows for c in r if c is not None)
        kept = all(s in flat for s in exp["pipeline_seed_ids"])
        opps_by_id = {o["id"]: o for o in st["opps"]}
        for dl in exp["deals"]:
            got = None
            for r in rows:
                if any(re.search(dl["rx"], str(c or ""), re.I) for c in r):
                    for c in r:
                        m = re.search(r"OPP-?\s*(\d+)", str(c or ""), re.I)
                        if m:
                            got = int(m.group(1))
            o = opps_by_id.get(got)
            ok = bool(o) and o["account_id"] == resolve(dl["account"]) and re.search(dl["rx"], o["name"] or "", re.I) is not None
            sheet_ok += ok
            checks[f"sheet {dl['name']}"] = f"OPP-{got}: {'ok' if ok else 'not a matching new opportunity'}" if got else "no CRM ID"
        checks["sheet_old_ids_kept"] = kept
    else:
        checks["sheet"] = "missing"
    nc, nd = len(exp["calls"]), len(exp["deals"])
    score = (0.30 * calls_pts / nc + 0.30 * opp_pts / nd + 0.15 * sum(items) / len(items) + 0.20 * sheet_ok / nd
             + 0.05 * (kept and sheet_ok > 0))
    if not any_call:
        score = min(score, 0.3)
        checks["cap"] = "no Call activity logged at all -> score <= 0.3"
    return result(score, checks)
