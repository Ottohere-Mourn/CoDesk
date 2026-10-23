"""W1 - Larkspur Support Desk (help desk web app on :8101): triaging the untriaged queue vs answering six tickets.

build(out_dir) writes: app_w1.py, seed.db, Support_Triage_Policy_v3.pdf, support_log_2026-10.xlsx, expected.json.
"""
from __future__ import annotations

import io
import json
import os
import random
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from PIL import Image, ImageDraw  # noqa: E402

import lib_w  # noqa: E402
from lib_w import font  # noqa: E402

PASSWORD = "Larkspur-Desk-26"

# id, username, name, team, title, ooo_until, backup id
STAFF = [
    (1, "sam.rivera", "Sam Rivera", "Lead", "Support Team Lead", None, None),
    (2, "nora.patel", "Nora Patel", "Billing", "Senior Support Engineer", None, 3),
    (3, "diego.alvarez", "Diego Alvarez", "Billing", "Support Engineer", None, 2),
    (4, "mei.tanaka", "Mei Tanaka", "Technical", "Senior Support Engineer", None, 6),
    (5, "oskar.lind", "Oskar Lind", "Technical", "Support Engineer", "30 Oct 2026", 6),
    (6, "femi.adeyemi", "Femi Adeyemi", "Technical", "Support Engineer", None, 4),
    (7, "hannah.cole", "Hannah Cole", "Accounts", "Senior Support Engineer", None, 8),
    (8, "raj.iyer", "Raj Iyer", "Accounts", "Support Engineer", "23 Oct 2026", 7),
]
SELF = "Self-serve (no account manager)"
# id, name, plan, domain, phone, account manager, dedicated engineer id, customer since, notes
ORGS = [
    (1, "Brightwater Logistics", "Enterprise", "brightwater-logistics.com", "+44 20 7946 0331", "Marcus Webb", 2, "2021-03", "UK entities: Brightwater Ltd, Brightwater Freight Ltd. Pays by bank transfer."),
    (2, "Quarry Hill Health", "Enterprise", "quarryhillhealth.org", "+1 503 555 0148", "Priya Shah", 5, "2022-07", "Uses SAML SSO (Azure AD). Billing office of 35 users."),
    (3, "Halvorsen Freight", "Enterprise", "halvorsenfreight.no", "+47 22 55 01 90", "Marcus Webb", 4, "2020-11", "Nightly ERP sync through the REST API."),
    (4, "Castellan Retail Group", "Enterprise", "castellanretail.com", "+1 312 555 0102", "Lena Brooks", 7, "2019-05", "Contracted seats: 260. Seasonal ramp-up every November."),
    (5, "Meridian Analytics", "Enterprise", "meridian-analytics.io", "+1 415 555 0190", "Lena Brooks", 6, "2023-01", ""),
    (6, "Pinecrest Dental", "Business", "pinecrestdental.com", "+1 206 555 0117", "Priya Shah", None, "2022-02", "Three clinics."),
    (7, "Tidewater Marine Supply", "Business", "tidewatermarine.com", "+1 757 555 0164", "Marcus Webb", None, "2026-09", "Migrating from QuickBooks Desktop."),
    (8, "Ardent Studio", "Business", "ardentstudio.co", "+44 161 555 0122", "Lena Brooks", None, "2024-04", "Uses Zapier integration."),
    (9, "Kessler & Moss LLP", "Business", "kesslermoss.law", "+1 617 555 0176", "Priya Shah", None, "2021-10", "Month-end close in the last week of each month."),
    (10, "Northgate Schools", "Business", "northgateschools.org", "+1 651 555 0133", "Marcus Webb", None, "2023-08", "Pays by card."),
    (11, "Oakline Bakery Co.", "Business", "oaklinebakery.com", "+1 773 555 0181", "Lena Brooks", None, "2024-01", "Four shops, central office."),
    (12, "Juniper Yoga", "Starter", "juniperyoga.studio", "+1 971 555 0109", SELF, None, "2025-02", ""),
    (13, "Fernwood Bikes", "Starter", "fernwoodbikes.com", "+1 541 555 0155", SELF, None, "2024-06", ""),
    (14, "Lumen Photography", "Starter", "lumenphoto.co", "+1 720 555 0128", SELF, None, "2025-08", ""),
    (15, "Redbird Cafe", "Starter", "redbirdcafe.com", "+1 512 555 0193", SELF, None, "2025-11", ""),
]

# ------------------------------------------------------------------ the untriaged queue (task A) incl. two of task B's
# id: (org, requester, email, channel, impact, created, subject, body, attachment, expected {team, priority, assignee} | kind)
TRIAGE = {
    1046: (1, "Tomasz Wilk", "t.wilk@brightwater-logistics.com", "Email", "Entire organization", "2026-10-20 16:42",
           "Account suspended for non-payment - but invoice INV-88213 was paid",
           "Hello,\n\nThis afternoon our whole company was locked out of Larkspur with the message \"Your subscription is "
           "suspended for non-payment\". Our October invoice INV-88213 (GBP 9,840.00) was paid by bank transfer on "
           "2 October, reference BWL-OCT-INV88213.\n\nNobody in our finance team (14 users) can send invoices. Please "
           "restore access as soon as possible.\n\nTomasz Wilk\nFinance Systems Lead, Brightwater Logistics", None,
           {"team": "Billing", "priority": "Urgent", "assignee": "Nora Patel"}),
    1047: (12, "Alicia Moreno", "alicia@juniperyoga.studio", "Web form", "Single user", "2026-10-20 19:05",
           "Sign-in alert from Sao Paulo - that wasn't me",
           "Hi,\n\nI just got an email saying \"New sign-in to your Larkspur account from Sao Paulo, Brazil\". I'm in "
           "Portland and nobody else uses my login. I changed my password right away but I'm worried someone has seen "
           "our client invoices. What should I do?\n\nAlicia", None,
           {"team": "Accounts", "priority": "Urgent", "assignee": None}),
    1049: (2, "Grace Lin", "g.lin@quarryhillhealth.org", "Email", "Entire organization", "2026-10-21 07:58",
           "SSO login fails for all staff after our certificate rotation",
           "Hello support,\n\nWe rotated our Azure AD SAML signing certificate last night. Since then nobody at Quarry "
           "Hill can sign in to Larkspur with SSO - the login page shows \"SAML response signature invalid\". We "
           "uploaded the new certificate under Admin > Security > SSO but it still fails.\n\nThis blocks our whole "
           "billing office (35 users).\n\nGrace Lin\nIT Manager, Quarry Hill Health", None,
           {"team": "Accounts", "priority": "Urgent", "assignee": "Femi Adeyemi"}),
    1050: (13, "Owen Price", "owen@fernwoodbikes.com", "Web form", "Single user", "2026-10-21 09:12",
           "How do I update the credit card we pay with?",
           "Hi,\n\nOur company card expires at the end of the month. Where do I change the card that Larkspur charges "
           "for our subscription?\n\nThanks,\nOwen", None,
           {"team": "Billing", "priority": "Low", "assignee": None}),
    1052: (None, "Kevin Marsh", "kevin@rankboost-seo.net", "Email", "Single user", "2026-10-21 10:31",
           "Get larkspur.io to #1 on Google - 70% off this week only",
           "Hi there,\n\nI was looking at larkspur.io and noticed you are missing out on a lot of organic traffic. Our "
           "white-hat SEO packages start at just $299/month and we guarantee first-page rankings in 60 days.\n\nReply "
           "YES for a free audit!\n\nKevin Marsh | RankBoost SEO", None, "spam"),
    1053: (8, "Priya Nair", "priya@ardentstudio.co", "Web form", "Multiple users", "2026-10-21 11:47",
           "Zapier webhooks stopped firing",
           "Hi,\n\nSince Saturday our \"invoice paid\" webhook to Zapier no longer fires, so our project tracker is not "
           "updated any more. The webhook still shows as Active under Settings > Integrations. Three of us rely on "
           "this every day.\n\nPriya Nair\nStudio Manager, Ardent Studio", None,
           {"team": "Technical", "priority": "High", "assignee": None}),
    1054: (10, "Mark Ellison", "m.ellison@northgateschools.org", "Email", "Single user", "2026-10-21 13:20",
           "Problem with our account - screenshot attached",
           "Hello,\n\nI got this message when I logged in this morning, see the screenshot. What do I need to do?\n\n"
           "Mark Ellison\nBursar, Northgate Schools", "payment_failed.png",
           {"team": "Billing", "priority": "Normal", "assignee": None}),
    1055: (3, "Ingrid Halvorsen", "ingrid@halvorsenfreight.no", "Web form", "Multiple users", "2026-10-21 14:03",
           "Dashboard takes 30+ seconds to load",
           "Hi,\n\nFor the last two days the main dashboard takes 30-40 seconds to load for our dispatch team (6 people). "
           "Other pages are fine. We are in Oslo and use Chrome.\n\nIngrid", None,
           {"team": "Technical", "priority": "High", "assignee": "Mei Tanaka"}),
    1056: (7, "Carl Jensen", "carl.jensen@tidewatermarine.com", "Web form", "Single user", "2026-10-21 15:36",
           "CSV import keeps failing",
           "Hi,\n\nI'm trying to import last month's invoices from our old system with the CSV import and it keeps "
           "failing. Screenshot attached - what are we doing wrong?\n\nCarl", "import_error.png",
           {"team": "Technical", "priority": "Normal", "assignee": None}),
    1057: (8, "Priya Nair", "priya@ardentstudio.co", "Email", "Multiple users", "2026-10-22 08:14",
           "Webhooks to Zapier not working",
           "Hello again,\n\nI reported this yesterday but haven't heard back. Our invoice-paid webhook to Zapier still "
           "isn't firing since Saturday. Can someone look at it please?\n\nPriya", None, "duplicate"),
    1058: (14, "Dev Malhotra", "dev@lumenphoto.co", "Web form", "Single user", "2026-10-22 08:40",
           "Remove a former employee's user",
           "Hi,\n\nOur assistant Jess left last week. How do I remove her user so she can't sign in any more? I'm the "
           "account owner.\n\nDev", None,
           {"team": "Accounts", "priority": "Low", "assignee": None}),
    1059: (4, "Monica Reyes", "monica.reyes@castellanretail.com", "Email", "Multiple users", "2026-10-22 09:05",
           "Need 40 more seats before Black Friday",
           "Hi team,\n\nWe are bringing 40 seasonal staff into our billing back office from 9 November for the Black "
           "Friday period. Can you add 40 seats to our account? They will need access until early January.\n\n"
           "Monica Reyes\nDirector of Finance Operations, Castellan Retail Group", None,
           {"team": "Accounts", "priority": "High", "assignee": "Hannah Cole"}),
    1060: (None, "Lucas Grant", "lucas@coinbridge-pay.io", "Email", "Single user", "2026-10-22 09:51",
           "Partnership opportunity: let your customers pay invoices in crypto",
           "Hello Larkspur team,\n\nCoinBridge Pay lets your users accept BTC, ETH and USDT on every invoice with a "
           "2-line integration. We'd love to set up a 20-minute call to discuss a revenue-share partnership.\n\n"
           "Best,\nLucas Grant\nHead of Partnerships, CoinBridge Pay", None, "spam"),
    1061: (11, "Sofia Bianchi", "sofia@oaklinebakery.com", "Web form", "Entire organization", "2026-10-22 10:18",
           "URGENT - see screenshot",
           "This is happening to everyone in our office since this morning. Screenshot attached. Please help!\n\nSofia",
           "server_error.png", {"team": "Technical", "priority": "High", "assignee": None}),
}
DUPLICATE_OF = {1057: 1053}

# ------------------------------------------------------------------ task B's six tickets
# id: (org, requester, email, created, subject, body, team, priority, status in seed)
B_SEED = {
    1037: (9, "Daniel Moss", "d.moss@kesslermoss.law", "2026-10-19 10:22", "Excel export of Aged Receivables is empty",
           "Hello,\n\nSince your update last week, exporting the Aged Receivables report to Excel gives us an .xlsx "
           "file with only the header row - no data at all. The report looks fine on screen. We need it for our "
           "month-end close next week. Is there a fix?\n\nDaniel Moss\nPractice Manager, Kessler & Moss LLP",
           "Technical", "Normal"),
    1043: (6, "Laura Kim", "laura@pinecrestdental.com", "2026-10-20 11:05", "Where can I download our 2025 invoices as PDF?",
           "Hi!\n\nOur accountant is asking for PDF copies of every invoice you sent us in 2025. Is there a way to "
           "download them all at once rather than one by one from the emails?\n\nThanks,\nLaura", "Billing", "Low"),
    1048: (15, "Hector Ruiz", "hector@redbirdcafe.com", "2026-10-21 09:40", "Charged twice for October",
           "Hi,\n\nI see two charges of $29.00 from LARKSPUR SOFTWARE on our bank statement for October. We only have "
           "one account. Please refund the extra one.\n\nHector", "Billing", "Normal"),
    1051: (3, "Lars Berg", "lars.berg@halvorsenfreight.no", "2026-10-21 12:30", "HTTP 429 errors from the API during nightly sync",
           "Hello,\n\nOur nightly ERP sync started failing with \"HTTP 429 Too Many Requests\" after a few minutes. We "
           "send about 4,500 requests per minute at peak. What is the rate limit for our account, and how should our "
           "client handle the 429 responses?\n\nLars Berg\nIntegration Developer, Halvorsen Freight", "Technical", "High"),
}
# what task B must produce (regexes are matched case-insensitively against the new public replies / internal notes)
B_EXPECT = {
    1037: {"kb": "KB-140", "status": "On-hold", "facts": {"csv workaround": r"\bcsv\b", "fix version/date": r"5\.2\.1|28\s*oct|oct\w*\s*28|2026-10-28|10/28"},
           "note": {"engineering ticket": r"ENG[\s-]?2291"}},
    1043: {"kb": "KB-104", "status": "Solved", "facts": {"invoice history": r"invoice\s+history"}},
    1048: {"kb": "KB-110", "status": "Pending", "facts": {"asks last four digits": r"last\s+(4|four)\b",
                                                           "5-7 business days": r"5\s*(-|–|—|to)\s*7"}},
    1051: {"kb": "KB-130", "status": "Solved", "facts": {"enterprise limit 3,000/min": r"3[,.\s]?000",
                                                          "retry-after header": r"retry[\s-]?after"}},
    1056: {"kb": "KB-121", "status": "Solved", "facts": {"date format": r"yyyy-mm-dd"}},
    1059: {"kb": "KB-115", "status": "Pending", "facts": {"names account manager": r"lena\s+brooks"},
           "note": {"hand-off to account manager": r"lena\s+brooks"}},
}

KB = [
    ("KB-001", "Ticket status guide", "Process",
     "How we use ticket statuses at Larkspur Support. Set the status in the ticket form when you submit your reply.\n\n"
     "| Status | Use it when |\n| New | Nobody has looked at the ticket yet (set by the system). |\n"
     "| Open | We are working on it and owe the customer an answer. |\n"
     "| Pending | We replied and are waiting for the customer to send something (information, a confirmation), OR we "
     "handed the request over to the customer's account manager. |\n"
     "| On-hold | The problem is a confirmed bug that is waiting for Engineering. |\n"
     "| Solved | Our reply fully answers the question or fixes the problem. |\n"
     "| Closed | Spam and duplicates only. Everything else is closed automatically 7 days after Solved. |\n\n"
     "## Internal notes that are required\n"
     "- On-hold: add an internal note with the engineering ticket number (ENG-....) from the Known issues article.\n"
     "- Hand-offs to an account manager (e.g. Enterprise seat or contract changes): add an internal note naming the "
     "account manager so they pick it up.\n\n"
     "## Replies\nAlways answer with a public reply and link the knowledge-base article you used (e.g. \"see KB-104\")."),
    ("KB-002", "Writing good replies", "Process",
     "Greet the customer by first name, answer the question in the first two sentences, then give the steps.\n\n"
     "- Cite the article number so the customer can find it again.\n- Never paste internal notes or engineering "
     "ticket numbers into a public reply.\n- Sign with your first name and \"Larkspur Support\"."),
    ("KB-104", "Downloading past invoices", "Billing",
     "Customers can download every invoice we have sent them as PDF.\n\n## Steps\n"
     "- Go to Settings > Billing > Invoice history.\n- Pick the year in the filter at the top.\n"
     "- Click \"Download all (ZIP)\" to get every PDF for that year at once, or the PDF icon on a single row.\n\n"
     "Invoices are kept for 7 years. Only account owners and billing admins can see Invoice history."),
    ("KB-105", "Customising invoice templates", "Billing",
     "Invoice templates (logo, colours, footer text) are edited under Settings > Branding > Invoice template. "
     "Changes apply to invoices created after saving; existing invoices keep their original look."),
    ("KB-108", "Updating the card on file", "Billing",
     "Account owners can change the card we charge under Settings > Billing > Payment method > Replace card. The "
     "new card is used from the next charge. If a payment has already failed, click \"Retry payment\" after "
     "replacing the card."),
    ("KB-110", "Duplicate or unexpected charges", "Billing",
     "Customers sometimes see two charges for the same period, e.g. after a card retry.\n\n## What to do\n"
     "- Ask the customer for the last four digits of the card and the dates of both charges (we cannot match bank "
     "statement lines without them).\n- Billing verifies the charges in the payment provider.\n"
     "- A confirmed duplicate charge is refunded to the same card within 5-7 business days after verification.\n\n"
     "Do not promise a refund before Billing has verified the charge."),
    ("KB-111", "Refunds after downgrades and cancellations", "Billing",
     "Annual plans: unused full months are refunded pro rata when a customer downgrades or cancels. Monthly plans are "
     "not refunded for the current month. Refunds take 5-10 business days."),
    ("KB-115", "Adding or removing seats", "Accounts",
     "## Starter and Business plans\nAccount owners change seats themselves under Admin > Users > Seats. The change is "
     "immediate and billed pro rata.\n\n## Enterprise plans\nSeats are part of the Enterprise contract. Support does not "
     "change seat counts. Tell the customer that their account manager (see the organization page) will contact them "
     "within 1 business day to update the contract, and hand the request over with an internal note naming the "
     "account manager."),
    ("KB-118", "Removing a user", "Accounts",
     "Account owners remove users under Admin > Users: open the user, click \"Deactivate\". Deactivated users cannot "
     "sign in; their invoices and history are kept. The seat becomes free immediately."),
    ("KB-121", "CSV import error codes", "Technical",
     "The invoice CSV import stops at the first invalid row and reports an error code.\n\n"
     "| Code | Meaning | Fix |\n| E-4017 | Unknown currency code | Use a 3-letter ISO 4217 code such as USD or EUR. |\n"
     "| E-4031 | Invalid date | Dates must be written as YYYY-MM-DD (for example 2026-10-21) in every date column. |\n"
     "| E-4052 | Duplicate invoice number | Invoice numbers must be unique within the account. |\n"
     "| E-4060 | File too large | Split the file; the limit is 10 MB or 5,000 rows. |\n\n"
     "After fixing the file, import it again; rows imported before the error are not duplicated."),
    ("KB-122", "Exporting reports", "Technical",
     "Every report can be exported as CSV, XLSX or PDF from the Export menu at the top right of the report. Large "
     "exports (over 50,000 rows) are emailed as a download link."),
    ("KB-130", "API rate limits", "Technical",
     "API requests are limited per account and per minute.\n\n| Plan | Limit |\n| Starter | 60 requests per minute |\n"
     "| Business | 600 requests per minute |\n| Enterprise | 3,000 requests per minute |\n\n"
     "When the limit is exceeded the API answers HTTP 429 Too Many Requests with a Retry-After header (in seconds). "
     "Clients must wait that long before retrying. For bulk syncs use the batch endpoints (up to 100 records per "
     "request) and spread the load over the sync window."),
    ("KB-131", "Webhook delivery and retries", "Technical",
     "Webhooks are retried 8 times with exponential backoff over 24 hours. After 8 failures the webhook is paused "
     "and shown as \"Paused\" under Settings > Integrations."),
    ("KB-135", "Setting up SAML single sign-on", "Accounts",
     "SSO is available on Enterprise plans. Upload the identity provider metadata under Admin > Security > SSO. When "
     "the IdP signing certificate is rotated, the new certificate must be uploaded before the old one expires."),
    ("KB-140", "Known issues - October 2026", "Technical",
     "Confirmed bugs that Engineering is working on. Tickets about these issues go On-hold with an internal note "
     "naming the ENG ticket.\n\n| ENG ticket | Problem | Status | Workaround |\n"
     "| ENG-2291 | XLSX export of reports produces a file with only the header row (since release 5.2.0) | Fix in "
     "release 5.2.1, scheduled for 28 Oct 2026 | Export as CSV and open the CSV in Excel |\n"
     "| ENG-2304 | Dashboard charts stay blank in Safari 17 | Investigating | Use Chrome or Firefox |\n"
     "| ENG-2310 | Webhooks delayed by up to 15 minutes | Fixed on 18 Oct 2026 | - |"),
    ("KB-150", "Resetting two-factor authentication", "Accounts",
     "If a user lost their 2FA device, the account owner can reset it under Admin > Users > (user) > Reset 2FA. "
     "Support can only reset 2FA for the account owner after an identity check by phone."),
]

# filler history: (org, requester first name, subject, body, team, priority)
FILLER = [
    (6, "Laura", "Invoice email not received by our accountant", "Our accountant says she never gets the invoice emails. Can you check?", "Billing", "Normal"),
    (3, "Lars", "API key rotation question", "How do we rotate API keys without downtime?", "Technical", "Normal"),
    (11, "Sofia", "Change company address on invoices", "We moved our central office. How do we change the address on our invoices?", "Billing", "Normal"),
    (1, "Tomasz", "Add a second legal entity", "We need invoices for Brightwater Freight Ltd as a separate entity.", "Billing", "High"),
    (13, "Owen", "Password reset email not arriving", "I asked for a password reset twice but nothing arrives.", "Accounts", "Low"),
    (9, "Daniel", "Trust account invoices", "Can we mark some invoices as paid from the trust account?", "Billing", "Normal"),
    (2, "Grace", "SSO metadata URL changed", "Our Azure tenant was migrated, the metadata URL changed. Anything to do on your side?", "Accounts", "High"),
    (8, "Priya", "Recurring invoice skipped September", "The monthly retainer invoice for one client was not generated in September.", "Technical", "High"),
    (4, "Monica", "Bulk export of customer list", "We need all customers with VAT numbers as a spreadsheet.", "Technical", "High"),
    (12, "Alicia", "Discount codes for class packs", "Can I add a discount code on invoices?", "Billing", "Low"),
    (10, "Mark", "Purchase order number on invoices", "The district needs our PO number on every invoice.", "Billing", "Normal"),
    (7, "Carl", "Help migrating from QuickBooks Desktop", "What is the best way to move our customers and open invoices from QuickBooks Desktop?", "Technical", "Normal"),
    (5, "Ana", "SCIM provisioning for new hires", "Can new hires be created automatically from Okta?", "Accounts", "High"),
    (14, "Dev", "Watermark on PDF invoices", "Is it possible to remove the 'Created with Larkspur' line on PDFs?", "Billing", "Low"),
    (15, "Hector", "Tip line on invoices", "Can customers add a tip when they pay an invoice online?", "Billing", "Low"),
    (3, "Ingrid", "Dispatch user cannot see invoices", "One of our dispatch users sees an empty invoice list.", "Accounts", "High"),
    (1, "Tomasz", "Credit note numbering", "Credit notes should use their own number range (CN-).", "Billing", "High"),
    (6, "Laura", "Two-factor authentication for all users", "Can we force 2FA for everyone in the practice?", "Accounts", "Normal"),
    (11, "Sofia", "Wrong logo on the payment page", "The online payment page still shows our old logo.", "Technical", "Normal"),
    (9, "Daniel", "Report: fees by matter", "Is there a report of fees invoiced per matter number?", "Technical", "Normal"),
    (4, "Monica", "Seat audit before renewal", "Please send us the list of active users for our renewal review.", "Accounts", "High"),
    (2, "Grace", "Duplicate patient account invoices", "Some invoices were created twice by our integration.", "Technical", "High"),
    (5, "Ana", "Data residency question", "Is our data stored in the EU?", "Accounts", "High"),
    (8, "Priya", "Xero sync stuck", "The Xero sync shows 'in progress' for 2 days.", "Technical", "High"),
    (13, "Owen", "Cancel the add-on we don't use", "Please remove the inventory add-on from our plan.", "Billing", "Low"),
    (7, "Carl", "Opening balances for customers", "How do we enter what each customer owed us before we switched to Larkspur?", "Technical", "Normal"),
    (10, "Mark", "Switch from monthly to annual billing", "Can we move to annual billing to get the discount?", "Billing", "Normal"),
    (1, "Tomasz", "User locked out after too many attempts", "One of our accountants is locked out after typing the wrong password.", "Accounts", "High"),
    (3, "Ingrid", "Webhook for credit notes", "Is there a webhook event when a credit note is issued?", "Technical", "Normal"),
    (12, "Alicia", "Change the account owner", "I'm taking over the studio, how do I become the account owner?", "Accounts", "Low"),
    (15, "Hector", "Receipt for our September payment", "Our bookkeeper needs a receipt for the September subscription.", "Billing", "Low"),
    (6, "Laura", "Statement of account per patient", "Can we send a statement listing all open invoices to one customer?", "Technical", "Normal"),
    (4, "Monica", "Contract renewal quote", "Please send the renewal quote for next year.", "Billing", "High"),
    (11, "Sofia", "Second user can't see Reports", "Our manager can sign in but the Reports menu is missing.", "Accounts", "Normal"),
    (5, "Ana", "Audit log export", "We need the audit log for the last 90 days for our SOC 2 review.", "Technical", "High"),
    (8, "Priya", "VAT reverse charge wording", "Our EU clients need 'reverse charge' on the invoice. Which setting is it?", "Technical", "Normal"),
    (14, "Dev", "Card was declined but it is valid", "Our card was declined for the subscription, the bank says it's fine.", "Billing", "Low"),
    (9, "Daniel", "Matter number column in invoice list", "Can we add a column for the matter number to the invoice list?", "Technical", "Normal"),
    (2, "Grace", "New billing admin needs access", "Please confirm how to give our new billing admin full access.", "Accounts", "High"),
    (13, "Owen", "Invoice reminders sent twice", "Some customers got the payment reminder twice yesterday.", "Technical", "Low"),
    (10, "Mark", "Tax exempt status for the district", "We are tax exempt - how do we stop tax being added to our subscription?", "Billing", "Normal"),
    (3, "Lars", "Sandbox account for API testing", "Can we get a sandbox account to test the new sync before go-live?", "Technical", "Normal"),
    (1, "Tomasz", "Remove a departed finance user", "Please confirm our former controller no longer has access.", "Accounts", "High"),
]


def payment_failed_png(path):
    def body(dr, im, top):
        dr.rectangle((0, top, im.width, top + 52), fill=(24, 50, 79))
        dr.text((24, top + 14), "Larkspur Invoicing", fill="white", font=font(20, bold=True))
        dr.text((260, top + 18), "Dashboard    Invoices    Customers    Reports    Settings", fill=(200, 215, 235), font=font(15))
        y = top + 90
        dr.rectangle((40, y, im.width - 40, y + 150), fill=(253, 236, 234), outline=(214, 69, 65), width=2)
        dr.text((64, y + 18), "Payment failed", fill=(176, 32, 28), font=font(24, bold=True))
        dr.text((64, y + 60), "We couldn't charge your Visa card ending in 4417 for your Larkspur subscription", fill=(60, 20, 20), font=font(17))
        dr.text((64, y + 84), "(card declined, code 51). Update your payment method before 29 Oct 2026 to", fill=(60, 20, 20), font=font(17))
        dr.text((64, y + 108), "avoid suspension of your account.", fill=(60, 20, 20), font=font(17))
        dr.rounded_rectangle((64, y + 170, 330, y + 210), 6, fill=(29, 95, 191))
        dr.text((84, y + 180), "Update payment method", fill="white", font=font(16, bold=True))
        dr.text((40, y + 250), "Northgate Schools  ·  Business plan  ·  Next invoice: 1 Nov 2026", fill=(110, 110, 110), font=font(15))
        dr.rectangle((40, y + 290, im.width - 40, y + 420), outline=(220, 224, 230))
        dr.text((60, y + 305), "Recent invoices", fill=(40, 40, 40), font=font(17, bold=True))
        for i, (n, c, a) in enumerate([("INV-2026-0912", "Ridgeview PTA", "$1,240.00"), ("INV-2026-0911", "Lakeside Sports Club", "$860.00"),
                                       ("INV-2026-0910", "County Music Service", "$2,150.00")]):
            dr.text((60, y + 340 + i * 24), f"{n}    {c}", fill=(60, 60, 60), font=font(15))
            dr.text((im.width - 160, y + 340 + i * 24), a, fill=(60, 60, 60), font=font(15))
    lib_w.browser_shot(path, "app.larkspur.io/dashboard", "Dashboard - Larkspur", body, (1000, 640))


def import_error_png(path):
    def body(dr, im, top):
        dr.rectangle((0, top, im.width, top + 52), fill=(24, 50, 79))
        dr.text((24, top + 14), "Larkspur Invoicing", fill="white", font=font(20, bold=True))
        y = top + 80
        dr.text((40, y), "Import invoices from CSV", fill=(30, 30, 30), font=font(24, bold=True))
        dr.text((40, y + 44), "File: invoices_sept_export.csv   (214 rows)", fill=(90, 90, 90), font=font(16))
        y += 90
        dr.rectangle((40, y, im.width - 40, y + 120), fill=(253, 236, 234), outline=(214, 69, 65), width=2)
        dr.text((62, y + 16), "Import failed - 0 of 214 rows imported", fill=(176, 32, 28), font=font(20, bold=True))
        dr.text((62, y + 56), "Error E-4031 at row 18, column \"invoice_date\".", fill=(60, 20, 20), font=font(17))
        dr.text((62, y + 82), "See the help center for the list of import error codes.", fill=(60, 20, 20), font=font(17))
        y += 150
        dr.text((40, y), "Preview (rows 16-19)", fill=(40, 40, 40), font=font(16, bold=True))
        cols = [("invoice_no", 40), ("customer", 200), ("invoice_date", 470), ("amount", 640), ("currency", 780)]
        rows = [("TMS-30415", "Harbor Point Marina", "29/09/2026", "1,284.50", "USD"),
                ("TMS-30416", "Bayside Yacht Club", "30/09/2026", "312.00", "USD"),
                ("TMS-30417", "Chesapeake Charters", "30/09/2026", "2,045.75", "USD"),
                ("TMS-30418", "Norfolk Boatworks", "01/10/2026", "97.20", "USD")]
        dr.rectangle((40, y + 28, im.width - 40, y + 56), fill=(238, 241, 245))
        for name, x in cols:
            dr.text((x + 8, y + 33), name, fill=(40, 40, 40), font=font(15, bold=True))
        for i, row in enumerate(rows):
            for (name, x), v in zip(cols, row):
                dr.text((x + 8, y + 64 + i * 28), v, fill=(50, 50, 50), font=font(15, mono=True))
    lib_w.browser_shot(path, "app.larkspur.io/import/invoices", "Import - Larkspur", body, (1000, 600))


def server_error_png(path):
    def body(dr, im, top):
        y = top + 70
        dr.text((60, y), "500", fill=(200, 205, 212), font=font(96, bold=True))
        dr.text((60, y + 130), "Internal Server Error", fill=(30, 30, 30), font=font(30, bold=True))
        dr.text((60, y + 185), "Something went wrong while generating the report \"Sales by product (monthly)\".", fill=(70, 70, 70), font=font(18))
        dr.text((60, y + 213), "Our engineers have been notified. Please try again later.", fill=(70, 70, 70), font=font(18))
        dr.text((60, y + 265), "Request ID: 7f3a-91c2-4be0    ·    2026-10-22 08:57:14 UTC", fill=(120, 120, 120), font=font(15, mono=True))
        dr.rounded_rectangle((60, y + 310, 240, y + 348), 6, outline=(29, 95, 191), width=2)
        dr.text((84, y + 319), "Back to Reports", fill=(29, 95, 191), font=font(16))
    lib_w.browser_shot(path, "app.larkspur.io/reports/sales-by-product/export?format=pdf", "Error - Larkspur", body, (1000, 560))


def policy_pdf(path):
    t = [["", "Single user", "Multiple users", "Entire organization"],
         ["Enterprise", "High", "High", "Urgent"], ["Business", "Normal", "High", "High"],
         ["Starter", "Low", "Normal", "Normal"]]
    blocks = [
        ("title", "Ticket triage & SLA policy"),
        ("p", "Larkspur Software - Customer Support. Version 3, effective 1 October 2026. Owner: Sam Rivera, Support Team Lead."),
        ("h", "1. Scope"),
        ("p", "Triage means giving every new ticket a team, a priority, an assignee decision and a status. Triage every "
              "ticket that has no team yet (the \"Untriaged\" view in the Support Desk). Do not re-triage tickets that "
              "already have a team."),
        ("h", "2. Team"),
        ("p", "Choose the team by what the customer needs:"),
        ("li", "Billing - the customer's subscription with Larkspur: our invoices and charges to the customer, refunds, "
               "payment methods and failed payments, plan changes, account suspension for non-payment."),
        ("li", "Accounts - signing in, passwords, SSO, two-factor authentication, users and seats, and security "
               "incidents (suspected account compromise, unauthorized access)."),
        ("li", "Technical - errors and bugs in the product, imports and exports, reports, the API, webhooks and "
               "integrations, performance."),
        ("h", "3. Priority"),
        ("p", "Priority follows the customer's plan (shown on the organization page) and the impact the customer "
              "reported on the ticket form:"),
        ("table", t, [250, 250, 250, 270]),
        ("p", "Exception: a suspected account compromise or unauthorized access is always Urgent and goes to the "
              "Accounts team, whatever the plan."),
        ("h", "4. Assignee"),
        ("p", "Enterprise customers have a dedicated support engineer (organization page): assign their tickets to that "
              "engineer, whatever the team. If the engineer is out of office, assign the ticket to the backup listed "
              "on the Staff page. Tickets of Business and Starter customers stay unassigned in their team's queue."),
        ("h", "5. Status"),
        ("p", "Set new tickets to Open when you triage them. If a colleague has already answered the ticket and set it "
              "to Pending, On-hold or Solved, keep that status."),
        ("h", "6. Spam and sales pitches"),
        ("p", "Set the status to Closed and add the tag \"spam\". Do not give them a team or a priority."),
        ("h", "7. Duplicates"),
        ("p", "When the same requester opened a second ticket about the same problem, close the newer ticket (status "
              "Closed, tag \"duplicate\") and add an internal note \"Duplicate of #<number of the older ticket>\". "
              "Triage the older ticket normally."),
        ("h", "8. Handover"),
        ("p", "When the queue is done, post a triage handover on the Team board of the Support Desk: how many tickets "
              "you triaged per team (Billing, Technical, Accounts), the numbers of the Urgent tickets, and how many "
              "spam and duplicate tickets you closed."),
        ("h", "9. SLA targets (first response)"),
        ("p", "Urgent 1 hour (24/7) - High 4 hours - Normal 1 business day - Low 3 business days. Business hours are "
              "Mon-Fri 08:00-18:00 in the customer's time zone."),
    ]
    lib_w.doc_pdf(path, blocks, header="Larkspur Software · Customer Support · INTERNAL",
                  footer="Support_Triage_Policy_v3 · approved 28 Sep 2026")


LOG_KB = {"Billing": ["KB-108", "KB-104", "KB-111", "KB-105"], "Technical": ["KB-122", "KB-131", "KB-130", "KB-121"],
          "Accounts": ["KB-118", "KB-150", "KB-135", "KB-115"]}


def support_log(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Log"
    ws.append(["Date", "Ticket", "Customer", "KB article", "Status", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")
    for r in rows:
        ws.append(list(r))
    for col, w in zip("ABCDEF", (12, 10, 26, 12, 11, 34)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return rows


def build(out):
    os.makedirs(out, exist_ok=True)
    rng = random.Random(20261101)
    lib_w.copy_app(HERE, out, "app_w1.py")
    con, app = lib_w.fresh_db(HERE, "app_w1.py", f"{out}/seed.db")
    for sid, user, name, team, title, ooo, backup in STAFF:
        con.execute("INSERT INTO users VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (sid, user, PASSWORD if sid == 1 else "", name, f"{user}@larkspur.io", team, title, ooo, backup, int(sid == 1)))
    for o in ORGS:
        con.execute("INSERT INTO orgs VALUES (?,?,?,?,?,?,?,?,?)", o)
    staff_by_team = {"Billing": [2, 3], "Technical": [4, 5, 6], "Accounts": [7, 8]}
    tmp = os.path.join(out, "_img")
    os.makedirs(tmp, exist_ok=True)
    imgs = {"payment_failed.png": payment_failed_png, "import_error.png": import_error_png, "server_error.png": server_error_png}
    for n, fn in imgs.items():
        fn(f"{tmp}/{n}")

    def add_ticket(tid, org, req, email, channel, impact, created, subject, body, status, priority, team, assignee,
                   tags="", updated=None, att=None):
        con.execute("INSERT INTO tickets VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (tid, subject, req, email, org, channel, impact, status, priority, team, assignee, tags, created,
                     updated or created))
        cur = con.execute("INSERT INTO comments (ticket_id,author,user_id,public,body,created) VALUES (?,?,?,?,?,?)",
                          (tid, req, None, 1, body, created))
        if att:
            con.execute("INSERT INTO attachments (ticket_id,comment_id,filename,ctype,data) VALUES (?,?,?,?,?)",
                        (tid, cur.lastrowid, att, "image/png", open(f"{tmp}/{att}", "rb").read()))

    def agent_reply(tid, uid, text, when, public=1):
        name = next(s[2] for s in STAFF if s[0] == uid)
        con.execute("INSERT INTO comments (ticket_id,author,user_id,public,body,created) VALUES (?,?,?,?,?,?)",
                    (tid, name, uid, public, text, when))

    org_by_id = {o[0]: o for o in ORGS}
    # --- history: tickets 1001..1045 except task B's (1037, 1043)
    fill = list(FILLER)
    rng.shuffle(fill)
    hist_ids = [i for i in range(1001, 1046) if i not in B_SEED]
    plan_status, log_rows = {}, []
    for k, tid in enumerate(hist_ids):
        day = 28 + k * 24 // len(hist_ids)
        created = f"2026-{9 if day <= 30 else 10:02d}-{day if day <= 30 else day - 30:02d} {8 + rng.randrange(10):02d}:{rng.randrange(60):02d}"
        if tid >= 1039:
            status = ["Open", "Open", "Pending", "On-hold", "Open", "Pending", "Open"][tid - 1039]
        elif tid >= 1030:
            status = rng.choice(["Solved", "Pending", "Solved", "Open"])
        else:
            status = rng.choice(["Closed", "Closed", "Solved"])
        status = {1031: "Pending", 1040: "Open"}.get(tid, status)
        plan_status[tid] = status
        if tid == 1012:  # an old spam ticket, closed the way the policy says
            add_ticket(tid, None, "Nina Shaw", "nina@leadsprint.biz", "Email", "Single user", created,
                       "Buy verified B2B email lists", "Hi, we sell verified lists of 50,000 finance decision makers...",
                       "Closed", None, None, None, "spam")
            continue
        org, first, subj, text, team, prio = fill[k % len(fill)]
        o = org_by_id[org]
        dom = o[3]
        assignee = o[6] if o[2] == "Enterprise" else (rng.choice(staff_by_team[team]) if status != "New" else None)
        if tid in (1031, 1040):
            assignee = 1  # Sam's own (not part of task B)
        last = {"Fernwood Bikes": "Price", "Lumen Photography": "Malhotra"}.get(o[1], "")
        req = f"{first} {last}".strip() if last else first + " " + rng.choice(["Hughes", "Lopez", "Novak", "Okoro", "Berg", "Kim"])
        add_ticket(tid, org, req, f"{first.lower()}@{dom}", rng.choice(["Email", "Web form"]),
                   rng.choice(["Single user", "Single user", "Multiple users"]), created, subj,
                   f"Hello,\n\n{text}\n\nThanks,\n{first}", status, prio, team, assignee,
                   updated=created[:11] + f"{min(int(created[11:13]) + 2, 21):02d}" + created[13:])
        if status in ("Solved", "Closed", "Pending"):
            agent_reply(tid, assignee, f"Hi {first},\n\nThanks for reaching out - " +
                        ("this is now sorted, see the steps in the linked article." if status != "Pending" else
                         "could you send us a few more details so we can look into it?") + "\n\nLarkspur Support",
                        created[:11] + f"{min(int(created[11:13]) + 1, 20):02d}" + created[13:])
        if status == "On-hold":
            agent_reply(tid, assignee, "Reproduced on our side, escalated to engineering as ENG-2287.", created, public=0)
        if status in ("Solved", "Closed", "Pending") and assignee == 1 or (tid % 4 == 1 and status in ("Solved", "Closed")):
            log_rows.append((created[:10], f"#{tid}", o[1], LOG_KB[team][tid % 4], status if status != "Closed" else "Solved", ""))
    # --- task B's four pre-triaged tickets (assigned to Sam)
    for tid, (org, req, email, created, subj, body, team, prio) in B_SEED.items():
        add_ticket(tid, org, req, email, "Email", "Single user" if tid != 1051 else "Multiple users", created, subj, body,
                   "Open", prio, team, 1)
        agent_reply(tid, 1, "Taking this one myself. -Sam", created[:14] + "59", public=0)
    # --- the untriaged queue
    for tid, (org, req, email, channel, impact, created, subj, body, att, exp) in TRIAGE.items():
        add_ticket(tid, org, req, email, channel, impact, created, subj, body, "New", None, None, None, att=att)
    for code, title, cat, body in KB:
        con.execute("INSERT INTO kb (code,title,category,body,updated) VALUES (?,?,?,?,?)",
                    (code, title, cat, body, "2026-10-" + ("18" if code == "KB-140" else "02")))
    con.execute("INSERT INTO posts (user_id,title,body,created) VALUES (?,?,?,?)",
                (4, "Release 5.2.0 is live", "Release 5.2.0 went out on 13 Oct. Known issues are tracked in KB-140 - "
                 "please link it instead of promising dates that are not in the article.\n\nMei", "2026-10-13 17:30"))
    con.execute("INSERT INTO posts (user_id,title,body,created) VALUES (?,?,?,?)",
                (1, "Triage handover - 19 Oct", "Triaged: Billing 2, Technical 5, Accounts 1\nUrgent: #1033\n"
                 "Closed: 1 spam, 0 duplicates\n\nSam", "2026-10-19 17:55"))
    con.commit()
    # ground truth / seed snapshot for the checkers
    seed = {str(r[0]): {"status": r[1], "priority": r[2], "team": r[3], "assignee_id": r[4]}
            for r in con.execute("SELECT id, status, priority, team, assignee_id FROM tickets")}
    maxes = {k: con.execute(f"SELECT COALESCE(MAX(id),0) FROM {k}").fetchone()[0] for k in ("comments", "events", "posts")}
    con.close()
    policy_pdf(f"{out}/Support_Triage_Policy_v3.pdf")
    log_rows = [r for r in log_rows if r[0] >= "2026-10-01"]
    support_log(f"{out}/support_log_2026-10.xlsx", log_rows)
    for n in imgs:
        os.remove(f"{tmp}/{n}")
    os.rmdir(tmp)
    staff_names = {s[0]: s[2] for s in STAFF}
    triage = {}
    for tid, v in TRIAGE.items():
        exp = v[-1]
        triage[str(tid)] = exp if isinstance(exp, str) else dict(exp)
    team_counts = {t: sum(1 for v in triage.values() if isinstance(v, dict) and v["team"] == t) for t in ("Billing", "Technical", "Accounts")}
    urgent = sorted(int(k) for k, v in triage.items() if isinstance(v, dict) and v["priority"] == "Urgent")
    json.dump({"port": 8101, "user": "sam.rivera", "password": PASSWORD, "sam_id": 1, "staff": staff_names,
               "triage": triage, "duplicate_of": {str(k): v for k, v in DUPLICATE_OF.items()},
               "team_counts": team_counts, "urgent": urgent, "spam_count": 2, "dup_count": 1,
               "b_tickets": {str(k): v for k, v in B_EXPECT.items()}, "seed": seed, "seed_max": maxes,
               "log_rows": len(log_rows), "log_path": "Documents/support_log_2026-10.xlsx"},
              open(f"{out}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/w1_out")
