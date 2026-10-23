"""o1 - one Calc workbook: keying the scanned September invoices into Actuals vs applying the CFO's budget transfers
and building the live variance analysis (formulas, conditional formatting, chart, notes) in the same open file."""
from __future__ import annotations

import datetime as dt
import re

from evaluation_examples.task_class import pg_olib as L
from evaluation_examples.task_class.pg_common import HOME, expected, fetch, guest, result, sh, tree, upload

XLSX = f"{HOME}/Documents/Ops_Budget_2026Q3.xlsx"
INV = f"{HOME}/Documents/invoices"
MEMO = f"{HOME}/Documents/BT-2026-07_budget_transfer.pdf"


def setup(sc):
    exp = expected("o1")
    pairs = [("o1/Ops_Budget_2026Q3.xlsx", XLSX), ("o1/invoices/AP_coding_rules.txt", f"{INV}/AP_coding_rules.txt"),
             ("o1/BT-2026-07_budget_transfer.pdf", MEMO)]
    pairs += [(f"o1/invoices/inbox/{f}", f"{INV}/inbox/{f}") for f in sorted(exp["files"])]
    upload(sc, pairs)
    sh(sc, f"mkdir -p {INV}/booked")
    sc._open_setup(XLSX)


# ---------------------------------------------------------------- helpers
def _key(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[^A-Z0-9]", "", str(v or "").upper()).lstrip("0")


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 40000 < v < 60000:
        return (L.EXCEL_EPOCH + dt.timedelta(days=float(v))).date()
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d %B %Y", "%d %b %Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _cols(bk, sheet, want: dict, default: dict):
    """Find a header row (within the first 5 rows) and map logical names -> column letters by header regex."""
    ws = bk.wf[sheet]
    for r in range(1, 6):
        heads = {c.column_letter: L.norm(c.value) for c in ws[r] if c.value is not None}
        hits = {}
        for name, pat in want.items():
            col = next((k for k, h in heads.items() if re.search(pat, h)), None)
            if col:
                hits[name] = col
        if len(hits) >= max(3, len(want) // 2):
            return r, {**default, **hits}
    return 1, default


ACT_WANT = {"date": r"date", "no": r"invoice|inv\b|number|no\.?$", "vendor": r"vendor|supplier", "cc": r"cost ?cent",
            "gl": r"\bgl\b|account", "net": r"^net", "vat": r"^vat|tax", "gross": r"gross|total",
            "charge": r"charge|budget"}
ACT_DEF = dict(zip(["date", "no", "vendor", "desc", "cc", "gl", "net", "vat", "gross", "charge"], "ABCDEFGHIJ"))


def _actual_rows(bk, sheet):
    head, cols = _cols(bk, sheet, ACT_WANT, ACT_DEF)
    ws = bk.wf[sheet]
    rows = []
    for r in range(head + 1, ws.max_row + 1):
        rec = {k: bk.safe(sheet, f"{c}{r}") for k, c in cols.items() if k != "desc"}
        if rec.get("no") in (None, "") and L.num(rec.get("net")) is None and L.num(rec.get("charge")) is None:
            continue
        rec["row"] = r
        rows.append(rec)
    return rows, cols, head


# ---------------------------------------------------------------- task A: book the September invoices
@L.safe
def evaluate_invoices(env, dest: str) -> dict:
    exp = expected("o1")
    path = fetch(env, XLSX, dest + ".xlsx")
    if not path:
        return result(0.0, {"workbook": "missing"})
    try:
        bk = L.XBook(path)
    except Exception as e:  # noqa: BLE001
        return result(0.0, {"workbook": f"unreadable: {e}"})
    sheet = bk.sheet(r"actuals?")
    if sheet is None:
        return result(0.0, {"actuals_sheet": "missing"})
    rows, cols, _ = _actual_rows(bk, sheet)
    by_key = {}
    for rec in rows:
        by_key.setdefault(_key(rec.get("no")), []).append(rec)

    per, detail, present = [], {}, 0
    for b in exp["bookings"]:
        hits = by_key.get(_key(b["no"]), [])
        if not hits:
            per.append(0.0)
            detail[b["no"]] = "missing"
            continue
        rec = hits[0]
        amounts = all(L.close(L.num(rec.get(k)) if rec.get(k) not in (None, "") or b[k] else 0.0, b[k])
                      for k in ("net", "vat", "gross"))  # an empty VAT cell counts as 0 where no VAT is charged
        parts = {"cc": L.num(rec.get("cc")) == b["cc"], "gl": L.num(rec.get("gl")) == b["gl"], "amounts": amounts,
                 "charge": L.close(L.num(rec.get("charge")), b["charge"]),
                 "date": _date(rec.get("date")) == dt.date.fromisoformat(b["date"])}
        present += L.close(L.num(rec.get("net")), b["net"])
        w = {"cc": 0.2, "gl": 0.15, "amounts": 0.25, "charge": 0.1, "date": 0.1}
        per.append(0.2 + sum(w[k] for k, ok in parts.items() if ok))
        bad = [k for k, ok in parts.items() if not ok]
        detail[b["no"]] = "ok" if not bad else "wrong " + ",".join(bad) + f" (row {rec['row']})"
    enough = present >= 6

    dup_no = next(f["no"] for f in exp["files"].values() if f["kind"] == "duplicate")
    q4_no = next(f["no"] for f in exp["files"].values() if f["kind"] == "q4")
    quote_no = next(f["no"] for f in exp["files"].values() if f["kind"] == "quote")
    no_dup = len(by_key.get(_key(dup_no), [])) <= 1
    no_q4 = not by_key.get(_key(q4_no))
    no_quote = not by_key.get(_key(quote_no))
    old_ok = all(any(L.close(L.num(r.get("net")), o["net"]) and L.num(r.get("cc")) == o["cc"]
                     for r in by_key.get(_key(o["no"]), [])) for o in exp["booked_before"])
    known = {_key(b["no"]) for b in exp["bookings"]} | {_key(o["no"]) for o in exp["booked_before"]}
    extra = sorted(str(r.get("no")) for r in rows if _key(r.get("no")) not in known)

    # files: booked PDFs moved to invoices/booked, the rest stay in the inbox
    files = tree(env, [INV.replace(HOME, "~")])
    booked = {m for p, m in files.items() if p.startswith("Documents/invoices/booked/")}
    inbox = {m for p, m in files.items() if p.startswith("Documents/invoices/inbox/")}
    fmeta = exp["files"]
    dup_md5 = {fmeta[f]["md5"] for f in fmeta if fmeta[f]["no"] == dup_no}
    moved = 0
    for b in exp["bookings"]:
        md5s = dup_md5 if b["no"] == dup_no else {fmeta[b["file"]]["md5"]}
        moved += bool(md5s & booked)
    stays = [fmeta[exp["not_booked"]["q4"]]["md5"] in inbox - booked,
             fmeta[exp["not_booked"]["quote"]]["md5"] in inbox - booked,
             len(dup_md5 & booked) == 1]
    files_score = 0.7 * moved / len(exp["bookings"]) + 0.3 * sum(stays) / 3

    note = guest(env, f"cat {INV}/inbox/NOT_BOOKED.txt 2>/dev/null || cat {INV}/NOT_BOOKED.txt 2>/dev/null")
    low = note.lower()
    items = {
        "duplicate": (any(s in low for s in ("scan_20260911_0007", "scan_20260909_0003", dup_no.lower())),
                      bool(re.search(r"duplicat|twice|copy|already|same invoice", low))),
        "q4": ("scan_20261002_0009" in low or q4_no.lower() in low,
               bool(re.search(r"october|\boct\b|q4|later|next quarter|outside|after 30", low))),
        "quote": ("scan_20260929_0010" in low or quote_no.lower() in low,
                  bool(re.search(r"quot|not an invoice|not a vat invoice", low)))}
    note_score = sum(0.5 * a + 0.5 * b for a, b in items.values()) / 3

    score = 0.6 * sum(per) / len(per)
    if enough:
        score += 0.1 * (no_dup + no_q4 + no_quote) / 3 + 0.05 * old_ok
    score += 0.12 * files_score + 0.13 * note_score
    if not enough:
        score = min(score, 0.4)
    checks = {"bookings": detail, "bookings_with_correct_net": f"{present}/{len(exp['bookings'])}",
              "duplicate_booked_once": no_dup, "october_invoice_not_booked": no_q4, "quotation_not_booked": no_quote,
              "jul_aug_rows_intact": old_ok, "unexpected_rows": extra[:8],
              "pdfs_moved_to_booked": f"{moved}/{len(exp['bookings'])}", "unbooked_left_in_inbox": stays,
              "not_booked_note": {k: list(v) for k, v in items.items()}}
    return result(score, checks)


# ---------------------------------------------------------------- task B: transfers + variance analysis
def _mentions(bk, sheet, ref, target, depth=3, seen=None) -> bool:
    """Does the formula in sheet!ref (or a formula it references, up to `depth`) refer to sheet `target`?"""
    seen = seen or set()
    if depth < 0 or (sheet, ref) in seen:
        return False
    seen.add((sheet, ref))
    f = bk.formula(sheet, ref)
    if not f:
        return False
    if re.search(rf"(^|[^A-Za-z]){target}'?!", f, re.I):
        return True
    for tok in re.findall(r"(?<![A-Za-z!$'])\$?([A-Z]{1,3})\$?(\d+)(?![\d(:])", f.replace("'", "")):
        if _mentions(bk, sheet, f"{tok[0]}{tok[1]}", target, depth - 1, seen):
            return True
    return False


def _shift(formula: str, dr: int, dc: int) -> str:
    """Shift relative references of a conditional-format expression from the range's top-left to another cell."""
    out = []
    for kind, tok in L._tokens(formula):
        if kind == "ref" and "!" not in tok and ":" not in tok:
            m = re.fullmatch(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)", tok)
            if m:
                c = m.group(2) if m.group(1) else L.n2col(L.col2n(m.group(2)) + dc)
                r = m.group(4) if m.group(3) else str(int(m.group(4)) + dr)
                tok = f"{m.group(1)}{c}{m.group(3)}{r}"
        out.append(tok)
    return "".join(out)


def _rule_fires(bk, sheet, sqref, rule, ref):
    from openpyxl.utils.cell import coordinate_from_string, range_boundaries
    v = L.num(bk.safe(sheet, ref))
    if rule.type == "cellIs":
        try:
            vals = [L.num(L._scalar(bk.eval(f, sheet))) for f in (rule.formula or [])]
        except Exception:  # noqa: BLE001
            return None
        if v is None or not vals or vals[0] is None:
            return False
        a = vals[0]
        b = vals[1] if len(vals) > 1 else None
        op = rule.operator or ""
        return {"greaterThan": v > a, "greaterThanOrEqual": v >= a, "lessThan": v < a, "lessThanOrEqual": v <= a,
                "equal": v == a, "notEqual": v != a,
                "between": b is not None and min(a, b) <= v <= max(a, b),
                "notBetween": b is not None and not (min(a, b) <= v <= max(a, b))}.get(op)
    if rule.type == "expression" and rule.formula:
        c0, r0, _, _ = range_boundaries(sqref.split()[0])
        col, row = coordinate_from_string(ref)
        try:
            return L._truth(bk.eval(_shift(rule.formula[0], row - r0, L.col2n(col) - c0), sheet))
        except Exception:  # noqa: BLE001
            return None
    return None


@L.safe
def evaluate_variance(env, dest: str) -> dict:
    exp = expected("o1")
    path = fetch(env, XLSX, dest + ".xlsx")
    if not path:
        return result(0.0, {"workbook": "missing"})
    try:
        bk = L.XBook(path)
    except Exception as e:  # noqa: BLE001
        return result(0.0, {"workbook": f"unreadable: {e}"})
    ccs = [int(c) for c in exp["ccs"]]
    checks = {}

    # 1. budget transfers (Budget sheet: original unchanged, revised = original + approved transfers)
    bsh = bk.sheet(r"budget")
    budget_ok = {}
    if bsh:
        head, cols = _cols(bk, bsh, {"cc": r"cost ?cent", "orig": r"original", "transfer": r"transfer|viremen",
                                     "revised": r"revised"}, {"cc": "A", "orig": "D", "transfer": "E", "revised": "F"})
        for r in range(head + 1, bk.wf[bsh].max_row + 1):
            code = L.num(bk.safe(bsh, f"{cols['cc']}{r}"))
            if code in ccs and int(code) not in budget_ok:
                c = str(int(code))
                budget_ok[int(code)] = (L.close(L.num(bk.safe(bsh, f"{cols['revised']}{r}")), exp["revised"][c], 0.5)
                                        and L.close(L.num(bk.safe(bsh, f"{cols['orig']}{r}")), exp["original"][c], 0.5))
    checks["revised_budget_ok"] = {str(c): budget_ok.get(c, False) for c in ccs}
    s_budget = sum(budget_ok.get(c, False) for c in ccs) / len(ccs)

    # 2. truth for actuals = Budget charge per cost centre in the FINAL Actuals sheet (whatever has been keyed in)
    ash = bk.sheet(r"actuals?")
    actual = {c: 0.0 for c in ccs}
    if ash:
        for rec in _actual_rows(bk, ash)[0]:
            cc, ch = L.num(rec.get("cc")), L.num(rec.get("charge"))
            if cc in actual and ch is not None:
                actual[int(cc)] += ch
    revised = {c: exp["revised"][str(c)] for c in ccs}
    want_flag = sorted(c for c in ccs if actual[c] > 1.1 * revised[c] + 1e-9)
    checks["actual_from_final_actuals"] = {str(c): round(v, 2) for c, v in actual.items()}

    vsh = bk.sheet(r"variance.*")
    rows_ok, pct_cells, s_var = {}, {}, 0.0
    if vsh:
        ws = bk.wf[vsh]
        for c in ccs:
            row = None
            for r in range(1, ws.max_row + 1):
                for col in range(1, min(4, ws.max_column) + 1):
                    v = bk.safe(vsh, f"{L.n2col(col)}{r}")
                    if L.num(v) == c or (isinstance(v, str) and re.match(rf"\s*{c}\b", v)):
                        row = r
                        break
                if row:
                    break
            if not row:
                rows_ok[str(c)] = "row missing"
                continue
            cells = {}
            for col in range(1, ws.max_column + 1):
                ref = f"{L.n2col(col)}{row}"
                x = L.num(bk.safe(vsh, ref))
                if x is not None and not (col <= 3 and x == c):
                    cells[ref] = x
            a, b = actual[c], revised[c]
            act_ref = next((r for r, x in cells.items() if L.close(x, a, 0.02)), None)
            live = bool(act_ref) and _mentions(bk, vsh, act_ref, "actuals")
            var_ok = any(L.close(abs(x), abs(a - b), 0.02) for x in cells.values())
            pct = (a - b) / b
            pct_ref = next((r for r, x in cells.items() if r != act_ref and
                            (L.close(abs(x), abs(pct), 0.0006) or L.close(abs(x), abs(pct * 100), 0.06))), None)
            if pct_ref:
                pct_cells[c] = pct_ref
            part = 0.5 * (bool(act_ref) and live) + 0.25 * var_ok + 0.25 * bool(pct_ref)
            s_var += part / len(ccs)
            rows_ok[str(c)] = {"row": row, "actual": act_ref, "actual_is_live_formula": live, "variance": var_ok,
                               "pct": pct_ref}
        # rows whose % could not be matched by value (e.g. stale numbers): take the cell under a '%' header
        head_pct = None
        for r in range(1, 6):
            for col in range(1, ws.max_column + 1):
                h = L.norm(ws.cell(row=r, column=col).value)
                if re.search(r"%|percent|pct", h):
                    head_pct = col
            if head_pct:
                break
        for c in ccs:
            ro = rows_ok.get(str(c))
            if head_pct and isinstance(ro, dict) and c not in pct_cells:
                pct_cells[c] = f"{L.n2col(head_pct)}{ro['row']}"
    checks["variance_sheet"] = vsh
    checks["variance_rows"] = rows_ok

    # 3. conditional formatting: a rule over the variance % cells that flags exactly the >10% overruns
    s_cf, cf_info = 0.0, []
    if vsh and len(pct_cells) == len(ccs):
        for sqref, rule in L.cf_rules(bk.wf[vsh]):
            if not all(L.in_ranges(sqref, ref) for ref in pct_cells.values()):
                continue
            fired = {c: _rule_fires(bk, vsh, sqref, rule, ref) for c, ref in pct_cells.items()}
            flagged = sorted(c for c, f in fired.items() if f)
            styled = rule.dxf is not None and (rule.dxf.fill is not None or rule.dxf.font is not None)
            cf_info.append({"range": sqref, "type": rule.type, "operator": rule.operator, "formula": rule.formula,
                            "flags": flagged, "styled": styled})
            s_cf = max(s_cf, 0.33 + 0.67 * (flagged == want_flag and styled))
    checks["conditional_formats"] = cf_info
    checks["should_flag"] = want_flag

    # 4. chart on the Variance sheet
    charts = L.xlsx_charts(path)
    n_chart = len(charts.get(vsh, [])) if vsh else 0
    s_chart = 1.0 if n_chart else (0.5 if any(charts.values()) else 0.0)
    checks["charts"] = {k: len(v) for k, v in charts.items() if v}

    # 5. notes for the meeting
    nsh = bk.sheet(r"notes?")
    text = ""
    if nsh:
        text = " ".join(str(c.value) for row in bk.wf[nsh].iter_rows() for c in row if isinstance(c.value, str))
    words = len(text.split())
    names = exp["ccs"]
    over = all(re.search(rf"\b{c}\b", text) or names[str(c)].lower() in text.lower() for c in want_flag)
    ref_ok = bool(re.search(r"BT\W*2026\W*0?7", text, re.I))
    s_notes = (words >= 25) / 3 + (bool(want_flag) and over and words >= 10) / 3 + ref_ok / 3
    checks["notes"] = {"sheet": nsh, "words": words, "names_overrun": over, "memo_ref": ref_ok}

    score = 0.2 * s_budget + 0.4 * s_var + 0.15 * s_cf + 0.1 * s_chart + 0.15 * s_notes
    if not vsh:
        score = min(score, 0.3)
    return result(score, checks)
