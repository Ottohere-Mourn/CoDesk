"""F2 assets - a bakery's scanned paperwork: filing the scans vs the accounts-payable three-way match.

Writes into out_dir:
  scans/scan_00NN.pdf        18 image-only PDFs (sheet-feeder scans: slight skew, noise, JPEG) - no text layer
  bookkeeping_2026-10.xlsx   shared workbook: "Document register" (task A), "AP check" (task B), "PO list" (given)
  f2.tar.gz                  scans + workbook, extracted once by setup
  expected.json              ground truth - host only
"""
from __future__ import annotations

import json
import os
import random
import tarfile

from PIL import Image, ImageDraw

import lib_f
from lib_f import font, hand_text, md5

PW, PH = 1240, 1754  # A4 at 150 dpi
BAKERY = ["Harbour Street Bakery", "12 Harbour Street", "Falmouth TR11 2AB"]

# vendor: (full name, address lines, accent colour, date format)
VENDORS = {
    "Millbrook": ("Millbrook Flour Co.", ["Unit 4, Mill Lane", "Stroud GL5 3AB", "VAT GB 214 7781 02"], (124, 84, 40), "%d/%m/%Y"),
    "Hollow": ("Hollow Farm Dairy", ["Hollow Farm, Lanner", "Redruth TR16 6HX", "VAT GB 302 1190 45"], (40, 110, 70), "%-d %B %Y"),
    "PackRight": ("PackRight Packaging Ltd", ["14 Trade Park Way", "Plymouth PL6 7RT", "VAT GB 188 4420 17"], (200, 90, 20), "%Y-%m-%d"),
    "Cedar": ("Cedar & Sons Produce", ["Market Yard 3", "Truro TR1 2QQ", "VAT GB 276 5512 88"], (70, 100, 40), "%d.%m.%y"),
    "BrightClean": ("BrightClean Supplies", ["Kernick Ind. Est. 22", "Penryn TR10 9EP", "VAT GB 355 0021 63"], (20, 120, 180), "%d-%b-%Y"),
    "Harbour": ("Harbour Print Studio", ["5 Arwenack Street", "Falmouth TR11 3JA", "VAT GB 401 2287 10"], (160, 40, 120), "%d/%m/%Y"),
}
POS = {  # PO no: (vendor, date, [(item, qty, unit price)])
    "PO-1041": ("Millbrook", "2026-09-28", [("Strong white flour 16kg", 20, 13.80), ("Wholemeal flour 16kg", 8, 12.60),
                                            ("Dark rye flour 10kg", 4, 11.40)]),
    "PO-1042": ("Hollow", "2026-09-30", [("Unsalted butter 2.5kg block", 30, 9.85), ("Whole milk 20L", 12, 14.40),
                                         ("Double cream 5L", 6, 17.90)]),
    "PO-1043": ("PackRight", "2026-10-01", [("Kraft bread bags (1000)", 3, 38.50), ("Cake boxes 10in (100)", 4, 21.75),
                                            ("Greaseproof sheets (500)", 2, 16.20)]),
    "PO-1044": ("Cedar", "2026-10-02", [("Bramley apples (crate)", 12, 18.00), ("Strawberries (flat)", 10, 13.90),
                                        ("Free-range eggs (tray of 30)", 15, 6.75)]),
    "PO-1045": ("BrightClean", "2026-10-05", [("Food-safe sanitiser 5L", 4, 12.95), ("Blue roll (6 pack)", 5, 9.40)]),
}
# invoices in scope for the AP check: (vendor, inv no, date, PO, vat rate, [(item, qty, price)])
INVOICES = [
    ("Millbrook", "INV-20417", "2026-10-03", "PO-1041", 0.0,
     [("Strong white flour 16kg", 20, 14.20), ("Wholemeal flour 16kg", 8, 12.60), ("Dark rye flour 10kg", 4, 11.40)]),
    ("Hollow", "HF-5532", "2026-10-06", "PO-1042", 0.0,
     [("Unsalted butter 2.5kg block", 30, 9.85), ("Whole milk 20L", 12, 14.40), ("Double cream 5L", 6, 17.90)]),
    ("PackRight", "PR-77310", "2026-10-07", "PO-1043", 0.2,
     [("Kraft bread bags (1000)", 3, 38.50), ("Cake boxes 10in (100)", 4, 21.75), ("Greaseproof sheets (500)", 2, 16.20)]),
    ("Cedar", "2026-1188", "2026-10-08", "PO-1044", 0.0,
     [("Bramley apples (crate)", 12, 18.00), ("Strawberries (flat)", 10, 14.50), ("Free-range eggs (tray of 30)", 15, 6.75)]),
    ("BrightClean", "BC-3092", "2026-10-12", "PO-1045", 0.2,
     [("Food-safe sanitiser 5L", 4, 12.95), ("Blue roll (6 pack)", 5, 9.40), ("Heavy-duty degreaser 5L", 2, 15.80)]),
]
# delivery notes: (vendor, DN no, date, PO, [(item, printed qty, received qty or None)], handwritten note)
DNOTES = [
    ("Millbrook", "DN-88213", "2026-10-02", "PO-1041",
     [("Strong white flour 16kg", 20, None), ("Wholemeal flour 16kg", 8, None), ("Dark rye flour 10kg", 4, None)],
     "All rec'd OK - Kim"),
    ("Hollow", "HFD-0917", "2026-10-05", "PO-1042",
     [("Unsalted butter 2.5kg block", 30, 26), ("Whole milk 20L", 12, None), ("Double cream 5L", 6, None)],
     "Butter: only 26 rec'd, 4 short! - KM"),
    ("Cedar", "CS-4471", "2026-10-08", "PO-1044",
     [("Bramley apples (crate)", 12, None), ("Strawberries (flat)", 10, None), ("Free-range eggs (tray of 30)", 15, 14)],
     "1 tray eggs smashed - sent back with driver. 14 kept. KM"),
]
KEYWORDS = {"Strong white flour 16kg": "strong", "Wholemeal flour 16kg": "wholemeal", "Dark rye flour 10kg": "rye",
            "Unsalted butter 2.5kg block": "butter", "Whole milk 20L": "milk", "Double cream 5L": "cream",
            "Kraft bread bags (1000)": "bag", "Cake boxes 10in (100)": "box", "Greaseproof sheets (500)": "greaseproof",
            "Bramley apples (crate)": "apple", "Strawberries (flat)": "strawberr", "Free-range eggs (tray of 30)": "egg",
            "Food-safe sanitiser 5L": "sanitis", "Blue roll (6 pack)": "roll", "Heavy-duty degreaser 5L": "degreas"}


def fmt_date(iso: str, f: str) -> str:
    import datetime as dt
    d = dt.date.fromisoformat(iso)
    return d.strftime(f.replace("%-d", str(d.day)))


def money(v: float) -> str:
    return f"£{v:,.2f}"


# ---------------------------------------------------------------- page renderers (white page, black/colour ink)
def letterhead(d: ImageDraw.ImageDraw, vendor: str, title: str):
    name, addr, col, _ = VENDORS[vendor]
    d.rectangle((0, 0, PW, 26), fill=col)
    d.ellipse((90, 70, 190, 170), fill=col)
    d.text((140, 120), name[0], fill="white", font=font("serif-bold", 64), anchor="mm")
    d.text((215, 78), name, fill=col, font=font("lato-black", 46))
    for i, ln in enumerate(addr):
        d.text((217, 136 + i * 28), ln, fill=(70, 70, 70), font=font("lato", 22))
    d.text((PW - 90, 86), title, fill=(30, 30, 30), font=font("lato-black", 54), anchor="ra")


def bill_to(d, y=300):
    d.text((90, y), "Bill to / Deliver to", fill=(90, 90, 90), font=font("lato-bold", 22))
    for i, ln in enumerate(BAKERY):
        d.text((90, y + 32 + i * 30), ln, fill=(20, 20, 20), font=font("lato", 26))


def meta(d, rows, x=720, y=300):
    for i, (k, v) in enumerate(rows):
        d.text((x, y + i * 36), k, fill=(90, 90, 90), font=font("lato-bold", 24))
        d.text((x + 200, y + i * 36), v, fill=(15, 15, 15), font=font("lato", 26))


def items_table(d, y, head, rows, widths):
    x0 = 90
    xs = [x0]
    for w in widths:
        xs.append(xs[-1] + w)
    d.rectangle((x0, y, xs[-1], y + 46), fill=(232, 232, 232))
    for j, h in enumerate(head):
        d.text((xs[j] + 12 if j == 0 else xs[j + 1] - 12, y + 10), h, fill=(20, 20, 20), font=font("lato-bold", 24),
               anchor="la" if j == 0 else "ra")
    y += 46
    for r in rows:
        for j, v in enumerate(r):
            d.text((xs[j] + 12 if j == 0 else xs[j + 1] - 12, y + 12), str(v), fill=(15, 15, 15), font=font("lato", 25),
                   anchor="la" if j == 0 else "ra")
        y += 50
        d.line((x0, y, xs[-1], y), fill=(200, 200, 200), width=2)
    d.rectangle((x0, y - 50 * len(rows) - 46, xs[-1], y), outline=(120, 120, 120), width=2)
    return y


def invoice_page(vendor, no, date, po, vat, lines, rng, title="INVOICE", note=None):
    im = Image.new("RGB", (PW, PH), "white")
    d = ImageDraw.Draw(im)
    letterhead(d, vendor, title)
    bill_to(d)
    m = [("Invoice no." if title == "INVOICE" else "Credit note", no), ("Date", fmt_date(date, VENDORS[vendor][3])),
         ("Account", "HSB-0" + str(sum(map(ord, vendor)) % 900 + 100))]
    if po:
        m.append(("Your PO", po))
    meta(d, m)
    rows = [(it, q, f"{p:.2f}", f"{q * p:,.2f}") for it, q, p in lines]
    y = items_table(d, 520, ["Description", "Qty", "Unit price £", "Amount £"], rows, [560, 120, 190, 190])
    net = round(sum(q * p for _, q, p in lines), 2)
    v = round(net * vat, 2)
    tot = round(net + v, 2)
    for i, (k, val) in enumerate([("Subtotal", net), (f"VAT {int(vat * 100)}%", v), ("Total due" if title == "INVOICE" else "Total credit", tot)]):
        f = font("lato-black" if i == 2 else "lato", 28)
        d.text((880, y + 30 + i * 44), k, fill=(20, 20, 20), font=f, anchor="ra")
        d.text((1150, y + 30 + i * 44), money(val), fill=(20, 20, 20), font=f, anchor="ra")
    d.text((90, PH - 230), "Payment terms: 30 days net. Please quote the invoice number with your payment.",
           fill=(70, 70, 70), font=font("lato", 22))
    d.text((90, PH - 196), f"Bank: Lloyds  Sort code 30-94-{rng.randint(10, 99)}  Account {rng.randint(10000000, 99999999)}",
           fill=(70, 70, 70), font=font("lato", 22))
    if note:
        note(im, d, y)
    return im, tot


def dn_page(vendor, no, date, po, lines, handnote, rng):
    im = Image.new("RGB", (PW, PH), "white")
    d = ImageDraw.Draw(im)
    letterhead(d, vendor, "DELIVERY NOTE")
    bill_to(d)
    meta(d, [("Delivery note", no), ("Date", fmt_date(date, VENDORS[vendor][3])), ("Your PO", po), ("Driver", "R. Penrose")])
    rows = [(it, q, "") for it, q, _ in lines]
    y0 = 520
    y = items_table(d, y0, ["Description", "Qty", "Checked"], rows, [640, 160, 260])
    for i, (it, q, rec) in enumerate(lines):
        ry = y0 + 46 + i * 50
        if rec is None:
            hand_text(im, (1010, ry + 2), "✓" if False else "ok", 34, rng, fill=(20, 30, 140))
        else:
            d.line((850, ry + 26, 920, ry + 20), fill=(20, 30, 140), width=4)
            hand_text(im, (940, ry - 2), str(rec), 40, rng, fill=(20, 30, 140), bold=True)
    hand_text(im, (110, y + 60), handnote, 40, rng, fill=(20, 30, 140))
    d.text((90, PH - 300), "Received in good condition (unless noted):", fill=(60, 60, 60), font=font("lato", 24))
    d.line((90, PH - 200, 600, PH - 200), fill=(0, 0, 0), width=2)
    hand_text(im, (120, PH - 262), "K. Morgan", 46, rng, fill=(20, 30, 140))
    return im


def receipt_page(shop, lines, no, date_s, total, rng, extra=()):
    w, h = 560, 1250
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    f = font("mono", 22)
    fb = font("mono-bold", 26)
    y = 40
    d.text((w // 2, y), shop[0], fill=(10, 10, 10), font=fb, anchor="ma")
    y += 40
    for ln in shop[1:]:
        d.text((w // 2, y), ln, fill=(30, 30, 30), font=f, anchor="ma")
        y += 30
    y += 20
    d.text((30, y), f"{date_s}   {no}", fill=(20, 20, 20), font=f)
    y += 50
    for it, amt in lines:
        d.text((30, y), it, fill=(20, 20, 20), font=f)
        d.text((w - 30, y), f"{amt:.2f}", fill=(20, 20, 20), font=f, anchor="ra")
        y += 32
    y += 16
    d.line((30, y, w - 30, y), fill=(40, 40, 40), width=2)
    y += 14
    d.text((30, y), "TOTAL", fill=(0, 0, 0), font=fb)
    d.text((w - 30, y), f"£{total:.2f}", fill=(0, 0, 0), font=fb, anchor="ra")
    y += 50
    for ln in extra:
        d.text((30, y), ln, fill=(40, 40, 40), font=f)
        y += 30
    d.text((w // 2, h - 80), "THANK YOU - PLEASE KEEP", fill=(30, 30, 30), font=f, anchor="ma")
    # thermal paper: slightly grey, faded
    return im.point(lambda v: 60 + v * 0.75 if v < 200 else v)


def utility_page(company, col, title, rows, amount_due, rng):
    im = Image.new("RGB", (PW, PH), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, PW, 180), fill=col)
    d.text((90, 50), company, fill="white", font=font("lato-black", 60))
    d.text((PW - 90, 70), title, fill="white", font=font("lato-bold", 36), anchor="ra")
    bill_to(d, 240)
    meta(d, rows, 700, 240)
    d.rectangle((90, 560, PW - 90, 700), outline=col, width=4)
    d.text((120, 590), "Amount due", fill=(30, 30, 30), font=font("lato-bold", 34))
    d.text((PW - 120, 590), money(amount_due), fill=col, font=font("lato-black", 64), anchor="ra")
    y = 760
    for i in range(6):
        d.rectangle((120 + i * 160, 1100 - (i * 37 % 5 + 2) * 50, 220 + i * 160, 1100), fill=col)
    d.text((90, 1140), "Your usage over the last six months", fill=(80, 80, 80), font=font("lato", 24))
    d.text((90, 1250), "Pay by Direct Debit, online or by phone. Quote your account number.", fill=(80, 80, 80), font=font("lato", 22))
    del y
    return im


def letter_page(rng):
    im = Image.new("RGB", (PW, PH), "white")
    d = ImageDraw.Draw(im)
    d.text((90, 90), "HARBOUR BANK", fill=(0, 60, 110), font=font("serif-bold", 56))
    d.text((92, 160), "Business Banking, PO Box 410, Exeter EX1 9QQ", fill=(80, 80, 80), font=font("lato", 22))
    for i, ln in enumerate(["Ms P. Nair", "Harbour Street Bakery", "12 Harbour Street", "Falmouth TR11 2AB"]):
        d.text((90, 280 + i * 32), ln, fill=(20, 20, 20), font=font("lato", 26))
    d.text((PW - 90, 280), "2 October 2026", fill=(20, 20, 20), font=font("lato", 26), anchor="ra")
    d.text((PW - 90, 316), "Our ref: HB/4471", fill=(20, 20, 20), font=font("lato", 26), anchor="ra")
    d.text((90, 480), "Changes to your business account tariff", fill=(0, 0, 0), font=font("lato-bold", 32))
    body = ["Dear Ms Nair,", "", "From 1 December 2026 the monthly fee for your Business Current Account",
            "will change from £6.50 to £8.00. Cash deposits over £1,000 per month will be",
            "charged at 0.5%. You do not need to do anything. If you would like to discuss",
            "a different tariff, please call your relationship manager on 0345 600 2100.", "",
            "Yours sincerely,", "", "", "Daniel Ashworth", "Relationship Manager"]
    for i, ln in enumerate(body):
        d.text((90, 560 + i * 38), ln, fill=(25, 25, 25), font=font("serif", 26))
    hand_text(im, (90, 560 + 8 * 38), "D. Ashworth", 52, rng, fill=(10, 10, 60))
    return im


def page_jpeg(im, rng, receipt=False):
    if receipt:  # a narrow receipt fed through the scanner: grey margins around it
        canvas = Image.new("RGB", (im.size[0] + 80, im.size[1] + 80), (205, 205, 205))
        canvas.paste(im, (40, 40))
        im = canvas
    return lib_f.scan(im, rng, quality=62, sigma=3.5)


# ---------------------------------------------------------------- build
def build(d: str):
    rng = random.Random(20261015)
    os.makedirs(f"{d}/scans", exist_ok=True)
    docs = []  # each: dict(kind, type, vendor_key, docno, date, total, jpeg, page_size)

    def add(kind, typ, vkey, docno, date, total, jpeg, size=(595.28, 841.89), **kw):
        docs.append(dict(kind=kind, type=typ, vendor=vkey, docno=docno, date=date, total=total, jpeg=jpeg, size=size, **kw))

    for vendor, no, date, po, vat, lines in INVOICES:
        im, tot = invoice_page(vendor, no, date, po, vat, lines, rng)
        add("doc", "Invoices", vendor, no, date, tot, page_jpeg(im, rng))
        if vendor == "Hollow":  # scanned twice by mistake (different skew and noise: same paper, different bytes)
            add("dup", "Invoices", vendor, no, date, tot, page_jpeg(im, rng), dup_of=no)

    def paid(im, dd, y):
        hand_text(im, (700, y + 200), "PAID 30/09 by card - PN", 44, rng, fill=(170, 20, 20), bold=True)
    im, tot = invoice_page("Harbour", "HPS-0412", "2026-09-29", None, 0.2,
                           [("A1 menu boards, matt laminate", 3, 24.00), ("Loyalty cards (500)", 1, 38.00)], rng, note=paid)
    add("doc", "Invoices", "Harbour", "HPS-0412", "2026-09-29", tot, page_jpeg(im, rng))
    im, tot = invoice_page("Millbrook", "CN-1180", "2026-10-09", None, 0.0,
                           [("Damaged flour sacks (INV-20211)", 2, 13.80)], rng, title="CREDIT NOTE")
    add("doc", "Credit notes", "Millbrook", "CN-1180", "2026-10-09", tot, page_jpeg(im, rng))
    for vendor, no, date, po, lines, note in DNOTES:
        add("doc", "Delivery notes", vendor, no, date, None, page_jpeg(dn_page(vendor, no, date, po, lines, note, rng), rng))
    r = receipt_page(["COSTWAY CASH & CARRY", "Kernick Road, Penryn", "VAT 199 2231 76"],
                     [("CASTER SUGAR 25KG", 21.49), ("VANILLA EXTRACT 500ML", 14.99), ("BAKING PARCHMENT X2", 9.89)],
                     "RCPT 004512", "04/10/2026 08:12", 46.37, rng, ["CARD  ****4410", "VAT INCL  £1.98"])
    add("doc", "Receipts", "Costway", "004512", "2026-10-04", 46.37, page_jpeg(r, rng, True), size=(198.0, 425.0))
    r = receipt_page(["Shell Harbour Road", "Falmouth TR11 4HX"], [("PUMP 3 DIESEL 36.42L", 58.20)],
                     "TRN 7781-22", "10/10/26 17:45", 58.20, rng, ["VAT 20%  £9.70", "FLEET CARD  ****0093"])
    add("doc", "Receipts", "Shell", "7781-22", "2026-10-10", 58.20, page_jpeg(r, rng, True), size=(198.0, 425.0))
    r = receipt_page(["OFFICE HUB", "Church St, Falmouth"], [("TILL ROLLS 57MM X10", 7.99), ("A4 LABELS 100", 4.49)],
                     "TILL 3 RCPT 1129", "13 OCT 2026 11:02", 12.48, rng, ["CASH  20.00", "CHANGE  7.52"])
    add("doc", "Receipts", "Office", "1129", "2026-10-13", 12.48, page_jpeg(r, rng, True), size=(198.0, 425.0))
    im = utility_page("Northgrid Energy", (0, 120, 90), "Business electricity bill",
                      [("Bill number", "NG-993201"), ("Bill date", "01/10/2026"), ("Account", "8812 3301"),
                       ("Period", "01/09 - 30/09/2026")], 412.66, rng)
    add("doc", "Utilities", "Northgrid", "NG-993201", "2026-10-01", 412.66, page_jpeg(im, rng))
    im = utility_page("AquaWest Water", (20, 90, 170), "Water & sewerage bill",
                      [("Bill no.", "AW-55021"), ("Date of bill", "30 Sep 2026"), ("Customer", "40 1182 7"),
                       ("Period", "Apr - Sep 2026")], 96.10, rng)
    add("doc", "Utilities", "AquaWest", "AW-55021", "2026-09-30", 96.10, page_jpeg(im, rng))
    add("doc", "Letters", "Harbour", "HB/4471", "2026-10-02", None, page_jpeg(letter_page(rng), rng))
    add("blank", None, None, None, None, None, page_jpeg(Image.new("RGB", (PW, PH), "white"), rng))
    # scan order: shuffled, the duplicate comes after its original
    order = docs[:]
    rng.shuffle(order)
    dup = next(x for x in order if x["kind"] == "dup")
    orig = next(x for x in order if x["kind"] == "doc" and x["docno"] == dup["dup_of"])
    order.remove(dup)
    order.insert(min(len(order), order.index(orig) + 3), dup)
    for i, x in enumerate(order, 1):
        x["scan"] = f"scan_{i:04d}.pdf"
        lib_f.image_pdf([x.pop("jpeg")], f"{d}/scans/{x['scan']}", x["size"])
        x["md5"] = md5(f"{d}/scans/{x['scan']}")
        if x["kind"] == "doc":
            safe = x["docno"].replace("/", "-")
            x["new_name"] = f"{x['date']}_{x['vendor']}_{safe}.pdf"
    # ---- AP ground truth
    dn_rec = {dn[3]: {it: (rec if rec is not None else q) for it, q, rec in dn[4]} for dn in DNOTES}
    ap_lines, memo = [], {}
    for vendor, no, date, po, vat, lines in INVOICES:
        po_items = {it: (q, p) for it, q, p in POS[po][2]}
        over = 0.0
        for it, q, p in lines:
            issues = []
            if it not in po_items:
                issues.append("not on PO")
                ob = q * p
            else:
                pq, pp = po_items[it]
                rec = dn_rec.get(po, {}).get(it, pq)
                if abs(p - pp) > 1e-9:
                    issues.append("price")
                if q != rec:
                    issues.append("qty")
                ob = max(0.0, q * p - rec * pp)
            over += ob
            ap_lines.append({"invoice": no, "po": po, "item": it, "kw": KEYWORDS[it], "qty": q, "price": p,
                             "total": round(q * p, 2), "issues": issues, "overbilled": round(ob, 2)})
        memo[no] = round(over, 2)
    # ---- workbook
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    hdr = PatternFill("solid", fgColor="E3ECD9")
    ws = wb.active
    ws.title = "Document register"
    ws.append(["File name", "Type", "Date", "Vendor", "Doc no.", "Total (£)"])
    ws2 = wb.create_sheet("AP check")
    ws2.append(["Invoice no.", "PO no.", "Item", "Qty", "Unit price (£)", "Line total (£)", "Issue"])
    ws3 = wb.create_sheet("PO list")
    ws3.append(["PO no.", "Supplier", "PO date", "Item", "Qty ordered", "Agreed unit price (£)"])
    for po, (vendor, date, items) in POS.items():
        for it, q, p in items:
            ws3.append([po, VENDORS[vendor][0], date, it, q, p])
    for w, widths in ((ws, (34, 16, 12, 24, 14, 12)), (ws2, (14, 10, 32, 8, 14, 14, 16)), (ws3, (10, 24, 12, 32, 12, 20))):
        for c, wd in zip("ABCDEFG", widths):
            w.column_dimensions[c].width = wd
        for cell in w[1]:
            cell.font = Font(bold=True)
            cell.fill = hdr
        w.freeze_panes = "A2"
    wb.save(f"{d}/bookkeeping_2026-10.xlsx")
    with tarfile.open(f"{d}/f2.tar.gz", "w:gz") as tf:
        tf.add(f"{d}/scans", arcname="Documents/Scans/2026-10")
        tf.add(f"{d}/bookkeeping_2026-10.xlsx", arcname="Documents/Bookkeeping/bookkeeping_2026-10.xlsx")
    json.dump({"docs": order, "ap_lines": ap_lines, "memo": memo, "invoices": [i[1] for i in INVOICES],
               "po_list": [[po, VENDORS[v][0], dt_, it, q, p] for po, (v, dt_, items) in POS.items() for it, q, p in items]},
              open(f"{d}/expected.json", "w"), indent=1)
