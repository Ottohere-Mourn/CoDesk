"""Shared grading helpers for the F1 office-document pairs o1-o4 (deployed next to pg_o*.py in task_class/).

- XBook: read an .xlsx and evaluate its formulas (SUM/SUMIF(S)/IF/VLOOKUP/... subset). A file last written by
  openpyxl has no cached values and a file last written by LibreOffice has them; evaluating the formula against the
  final cells is what the user would see when opening the file, so we evaluate and fall back to the cache.
- docx helpers: visible text (pending insertions shown, pending deletions hidden), tracked revisions, effective
  run formatting through the style chain, headers/footers, fields.
- PDF helpers (PyMuPDF) and fuzzy text matching.
"""
from __future__ import annotations

import datetime as dt
import difflib
import functools
import re
import traceback
import zipfile


def safe(fn):
    """Evaluator wrapper: an unexpected file shape must give score 0 with a reason, never crash the run."""
    @functools.wraps(fn)
    def wrapper(env, dest):
        try:
            return fn(env, dest)
        except Exception as e:  # noqa: BLE001
            return {"score": 0.0, "checks": {"evaluator_error": repr(e)[:300],
                                              "trace": traceback.format_exc().splitlines()[-4:]}}
    return wrapper

# ====================================================================== text utilities


def norm(s) -> str:
    """Lower-case, unify quotes/dashes, collapse whitespace."""
    s = str(s or "")
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    s = s.replace("–", "-").replace("—", "-").replace(" ", " ").replace("‑", "-")
    return re.sub(r"\s+", " ", s).strip().lower()


def alnum(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", norm(s)).strip()


def similar(a, b) -> float:
    a, b = alnum(a), alnum(b)
    if not a or not b:
        return 0.0
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()


def best_window(hay: str, needle: str) -> float:
    """Best similarity of `needle` against any equally long word window of `hay` (both alnum-normalised)."""
    h, n = alnum(hay).split(), alnum(needle).split()
    if not n or not h:
        return 0.0
    k = len(n)
    target = " ".join(n)
    best = 0.0
    first = set(n[:3])
    for i in range(0, max(1, len(h) - k + 1)):
        if h[i] not in first and i % 4:
            continue
        for kk in (k - 2, k, k + 2):
            if kk <= 0:
                continue
            cand = " ".join(h[i:i + kk])
            r = difflib.SequenceMatcher(None, cand, target, autojunk=False).ratio()
            best = max(best, r)
    return best


def num(v):
    """Parse a number from a cell value or a string like '£1,234.50', '(400.00)', '-1 080,00', '12%'."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("−", "-").replace(" ", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace("£", "").replace("€", "").replace("$", "").replace("EUR", "").replace("GBP", "")
    pct = s.endswith("%")
    s = s.rstrip("%")
    if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s) or re.fullmatch(r"-?\d+(\.\d+)?", s):
        x = float(s.replace(",", ""))
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s) or re.fullmatch(r"-?\d+,\d{1,2}", s):
        x = float(s.replace(".", "").replace(",", "."))
    else:
        return None
    x = -x if neg else x
    return x / 100 if pct else x


def close(a, b, tol=0.011) -> bool:
    return a is not None and b is not None and abs(float(a) - float(b)) <= tol


# ====================================================================== xlsx formula evaluation


class XLError(Exception):
    pass


class Unsupported(Exception):
    pass


def col2n(c: str) -> int:
    n = 0
    for ch in c.upper():
        n = n * 26 + ord(ch) - 64
    return n


def n2col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def _serial(v):
    if isinstance(v, dt.datetime):
        return (v - EXCEL_EPOCH).total_seconds() / 86400
    if isinstance(v, dt.date):
        return float((v - EXCEL_EPOCH.date()).days)
    if isinstance(v, dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v


class Arr:
    """2-D array of values (a range or an array result)."""

    def __init__(self, rows):
        self.rows = rows

    def flat(self):
        return [v for r in self.rows for v in r]

    @property
    def shape(self):
        return len(self.rows), (len(self.rows[0]) if self.rows else 0)


_TOK = re.compile(r"""
 (?P<ws>\s+)
|(?P<str>"(?:[^"]|"")*")
|(?P<err>\#(?:DIV/0!|N/A|VALUE!|REF!|NAME\?|NUM!|NULL!))
|(?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?(?:\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?\d+:\$?\d+)(?![\w(]))
|(?P<num>(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)
|(?P<func>[A-Za-z_][\w.]*(?=\())
|(?P<name>[A-Za-z_][\w.]*)
|(?P<op><=|>=|<>|[-+*/^&=<>%(),;{}])
""", re.X)


def _tokens(f: str):
    out, pos = [], 0
    while pos < len(f):
        m = _TOK.match(f, pos)
        if not m:
            raise Unsupported(f"token at {f[pos:pos + 10]!r}")
        pos = m.end()
        kind = m.lastgroup
        if kind != "ws":
            out.append((kind, m.group(kind)))
    return out


def _crit(c):
    """SUMIF-style criterion -> predicate."""
    if isinstance(c, Arr):
        c = c.rows[0][0] if c.rows else None
    if isinstance(c, (int, float)) and not isinstance(c, bool):
        return lambda v: _eqnum(v, c)
    s = "" if c is None else str(c)
    m = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", s, re.S)
    op, val = m.group(1) or "=", m.group(2)
    x = num(val) if val.strip() else None
    if x is not None and op in ("<", ">", "<=", ">="):
        f = {"<": lambda a: a < x, ">": lambda a: a > x, "<=": lambda a: a <= x, ">=": lambda a: a >= x}[op]
        return lambda v: (num(_serial(v)) is not None and not isinstance(v, bool)) and f(num(_serial(v)))
    if x is not None and re.fullmatch(r"\s*-?[\d.,]+\s*", val):
        eq = lambda v: _eqnum(v, x)  # noqa: E731
    elif val == "":
        eq = lambda v: v is None or v == ""  # noqa: E731
    else:
        pat = re.compile("^" + re.escape(val.lower()).replace(r"\*", ".*").replace(r"\?", ".") + "$", re.S)
        eq = lambda v: v is not None and bool(pat.match(str(v).strip().lower()))  # noqa: E731
    return eq if op == "=" else (lambda v: not eq(v))


def _eqnum(v, x):
    if isinstance(v, bool) or v is None:
        return False
    y = num(_serial(v)) if not isinstance(v, (int, float)) else float(v)
    return y is not None and abs(y - float(x)) < 1e-9


def _n(v):
    """Coerce to number for arithmetic (Excel rules: blank -> 0, bool -> 0/1, numeric text -> number)."""
    v = _serial(v)
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    x = num(v)
    if x is None:
        raise XLError("#VALUE!")
    return x


def _nums(args):
    """Numbers for SUM/AVERAGE/MIN/MAX: ranges contribute only numeric cells, direct args are coerced."""
    out = []
    for a in args:
        if isinstance(a, Arr):
            out += [float(_serial(v)) for v in a.flat()
                    if isinstance(_serial(v), (int, float)) and not isinstance(v, bool)]
        else:
            out.append(_n(a))
    return out


def _scalar(v):
    if isinstance(v, Arr):
        return v.rows[0][0] if v.rows and v.rows[0] else None
    return v


def _truth(v):
    v = _scalar(v)
    if isinstance(v, str):
        if v.upper() in ("TRUE", "FALSE"):
            return v.upper() == "TRUE"
        raise XLError("#VALUE!")
    return bool(_n(v))


def _cmp(op, a, b):
    a, b = _serial(a), _serial(b)
    if a is None:
        a = "" if isinstance(b, str) else 0.0
    if b is None:
        b = "" if isinstance(a, str) else 0.0
    if isinstance(a, str) and isinstance(b, str):
        a, b = a.lower(), b.lower()
    elif isinstance(a, str) or isinstance(b, str):  # text > numbers in Excel ordering
        ta, tb = isinstance(a, str), isinstance(b, str)
        return {"=": False, "<>": True, "<": tb, ">": ta, "<=": tb, ">=": ta}[op]
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


def _bin(op, a, b):
    if isinstance(a, Arr) or isinstance(b, Arr):
        A = a if isinstance(a, Arr) else None
        B = b if isinstance(b, Arr) else None
        r, c = (A or B).shape
        get = lambda X, x, i, j: (X.rows[min(i, X.shape[0] - 1)][min(j, X.shape[1] - 1)] if X else x)  # noqa: E731
        return Arr([[_bin(op, get(A, a, i, j), get(B, b, i, j)) for j in range(c)] for i in range(r)])
    if op in ("=", "<>", "<", ">", "<=", ">="):
        return _cmp(op, a, b)
    if op == "&":
        f = lambda v: "" if v is None else ("TRUE" if v is True else "FALSE" if v is False else  # noqa: E731
                                            (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)))
        return f(a) + f(b)
    x, y = _n(a), _n(b)
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if op == "/":
        if y == 0:
            raise XLError("#DIV/0!")
        return x / y
    if op == "^":
        return x ** y
    raise Unsupported(op)


class XBook:
    """Formulas from openpyxl (data_only=False), cached values from a second load (data_only=True)."""

    def __init__(self, path):
        from openpyxl import load_workbook
        self.wf = load_workbook(path, data_only=False)
        self.wv = load_workbook(path, data_only=True)
        self.path = path
        self.memo, self.stack = {}, set()
        self.names = {n.lower(): n for n in self.wf.sheetnames}

    def sheet(self, pattern):
        """First sheet whose name matches the regex (case-insensitive), or None."""
        for n in self.wf.sheetnames:
            if re.fullmatch(pattern, n.strip(), re.I):
                return n
        for n in self.wf.sheetnames:
            if re.search(pattern, n, re.I):
                return n
        return None

    def raw(self, sheet, ref):
        return self.wf[sheet][ref].value

    def cached(self, sheet, ref):
        return self.wv[sheet][ref].value

    def is_formula(self, sheet, ref):
        v = self.wf[sheet][ref].value
        return isinstance(v, str) and v.startswith("=") or type(v).__name__ in ("ArrayFormula", "DataTableFormula")

    def formula(self, sheet, ref) -> str:
        v = self.wf[sheet][ref].value
        if type(v).__name__ == "ArrayFormula":
            v = v.text
        return v if isinstance(v, str) and v.startswith("=") else ""

    def value(self, sheet, ref):
        """Evaluated value of a cell (formula evaluated; cached value if evaluation is unsupported)."""
        key = (sheet, ref.replace("$", "").upper())
        if key in self.memo:
            v = self.memo[key]
            if isinstance(v, XLError):
                raise v
            return v
        raw = self.wf[sheet][key[1]].value
        f = raw.text if type(raw).__name__ == "ArrayFormula" else raw
        if isinstance(f, str) and f.startswith("="):
            if key in self.stack:
                raise XLError("#CIRC")
            self.stack.add(key)
            try:
                v = self.eval(f[1:], sheet)
                v = _scalar(v)
            except XLError as e:
                self.memo[key] = e
                raise
            except (Unsupported, RecursionError, IndexError, ValueError, TypeError, KeyError, ZeroDivisionError,
                    AttributeError):
                v = self.wv[sheet][key[1]].value
            finally:
                self.stack.discard(key)
            self.memo[key] = v
            return v
        self.memo[key] = raw
        return raw

    def safe(self, sheet, ref):
        try:
            return self.value(sheet, ref)
        except XLError:
            return None

    # ------------------------------------------------------------------ parser
    def eval(self, text, sheet):
        saved = (getattr(self, "toks", None), getattr(self, "i", 0), getattr(self, "cur", None))
        try:
            self.toks, self.i, self.cur = _tokens(text), 0, sheet
            v = self._expr()
            if self.i != len(self.toks):
                raise Unsupported("trailing tokens")
            return v
        finally:
            self.toks, self.i, self.cur = saved

    def _peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def _take(self):
        t = self._peek()
        self.i += 1
        return t

    def _expr(self):
        a = self._concat()
        while self._peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self._take()[1]
            a = _bin(op, a, self._concat())
        return a

    def _concat(self):
        a = self._add()
        while self._peek()[1] == "&":
            self._take()
            a = _bin("&", a, self._add())
        return a

    def _add(self):
        a = self._mul()
        while self._peek()[1] in ("+", "-"):
            op = self._take()[1]
            a = _bin(op, a, self._mul())
        return a

    def _mul(self):
        a = self._pow()
        while self._peek()[1] in ("*", "/"):
            op = self._take()[1]
            a = _bin(op, a, self._pow())
        return a

    def _pow(self):
        a = self._unary()
        while self._peek()[1] == "^":
            self._take()
            a = _bin("^", a, self._unary())
        return a

    def _unary(self):
        if self._peek()[1] in ("-", "+"):
            op = self._take()[1]
            v = self._unary()
            return _bin("*", -1.0 if op == "-" else 1.0, v)
        v = self._primary()
        while self._peek()[1] == "%":
            self._take()
            v = _bin("/", v, 100.0)
        return v

    def _primary(self):
        kind, tok = self._take()
        if kind == "num":
            return float(tok)
        if kind == "str":
            return tok[1:-1].replace('""', '"')
        if kind == "err":
            raise XLError(tok)
        if kind == "ref":
            return self._ref(tok)
        if kind == "name":
            if tok.upper() in ("TRUE", "FALSE"):
                return tok.upper() == "TRUE"
            raise Unsupported(f"name {tok}")
        if kind == "func":
            return self._call(tok)
        if tok == "(":
            v = self._expr()
            if self._take()[1] != ")":
                raise Unsupported("paren")
            return v
        raise Unsupported(f"unexpected {tok}")

    def _ref(self, tok):
        sheet = self.cur
        if "!" in tok:
            sh, tok = tok.rsplit("!", 1)
            sh = sh[1:-1].replace("''", "'") if sh.startswith("'") else sh
            sheet = self.names.get(sh.lower())
            if sheet is None:
                raise XLError("#REF!")
        tok = tok.replace("$", "").upper()
        ws = self.wf[sheet]
        if ":" not in tok:
            return self.value(sheet, tok)
        a, b = tok.split(":")
        ma, mb = re.fullmatch(r"([A-Z]*)(\d*)", a), re.fullmatch(r"([A-Z]*)(\d*)", b)
        c0 = col2n(ma.group(1)) if ma.group(1) else 1
        c1 = col2n(mb.group(1)) if mb.group(1) else ws.max_column
        r0 = int(ma.group(2)) if ma.group(2) else 1
        r1 = int(mb.group(2)) if mb.group(2) else ws.max_row
        r1, c1 = min(r1, max(ws.max_row, r0)), min(c1, max(ws.max_column, c0))
        if (r1 - r0 + 1) * (c1 - c0 + 1) > 200000:
            raise Unsupported("range too large")
        return Arr([[self._cellval(sheet, r, c) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)])

    def _cellval(self, sheet, r, c):
        try:
            return self.value(sheet, f"{n2col(c)}{r}")
        except XLError:
            return None

    def _args(self):
        """Parse an argument list after the function name; arguments are evaluated lazily (thunks)."""
        if self._take()[1] != "(":
            raise Unsupported("call")
        args = []
        if self._peek()[1] == ")":
            self._take()
            return args
        while True:
            start = self.i
            depth = 0
            while True:
                k, t = self._peek()
                if k is None:
                    raise Unsupported("eof")
                if t in ("(", "{"):
                    depth += 1
                elif t in (")", "}"):
                    if depth == 0:
                        break
                    depth -= 1
                elif t in (",", ";") and depth == 0:
                    break
                self.i += 1
            args.append(self.toks[start:self.i])
            k, t = self._take()
            if t == ")":
                return args
        return args

    def _ev(self, toks):
        if not toks:
            return None
        saved = (self.toks, self.i)
        self.toks, self.i = toks, 0
        try:
            v = self._expr()
            if self.i != len(self.toks):
                raise Unsupported("arg")
            return v
        finally:
            self.toks, self.i = saved

    def _call(self, name):
        name = re.sub(r"^(_xlfn\.|_xlws\.|com\.microsoft\.)", "", name, flags=re.I).upper()
        raw = self._args()
        ev = self._ev
        if name == "IF":
            c = _truth(ev(raw[0]))
            if c:
                return ev(raw[1]) if len(raw) > 1 else True
            return ev(raw[2]) if len(raw) > 2 else False
        if name in ("IFERROR", "IFNA"):
            try:
                v = ev(raw[0])
                _scalar(v)
                return v
            except XLError:
                return ev(raw[1])
        args = [ev(a) for a in raw]
        if name == "SUM":
            return sum(_nums(args))
        if name in ("AVERAGE", "MIN", "MAX", "COUNT"):
            xs = _nums(args)
            if name == "COUNT":
                return float(len(xs))
            if not xs:
                if name == "AVERAGE":
                    raise XLError("#DIV/0!")
                return 0.0
            return {"AVERAGE": sum(xs) / len(xs), "MIN": min(xs), "MAX": max(xs)}[name]
        if name == "COUNTA":
            return float(sum(1 for a in args for v in (a.flat() if isinstance(a, Arr) else [a]) if v not in (None, "")))
        if name in ("SUMIF", "AVERAGEIF", "COUNTIF"):
            rng = args[0] if isinstance(args[0], Arr) else Arr([[args[0]]])
            pred = _crit(args[1])
            if name == "COUNTIF":
                return float(sum(1 for v in rng.flat() if pred(v)))
            srng = args[2] if len(args) > 2 and isinstance(args[2], Arr) else rng
            vals = [s for v, s in zip(rng.flat(), self._shaped(srng, rng)) if pred(v)]
            xs = [float(_serial(s)) for s in vals if isinstance(_serial(s), (int, float)) and not isinstance(s, bool)]
            if name == "AVERAGEIF":
                if not xs:
                    raise XLError("#DIV/0!")
                return sum(xs) / len(xs)
            return sum(xs)
        if name in ("SUMIFS", "COUNTIFS", "AVERAGEIFS"):
            if name == "COUNTIFS":
                pairs, srng = args, None
            else:
                srng, pairs = args[0], args[1:]
            base = pairs[0]
            ok = [True] * len(base.flat())
            for k in range(0, len(pairs), 2):
                rng = pairs[k]
                pred = _crit(pairs[k + 1])
                for i, v in enumerate(self._shaped(rng, base)):
                    ok[i] = ok[i] and pred(v)
            if name == "COUNTIFS":
                return float(sum(ok))
            vals = [s for s, o in zip(self._shaped(srng, base), ok) if o]
            xs = [float(_serial(s)) for s in vals if isinstance(_serial(s), (int, float)) and not isinstance(s, bool)]
            if name == "AVERAGEIFS":
                if not xs:
                    raise XLError("#DIV/0!")
                return sum(xs) / len(xs)
            return sum(xs)
        if name == "SUMPRODUCT":
            arrs = [a if isinstance(a, Arr) else Arr([[a]]) for a in args]
            n = len(arrs[0].flat())
            tot = 0.0
            for i in range(n):
                p = 1.0
                for a in arrs:
                    v = a.flat()[i] if i < len(a.flat()) else 0
                    v = _serial(v)
                    p *= float(v) if isinstance(v, (int, float)) else 0.0
                tot += p
            return tot
        if name == "ABS":
            return abs(_n(_scalar(args[0])))
        if name in ("ROUND", "ROUNDUP", "ROUNDDOWN"):
            x, d = _n(_scalar(args[0])), int(_n(_scalar(args[1]))) if len(args) > 1 else 0
            if name == "ROUND":
                import decimal
                q = decimal.Decimal(1).scaleb(-d)
                return float(decimal.Decimal(repr(x)).quantize(q, rounding=decimal.ROUND_HALF_UP))
            import math
            f = 10 ** d
            return (math.ceil if (name == "ROUNDUP") == (x >= 0) else math.floor)(x * f) / f
        if name in ("AND", "OR"):
            vals = [_truth(v) for a in args for v in (a.flat() if isinstance(a, Arr) else [a]) if v is not None]
            return all(vals) if name == "AND" else any(vals)
        if name == "NOT":
            return not _truth(args[0])
        if name in ("VLOOKUP", "HLOOKUP"):
            key, table, idx = _scalar(args[0]), args[1], int(_n(_scalar(args[2])))
            exact = len(args) > 3 and not _truth(args[3])
            rows = table.rows if name == "VLOOKUP" else [list(c) for c in zip(*table.rows)]
            hit = None
            for r in rows:
                if r and (_cmp("=", r[0], key) or (num(r[0]) is not None and num(key) is not None
                                                    and abs(num(r[0]) - num(key)) < 1e-9)):
                    hit = r
                    break
            if hit is None and not exact:
                for r in rows:
                    if r and r[0] is not None and _cmp("<=", r[0], key):
                        hit = r
            if hit is None:
                raise XLError("#N/A")
            return hit[idx - 1]
        if name == "INDEX":
            arr, r = args[0], int(_n(_scalar(args[1])))
            c = int(_n(_scalar(args[2]))) if len(args) > 2 else 1
            if arr.shape[0] == 1 and len(args) == 2:
                r, c = 1, r
            return arr.rows[r - 1][c - 1]
        if name == "MATCH":
            key, arr = _scalar(args[0]), args[1].flat()
            for i, v in enumerate(arr, 1):
                if _cmp("=", v, key) or (num(v) is not None and num(key) is not None and abs(num(v) - num(key)) < 1e-9):
                    return float(i)
            raise XLError("#N/A")
        if name in ("CONCATENATE", "CONCAT"):
            out = ""
            for a in args:
                for v in (a.flat() if isinstance(a, Arr) else [a]):
                    out = _bin("&", out, v)
            return out
        if name == "TEXT":
            return str(_scalar(args[0]))
        if name == "VALUE":
            return _n(_scalar(args[0]))
        if name == "N":
            v = _scalar(args[0])
            return float(v) if isinstance(v, (int, float)) else 0.0
        if name == "ISBLANK":
            return _scalar(args[0]) in (None, "")
        if name in ("ISNUMBER", "ISTEXT"):
            v = _scalar(args[0])
            return isinstance(v, (int, float)) and not isinstance(v, bool) if name == "ISNUMBER" else isinstance(v, str)
        raise Unsupported(name)

    @staticmethod
    def _shaped(rng, like):
        flat = rng.flat() if isinstance(rng, Arr) else [rng]
        n = len(like.flat()) if isinstance(like, Arr) else 1
        return (flat + [None] * n)[:n]


# ====================================================================== xlsx structure helpers (zip level)

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels(z, part):
    from lxml import etree
    d, f = part.rsplit("/", 1)
    rp = f"{d}/_rels/{f}.rels"
    if rp not in z.namelist():
        return {}
    root = etree.fromstring(z.read(rp))
    out = {}
    for r in root:
        tgt = r.get("Target")
        if r.get("TargetMode") == "External":
            continue
        if tgt.startswith("/"):
            full = tgt[1:]
        else:
            parts = (d + "/" + tgt).split("/")
            st = []
            for p in parts:
                if p == "..":
                    st.pop()
                elif p and p != ".":
                    st.append(p)
            full = "/".join(st)
        out[r.get("Id")] = (r.get("Type").rsplit("/", 1)[-1], full)
    return out


def xlsx_charts(path) -> dict:
    """{sheet name: [chart xml strings]} for charts anchored on each worksheet."""
    from lxml import etree
    out = {}
    with zipfile.ZipFile(path) as z:
        wb = etree.fromstring(z.read("xl/workbook.xml"))
        rels = _rels(z, "xl/workbook.xml")
        for s in wb.iter(f"{{{NS_MAIN}}}sheet"):
            rid = s.get(f"{{{NS_R}}}id")
            if rid not in rels:
                continue
            part = rels[rid][1]
            charts = []
            for typ, drawing in _rels(z, part).values():
                if typ != "drawing" or drawing not in z.namelist():
                    continue
                for typ2, chart in _rels(z, drawing).values():
                    if typ2 == "chart" and chart in z.namelist():
                        charts.append(z.read(chart).decode("utf8", "replace"))
            out[s.get("name")] = charts
    return out


def cf_rules(ws):
    """[(cell ranges string, rule)] for the openpyxl worksheet's conditional formats."""
    out = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            out.append((str(cf.sqref), rule))
    return out


def in_ranges(sqref: str, ref: str) -> bool:
    from openpyxl.utils.cell import range_boundaries
    m = re.fullmatch(r"([A-Z]+)(\d+)", ref.replace("$", "").upper())
    c, r = col2n(m.group(1)), int(m.group(2))
    for part in sqref.split():
        try:
            c0, r0, c1, r1 = range_boundaries(part)
        except ValueError:
            continue
        c0, r0 = c0 or 1, r0 or 1
        c1, r1 = c1 or 16384, r1 or 1048576
        if c0 <= c <= c1 and r0 <= r <= r1:
            return True
    return False


# ====================================================================== docx helpers

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WNS = {"w": W}


def q(tag):
    p, t = tag.split(":")
    return f"{{{W}}}{t}"


class Docx:
    def __init__(self, path):
        from lxml import etree
        self.path = path
        self.z = zipfile.ZipFile(path)
        self.names = self.z.namelist()
        self.doc = etree.fromstring(self.z.read("word/document.xml"))
        self.body = self.doc.find(q("w:body"))
        self.styles = etree.fromstring(self.z.read("word/styles.xml")) if "word/styles.xml" in self.names else None
        self._style_by_id = {}
        if self.styles is not None:
            for s in self.styles.findall(q("w:style")):
                self._style_by_id[s.get(q("w:styleId"))] = s
        self.rels = _rels(self.z, "word/document.xml")

    def part(self, name):
        from lxml import etree
        return etree.fromstring(self.z.read(name)) if name in self.names else None

    # ---------------------------------------------------------------- styles
    def style_name(self, sid):
        s = self._style_by_id.get(sid)
        if s is None:
            return sid or ""
        n = s.find(q("w:name"))
        return n.get(q("w:val")) if n is not None else sid

    def para_style(self, p):
        ps = p.find("w:pPr/w:pStyle", WNS)
        return ps.get(q("w:val")) if ps is not None else None

    def para_style_name(self, p):
        sid = self.para_style(p)
        if sid is None:
            # default paragraph style
            for s in self._style_by_id.values():
                if s.get(q("w:type")) == "paragraph" and s.get(q("w:default")) in ("1", "true"):
                    return self.style_name(s.get(q("w:styleId")))
            return "Normal"
        return self.style_name(sid)

    def _style_chain(self, sid):
        seen = []
        while sid and sid in self._style_by_id and sid not in seen:
            seen.append(sid)
            b = self._style_by_id[sid].find(q("w:basedOn"))
            sid = b.get(q("w:val")) if b is not None else None
        return [self._style_by_id[s] for s in seen]

    def _default_pstyle(self):
        for s in self._style_by_id.values():
            if s.get(q("w:type")) == "paragraph" and s.get(q("w:default")) in ("1", "true"):
                return s.get(q("w:styleId"))
        return None

    def _rpr_chain(self, p, r):
        """rPr elements from most to least specific for run r in paragraph p."""
        chain = []
        if r is not None:
            rpr = r.find(q("w:rPr"))
            if rpr is not None:
                chain.append(rpr)
                rs = rpr.find(q("w:rStyle"))
                if rs is not None:
                    chain += [s.find(q("w:rPr")) for s in self._style_chain(rs.get(q("w:val")))]
        sid = self.para_style(p) or self._default_pstyle()
        chain += [s.find(q("w:rPr")) for s in self._style_chain(sid)]
        if self.styles is not None:
            d = self.styles.find("w:docDefaults/w:rPrDefault/w:rPr", WNS)
            chain.append(d)
        return [c for c in chain if c is not None]

    def _ppr_chain(self, p):
        chain = []
        ppr = p.find(q("w:pPr"))
        if ppr is not None:
            chain.append(ppr)
        sid = self.para_style(p) or self._default_pstyle()
        chain += [s.find(q("w:pPr")) for s in self._style_chain(sid)]
        if self.styles is not None:
            chain.append(self.styles.find("w:docDefaults/w:pPrDefault/w:pPr", WNS))
        return [c for c in chain if c is not None]

    def run_size(self, p, r):
        for e in self._rpr_chain(p, r):
            sz = e.find(q("w:sz"))
            if sz is not None and sz.get(q("w:val")):
                return int(float(sz.get(q("w:val")))) / 2
        return 10.0

    def run_font(self, p, r):
        for e in self._rpr_chain(p, r):
            f = e.find(q("w:rFonts"))
            if f is not None and (f.get(q("w:ascii")) or f.get(q("w:hAnsi"))):
                return f.get(q("w:ascii")) or f.get(q("w:hAnsi"))
        return ""

    def run_bold(self, p, r):
        for e in self._rpr_chain(p, r):
            b = e.find(q("w:b"))
            if b is not None:
                return b.get(q("w:val")) not in ("0", "false", "off")
        return False

    def ppr_val(self, p, path, attr):
        """First value of attribute `attr` on element `path` (e.g. 'w:spacing', 'w:line') along the pPr chain."""
        for e in self._ppr_chain(p):
            el = e.find(path, WNS)
            if el is not None and el.get(q(attr)) is not None:
                return el.get(q(attr))
        return None

    def ppr_el(self, p, path):
        for e in self._ppr_chain(p):
            el = e.find(path, WNS)
            if el is not None:
                return el
        return None

    # ---------------------------------------------------------------- text
    @staticmethod
    def ptext(p) -> str:
        """Visible text of a paragraph: pending insertions included, pending deletions excluded."""
        out = []
        for el in p.iter(q("w:t"), q("w:tab"), q("w:br"), q("w:delText"), q("w:noBreakHyphen")):
            tag = el.tag.rsplit("}", 1)[1]
            if tag == "delText":
                continue
            if any(a.tag in (q("w:del"), q("w:moveFrom")) for a in el.iterancestors()):
                continue
            if tag == "t":
                out.append(el.text or "")
            elif tag == "tab":
                out.append("\t")
            elif tag == "noBreakHyphen":
                out.append("-")
            else:
                out.append("\n")
        return "".join(out)

    def paragraphs(self, root=None, tables=True):
        """Body paragraphs in document order (descending into tables / sdt / text boxes when tables=True)."""
        root = self.body if root is None else root
        if tables:
            return [p for p in root.iter(q("w:p"))]
        return [p for p in root if p.tag == q("w:p")] + []

    def body_paragraphs(self):
        """Top-level paragraphs (incl. inside sdt blocks such as a TOC), not inside tables."""
        out = []
        for el in self.body.iter(q("w:p")):
            if any(a.tag == q("w:tbl") for a in el.iterancestors()):
                continue
            if any(a.tag in (q("w:txbxContent"),) for a in el.iterancestors()):
                continue
            out.append(el)
        return out

    def tables(self):
        return [t for t in self.body.iter(q("w:tbl")) if not any(a.tag == q("w:tbl") for a in t.iterancestors())]

    @classmethod
    def table_rows(cls, tbl):
        rows = []
        for tr in tbl.findall(q("w:tr")):
            rows.append([" ".join(cls.ptext(p) for p in tc.iter(q("w:p"))).strip() for tc in tr.findall(q("w:tc"))])
        return rows

    def full_text(self) -> str:
        return "\n".join(self.ptext(p) for p in self.body.iter(q("w:p")))

    def revisions(self):
        """[(kind ins/del, author, text)] for pending tracked changes in the body."""
        out = []
        for el in self.body.iter(q("w:ins"), q("w:del"), q("w:moveFrom"), q("w:moveTo")):
            kind = el.tag.rsplit("}", 1)[1]
            txt = "".join((t.text or "") for t in el.iter(q("w:t"), q("w:delText")))
            out.append((kind, el.get(q("w:author")) or "", txt))
        return out

    def comments(self):
        """[(id, author, text)] for comments that are still anchored in the document body."""
        ids = {c.get(q("w:id")) for c in self.body.iter(q("w:commentReference"))}
        ids |= {c.get(q("w:id")) for c in self.body.iter(q("w:commentRangeStart"))}
        out = []
        for rel in self.rels.values():
            if rel[0] == "comments":
                root = self.part(rel[1])
                if root is None:
                    continue
                for c in root.findall(q("w:comment")):
                    if c.get(q("w:id")) in ids:
                        out.append((c.get(q("w:id")), c.get(q("w:author")) or "",
                                    " ".join(self.ptext(p) for p in c.iter(q("w:p")))))
        return out

    def header_footer(self, kind="footer"):
        """Texts and field instructions of all headers or footers referenced from section properties."""
        texts, instrs, roots = [], [], []
        for rel in self.rels.values():
            if rel[0] == kind:
                root = self.part(rel[1])
                if root is None:
                    continue
                roots.append(root)
                texts.append("\n".join(self.ptext(p) for p in root.iter(q("w:p"))))
                instrs += [(e.get(q("w:instr")) or "") for e in root.iter(q("w:fldSimple"))]
                instrs += [(e.text or "") for e in root.iter(q("w:instrText"))]
        return texts, instrs, roots

    def sections(self):
        return list(self.body.iter(q("w:sectPr")))


def twips(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


# ====================================================================== pdf helpers


def pdf_pages(path):
    """[page text] with PyMuPDF; None if unreadable."""
    try:
        import fitz
        with fitz.open(path) as d:
            return [p.get_text() for p in d]
    except Exception:
        return None


def pdf_lines(path):
    """[[(y, text of the visual line)]] per page: spans grouped by baseline, left to right."""
    import fitz
    out = []
    with fitz.open(path) as d:
        for page in d:
            rows = {}
            for b in page.get_text("dict")["blocks"]:
                for ln in b.get("lines", []):
                    for sp in ln["spans"]:
                        if not sp["text"].strip():
                            continue
                        y = round(sp["origin"][1] / 3)
                        key = next((k for k in rows if abs(k - y) <= 1), y)
                        rows.setdefault(key, []).append((sp["origin"][0], sp["text"]))
            out.append([(k * 3, " ".join(t for _, t in sorted(v)).strip()) for k, v in sorted(rows.items())])
    return out


def pdf_meta(path):
    try:
        import fitz
        with fitz.open(path) as d:
            return {"pages": d.page_count, "sizes": [(round(p.rect.width), round(p.rect.height)) for p in d]}
    except Exception:
        return None
