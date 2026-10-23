"""F3 assets - a community fair's photos: the thank-you slideshow + photo album vs next year's poster.

Writes into out_dir:
  photos/IMG-20261017-WA00NN.jpg   22 event photos (1600x1200, as received via a messenger: no EXIF times)
  programme_scan.pdf               image-only scan of the printed programme with handwritten changes
  committee_note.jpg               phone photo of a handwritten sticky note (next year's date/venue)
  MakersFair_2026_recap.pptx       the slideshow template (title + closing slide)
  poster_2027.docx                 the poster template (placeholders)
  f3.tar.gz                        everything above in Guest layout, extracted once by setup
  expected.json                    ground truth - host only
"""
from __future__ import annotations

import json
import os
import random
import tarfile

import numpy as np
from PIL import Image, ImageDraw, ImageFilter

import lib_f
from lib_f import font, hand_text, md5, noise

W, H = 1600, 1200
# programme: key, printed time, actual time (None = cancelled), title, who, place
PROGRAMME = [
    ("opening", "10:00", "10:00", "Official opening", "Mayor Cllr Jenna Treloar", "Main stage"),
    ("pottery", "10:30", "10:30", "Pottery wheel demo", "Penhallow Pottery", "Stall 4"),
    ("choir", "11:15", "11:15", "Children's choir", "St Mary's Primary School", "Main stage"),
    ("forge", "12:00", None, "Blacksmith demo", "Kernow Forge", "Old yard"),
    ("ukulele", "13:00", "13:00", "Ukulele set", "Falmouth Ukulele Club", "Main stage"),
    ("bread", "14:00", "14:30", "Bread masterclass", "Harbour Street Bakery", "Tent B"),
    ("repair", "15:00", "15:00", "Repair cafe", "Fix-it Falmouth volunteers", "Church hall"),
    ("raffle", "16:00", "16:00", "Grand raffle draw & close", "Riverside Makers Collective", "Main stage"),
]
# photos per item: list of umbrella flags (True = the family with the red umbrella is in the picture)
PHOTOS = {"opening": [True, False], "pottery": [False, True, False], "choir": [False, False, False],
          "ukulele": [False, True, False], "bread": [False, False, False], "repair": [True, False],
          "raffle": [False, False], "general": [False, False, False], "forge": [False]}
POSTER_ACTS = ["pottery", "choir", "ukulele"]
SKIN = [(241, 194, 160), (224, 172, 130), (198, 134, 92), (141, 85, 56), (255, 219, 186)]
SHIRTS = [(200, 60, 60), (60, 90, 170), (240, 200, 60), (90, 150, 90), (150, 80, 160), (240, 140, 50), (60, 60, 60),
          (230, 230, 230), (40, 140, 160)]


# ---------------------------------------------------------------- drawing primitives
def person(d, x, feet, h, rng, shirt=None, child=False):
    h = h * (0.62 if child else 1.0)
    skin = rng.choice(SKIN)
    shirt = shirt or rng.choice(SHIRTS)
    trous = rng.choice([(50, 60, 90), (40, 40, 40), (120, 100, 80), (70, 90, 130)])
    hw = h * 0.16
    d.rectangle((x - hw * 0.8, feet - h * 0.45, x - hw * 0.1, feet), fill=trous)
    d.rectangle((x + hw * 0.1, feet - h * 0.45, x + hw * 0.8, feet), fill=trous)
    d.rounded_rectangle((x - hw, feet - h * 0.82, x + hw, feet - h * 0.42), radius=int(hw * 0.4), fill=shirt)
    r = h * 0.1
    d.ellipse((x - r, feet - h * 0.82 - 2 * r, x + r, feet - h * 0.82), fill=skin)
    d.chord((x - r, feet - h * 0.82 - 2 * r - 2, x + r, feet - h * 0.82 - r * 0.6), 180, 360,
            fill=rng.choice([(40, 30, 20), (90, 60, 30), (200, 170, 90), (20, 20, 20), (150, 150, 150)]))


def crowd(d, rng, n, y0, y1, x0=0, x1=W, scale=1.0):
    pts = sorted([(rng.uniform(x0, x1), rng.uniform(y0, y1)) for _ in range(n)], key=lambda p: p[1])
    for x, y in pts:
        person(d, x, y, (180 + (y - y0) * 0.9) * scale, rng, child=rng.random() < 0.2)


def umbrella_family(d, x, feet, h, rng):
    for i, dx in enumerate((-55, 0, 45)):
        person(d, x + dx, feet + i * 6, h * (0.62 if i == 2 else 1.0), rng,
               shirt=[(80, 80, 160), (230, 230, 230), (240, 200, 60)][i])
    top = feet - h * 1.25
    d.line((x, top + 20, x, feet - h * 0.5), fill=(40, 40, 40), width=6)
    d.chord((x - h * 0.75, top - h * 0.3, x + h * 0.75, top + h * 0.45), 180, 360, fill=(214, 24, 30))
    for k in range(-2, 3):
        d.line((x, top - h * 0.3 + 4, x + k * h * 0.3, top + h * 0.07), fill=(150, 10, 15), width=3)


def sky_grass(rng, indoor=False):
    if indoor:
        a = np.zeros((H, W, 3), np.float32)
        a[:] = (224, 214, 190)
        a[int(H * 0.62):] = (165, 120, 80)
        im = Image.fromarray(a.astype(np.uint8))
        d = ImageDraw.Draw(im)
        for i in range(4):
            x = 120 + i * 400 + rng.randint(-20, 20)
            d.rectangle((x, 120, x + 220, 420), fill=(180, 210, 235), outline=(255, 255, 255), width=10)
        for i in range(0, W, 60):
            d.line((i, int(H * 0.62), i - 200, H), fill=(150, 108, 70), width=3)
        return im
    y = np.linspace(0, 1, H)[:, None, None]
    sky = np.array([110, 165, 225], np.float32) * (1 - y) + np.array([200, 225, 245], np.float32) * y
    a = np.broadcast_to(sky, (H, W, 3)).copy()
    hz = int(H * 0.55)
    g = np.linspace(0, 1, H - hz)[:, None, None]
    a[hz:] = np.array([110, 160, 80], np.float32) * (1 - g) + np.array([80, 135, 60], np.float32) * g
    im = Image.fromarray(a.astype(np.uint8))
    d = ImageDraw.Draw(im)
    for i in range(14):
        x = rng.randint(-50, W + 50)
        r = rng.randint(60, 130)
        d.ellipse((x - r, hz - r * 1.6, x + r, hz + 10), fill=(50 + rng.randint(0, 30), 100 + rng.randint(0, 40), 50))
    for i in range(3):
        cx, cy = rng.randint(0, W), rng.randint(40, 200)
        for k in range(4):
            d.ellipse((cx + k * 50 - 60, cy - 30 + (k % 2) * 15, cx + k * 50 + 40, cy + 40), fill=(250, 250, 252))
    return im


def bunting(d, y, rng, dx=0):
    d.line((0, y, W, y + 30), fill=(90, 90, 90), width=3)
    for i in range(0, W, 70):
        c = rng.choice([(220, 50, 50), (250, 200, 40), (40, 120, 200), (60, 170, 90), (240, 240, 240)])
        yy = y + i * 30 / W
        d.polygon([(i + dx % 70, yy), (i + 50 + dx % 70, yy), (i + 25 + dx % 70, yy + 55)], fill=c)


def banner(d, box, text, fill=(255, 255, 255), ink=(30, 40, 90), size=56):
    d.rectangle(box, fill=fill, outline=(60, 60, 60), width=3)
    cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
    f = font("lato-black", size)
    while f.getlength(text) > (box[2] - box[0]) - 30 and size > 18:
        size -= 2
        f = font("lato-black", size)
    d.text((cx, cy), text, fill=ink, font=f, anchor="mm")


def stage(d, ox, text, rng):
    d.rectangle((ox + 150, 250, ox + 1450, 700), fill=(30, 40, 80))
    d.polygon([(ox + 100, 250), (ox + 1500, 250), (ox + 1450, 190), (ox + 150, 190)], fill=(200, 200, 205))
    banner(d, (ox + 250, 270, ox + 1350, 370), text)
    d.rectangle((ox + 120, 700, ox + 1480, 790), fill=(120, 80, 50))
    for k in range(6):
        d.line((ox + 120 + k * 270, 700, ox + 120 + k * 270, 790), fill=(90, 60, 40), width=4)
    d.rectangle((ox + 170, 500, ox + 260, 700), fill=(20, 20, 20))
    d.rectangle((ox + 1340, 500, ox + 1430, 700), fill=(20, 20, 20))


def gazebo(d, ox, text, col, rng):
    d.polygon([(ox + 300, 330), (ox + 1300, 330), (ox + 1200, 230), (ox + 400, 230)], fill=col)
    for k in range(8):
        x = ox + 300 + k * 125
        d.polygon([(x, 330), (x + 62, 330), (x + 50, 360), (x + 12, 360)], fill=col if k % 2 else (245, 245, 245))
    d.line((ox + 320, 330, ox + 320, 760), fill=(80, 80, 80), width=10)
    d.line((ox + 1280, 330, ox + 1280, 760), fill=(80, 80, 80), width=10)
    banner(d, (ox + 430, 245, ox + 1170, 325), text, fill=(250, 245, 230), ink=(90, 40, 20), size=48)
    d.rectangle((ox + 380, 620, ox + 1220, 650), fill=(240, 240, 235))
    d.rectangle((ox + 390, 650, ox + 1210, 760), fill=(200, 60, 60))


def marquee(d, ox, text):
    d.polygon([(ox + 200, 360), (ox + 800, 160), (ox + 1400, 360)], fill=(245, 245, 245), outline=(180, 180, 180))
    d.rectangle((ox + 200, 360, ox + 1400, 780), fill=(240, 240, 240), outline=(180, 180, 180), width=3)
    d.polygon([(ox + 780, 160), (ox + 780, 60), (ox + 870, 85), (ox + 780, 110)], fill=(220, 40, 40))
    d.text((ox + 812, 85), "B", fill="white", font=font("lato-black", 34), anchor="mm")
    banner(d, (ox + 380, 372, ox + 1220, 452), text, fill=(255, 250, 235), ink=(120, 70, 20), size=50)
    d.rectangle((ox + 450, 470, ox + 1150, 780), fill=(70, 50, 40))
    d.rectangle((ox + 300, 640, ox + 1300, 670), fill=(230, 220, 200))


def props(d, key, ox, rng):
    if key == "opening":
        person(d, ox + 800, 690, 260, rng, shirt=(30, 30, 30))
        d.arc((ox + 765, 480, ox + 835, 550), 20, 160, fill=(230, 190, 40), width=8)
        d.polygon([(ox + 700, 600), (ox + 900, 600), (ox + 880, 700), (ox + 720, 700)], fill=(120, 70, 40))
        d.line((ox + 770, 600, ox + 780, 560), fill=(60, 60, 60), width=4)
    elif key == "choir":
        for row in range(2):
            for k in range(7):
                person(d, ox + 380 + k * 140 + row * 70, 690 - row * 40, 260, rng, shirt=(40, 80, 170), child=True)
    elif key == "ukulele":
        for k in range(4):
            x = ox + 420 + k * 250
            person(d, x, 700, 290, rng)
            d.ellipse((x - 10, 560, x + 60, 610), fill=(170, 110, 50))
            d.line((x + 40, 575, x + 120, 520), fill=(90, 60, 30), width=8)
    elif key == "raffle":
        d.ellipse((ox + 650, 470, ox + 950, 690), fill=(240, 190, 30), outline=(150, 110, 0), width=6)
        d.text((ox + 800, 580), "RAFFLE", fill=(120, 30, 30), font=font("lato-black", 54), anchor="mm")
        person(d, ox + 1050, 700, 300, rng)
    elif key == "pottery":
        d.ellipse((ox + 600, 590, ox + 860, 650), fill=(150, 150, 150))
        d.rounded_rectangle((ox + 680, 510, ox + 780, 610), radius=30, fill=(170, 110, 70))
        person(d, ox + 920, 720, 260, rng, shirt=(110, 110, 120))
        for k in range(5):
            d.rounded_rectangle((ox + 420 + k * 40, 580, ox + 450 + k * 40, 620), radius=8, fill=(100 + k * 20, 140, 160))
    elif key == "bread":
        for k in range(7):
            d.ellipse((ox + 330 + k * 130, 600, ox + 430 + k * 130, 650), fill=(190, 130, 60), outline=(140, 90, 40))
        person(d, ox + 800, 700, 220, rng, shirt=(250, 250, 250))
    elif key == "repair":
        for k in range(3):
            x = ox + 250 + k * 420
            d.rectangle((x, 640, x + 320, 670), fill=(150, 110, 70))
            d.rectangle((x + 40, 590, x + 120, 640), fill=(180, 180, 190))
            d.polygon([(x + 200, 640), (x + 230, 560), (x + 290, 560), (x + 260, 640)], fill=(240, 220, 120))
            person(d, x + 160, 760, 300, rng)
    elif key == "forge":
        d.rectangle((ox + 600, 560, ox + 900, 610), fill=(50, 50, 55))
        d.rectangle((ox + 700, 610, ox + 800, 720), fill=(60, 60, 65))
        d.rectangle((ox + 1000, 400, ox + 1400, 560), fill=(250, 250, 250), outline=(30, 30, 30), width=4)
        d.text((ox + 1200, 440), "Blacksmith demo", fill=(30, 30, 30), font=font("lato-black", 36), anchor="mm")
        d.text((ox + 1200, 500), "CANCELLED - sorry!", fill=(200, 20, 20), font=font("lato-black", 40), anchor="mm")
        d.line((ox + 1200, 560, ox + 1200, 760), fill=(90, 70, 50), width=10)


def scene(key, idx, umbrella, rng):
    ox = rng.randint(-150, 150)
    indoor = key == "repair"
    im = sky_grass(rng, indoor)
    d = ImageDraw.Draw(im)
    if not indoor:
        bunting(d, rng.randint(60, 140), rng, ox)
    if key in ("opening", "choir", "ukulele", "raffle"):
        text = {"opening": "Riverside Makers Fair 2026 - Welcome!", "choir": "St Mary's Primary School Choir",
                "ukulele": "Falmouth Ukulele Club", "raffle": "Grand Raffle Draw"}[key]
        stage(d, ox, text, rng)
        props(d, key, ox, rng)
        crowd(d, rng, 12, 860, 1150)
    elif key == "pottery":
        gazebo(d, ox, "Stall 4 - Penhallow Pottery", (60, 120, 140), rng)
        props(d, key, ox, rng)
        crowd(d, rng, 8, 880, 1150)
    elif key == "bread":
        marquee(d, ox, "Harbour Street Bakery - Bread masterclass")
        props(d, key, ox, rng)
        crowd(d, rng, 9, 870, 1150)
    elif key == "repair":
        banner(d, (ox + 400, 300, ox + 1200, 390), "REPAIR CAFE - Fix-it Falmouth", fill=(255, 255, 255), ink=(30, 110, 60))
        props(d, key, ox, rng)
        crowd(d, rng, 5, 900, 1150)
    elif key == "forge":
        props(d, key, ox, rng)
        crowd(d, rng, 3, 880, 1150)
    elif key == "general":
        if idx == 0:
            d.rectangle((ox + 300, 200, ox + 340, 900), fill=(110, 80, 50))
            d.rectangle((ox + 1260, 200, ox + 1300, 900), fill=(110, 80, 50))
            banner(d, (ox + 280, 200, ox + 1320, 320), "Riverside Makers Fair - Saturday 17 October 2026",
                   fill=(250, 220, 60), ink=(40, 40, 40))
            crowd(d, rng, 14, 820, 1150)
        elif idx == 1:
            for k, col in enumerate([(230, 90, 60), (60, 150, 200)]):
                x = ox + 200 + k * 650
                d.rounded_rectangle((x, 420, x + 560, 760), radius=30, fill=col)
                d.rectangle((x + 60, 480, x + 400, 600), fill=(40, 40, 50))
                d.ellipse((x + 60, 720, x + 160, 820), fill=(30, 30, 30))
                d.ellipse((x + 400, 720, x + 500, 820), fill=(30, 30, 30))
                d.text((x + 280, 650), ["Pasty Shack", "Kernow Coffee"][k], fill="white", font=font("lato-black", 46), anchor="mm")
            crowd(d, rng, 10, 860, 1150)
        else:
            crowd(d, rng, 26, 700, 1150, scale=0.9)
    if umbrella:
        umbrella_family(d, rng.choice([250, 1300]) + ox * 0.3, 1120, 330, rng)
    im = im.filter(ImageFilter.GaussianBlur(0.8))
    return noise(im, rng, 2.5)


# ---------------------------------------------------------------- documents
def programme_page(rng):
    PW, PH = 1240, 1754
    im = Image.new("RGB", (PW, PH), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, PW, 260), fill=(250, 220, 60))
    d.text((PW // 2, 90), "Riverside Makers Fair", fill=(40, 40, 40), font=font("lato-black", 84), anchor="mm")
    d.text((PW // 2, 190), "Saturday 17 October 2026  -  Riverside Park, Falmouth", fill=(40, 40, 40),
           font=font("lato-bold", 38), anchor="mm")
    d.text((110, 330), "Programme", fill=(30, 30, 30), font=font("lato-black", 56))
    y = 440
    rows = []
    for key, t, actual, title, who, place in PROGRAMME:
        d.text((110, y), t, fill=(30, 30, 30), font=font("lato-black", 40))
        d.text((300, y), title, fill=(30, 30, 30), font=font("lato-bold", 40))
        d.text((300, y + 52), f"{who}  -  {place}", fill=(90, 90, 90), font=font("lato", 30))
        rows.append((key, y, t, actual))
        y += 140
    for key, yy, t, actual in rows:
        if actual is None:
            d.line((100, yy + 30, 1000, yy + 50), fill=(20, 30, 140), width=6)
            hand_text(im, (780, yy + 60), "CANCELLED - rain!", 46, rng, fill=(20, 30, 140), bold=True)
        elif actual != t:
            d.line((100, yy + 22, 230, yy + 18), fill=(20, 30, 140), width=6)
            hand_text(im, (90, yy - 50), actual, 48, rng, fill=(20, 30, 140), bold=True)
            hand_text(im, (850, yy + 4), "(moved, oven late)", 36, rng, fill=(20, 30, 140))
    d.text((110, PH - 170), "Free entry. Stalls open 10:00-16:30. Food & drink all day.", fill=(80, 80, 80),
           font=font("lato", 30))
    d.text((110, PH - 120), "Organised by the Riverside Makers Collective - riversidemakers.org.uk", fill=(80, 80, 80),
           font=font("lato", 30))
    return im


def sticky_note(rng):
    im = Image.new("RGB", (1400, 1050), (120, 85, 55))
    d = ImageDraw.Draw(im)
    for i in range(0, 1050, 7):
        d.line((0, i, 1400, i + rng.randint(-3, 3)), fill=(112 + rng.randint(-8, 8), 80, 52), width=3)
    note = Image.new("RGB", (900, 820), (252, 236, 120))
    nd = ImageDraw.Draw(note)
    nd.rectangle((0, 0, 900, 90), fill=(245, 225, 100))
    hand_text(note, (60, 120), "Makers Fair 2027 - BOOKED", 62, rng, fill=(25, 25, 30), bold=True)
    hand_text(note, (60, 240), "Sat 9 Oct", 58, rng, fill=(25, 25, 30))
    nd.line((50, 285, 330, 275), fill=(25, 25, 30), width=6)
    hand_text(note, (360, 240), "-> Sat 16 Oct 2027", 58, rng, fill=(25, 25, 30), bold=True)
    hand_text(note, (60, 360), "Princess Pavilion", 62, rng, fill=(25, 25, 30), bold=True)
    hand_text(note, (60, 450), "(main hall + gardens)", 50, rng, fill=(25, 25, 30))
    hand_text(note, (60, 560), "10am - 4pm, free entry", 56, rng, fill=(25, 25, 30))
    hand_text(note, (520, 700), "- Sam", 52, rng, fill=(25, 25, 30))
    note = note.rotate(4, expand=True, fillcolor=(0, 0, 0))
    mask = Image.new("L", (900, 820), 255).rotate(4, expand=True)
    im.paste(note, (250, 110), mask)
    im = im.filter(ImageFilter.GaussianBlur(1.1))
    return noise(im, rng, 3.0)


def recap_deck(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Riverside Makers Fair 2026"
    s.placeholders[1].text = "Thank you, volunteers! - photo recap, Saturday 17 October 2026"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "See you next year!"
    s.placeholders[1].text = "Riverside Makers Collective - riversidemakers.org.uk"
    for sl in prs.slides:
        for r in sl.shapes.title.text_frame.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0x2B, 0x3A, 0x67)
            r.font.size = Pt(40)
    prs.save(path)


def poster_doc(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))

    def para(text, size, bold=False, color=(0x2B, 0x3A, 0x67)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = RGBColor(*color)
        return p
    para("SAVE THE DATE", 20, True, (0xC0, 0x39, 0x2B))
    para("Riverside Makers Fair 2027", 34, True)
    para("[DATE]  |  [TIME]  |  [VENUE]", 18, True)
    para("Makers' stalls, live music, demos and food - free entry, all welcome!", 13, False, (0x44, 0x44, 0x44))
    para("Highlights from 2026", 16, True)
    t = doc.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for j in range(3):
        t.cell(0, j).text = "[photo]"
        t.cell(1, j).text = "[caption]"
    para("Want a stall? Email stalls@riversidemakers.org.uk", 13, False, (0x44, 0x44, 0x44))
    doc.save(path)


# ---------------------------------------------------------------- build
def build(d: str):
    rng = random.Random(20261017)
    os.makedirs(f"{d}/photos", exist_ok=True)
    photos = []
    order = []
    for key, flags in PHOTOS.items():
        for i, umb in enumerate(flags):
            order.append((key, i, umb))
    rng.shuffle(order)
    for n, (key, i, umb) in enumerate(order, 1):
        im = scene(key, i, umb, rng)
        name = f"IMG-20261017-WA{n + 11:04d}.jpg"
        im.save(f"{d}/photos/{name}", "JPEG", quality=82)
        photos.append({"name": name, "item": key, "umbrella": umb, "md5": md5(f"{d}/photos/{name}")})
    lib_f.image_pdf([lib_f.scan(programme_page(rng), rng, quality=65)], f"{d}/programme_scan.pdf")
    sticky_note(rng).save(f"{d}/committee_note.jpg", "JPEG", quality=85)
    recap_deck(f"{d}/MakersFair_2026_recap.pptx")
    poster_doc(f"{d}/poster_2027.docx")
    with tarfile.open(f"{d}/f3.tar.gz", "w:gz") as tf:
        tf.add(f"{d}/photos", arcname="Pictures/MakersFair_2026")
        for f in ("programme_scan.pdf", "committee_note.jpg", "MakersFair_2026_recap.pptx", "poster_2027.docx"):
            tf.add(f"{d}/{f}", arcname=f"Documents/MakersFair/{f}")
    items = [{"key": k, "printed": t, "time": a, "title": ti, "who": w, "place": p} for k, t, a, ti, w, p in PROGRAMME]
    json.dump({"photos": photos, "programme": items, "poster_acts": POSTER_ACTS,
               "poster": {"date": "2027-10-16", "venue": "Princess Pavilion", "time": "10am-4pm"},
               "deck_md5": md5(f"{d}/MakersFair_2026_recap.pptx"), "poster_md5": md5(f"{d}/poster_2027.docx")},
              open(f"{d}/expected.json", "w"), indent=1)
