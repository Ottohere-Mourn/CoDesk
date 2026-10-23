"""W3 - Bluepeak CRM (sales CRM web app on :8103): merging duplicate accounts vs logging the week's calls.

build(out_dir) writes: app_w3.py, seed.db, call_notes_wk43.docx, pipeline_q4.xlsx, card_ben_alvarez.png,
card_grace_liu.png, expected.json.
"""
from __future__ import annotations

import json
import os
import random
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from PIL import Image, ImageDraw  # noqa: E402

import lib_w  # noqa: E402
from lib_w import font  # noqa: E402

PASSWORD = "Bluepeak-CRM-26"
# id, username, name, title, region, active, left_on
USERS = [
    (1, "morgan.lee", "Morgan Lee", "Sales Operations Manager", "All regions", 1, None),
    (2, "jamie.ortega", "Jamie Ortega", "Account Executive", "US West", 1, None),
    (3, "priya.desai", "Priya Desai", "Account Executive", "US East", 1, None),
    (4, "tom.becker", "Tom Becker", "Account Executive", "DACH", 1, None),
    (5, "aisha.grant", "Aisha Grant", "Account Executive", "Canada", 1, None),
    (6, "liam.walsh", "Liam Walsh", "Account Executive", "DACH", 0, "2026-08-31"),
    (7, "sofia.romero", "Sofia Romero", "Account Executive", "Canada", 0, "2026-09-15"),
    (8, "ken.watanabe", "Ken Watanabe", "Account Executive", "UK & Ireland", 1, None),
    (9, "hannah.price", "Hannah Price", "Inside Sales Representative", "US Central", 1, None),
]
US, CA, UK, DE, AT, CH = "United States", "Canada", "United Kingdom", "Germany", "Austria", "Switzerland"
COLS = ["name", "website", "tax_id", "erp_no", "phone", "industry", "street", "city", "postcode", "country", "owner_id", "created"]

# key: (name, website, tax_id, erp_no, phone, industry, street, city, postcode, country, owner, created)
PAIRS = {  # (master, duplicate) - the master follows the data quality guide
    "P1": (("Trailhead Outfitters", "trailheadoutfitters.com", "82-4410937", "ERP-40112", "", "Outdoor Retail", "1450 W Main St", "Boise", "83702", US, 2, "2019-04-02"),
           ("Trailhead Outfitters, Inc.", "www.trailheadoutfitters.com", "82-4410937", "", "+1 (208) 555-0144", "Outdoor Retail", "1450 West Main Street", "Boise", "83702", US, 9, "2025-08-14")),
    "P2": (("Summit & Pine Supply Co.", "summitandpine.com", "47-2290318", "ERP-40388", "+1 (802) 555-0170", "", "88 Pine St", "Burlington", "05401", US, 3, "2019-09-23"),
           ("Summit and Pine Supply", "", "47-2290318", "", "+1 (802) 555-0171", "Sporting Goods Retail", "88 Pine Street", "Burlington", "05401", US, 3, "2025-11-03")),
    "P3": (("Alpenglow Sport GmbH", "alpenglow-sport.de", "DE284117903", "ERP-50071", "+49 89 5550 2231", "Ski & Snowboard", "Leopoldstr. 61", "München", "80802", DE, 6, "2022-03-15"),
           ("Alpenglow Sports", "https://www.alpenglow-sport.de/shop", "DE 284 117 903", "", "+49 89 5550 2230", "Ski & Snowboard", "Leopoldstraße 61", "Munich", "80802", DE, 4, "2017-06-01")),
    "P4": (("Cascade Paddle Works", "cascadepaddle.com", "91-3376402", "", "+1 (541) 555-0112", "Paddle Sports", "", "", "", US, 2, "2021-06-10"),
           ("Cascade Paddleworks LLC", "cascadepaddle.com", "", "", "+1 (541) 555-0113", "Paddle Sports", "220 River Rd", "Bend", "97701", US, 2, "2023-02-01")),
    "P5": (("Northern Lights Camping Ltd", "northernlightscamping.ca", "819245637RT0001", "ERP-60219", "+1 (604) 555-0150", "Camping", "2100 Main St", "Vancouver", "V5T 3C9", CA, 7, "2020-05-12"),
           ("Northern Lights Camping", "https://northernlightscamping.ca", "", "", "+1 (604) 555-0151", "Camping", "2100 Main Street", "Vancouver", "V5T 3C9", CA, 5, "2025-09-30")),
    "P6": (("Redrock Climbing Co.", "redrockclimbing.com", "86-1029384", "", "", "Climbing & Mountaineering", "77 Canyon Rd", "Moab", "84532", US, 2, "2020-03-02"),
           ("Red Rock Climbing Company", "redrockclimbingco.com", "86-1029384", "", "+1 (435) 555-0166", "Climbing & Mountaineering", "77 Canyon Road", "Moab", "84532", US, 9, "2024-05-20")),
    "P7": (("Harborview Marine", "harborviewmarine.com", "", "ERP-40977", "+1 (207) 555-0133", "", "5 Commercial St", "Portland", "04101", US, 3, "2024-01-18"),
           ("Harbor View Marine Supply", "www.harborviewmarine.com", "01-0563318", "", "+1 (207) 555-0134", "Marine Supply", "5 Commercial Street", "Portland", "04101", US, 3, "2018-07-09")),
}
LOOKALIKES = {  # similar names, different companies: must stay separate
    "L1": (("Summit Gear Co.", "summitgear.com", "84-6612390", "ERP-40551", "+1 (303) 555-0190", "Outdoor Retail", "1800 Pearl St", "Boulder", "80302", US, 9, "2018-02-11"),
           ("Summit Gear Ltd", "summitgear.co.uk", "GB 291 4471 06", "ERP-70114", "+44 20 7946 0712", "Outdoor Retail", "14 Long Acre", "London", "WC2E 9LH", UK, 8, "2021-10-04")),
    "L2": (("Alpine Edge Outfitters", "alpineedge.com", "84-3019276", "", "+1 (720) 555-0145", "Ski & Snowboard", "910 Broadway", "Denver", "80203", US, 9, "2020-11-30"),
           ("Alpine Edge Sports", "alpineedgesports.ca", "703118254RT0001", "ERP-60433", "+1 (403) 555-0181", "Ski & Snowboard", "333 17 Ave SW", "Calgary", "T2S 0A8", CA, 5, "2022-04-21")),
    "L3": (("Blue Ridge Trail Supply", "blueridgetrailsupply.com", "54-2210987", "ERP-40220", "+1 (540) 555-0108", "Outdoor Retail", "310 Market St", "Roanoke", "24011", US, 3, "2019-01-15"),
           ("Blue Ridge Trail Supply", "brtrailsupply.com", "58-7716023", "", "+1 (706) 555-0129", "Camping", "45 E Main St", "Blue Ridge", "30513", US, 3, "2024-09-02")),
}
B_ACCOUNTS = {  # task B's ordinary accounts
    "GRANITE": ("Granite Peak Sports", "granitepeaksports.com", "84-2205511", "ERP-40610", "+1 (303) 555-0180", "Outdoor Retail", "1120 Washington Ave", "Golden", "80401", US, 9, "2020-08-17"),
    "COASTAL": ("Coastal Kayak Centers", "coastalkayakcenters.com", "06-1188240", "ERP-40733", "+1 (860) 555-0100", "Paddle Sports", "12 Water St", "Stonington", "06378", US, 3, "2021-04-06"),
    "PINECONE": ("Pinecone Family Camping", "pineconecamping.com", "38-3390127", "", "+1 (231) 555-0119", "Camping", "402 Front St", "Traverse City", "49684", US, 9, "2023-05-25"),
    "BLUEWATER": ("Bluewater Expeditions", "bluewaterexpeditions.com", "91-2276654", "ERP-40852", "+1 (206) 555-0160", "Rental & Guiding", "Pier 66", "Seattle", "98121", US, 2, "2019-12-03"),
}
# other customers: name, city, postcode, country, industry
OTHERS = [
    ("Ridgeline Sports", "Missoula", "59801", US, "Outdoor Retail"), ("Basecamp Supply Co.", "Flagstaff", "86001", US, "Camping"),
    ("Timberline Bike & Ski", "Hood River", "97031", US, "Cycling"), ("Big Sky Outfitters", "Bozeman", "59715", US, "Outdoor Retail"),
    ("Lakeshore Paddle Co.", "Duluth", "55802", US, "Paddle Sports"), ("Mesa Trail Running", "Santa Fe", "87501", US, "Sporting Goods Retail"),
    ("Evergreen Camp Store", "Asheville", "28801", US, "Camping"), ("Sierra Summit Rentals", "Bishop", "93514", US, "Rental & Guiding"),
    ("Prairie Wind Cycles", "Omaha", "68102", US, "Cycling"), ("Cape Cod Surf & Paddle", "Hyannis", "02601", US, "Paddle Sports"),
    ("Ozark Trail Outfitters", "Fayetteville", "72701", US, "Outdoor Retail"), ("Wildflower Climbing Gym", "Austin", "78702", US, "Climbing & Mountaineering"),
    ("Great Lakes Outdoor", "Madison", "53703", US, "Outdoor Retail"), ("Desert Rose Adventure", "Tucson", "85701", US, "Rental & Guiding"),
    ("Moose Lake Marine", "Ely", "55731", US, "Marine Supply"), ("Tahoe Board Shop", "Truckee", "96161", US, "Ski & Snowboard"),
    ("Peak Performance Ski", "Park City", "84060", US, "Ski & Snowboard"), ("Gulf Coast Kayak", "Pensacola", "32502", US, "Paddle Sports"),
    ("Appalachian Gear Exchange", "Boone", "28607", US, "Online Retail"), ("Northwest Trail Co.", "Bellingham", "98225", US, "Outdoor Retail"),
    ("Granite State Outdoors", "Concord", "03301", US, "Outdoor Retail"), ("Coastal Outfitters", "Charleston", "29401", US, "Outdoor Retail"),
    ("Pine Cone Hardware & Camping", "Bend", "97702", US, "Camping"), ("Blue Water Dive Center", "Key West", "33040", US, "Marine Supply"),
    ("Rocky Mountain Rentals", "Banff", "T1L 1A1", CA, "Rental & Guiding"), ("Maritime Paddle Supply", "Halifax", "B3J 1S9", CA, "Paddle Sports"),
    ("Laurentian Ski Shop", "Mont-Tremblant", "J8E 1T1", CA, "Ski & Snowboard"), ("Muskoka Trail Supply", "Huntsville", "P1H 1A7", CA, "Camping"),
    ("Peak District Outdoor", "Sheffield", "S1 2HE", UK, "Outdoor Retail"), ("Highland Trek Supplies", "Fort William", "PH33 6AN", UK, "Outdoor Retail"),
    ("Lakeland Climbing Centre", "Kendal", "LA9 4PU", UK, "Climbing & Mountaineering"), ("Bergsport Huber", "Garmisch-Partenkirchen", "82467", DE, "Climbing & Mountaineering"),
    ("Alpinhaus Innsbruck", "Innsbruck", "6020", AT, "Ski & Snowboard"), ("Gipfelwerk AG", "Zürich", "8001", CH, "Climbing & Mountaineering"),
    ("Nordwand Outdoor GmbH", "Berlin", "10115", DE, "Online Retail"), ("Tiroler Bergladen", "Kitzbühel", "6370", AT, "Ski & Snowboard"),
]
OWNER_BY_COUNTRY = {US: [2, 3, 9, 2, 3], CA: [5, 5, 7], UK: [8], DE: [4, 4, 6], AT: [4], CH: [4]}

# named contacts: account key -> [(first, last, title, email, phone)]
NAMED_CONTACTS = {
    "P1.dup": [("Dana", "Fox", "Purchasing Manager", "dana.fox@trailheadoutfitters.com", "+1 (208) 555-0147")],
    "P1.master": [("Greg", "Holt", "Owner", "greg@trailheadoutfitters.com", "+1 (208) 555-0140")],
    "P5.dup": [("Marc", "Tremblay", "Store Manager", "marc@northernlightscamping.ca", "+1 (604) 555-0152")],
    "P5.master": [("Helen", "Park", "Controller", "helen.park@northernlightscamping.ca", "+1 (604) 555-0155")],
    "P3.dup": [("Katrin", "Vogel", "Head Buyer", "k.vogel@alpenglow-sport.de", "+49 89 5550 2240")],
    "P3.master": [("Stefan", "Maier", "Geschäftsführer", "s.maier@alpenglow-sport.de", "+49 89 5550 2232")],
    "L1.b": [("Oliver", "Hughes", "Buying Director", "oliver.hughes@summitgear.co.uk", "+44 20 7946 0715")],
    "L1.a": [("Oliver", "Hayes", "Store Manager", "oliver@summitgear.com", "+1 (303) 555-0193")],
    "GRANITE": [("Rachel", "Kim", "Owner", "rachel@granitepeaksports.com", "+1 (303) 555-0181")],
    "COASTAL": [("Ian", "Murphy", "General Manager", "ian.murphy@coastalkayakcenters.com", "+1 (860) 555-0101")],
    "PINECONE": [("Lucy", "Chen", "Owner", "lucy@pineconecamping.com", "+1 (231) 555-0120")],
    "BLUEWATER": [("Ahmed", "Saleh", "Founder", "ahmed@bluewaterexpeditions.com", "+1 (206) 555-0161")],
}
FIRST = ["Alex", "Jordan", "Casey", "Taylor", "Morgan", "Riley", "Sam", "Jesse", "Chris", "Dana", "Erin", "Kelly", "Pat",
         "Robin", "Drew", "Avery", "Quinn", "Lena", "Mark", "Nina", "Oscar", "Paula", "Victor", "Wendy", "Yusuf", "Zoe"]
LAST = ["Anders", "Brooks", "Carter", "Dalton", "Ellis", "Fischer", "Garcia", "Hale", "Ivers", "Jensen", "Kowalski",
        "Lund", "Moreau", "Nakamura", "Olsen", "Pereira", "Quist", "Russo", "Stein", "Turner", "Uddin", "Vance", "Weber"]
TITLES = ["Buyer", "Store Manager", "Owner", "Purchasing Manager", "Accounts Payable", "Marketing Manager", "Operations Manager"]
PRODUCTS = ["fall restock", "winter jackets", "tents 2027", "climbing hardware", "paddle accessories", "ski bindings",
            "base layers", "trail shoes", "camp kitchen", "backpacks", "rental fleet refresh", "kids range"]
SUBJECTS = ["Quarterly check-in", "Order status", "Catalog walkthrough", "Pricing question", "Delivery delay follow-up",
            "Trade show follow-up", "Payment terms", "New season preview"]

# ------------------------------------------------------------------ task B: what the notes ask for
CALLS = [  # account key, date, contact (first last), keyword regex for the summary
    ("P1.dup", "2026-10-19", "Dana Fox", r"tent|restock|winter"),
    ("GRANITE", "2026-10-19", "Rachel Kim", r"shoe|footwear|trail|ben"),
    ("P5.dup", "2026-10-20", "Marc Tremblay", r"q1|january|purchasing|promot|follow"),
    ("COASTAL", "2026-10-20", "Ian Murphy", r"office|move|deliver|address|fall"),
    ("P3.dup", "2026-10-21", "Katrin Vogel", r"ski|touring|pack|launch"),
    ("L1.b", "2026-10-21", "Oliver Hughes", r"rain|shell|spring|negotiat"),
    ("PINECONE", "2026-10-22", "Lucy Chen", r"tent|bundle|signed|holiday|order"),
    ("BLUEWATER", "2026-10-22", "Ahmed Saleh", r"grace|operations|order|quarter|hand"),
]
DEALS = [  # account key, name, amount, stage, close, regex identifying the opportunity name
    ("P1.dup", "Trailhead – Winter 2027 tents restock", 48500, "Proposal", "2026-12-15", r"winter.*tent|tent.*restock"),
    ("GRANITE", "Granite Peak – Trail running shoes 2027", 26750, "Proposal", "2026-12-01", r"trail\s*-?\s*running|running\s+shoe"),
    ("P3.dup", "Alpenglow – Ski touring packs EU launch", 72000, "Qualification", "2027-01-31", r"ski\s*-?\s*touring|touring\s+pack"),
    ("L1.b", "Summit Gear UK – Spring 2027 rain shells", 31200, "Negotiation", "2026-11-28", r"rain\s*-?\s*shell"),
    ("PINECONE", "Pinecone – Holiday family tent bundle", 18400, "Closed Won", "2026-10-22", r"holiday|family\s+tent"),
]
CONTACT_UPDATES = {"Marc Tremblay": {"title": r"head\s+of\s+purchasing", "phone": "6045550199"},
                   "Ian Murphy": {"phone": "8605550123"}}
NEW_CONTACTS = {"GRANITE": ("Ben", "Alvarez", "ben.alvarez@granitepeaksports.com", "3035550187"),
                "BLUEWATER": ("Grace", "Liu", "grace.liu@bluewaterexpeditions.com", "2065550164")}
COASTAL_ADDRESS = {"street": r"41\s+harbou?r\s+r(oa)?d", "city": r"^mystic$", "postcode": r"^0?6355$"}

NOTES = [
    ("Mon 19 Oct - Trailhead Outfitters, Inc. - call with Dana Fox (Purchasing Manager)",
     ["Strong October sell-through on 3- and 4-season tents; Dana wants to restock winter tents for the 2027 season. "
      "I promised a proposal by Friday.",
      "Deal: \"Trailhead – Winter 2027 tents restock\", $48,500, stage Proposal, expected close 15 Dec 2026."]),
    ("Mon 19 Oct - Granite Peak Sports - call with Rachel Kim (Owner)",
     ["Rachel introduced their new footwear buyer, Ben Alvarez - his business card is in my Pictures folder "
      "(card_ben_alvarez.png). Please add him as a contact on the account.",
      "They want our trail running shoes for 2027. Deal: \"Granite Peak – Trail running shoes 2027\", $26,750, stage "
      "Proposal, expected close 1 Dec 2026."]),
    ("Tue 20 Oct - Northern Lights Camping - call with Marc Tremblay",
     ["Marc has been promoted to Head of Purchasing - update his title. New direct line: +1 (604) 555-0199.",
      "No new order before Q1; follow up in January. No opportunity for now."]),
    ("Tue 20 Oct - Coastal Kayak Centers - call with Ian Murphy (General Manager)",
     ["They moved their head office: new address 41 Harbor Road, Mystic, CT 06355 - please update the account address. "
      "Ian's new mobile number is +1 (860) 555-0123.",
      "Happy with the fall deliveries; nothing new to order right now."]),
    ("Wed 21 Oct - Alpenglow Sports - call with Katrin Vogel (Head Buyer)",
     ["Katrin is planning the EU launch of ski touring backpacks for next winter and wants a first quote.",
      "Deal: \"Alpenglow – Ski touring packs EU launch\", $72,000, stage Qualification, expected close 31 Jan 2027."]),
    ("Wed 21 Oct - Summit Gear Ltd (London) - call with Oliver Hughes (Buying Director)",
     ["Negotiating volume pricing for spring rain shells; Oliver expects to sign by end of November.",
      "Deal: \"Summit Gear UK – Spring 2027 rain shells\", $31,200, stage Negotiation, expected close 28 Nov 2026."]),
    ("Thu 22 Oct - Pinecone Family Camping - call with Lucy Chen (Owner)",
     ["Lucy signed the order for the holiday family tent bundle on the call - it's won.",
      "Deal: \"Pinecone – Holiday family tent bundle\", $18,400, stage Closed Won, close date 22 Oct 2026."]),
    ("Thu 22 Oct - Bluewater Expeditions - call with Ahmed Saleh (Founder)",
     ["Ahmed has handed day-to-day ordering over to Grace Liu, their Operations Manager - business card in my Pictures "
      "folder (card_grace_liu.png). Please add her as a contact.",
      "No order this quarter."]),
]


GUIDE = """Version 2, September 2026. Owner: Sales Operations (Morgan Lee). Applies to every account in Bluepeak CRM.

## 1. When are two accounts duplicates?
Two accounts are the same customer when they have the same tax / VAT ID, or the same website domain (ignore http://, https://, "www." and anything after the domain). Spelling, punctuation and legal suffixes of the name (Inc., LLC, Ltd, GmbH, Co., "&" vs "and") often differ between duplicates.
Similar or even identical names are NOT enough: shops with similar names in different cities or countries, with a different domain and tax ID, are different customers. Never merge them.

## 2. Which record survives (the master)
- The record that has an ERP number (synced from SAP) is the master.
- If neither record has an ERP number, the older record (earlier Created date) is the master.

## 3. Field values on the master
- Keep the master's values.
- Where the master's field is empty, take the duplicate's value (typically phone, industry, street, city, postcode).

## 4. Owner
Keep the master's owner, unless that user is inactive (see Users). Then the duplicate's owner takes over the account.

## 5. How to merge
Use Merge on the account page (never delete or rename accounts by hand). Contacts, opportunities and activities move to the master automatically; the duplicate becomes a read-only record that points to the master.

## 6. Merge log
Record every merge in your merge log: the master's account ID and the merged account's ID."""


def business_card(path, seed, company, color, name, title, email, phone, address, web):
    im = Image.new("RGB", (1100, 700), (118, 104, 88))  # desk
    card = Image.new("RGB", (880, 500), (252, 251, 247))
    d = ImageDraw.Draw(card)
    d.rectangle((0, 0, 880, 16), fill=color)
    d.polygon([(48, 150), (100, 62), (126, 104), (146, 80), (196, 150)], fill=color)
    d.text((216, 82), company, fill=color, font=font(34, bold=True))
    d.text((48, 210), name, fill=(30, 30, 30), font=font(40, bold=True))
    d.text((48, 264), title, fill=(90, 90, 90), font=font(26))
    d.line((48, 315, 830, 315), fill=(200, 200, 200), width=2)
    for i, line in enumerate([email, phone, address, web]):
        d.text((48, 335 + i * 36), line, fill=(40, 40, 40), font=font(25))
    card = card.rotate(-3, expand=True, fillcolor=(118, 104, 88), resample=Image.BICUBIC)
    im.paste(card, (90, 80))
    lib_w.photo_noise(im, seed, 9)
    im.save(path)


def build(out):
    os.makedirs(out, exist_ok=True)
    rng = random.Random(20261103)
    lib_w.copy_app(HERE, out, "app_w3.py")
    con, _ = lib_w.fresh_db(HERE, "app_w3.py", f"{out}/seed.db")
    for u in USERS:
        con.execute("INSERT INTO users (id,username,password,name,email,title,region,active,left_on,can_login) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (u[0], u[1], PASSWORD if u[0] == 1 else "", u[2], f"{u[1]}@bluepeak-outdoor.com", u[3], u[4], u[5], u[6], int(u[0] == 1)))
    # ---- all accounts with their keys, ids assigned in a shuffled order so pairs are not neighbours by id
    accts = []
    for k, (m, d) in PAIRS.items():
        accts += [(f"{k}.master", m), (f"{k}.dup", d)]
    for k, (a, b) in LOOKALIKES.items():
        accts += [(f"{k}.a", a), (f"{k}.b", b)]
    accts += list(B_ACCOUNTS.items())
    used_tax = set()
    for i, (name, city, pc, country, ind) in enumerate(OTHERS):
        slug = "".join(ch for ch in name.lower().replace("&", "and") if ch.isalnum())[:24]
        tld = {US: "com", CA: "ca", UK: "co.uk", DE: "de", AT: "at", CH: "ch"}[country]
        while True:
            tax = {US: f"{rng.randint(10, 99)}-{rng.randint(1000000, 9999999)}", CA: f"{rng.randint(100000000, 999999999)}RT0001",
                   UK: f"GB {rng.randint(100, 999)} {rng.randint(1000, 9999)} {rng.randint(10, 99)}", DE: f"DE{rng.randint(100000000, 999999999)}",
                   AT: f"ATU{rng.randint(10000000, 99999999)}", CH: f"CHE-{rng.randint(100, 999)}.{rng.randint(100, 999)}.{rng.randint(100, 999)}"}[country]
            if tax not in used_tax:
                used_tax.add(tax)
                break
        cc = {US: "+1 ({}) 555-01{:02d}", CA: "+1 ({}) 555-01{:02d}"}.get(country)
        phone = (cc.format(rng.choice([406, 928, 541, 218, 505, 828, 760, 402, 508, 479, 512, 608, 520, 530, 435, 850, 828, 360, 603, 843, 541, 305]),
                           rng.randint(10, 99)) if cc else
                 {UK: f"+44 114 496 {rng.randint(1000, 9999)}", DE: f"+49 8821 {rng.randint(10000, 99999)}",
                  AT: f"+43 512 {rng.randint(100000, 999999)}", CH: f"+41 44 {rng.randint(100, 999)} {rng.randint(10, 99)} {rng.randint(10, 99)}"}[country])
        erp = f"ERP-{rng.randint(41000, 69999)}" if rng.random() < 0.7 else ""
        street = f"{rng.randint(2, 2900)} {rng.choice(['Main St', 'Market St', 'Lake Ave', 'Mountain Rd', 'High St', 'Bahnhofstr.', 'River Rd', 'Elm St'])}"
        year = rng.randint(2016, 2025)
        accts.append((f"O{i}", (name, f"{slug}.{tld}", tax, erp, phone, ind, street, city, pc, country,
                                rng.choice(OWNER_BY_COUNTRY[country]), f"{year}-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}")))
    order = list(range(len(accts)))
    rng.shuffle(order)
    ids = {}
    for pos, idx in enumerate(order):
        key, row = accts[idx]
        aid = 1001 + pos
        ids[key] = aid
        con.execute("INSERT INTO accounts (id,name,website,tax_id,erp_no,phone,industry,street,city,postcode,country,owner_id,created,updated,status) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,'active')", (aid, *row, row[11] + " 09:00:00"))
    rows = dict(accts)
    # ---- contacts
    contact_ids, used_email = {}, set()
    for key, (name, web, *_rest) in accts:
        aid = ids[key]
        dom = (web or "").replace("https://", "").replace("http://", "").replace("www.", "").split("/")[0] or "gmail.com"
        named = NAMED_CONTACTS.get(key, [])
        for c in named:
            cur = con.execute("INSERT INTO contacts (account_id,first,last,title,email,phone,created,updated) VALUES (?,?,?,?,?,?,?,?)",
                              (aid, *c, "2025-01-10 10:00:00", "2025-01-10 10:00:00"))
            contact_ids[f"{c[0]} {c[1]}"] = cur.lastrowid
        for _ in range(rng.randint(1, 3) - (1 if named else 0) + 1):
            while True:
                f, l = rng.choice(FIRST), rng.choice(LAST)
                email = f"{f.lower()}.{l.lower()}@{dom}"
                if email not in used_email and f"{f} {l}" not in contact_ids:
                    used_email.add(email)
                    break
            ph = rows[key][4][:-2] + f"{rng.randint(10, 99)}" if rows[key][4] else ""
            con.execute("INSERT INTO contacts (account_id,first,last,title,email,phone,created,updated) VALUES (?,?,?,?,?,?,?,?)",
                        (aid, f, l, rng.choice(TITLES), email, ph, "2025-03-02 10:00:00", "2025-03-02 10:00:00"))
    # ---- opportunities (OPP-2001...) and activities
    keys = [k for k, _ in accts]
    pipeline_rows = []
    for i in range(80):
        key = rng.choice(keys)
        stage = rng.choice(["Qualification", "Proposal", "Negotiation", "Closed Won", "Closed Won", "Closed Lost"])
        close = f"2026-{rng.randint(7, 12):02d}-{rng.randint(1, 28):02d}" if stage.startswith("Closed") else \
            f"{rng.choice(['2026-11', '2026-12', '2027-01', '2027-02'])}-{rng.randint(1, 28):02d}"
        short = rows[key][0].split()[0].rstrip(",")
        oname = f"{short} – {rng.choice(PRODUCTS)}"
        amount = float(rng.randrange(4000, 95000, 250))
        con.execute("INSERT INTO opps (id,account_id,name,amount,stage,close_date,owner_id,created,updated) VALUES (?,?,?,?,?,?,?,?,?)",
                    (2001 + i, ids[key], oname, amount, stage, close, rows[key][10], "2026-08-01 10:00:00", "2026-09-15 10:00:00"))
        if not stage.startswith("Closed") and len(pipeline_rows) < 8 and rows[key][10] in (2, 3, 9, 5):
            pipeline_rows.append((rows[key][0], oname, amount, stage, close, next(u[2] for u in USERS if u[0] == rows[key][10]), f"OPP-{2001 + i}"))
    for i in range(140):
        key = rng.choice(keys)
        aid = ids[key]
        cs = con.execute("SELECT id FROM contacts WHERE account_id=?", (aid,)).fetchall()
        typ = rng.choice(["Call", "Call", "Email", "Meeting", "Note"])
        date = f"2026-{rng.randint(7, 10):02d}-{rng.randint(1, 16):02d}"
        con.execute("INSERT INTO activities (account_id,contact_id,type,date,subject,notes,user_id,created) VALUES (?,?,?,?,?,?,?,?)",
                    (aid, rng.choice(cs)[0] if cs and typ != "Note" else None, typ, date, rng.choice(SUBJECTS),
                     rng.choice(["Discussed next season's range.", "Sent updated price list.", "Customer asked about lead times.",
                                 "Agreed to follow up after their inventory count.", "Shared marketing materials."]),
                     rng.choice([2, 3, 4, 5, 8, 9]), date + " 15:00:00"))
    con.execute("INSERT INTO pages VALUES (?,?,?,?)", ("data-quality", "Data quality guide: duplicate accounts", GUIDE, "2026-09-08"))
    con.commit()
    maxes = {t: con.execute(f"SELECT COALESCE(MAX(id),0) FROM {t}").fetchone()[0] for t in ("contacts", "opps", "activities", "merges")}
    con.close()
    # ---- expected final masters per the guide (computed, not hand-written)
    pairs = {}
    for k, (m, d) in PAIRS.items():
        mm, dd = dict(zip(COLS, m)), dict(zip(COLS, d))
        final = {c: (mm[c] or dd[c]) for c in COLS if c not in ("owner_id", "created")}
        active = {u[0]: u[5] for u in USERS}
        final["owner_id"] = mm["owner_id"] if active[mm["owner_id"]] else dd["owner_id"]
        pairs[k] = {"master": ids[f"{k}.master"], "dup": ids[f"{k}.dup"], "final": final,
                    "names": [mm["name"], dd["name"]]}
    # ---- documents for task B
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Liberation Sans"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Call notes - week 43 (19-23 Oct 2026)", level=1)
    doc.add_paragraph("Morgan Lee - to be entered in the CRM (activities, opportunities, contact changes).")
    for head, paras in NOTES:
        doc.add_heading(head, level=2)
        for p in paras:
            doc.add_paragraph(p)
    doc.save(f"{out}/call_notes_wk43.docx")
    business_card(f"{out}/card_ben_alvarez.png", 31, "GRANITE PEAK SPORTS", (120, 62, 30), "Ben Alvarez", "Footwear Buyer",
                  "ben.alvarez@granitepeaksports.com", "+1 (303) 555-0187", "1120 Washington Ave, Golden, CO 80401", "granitepeaksports.com")
    business_card(f"{out}/card_grace_liu.png", 32, "BLUEWATER EXPEDITIONS", (20, 92, 140), "Grace Liu", "Operations Manager",
                  "grace.liu@bluewaterexpeditions.com", "+1 (206) 555-0164", "Pier 66, Seattle, WA 98121", "bluewaterexpeditions.com")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Q4 pipeline"
    ws.append(["Account", "Opportunity", "Amount (USD)", "Stage", "Expected close", "Owner", "CRM ID"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9EAD3")
    for r in pipeline_rows:
        ws.append(list(r))
    owner_of = {"P1.dup": "Hannah Price", "GRANITE": "Hannah Price", "P3.dup": "Tom Becker", "L1.b": "Ken Watanabe", "PINECONE": "Hannah Price"}
    for key, name, amount, stage, close, _ in DEALS:
        ws.append([rows[key][0], name, amount, stage, close, owner_of[key], None])
    for col, w in zip("ABCDEFG", (30, 44, 14, 14, 15, 16, 12)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "#,##0"
    wb.save(f"{out}/pipeline_q4.xlsx")
    json.dump({
        "port": 8103, "user": "morgan.lee", "password": PASSWORD, "morgan_id": 1, "seed_max": maxes,
        "pairs": pairs, "lookalikes": {k: [ids[f"{k}.a"], ids[f"{k}.b"]] for k in LOOKALIKES},
        "calls": [{"account": ids[k], "name": rows[k][0], "date": dte, "contact": contact_ids[c], "contact_name": c, "kw": kw}
                  for k, dte, c, kw in CALLS],
        "deals": [{"account": ids[k], "name": n, "amount": a, "stage": s, "close": cl, "rx": rx} for k, n, a, s, cl, rx in DEALS],
        "contact_updates": {str(contact_ids[n]): dict(v, name=n) for n, v in CONTACT_UPDATES.items()},
        "new_contacts": [{"account": ids[k], "first": f, "last": l, "email": e, "phone": p} for k, (f, l, e, p) in NEW_CONTACTS.items()],
        "coastal": {"account": ids["COASTAL"], "fields": COASTAL_ADDRESS},
        "pipeline_seed_ids": [r[6] for r in pipeline_rows],
    }, open(f"{out}/expected.json", "w"), indent=1, ensure_ascii=False)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/w3_out")
