"""C2 - the Downloads folder: tidying it up vs collecting the receipts of one trip for an expense claim."""
from __future__ import annotations

import os

from evaluation_examples.task_class.pg_common import HOME, expected, fetch, result, sh, tree, upload

FORM = f"{HOME}/Documents/claims/expense_form.xlsx"


def setup(sc):
    exp = expected("c2")
    upload(sc, [(f"c2/Downloads/{f['name']}", f"{HOME}/Downloads/{f['name']}") for f in exp["files"]]
           + [("c2/claims/expense_form.xlsx", FORM)])
    sh(sc, "mkdir -p ~/Documents ~/Pictures")
    sc.launch(["nautilus", f"{HOME}/Downloads"])


def evaluate_tidy(env, dest: str) -> dict:
    exp = expected("c2")
    t = tree(env, ["~/Downloads", "~/Documents/Archive", "~/Pictures/Downloads"])
    leftover = [p for p in t if p.startswith("Downloads/") and not os.path.basename(p).startswith(".")]
    ok = 0
    for f in exp["files"]:
        if f["target"]:
            ok += t.get(f"{f['target']}/{f['name']}") == f["md5"]
        else:  # installers and duplicate copies must be gone from Downloads
            ok += f"Downloads/{f['name']}" not in t
    n = len(exp["files"])
    checks = {"files_in_place": f"{ok}/{n}", "downloads_empty": not leftover, "leftover": leftover[:8]}
    return result((ok / n) * 0.9 + 0.1 * (not leftover), checks)


def evaluate_claim(env, dest: str) -> dict:
    exp = expected("c2")
    want = sorted(r["amount"] for r in exp["receipts"].values())
    path = fetch(env, FORM, dest + ".xlsx")
    amounts, total = [], None
    if path:
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).worksheets[0]
        for row in range(6, 11):
            v = ws.cell(row, 4).value
            if isinstance(v, (int, float)):
                amounts.append(round(float(v), 2))
        tv = ws.cell(11, 4).value
        total = float(tv) if isinstance(tv, (int, float)) else None
    lines_ok = sorted(amounts) == want
    total_ok = total is not None and abs(total - exp["total"]) < 0.01
    t = tree(env, ["~/Documents/claims/2026-09-trip"])
    got = sorted(t.values())
    folder_ok = got == sorted(r["md5"] for r in exp["receipts"].values())
    checks = {"form_found": bool(path), "amounts": amounts, "total": total, "receipt_copies": len(got),
              "receipt_folder_exact": folder_ok}
    return result((lines_ok + total_ok + folder_ok) / 3, checks)
