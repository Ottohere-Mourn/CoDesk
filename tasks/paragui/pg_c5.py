"""C5 - one slide deck: adding the results slides vs applying the new brand guide to every slide."""
from __future__ import annotations

import io
import re

from evaluation_examples.task_class.pg_common import HOME, expected, fetch, result, upload

DECK = f"{HOME}/Documents/Q3_model_review.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def setup(sc):
    upload(sc, [("c5/Q3_model_review.pptx", DECK), ("c5/results.xlsx", f"{HOME}/Documents/results.xlsx"),
                ("c5/brand_guide.txt", f"{HOME}/Documents/brand_guide.txt"),
                ("c5/brand/northwind_logo_2026.png", f"{HOME}/Pictures/brand/northwind_logo_2026.png")])
    sc._open_setup(DECK)


def _text(slide) -> str:
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def evaluate_results(env, dest: str) -> dict:
    path = fetch(env, DECK, dest + ".pptx")
    if not path:
        return result(0.0, {"file": "missing"})
    from pptx import Presentation
    exp = expected("c5")
    slides = list(Presentation(path).slides)
    new = slides[exp["slides"]:]
    table = chart = concl = False
    for s in new:
        for sh in s.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                cells = " ".join(c.text for row in sh.table.rows for c in row.cells)
                models = sum(m.lower() in cells.lower() for m, _, _ in exp["results"])
                accs = sum(bool(re.search(rf"{a:.3f}|{a * 100:.1f}|{a:.2f}", cells)) for _, a, _ in exp["results"])
                table |= models == 4 and accs == 4
            chart |= bool(getattr(sh, "has_chart", False) and sh.has_chart)
        # the title may be a plain text box, and "ConvNeXt-T" may be written with another dash or as "ConvNeXt Tiny"
        flat = re.sub(r"[^a-z0-9]", "", _text(s).lower())
        concl |= "conclusion" in flat and re.sub(r"[^a-z0-9]", "", exp["best"].lower()) in flat
    checks = {"slides": len(slides), "results_table": table, "accuracy_chart": chart, "conclusions": concl,
              "new_slide_text": [_text(s)[:160] for s in new]}
    return result((table + chart + concl) / 3 if len(new) >= 3 else (table + chart + concl) / 6, checks)


def _attrs(el, acc):
    """Fill font attributes (bold, size, colour, typeface) from an a:rPr / a:defRPr element if not set yet."""
    if el is None:
        return
    if "b" not in acc and el.get("b") is not None:
        acc["b"] = el.get("b") in ("1", "true")
    if "sz" not in acc and el.get("sz") is not None:
        acc["sz"] = int(el.get("sz")) / 100
    clr = el.find("a:solidFill/a:srgbClr", NS)
    if "color" not in acc and clr is not None:
        acc["color"] = clr.get("val").upper()
    latin = el.find("a:latin", NS)
    if "font" not in acc and latin is not None and not latin.get("typeface", "").startswith("+"):
        acc["font"] = latin.get("typeface")


def _title_shape(slide):
    """The title placeholder, or -- for slides built with a plain text box as title -- the topmost text shape."""
    shape = slide.shapes.title
    if shape is not None and shape.text.strip():
        return shape
    boxes = [s for s in slide.shapes if s.has_text_frame and s.text.strip() and s.top is not None]
    return min(boxes, key=lambda s: s.top) if boxes else None


def _title_runs_ok(slide, master_title_style) -> tuple[bool, dict]:
    shape = _title_shape(slide)
    if shape is None:
        return False, {}
    is_ph = shape.is_placeholder
    layout_ph = next((p for p in slide.slide_layout.placeholders
                      if is_ph and p.placeholder_format.type == shape.placeholder_format.type), None)
    chain = [shape._element.find(".//a:lstStyle/a:lvl1pPr/a:defRPr", NS)]
    if layout_ph is not None:
        chain += [layout_ph._element.find(".//a:lstStyle/a:lvl1pPr/a:defRPr", NS), layout_ph._element.find(".//a:rPr", NS)]
    if is_ph:  # a text box does not inherit the master's title style
        chain.append(master_title_style)
    seen = {}
    ok = True
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if not run.text.strip():
                continue
            acc = {}
            _attrs(run._r.find("a:rPr", NS), acc)
            for el in chain:
                _attrs(el, acc)
            seen = acc
            col = acc.get("color", "000000")
            close = all(abs(int(col[i:i + 2], 16) - int("1F3864"[i:i + 2], 16)) <= 24 for i in (0, 2, 4))
            ok &= acc.get("font") == "Liberation Sans" and acc.get("b") is True and abs(acc.get("sz", 0) - 32) < 0.6 and close
    return ok, seen


def _logo_kind(image_blob) -> str:
    from PIL import Image
    im = Image.open(io.BytesIO(image_blob)).convert("RGBA").resize((40, 40))
    px = [p for p in im.getdata() if p[3] > 128 and sum(p[:3]) < 650]
    if not px:
        return "other"
    r, g, b = (sum(p[i] for p in px) / len(px) for i in range(3))
    return "old" if r > b + 60 else ("new" if b > r + 30 else "other")


def evaluate_brand(env, dest: str) -> dict:
    path = fetch(env, DECK, dest + ".pptx")
    if not path:
        return result(0.0, {"file": "missing"})
    from pptx import Presentation
    prs = Presentation(path)
    master_style = prs.slide_master._element.find(".//p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr", NS)
    per, detail = [], []
    for i, s in enumerate(prs.slides, 1):
        title_ok, seen = _title_runs_ok(s, master_style)
        kinds = []
        for holder in (s, s.slide_layout, s.slide_layout.slide_master):
            for sh in holder.shapes:
                if sh.shape_type == 13:  # picture
                    try:
                        kinds.append(_logo_kind(sh.image.blob))
                    except Exception:
                        pass
        logo_ok = "new" in kinds and "old" not in kinds
        xml = s._element.xml + s.slide_layout._element.xml
        conf_ok = "confidential" in xml.lower()
        per.append((title_ok + logo_ok + conf_ok) / 3)
        detail.append({"slide": i, "title": title_ok, "seen": seen, "logo": logo_ok, "confidential": conf_ok})
    score = sum(per) / len(per) if per else 0.0
    return result(score, {"slides": len(per), "per_slide": detail})
