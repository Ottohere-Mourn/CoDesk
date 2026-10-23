"""o1 - one Calc workbook (Harbourside Community Health operations budget, Q3 2026):
A keys the September supplier invoices (image-only scans) into the Actuals sheet per the AP coding rules;
B applies the CFO's budget-transfer memo (scan) and builds the live variance analysis + notes for the review."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import random
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_oscan import Page, save_scan  # noqa: E402

ORG = "Harbourside Community Health"
ADDR = "14 Quay Street, Portsmouth PO1 3AB"
CCS = [  # code, name, holder, original Q3 budget
    (110, "Front Office & Reception", "J. Patel", 18000),
    (210, "Facilities Maintenance", "R. Doyle", 24000),
    (220, "Cleaning & Hygiene", "R. Doyle", 14000),
    (310, "IT & Telecoms", "S. Lindgren", 15000),
    (410, "Clinical Supplies", "Dr A. Mensah", 32000),
    (510, "Training & Travel", "K. O'Neill", 9000),
]
DEPTS = {110: "Reception", 210: "Estates & Facilities", 220: "Housekeeping", 310: "IT Service Desk",
         410: "Pharmacy / Clinic Stores", 510: "Learning & Development (HR)"}
GLS = [(6100, "Repairs & maintenance (buildings, plant, electrical, plumbing, grounds, furniture)"),
       (6150, "Cleaning services & hygiene supplies"),
       (6300, "IT hardware, software & IT support"),
       (6350, "Telephone & broadband"),
       (6400, "Medical & clinical consumables"),
       (6500, "Office supplies, printing & stationery"),
       (6600, "Training courses & conference fees"),
       (6650, "Staff travel & accommodation")]
TRANSFERS = {110: -2000, 210: 6000, 220: -1500, 310: 1500, 410: 0, 510: -4000}
MEMO_REF = "BT-2026-07"

# Jul + Aug invoices already booked: date, invoice no, vendor, description, cc, gl, net, vat
BOOKED = [
    ("2026-07-03", "OD-54211", "OfficeDirect", "Printer toner and paper", 110, 6500, 742.50, 148.50),
    ("2026-07-03", "MW-77102", "MediWholesale UK", "Gloves, dressings, syringes", 410, 6400, 6850.00, 1370.00),
    ("2026-07-08", "HX-5391", "Hexagon Facilities Services", "Planned maintenance July", 210, 6100, 2180.00, 436.00),
    ("2026-07-10", "VX-JUL26-087", "Voxline Telecom", "Telephone & broadband July 2026", 310, 6350, 1210.00, 242.00),
    ("2026-07-15", "0092877", "Northline Electrical", "Fire alarm panel repair", 210, 6100, 3420.00, 684.00),
    ("2026-07-17", "CF-0921", "CareSkills Academy", "Safeguarding level 2 course (8 staff)", 510, 6600, 1600.00, 0.00),
    ("2026-07-21", "PR-3251", "Quayside Print Studio", "Appointment cards (10,000)", 110, 6500, 1180.00, 236.00),
    ("2026-07-24", "TS-10301", "TechServe Solutions", "12 laptops for clinical admin", 310, 6300, 9840.00, 1968.00),
    ("2026-07-30", "TS-10288", "TechServe Solutions", "Reception check-in kiosk support", 110, 6300, 3215.20, 643.04),
    ("2026-07-31", "SP-22650", "Sparkle Pro Cleaning Ltd", "Contract cleaning July", 220, 6150, 3450.00, 690.00),
    ("2026-08-02", "BH-7720", "Brennan Heating Ltd", "Emergency boiler replacement (Block B)", 210, 6100, 21500.00, 4300.00),
    ("2026-08-04", "RC-1180", "Calmwave Interiors", "Reception seating refurbishment", 110, 6100, 6950.00, 1390.00),
    ("2026-08-05", "MW-77215", "MediWholesale UK", "Gloves, dressings, catheters", 410, 6400, 7400.00, 1480.00),
    ("2026-08-10", "VX-AUG26-102", "Voxline Telecom", "Telephone & broadband August 2026", 310, 6350, 1210.00, 242.00),
    ("2026-08-11", "HX-5462", "Hexagon Facilities Services", "Planned maintenance August", 210, 6100, 2180.00, 436.00),
    ("2026-08-13", "TR-55810", "RailCo Travel", "Staff rail travel - NHS Confed Expo", 510, 6650, 1500.00, 0.00),
    ("2026-08-14", "HY-9013", "HygieneFirst Supplies", "Hand sanitiser and paper towels", 220, 6150, 1300.00, 260.00),
    ("2026-08-19", "OD-54790", "OfficeDirect", "Stationery and visitor badges", 110, 6500, 612.30, 122.46),
    ("2026-08-20", "SW-4471", "Softwise Licensing", "Antivirus licence renewal", 310, 6300, 1340.00, 268.00),
    ("2026-08-22", "PH-2244", "ClinicSource Ltd", "Sterile instrument packs", 410, 6400, 4500.00, 900.00),
    ("2026-08-25", "GW-3316", "Greenway Grounds", "Grounds maintenance Jul-Aug", 210, 6100, 1960.00, 392.00),
    ("2026-08-29", "AQ-2207", "AquaFlow Plumbing", "Water leak repair, staff kitchen", 210, 6100, 3560.00, 712.00),
    ("2026-08-31", "SP-22741", "Sparkle Pro Cleaning Ltd", "Contract cleaning August", 220, 6150, 3450.00, 690.00),
]

VENDORS = {
    "Sparkle Pro Cleaning Ltd": ("Unit 4, Fratton Trading Estate, Portsmouth PO4 8SX", "GB 284 1193 07", 1),
    "Hexagon Facilities Services": ("Hexagon House, 2 Northarbour Road, Cosham PO6 3TH", "GB 397 5521 64", 0),
    "Northline Electrical": ("88 London Road, North End, Portsmouth PO2 0LZ", "GB 166 2048 91", 2),
    "MediWholesale UK": ("Gateway Park, Basingstoke RG24 8GT", "GB 702 3319 45", 0),
    "ClinicSource Ltd": ("3 Mill Lane, Romsey SO51 8EQ", "GB 551 0742 18", 1),
    "TechServe Solutions": ("Lakeside North Harbour, Western Road, Portsmouth PO6 3EN", "GB 813 4407 22", 2),
    "Voxline Telecom": ("PO Box 7110, Reading RG1 9QX", "GB 948 2205 61", 0),
    "OfficeDirect": ("Kingsway Business Park, Southampton SO15 0AB", "GB 225 7781 39", 1),
    "Quayside Print Studio": ("21 Broad Street, Old Portsmouth PO1 2JE", "GB 340 9917 05", 2),
    "J. Okonkwo Training": ("7 Elm Grove, Southsea PO5 1JF", None, 1),
}

# The September post: file, kind, data. kind: invoice | credit | duplicate | q4 | quote
DOCS = [
    ("scan_20260903_0001.pdf", "invoice", dict(no="MW-77310", vendor="MediWholesale UK", date="2026-09-03", cc=410, gl=6400,
     lines=[("Nitrile examination gloves, box of 100", 240, 9.50), ("Wound dressings 10x10 cm, pack of 50", 80, 26.25),
            ("Syringes 5 ml luer lock, box of 100", 38, 65.00)], vat=0.2, datefmt="%d %B %Y")),
    ("Invoice_SP-22790.pdf", "invoice", dict(no="SP-22790", vendor="Sparkle Pro Cleaning Ltd", date="2026-09-05", cc=220, gl=6150,
     lines=[("Deep clean after flood - ground floor (2 operatives x 2 days)", 1, 1300.00)], vat=0.2, datefmt="%d/%m/%Y")),
    ("scan_20260908_0002.pdf", "invoice", dict(no="HX-5520", vendor="Hexagon Facilities Services", date="2026-09-08", cc=210, gl=6100,
     lines=[("Planned preventive maintenance - September", 1, 1460.00), ("Quarterly lift inspection (2 lifts)", 2, 360.00)],
     vat=0.2, datefmt="%d.%m.%Y")),
    ("scan_20260909_0003.pdf", "invoice", dict(no="TS-10483", vendor="TechServe Solutions", date="2026-09-09", cc=310, gl=6300,
     lines=[("48-port managed network switch", 1, 2340.00), ("On-site installation and configuration (hours)", 6, 100.00)],
     vat=0.2, datefmt="%d %b %Y")),
    ("Voxline_bill_Sep2026.pdf", "invoice", dict(no="VX-SEP26-118", vendor="Voxline Telecom", date="2026-09-10", cc=310, gl=6350,
     lines=[("Telephone lines & broadband - September 2026", 1, 1210.00)], vat=0.2, datefmt="%d/%m/%Y")),
    ("scan_20260911_0007.pdf", "duplicate", dict(of="scan_20260909_0003.pdf")),
    ("scan_20260912_0004.pdf", "invoice", dict(no="0093114", vendor="Northline Electrical", date="2026-09-12", cc=210, gl=6100,
     lines=[("Emergency call-out (Saturday)", 1, 280.00), ("LED corridor light fittings supplied and fitted", 16, 85.00)],
     vat=0.2, datefmt="%d %B %Y")),
    ("scan_20260915_0005.pdf", "credit", dict(no="CN-0412", vendor="Hexagon Facilities Services", date="2026-09-15", cc=210, gl=6100,
     lines=[("Credit: duplicate call-out charge on invoice HX-5391", 1, 400.00)], vat=0.2, datefmt="%d.%m.%Y")),
    ("OfficeDirect_OD-55102.pdf", "invoice", dict(no="OD-55102", vendor="OfficeDirect", date="2026-09-16", cc=110, gl=6500,
     lines=[("Toner cartridge HP 59A", 4, 112.35), ("A4 copier paper, box of 5 reams", 12, 18.95),
            ("Visitor badges with lanyards, pack of 100", 2, 94.30)], vat=0.2, datefmt="%d/%m/%Y")),
    ("scan_20260918_0006.pdf", "invoice", dict(no="PH-2291", vendor="ClinicSource Ltd", date="2026-09-18", cc=410, gl=6400,
     lines=[("Sterile instrument packs, type C", 40, 32.50)], vat=0.2, datefmt="%d %b %Y")),
    ("scan_20260922_0008.pdf", "invoice", dict(no="PR-3307", vendor="Quayside Print Studio", date="2026-09-22", cc=110, gl=6500,
     lines=[("Patient information leaflets, A5, 4pp (5,000)", 1, 1420.00)], vat=0.2, datefmt="%d %B %Y")),
    ("Okonkwo_invoice_114.pdf", "invoice", dict(no="2026-114", vendor="J. Okonkwo Training", date="2026-09-25", cc=510, gl=6600,
     lines=[("First Aid at Work course, 1 day, 12 delegates", 12, 90.00)], vat=0.0, datefmt="%d/%m/%Y")),
    ("Invoice_SP-22817.pdf", "invoice", dict(no="SP-22817", vendor="Sparkle Pro Cleaning Ltd", date="2026-09-30", cc=220, gl=6150,
     lines=[("Contract cleaning - September 2026", 1, 3450.00)], vat=0.2, datefmt="%d/%m/%Y")),
    ("scan_20261002_0009.pdf", "q4", dict(no="VX-OCT26-131", vendor="Voxline Telecom", date="2026-10-01", cc=310, gl=6350,
     lines=[("Telephone lines & broadband - October 2026", 1, 1210.00)], vat=0.2, datefmt="%d/%m/%Y")),
    ("scan_20260929_0010.pdf", "quote", dict(no="Q-8841", vendor="Hexagon Facilities Services", date="2026-09-29", cc=210, gl=6100,
     lines=[("Roof gutter repair, Block A (scaffold + labour)", 1, 3200.00)], vat=0.2, datefmt="%d.%m.%Y")),
]

POLICY = f"""{ORG} - Accounts payable: coding rules for the operations budget workbook
(Finance office, revised April 2026)

1. One row in the Actuals sheet per invoice (not per invoice line). Posting date = the invoice date.
   Invoice no. exactly as printed by the supplier.
2. Cost centre: taken from the department in the "Deliver to" block of the invoice -
      Reception ............................ 110
      Estates & Facilities ................. 210
      Housekeeping ......................... 220
      IT Service Desk ...................... 310
      Pharmacy / Clinic Stores ............. 410
      Learning & Development (HR) .......... 510
   The supplier does not decide the cost centre (e.g. TechServe works for both Reception and IT).
3. GL account: by what was bought, see the Codes sheet of the workbook.
4. Amounts: enter Net, VAT and Gross exactly as on the invoice.
   Budget charge = Net, because we recover the VAT - EXCEPT cost centre 410 (clinical, VAT-exempt
   activity): the VAT cannot be recovered there, so the budget is charged with the Gross amount.
   (Column J of the Actuals sheet already holds this formula for the existing rows.)
5. Credit notes are booked as their own row with negative Net, VAT, Gross and budget charge,
   under the credit note number.
6. Only invoices dated 1 July - 30 September 2026 belong to Q3. Anything dated later stays in the
   inbox for the Q4 workbook. Quotations, pro-formas and statements are not invoices: do not book.
   A document scanned twice is booked once.
7. After booking, move the PDF from invoices/inbox to invoices/booked. List every document you leave
   in the inbox in invoices/inbox/NOT_BOOKED.txt (file name + reason), so the clerk knows why.
"""


def money(x):
    return f"{x:,.2f}"


def _invoice(path, kind, d, seed):
    rng = random.Random(seed)
    addr, vatno, layout = VENDORS[d["vendor"]]
    day = dt.date.fromisoformat(d["date"])
    sign = -1 if kind == "credit" else 1
    title = {"credit": "CREDIT NOTE", "quote": "QUOTATION"}.get(kind, "INVOICE" if vatno else "INVOICE")
    if kind not in ("credit", "quote") and vatno:
        title = "TAX INVOICE" if layout == 2 else "INVOICE"
    p = Page()
    serif = layout == 1
    kb, kn = ("serifbold", "serif") if serif else ("bold", "sans")
    # letterhead
    if layout == 0:
        p.rect((80, 70, 1160, 190), outline=(60, 60, 60), w=3)
        p.text(110, 92, d["vendor"], 44, kb)
        p.text(110, 148, addr, 22, kn)
        p.right(1130, 100, title, 40, kb)
    elif layout == 1:
        p.text(620, 80, d["vendor"].upper(), 40, kb, anchor="ma")
        p.text(620, 135, addr, 22, kn, anchor="ma")
        p.hline(90, 1150, 180, 3)
        p.text(620, 200, title, 36, kb, anchor="ma")
    else:
        p.text(90, 80, d["vendor"], 46, kb)
        p.text(90, 140, addr, 22, kn)
        p.text(90, 205, title, 34, kb)
    y = 270
    label = {"credit": "Credit note no.", "quote": "Quotation no."}.get(kind, "Invoice no.")
    date_label = "Date" if kind != "quote" else "Quotation date"
    p.text(730, y, f"{label}:", 24, kb)
    p.text(985, y, d["no"], 24, kn)
    p.text(730, y + 40, f"{date_label}:", 24, kb)
    p.text(985, y + 40, day.strftime(d["datefmt"]), 24, kn)
    if vatno:
        p.text(730, y + 80, "VAT reg. no.:", 24, kb)
        p.text(985, y + 80, vatno, 24, kn)
    p.text(730, y + 120, "Account:", 24, kb)
    p.text(985, y + 120, "HCH-00" + str(100 + seed % 800), 24, kn)
    p.text(90, y, "Deliver to / Invoice to:", 24, kb)
    p.text(90, y + 40, ORG, 24, kn)
    p.text(90, y + 80, "Attn: " + DEPTS[d["cc"]], 24, kn)
    p.text(90, y + 120, ADDR, 24, kn)
    y += 210
    if kind == "credit":
        p.text(90, y, "Reason for credit: duplicate call-out charge raised on our invoice HX-5391 (July).", 24, kn)
        y += 50
    rows, net = [], 0.0
    for desc, qty, unit in d["lines"]:
        amt = round(qty * unit, 2)
        net += amt
        rows.append([desc, str(qty), money(unit), money(sign * amt)])
    net = round(net, 2)
    vat = round(net * d["vat"], 2)
    y = p.table(90, y, [("Description", 600), ("Qty", 90), ("Unit price", 180), ("Amount GBP", 190)], rows,
                size=22, head_kind=kb, kind=kn, aligns=["l", "r", "r", "r"])
    y += 10
    tot = [("Net amount", sign * net), (f"VAT @ {int(d['vat'] * 100)}%", sign * vat), ("TOTAL GBP", sign * (net + vat))]
    if not vatno:
        tot = [("Net amount", net), ("VAT", 0.0), ("TOTAL GBP", net)]
    for i, (lab, v) in enumerate(tot):
        k = kb if i == 2 else kn
        p.right(900, y, lab, 26, k)
        p.right(1150, y, money(v), 26, k)
        y += 46
    p.hline(700, 1150, y + 4, 2)
    y += 50
    if not vatno:
        y = p.para(90, y, "No VAT charged - the supplier is not registered for VAT.", 22, kn)
    if kind == "quote":
        y = p.para(90, y, "This is a quotation, valid for 60 days. It is NOT a VAT invoice and no payment is due. "
                          "Please sign and return to accept; an invoice will follow on completion of the works.", 22, kn)
    elif kind == "credit":
        y = p.para(90, y, "This credit will be offset against your next payment. No action required.", 22, kn)
    else:
        due = day + dt.timedelta(days=30)
        y = p.para(90, y, f"Payment terms: 30 days net. Please pay by {due.strftime('%d %B %Y')} quoting the "
                          f"{label.lower()} as reference.", 22, kn)
    p.text(90, 1560, "Bank: Solent Mutual Bank  Sort code 40-12-76  Account 6110" + str(2000 + seed % 7000), 20, kn)
    p.text(90, 1595, f"{d['vendor']} - registered in England and Wales", 20, kn)
    if rng.random() < 0.6 and kind != "quote":
        p.stamp(820, 1300, ["RECEIVED", day.strftime("%d %b %Y").upper(), "HCH FINANCE"], size=26, angle=rng.choice([-9, 7, -14]))
    return p, net, vat


def _memo(path):
    p = Page()
    p.text(90, 80, ORG, 40, "bold")
    p.text(90, 135, "Finance Directorate", 26)
    p.hline(90, 1150, 180, 3)
    p.text(90, 205, "MEMO - BUDGET TRANSFER APPROVAL", 34, "bold")
    y = 275
    for a, b in (("Reference:", MEMO_REF), ("Date:", "15 September 2026"), ("From:", "Martin Achebe, Chief Financial Officer"),
                 ("To:", "Elena Marsh, Operations Manager"), ("Subject:", "Q3 2026 operations budget - virement requests")):
        p.text(90, y, a, 25, "bold")
        p.text(300, y, b, 25)
        y += 44
    y += 20
    y = p.para(90, y, "I have reviewed the virement requests submitted on 9 September. The decisions below apply to the "
                      "Q3 2026 budget (July - September). Approved transfers are to be entered in the operations budget "
                      "workbook against the original budgets; declined requests must not be entered.", 25, width=1060)
    y += 20
    rows = [["1", "510 Training & Travel", "210 Facilities Maint.", "4,000", "APPROVED"],
            ["2", "110 Front Office & Rec.", "210 Facilities Maint.", "2,000", "APPROVED"],
            ["3", "220 Cleaning & Hygiene", "310 IT & Telecoms", "1,500", "APPROVED"],
            ["4", "510 Training & Travel", "410 Clinical Supplies", "3,000", "DECLINED"]]
    y = p.table(90, y, [("#", 50), ("From cost centre", 330), ("To cost centre", 330), ("GBP", 140), ("Decision", 210)],
                rows, size=23, aligns=["l", "l", "l", "r", "l"])
    y += 10
    y = p.para(90, y, "Notes: (1) and (2) fund the emergency boiler replacement in Block B (August). (3) covers the "
                      "replacement network switch. (4) is declined - Clinical Supplies is expected to stay within its "
                      "original allocation; please raise again in Q4 if needed.", 23, width=1060)
    y += 40
    p.hand(110, y, "M. Achebe", 46, color=(25, 35, 120), angle=4, shear=0.35)
    p.hline(90, 500, y + 75, 2)
    p.text(90, y + 85, "Martin Achebe, CFO", 22)
    p.stamp(780, y - 30, ["APPROVED", "15 SEP 2026"], color=(30, 110, 50), size=30, angle=-8)
    save_scan([p], path, seed=7101)


def build(d):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation  # noqa: F401
    os.makedirs(f"{d}/invoices/inbox", exist_ok=True)
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    # ---------------------------------------------------------------- Budget
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = f"{ORG} - Operations budget Q3 2026 (July - September), GBP"
    ws["A1"].font = Font(bold=True, size=14)
    heads = ["Cost centre", "Name", "Budget holder", "Original Q3 budget", "Approved transfers", "Revised Q3 budget"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font, c.fill = bold, head_fill
    for r, (code, name, holder, amt) in enumerate(CCS, 4):
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=holder)
        ws.cell(row=r, column=4, value=amt).number_format = "#,##0.00"
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        ws.cell(row=r, column=6, value=f"=D{r}+E{r}").number_format = "#,##0.00"
    tr = 4 + len(CCS)
    ws.cell(row=tr, column=2, value="Total").font = bold
    for col in "DEF":
        c = ws[f"{col}{tr}"]
        c.value = f"=SUM({col}4:{col}{tr - 1})"
        c.number_format, c.font = "#,##0.00", bold
    ws[f"A{tr + 2}"] = "Transfers (virements) need written CFO approval; enter them in column E against the cost centre."
    ws[f"A{tr + 2}"].font = Font(italic=True, size=9)
    for col, w in zip("ABCDEF", (12, 28, 16, 20, 20, 20)):
        ws.column_dimensions[col].width = w
    # ---------------------------------------------------------------- Actuals
    wa = wb.create_sheet("Actuals")
    ah = ["Posting date", "Invoice no.", "Vendor", "Description", "Cost centre", "GL account", "Net", "VAT", "Gross",
          "Budget charge"]
    for i, h in enumerate(ah, 1):
        c = wa.cell(row=1, column=i, value=h)
        c.font, c.fill = bold, head_fill
        c.alignment = Alignment(horizontal="center")
    for r, (day, no, vendor, desc, cc, gl, net, vat) in enumerate(BOOKED, 2):
        wa.cell(row=r, column=1, value=dt.datetime.fromisoformat(day)).number_format = "DD/MM/YYYY"
        wa.cell(row=r, column=2, value=no)
        wa.cell(row=r, column=3, value=vendor)
        wa.cell(row=r, column=4, value=desc)
        wa.cell(row=r, column=5, value=cc)
        wa.cell(row=r, column=6, value=gl)
        for col, v in ((7, net), (8, vat), (9, round(net + vat, 2))):
            wa.cell(row=r, column=col, value=v).number_format = "#,##0.00"
        wa.cell(row=r, column=10, value=f"=IF(E{r}=410,I{r},G{r})").number_format = "#,##0.00"
    wa.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJ", (13, 15, 28, 40, 12, 12, 12, 11, 12, 15)):
        wa.column_dimensions[col].width = w
    # ---------------------------------------------------------------- Codes
    wc = wb.create_sheet("Codes")
    wc["A1"], wc["A1"].font = "Cost centres", Font(bold=True, size=12)
    for i, h in enumerate(["Cost centre", "Name", "Invoice 'Deliver to' department"], 1):
        wc.cell(row=2, column=i, value=h).font = bold
    for r, (code, name, _, _) in enumerate(CCS, 3):
        wc.cell(row=r, column=1, value=code)
        wc.cell(row=r, column=2, value=name)
        wc.cell(row=r, column=3, value=DEPTS[code])
    r0 = 4 + len(CCS)
    wc.cell(row=r0, column=1, value="GL accounts").font = Font(bold=True, size=12)
    wc.cell(row=r0 + 1, column=1, value="GL").font = bold
    wc.cell(row=r0 + 1, column=2, value="Use for").font = bold
    for i, (gl, desc) in enumerate(GLS, r0 + 2):
        wc.cell(row=i, column=1, value=gl)
        wc.cell(row=i, column=2, value=desc)
    for col, w in zip("ABC", (12, 80, 34)):
        wc.column_dimensions[col].width = w
    wb.save(f"{d}/Ops_Budget_2026Q3.xlsx")

    # ---------------------------------------------------------------- scans
    open(f"{d}/invoices/AP_coding_rules.txt", "w").write(POLICY)
    expected_bookings, files = [], {}
    rendered = {}
    for i, (fname, kind, data) in enumerate(DOCS):
        if kind == "duplicate":
            src = rendered[data["of"]]
            p = src["page"]
            p.stamp(860, 120, ["COPY"], color=(90, 90, 90), size=40, angle=6)
            save_scan([p], f"{d}/invoices/inbox/{fname}", seed=9100 + i)
            files[fname] = {"kind": kind, "of": data["of"], "no": src["no"]}
            continue
        page, net, vat = _invoice(fname, kind, data, seed=100 + i * 17)
        save_scan([page], f"{d}/invoices/inbox/{fname}", seed=5100 + i)
        rendered[fname] = {"page": page, "no": data["no"]}
        files[fname] = {"kind": kind, "no": data["no"]}
        if kind in ("invoice", "credit"):
            s = -1 if kind == "credit" else 1
            gross = round(net + vat, 2)
            expected_bookings.append({
                "no": data["no"], "vendor": data["vendor"], "date": data["date"], "cc": data["cc"], "gl": data["gl"],
                "net": round(s * net, 2), "vat": round(s * vat, 2), "gross": round(s * gross, 2),
                "charge": round(s * (gross if data["cc"] == 410 else net), 2), "file": fname})
    for f in files:
        files[f]["md5"] = hashlib.md5(open(f"{d}/invoices/inbox/{f}", "rb").read()).hexdigest()
    _memo(f"{d}/{MEMO_REF}_budget_transfer.pdf")

    # ---------------------------------------------------------------- expected values
    orig = {code: amt for code, _, _, amt in CCS}
    revised = {code: orig[code] + TRANSFERS[code] for code in orig}
    ja = {code: 0.0 for code in orig}
    for row in BOOKED:
        ja[row[4]] += round(row[6] + row[7], 2) if row[4] == 410 else row[6]
    sept = {code: 0.0 for code in orig}
    for b in expected_bookings:
        sept[b["cc"]] += b["charge"]
    flags_pre = sorted(c for c in orig if ja[c] > 1.1 * revised[c])
    flags_post = sorted(c for c in orig if ja[c] + sept[c] > 1.1 * revised[c])
    assert flags_pre == flags_post == [210], (flags_pre, flags_post)
    assert len(expected_bookings) == 12
    json.dump({"bookings": expected_bookings, "files": files,
               "booked_before": [{"no": r[1], "cc": r[4], "net": r[6]} for r in BOOKED],
               "not_booked": {"duplicate": "scan_20260911_0007.pdf", "q4": "scan_20261002_0009.pdf",
                              "quote": "scan_20260929_0010.pdf"},
               "ccs": {str(c): n for c, n, _, _ in CCS}, "original": {str(k): v for k, v in orig.items()},
               "transfers": {str(k): v for k, v in TRANSFERS.items()}, "revised": {str(k): v for k, v in revised.items()},
               "actual_pre": {str(k): round(v, 2) for k, v in ja.items()},
               "actual_post": {str(k): round(ja[k] + sept[k], 2) for k in ja},
               "flagged": [210], "memo_ref": MEMO_REF},
              open(f"{d}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/o1")
