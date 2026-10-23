"""o4 - one Impress presentation: applying the advisor's handwritten handout comments (scan) + review PDF vs preparing
the deck for the conference AV team (16:9, speaker notes from a scanned script, acknowledgements slide, notes PDF)."""
from __future__ import annotations

import re

from evaluation_examples.task_class import pg_olib as L
from evaluation_examples.task_class.pg_common import HOME, expected, fetch, result, upload

DIR = f"{HOME}/Documents/Talk"
PPTX = f"{DIR}/StreamJoin_DSSW.pptx"
REVIEW_PDF = f"{DIR}/StreamJoin_DSSW_v2_for_Markus.pdf"
NOTES_PDF = f"{DIR}/StreamJoin_DSSW_notes.pdf"


def setup(sc):
    upload(sc, [("o4/Talk/StreamJoin_DSSW.pptx", PPTX), ("o4/Talk/IDS_talk_rules.txt", f"{DIR}/IDS_talk_rules.txt"),
                ("o4/Talk/handout_comments_MW.pdf", f"{DIR}/handout_comments_MW.pdf"),
                ("o4/Talk/speaker_script.pdf", f"{DIR}/speaker_script.pdf")])
    sc._open_setup(PPTX)


def _deck(env, dest):
    path = fetch(env, PPTX, dest + ".pptx")
    if not path:
        return None, "missing"
    try:
        from pptx import Presentation
        prs = Presentation(path)
    except Exception as e:  # noqa: BLE001
        return None, f"unreadable: {e}"
    slides = []
    for s in prs.slides:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = ""
        try:
            if s.shapes.title is not None:
                title = s.shapes.title.text_frame.text
        except Exception:  # noqa: BLE001
            title = ""
        if not title.strip() and texts:
            title = texts[0].splitlines()[0]
        notes = ""
        if s.has_notes_slide:
            ns = s.notes_slide
            tf = ns.notes_text_frame
            if tf is not None and tf.text.strip():
                notes = tf.text
            else:
                cands = [sh.text_frame.text for sh in ns.shapes if sh.has_text_frame
                         and not re.fullmatch(r"\s*(\d+|‹#›|<#>)?\s*", sh.text_frame.text)]
                notes = max(cands, key=len) if cands else ""
        slides.append({"title": title.strip(), "text": "\n".join(texts), "notes": notes.strip()})
    return {"slides": slides, "w": prs.slide_width, "h": prs.slide_height}, None


def _is_ack(s):
    return "acknowledg" in L.norm(s["title"] + " " + s["text"])


def _canon(title, names, thr=0.8):
    best, name = 0.0, None
    for n in names:
        r = L.similar(title, n)
        if r > best:
            best, name = r, n
    return name if best >= thr else None


# ---------------------------------------------------------------- task A: advisor's comments
@L.safe
def evaluate_comments(env, dest: str) -> dict:
    exp = expected("o4")
    deck, err = _deck(env, dest)
    if deck is None:
        return result(0.0, {"deck": err})
    final = exp["final_titles"]
    main = [s for s in deck["slides"] if not _is_ack(s)]
    names = final + ["Motivaton", "Thank You", "Backup: Parameter Sweep"]
    canon = [_canon(s["title"], names) or s["title"] for s in main]
    by = {c: s for c, s in zip(canon, main)}
    alltext = L.norm(" ".join(s["title"] + " " + s["text"] for s in main))

    def has(c, *needles, absent=()):
        t = L.norm(by[c]["text"]) if c in by else ""
        return bool(t) and all(L.norm(n) in t for n in needles) and not any(L.norm(a) in t for a in absent)
    checks = {
        "typo_motivation": "Motivation" in by and L.norm(by["Motivation"]["title"]) == "motivation"
                           and "motivaton" not in alltext,
        "single_node_bullet_removed": has("Problem Setting", "fixed memory budget", "minimise p99",
                                          absent=["single-node setting"]),
        "latency_gain_84.3": has("Key Idea: Adaptive Window Partitioning", "84.3%", absent=["82.1"]),
        "cpu_gold_6338": has("Experimental Setup", "gold 6338", absent=["6248"]),
        "throughput_before_latency": "Results: Throughput" in canon and "Results: Latency" in canon
                                    and canon.index("Results: Throughput") < canon.index("Results: Latency"),
        "backup_slide_deleted": "parameter sweep" not in alltext,
        "chen_vldb_2025_added": has("Related Work", "chen", "2025", "splitjoin"),
        "code_link_added": has("Limitations and Future Work", "github.com/lpark/streamjoin", "multi-way"),
        "closing_is_questions": bool(main) and L.norm(main[-1]["title"]).rstrip("?") == "questions"
                                and "thank you" not in L.norm(" ".join(s["title"] for s in main)),
    }
    order_ok = canon == final
    checks["slide_order"] = order_ok
    checks["titles"] = [s["title"][:40] for s in main]

    # review PDF for the advisor
    s_pdf, pinfo = 0.0, {"exists": False}
    pdf = fetch(env, REVIEW_PDF, dest + ".pdf")
    if pdf:
        pages, meta = L.pdf_pages(pdf), L.pdf_meta(pdf)
        if pages and meta:
            land = [t for t, (w, h) in zip(pages, meta["sizes"]) if w > h] or pages
            seq = []
            for t in land:
                nt = L.norm(t)
                if "acknowledg" in nt[:200]:
                    continue
                hits = [(nt.find(L.norm(x)), x) for x in final + ["Backup: Parameter Sweep", "Thank You", "Motivaton"]
                        if L.norm(x) in nt]
                if hits:
                    first = min(hits)[1]
                    if not seq or seq[-1] != first:
                        seq.append(first)
            flat = L.norm(" ".join(pages))
            content = all(x in flat for x in ("84.3%", "github.com/lpark/streamjoin", "questions?"))
            s_pdf = 0.05 + 0.15 * (seq == final) + 0.1 * content
            pinfo = {"exists": True, "slide_pages": len(land), "title_sequence_ok": seq == final, "content_ok": content}
    checks["review_pdf"] = pinfo
    edits = [v for k, v in checks.items() if isinstance(v, bool) and k != "slide_order"]
    score = 0.07 * sum(edits) + 0.07 * order_ok + s_pdf
    return result(score, checks)


# ---------------------------------------------------------------- task B: AV-team version
@L.safe
def evaluate_avprep(env, dest: str) -> dict:
    exp = expected("o4")
    deck, err = _deck(env, dest)
    if deck is None:
        return result(0.0, {"deck": err})
    ratio = deck["w"] / deck["h"] if deck["h"] else 0
    wide = abs(ratio - 16 / 9) <= 0.012
    notes_by_id = {nid: (aliases, text) for nid, aliases, text in exp["notes"]}
    alias = {a: nid for nid, aliases, _ in exp["notes"] for a in aliases}
    slides = deck["slides"]
    ok, detail, want_texts = 0, [], []
    for s in slides:
        nid = "ack" if _is_ack(s) else alias.get(_canon(s["title"], list(alias), 0.75) or "", None)
        if nid is None:
            good = len(s["notes"].split()) >= 10
            detail.append({"slide": s["title"][:30], "id": None, "notes": good})
        else:
            want = notes_by_id[nid][1]
            want_texts.append(want)
            sim = L.similar(s["notes"], want)
            good = sim >= 0.85
            detail.append({"slide": s["title"][:30], "id": nid, "sim": round(sim, 2)})
        ok += good
    s_notes = ok / len(slides) if slides else 0.0

    acks = [i for i, s in enumerate(slides) if _is_ack(s)]
    grant = re.sub(r"\s", "", exp["grant"].lower())
    ack_ok = bool(acks) and any(grant in re.sub(r"\s", "", L.norm(slides[i]["text"])) and
                                re.search(r"german research foundation|\bdfg\b", L.norm(slides[i]["text"]))
                                for i in acks)
    ack_pos = bool(acks) and acks[-1] == len(slides) - 2

    s_pdf, pinfo = 0.0, {"exists": False}
    pdf = fetch(env, NOTES_PDF, dest + ".pdf")
    if pdf:
        pages = L.pdf_pages(pdf)
        if pages:
            flat = " ".join(pages)
            found = sum(L.best_window(flat, w) >= 0.85 for w in want_texts)
            frac = found / len(want_texts) if want_texts else 0.0
            s_pdf = 0.2 * (len(pages) >= len(slides)) + 0.8 * frac
            pinfo = {"exists": True, "pages": len(pages), "notes_found": f"{found}/{len(want_texts)}"}

    score = 0.15 * wide + 0.45 * s_notes + 0.1 * ack_ok + 0.05 * ack_pos + 0.25 * s_pdf
    checks = {"aspect_16_9": wide, "ratio": round(ratio, 4), "slides_with_correct_notes": f"{ok}/{len(slides)}",
              "notes": detail, "ack_slide_content": ack_ok, "ack_before_closing": ack_pos, "notes_pdf": pinfo}
    return result(score, checks)
