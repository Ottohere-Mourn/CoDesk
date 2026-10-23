"""X1 - month-end close at Harbour & Pine Coffee Roasters (Thunderbird + Ledgerline web app in Chrome + LibreOffice).

A (pgx1recon): September bank reconciliation - scanned statement from the mail -> receipts posted in Ledgerline ->
   accountant's Calc template filled in -> reply to the accountant with the workbook attached (draft).
B (pgx1report): September management report - figures from Ledgerline's Reports page (period selector) -> Calc table +
   chart -> Writer report + PDF -> PDF uploaded as board pack in Ledgerline -> reply to the director with the PDF (draft).
Shared live state: one Thunderbird profile/window, one Chrome window with one Ledgerline login session whose period
selector and list filters are stored per session (B switching the period re-scopes A's invoice list and vice versa).
"""
from __future__ import annotations

import re

from evaluation_examples.task_class import pg_common as pc
from evaluation_examples.task_class import pg_xcommon as xc

APP, PORT = "ledgerline", 8121
REC = f"{pc.HOME}/Documents/Finance/Bank_Rec_2026-09.xlsx"
PACK = f"{pc.HOME}/Documents/Reports/Mgmt_Pack_2026-09.xlsx"
DOCX = f"{pc.HOME}/Documents/Reports/Management_Report_2026-09.docx"
PDF = f"{pc.HOME}/Documents/Reports/Management_Report_2026-09.pdf"
CHANNELS = ["Wholesale", "Online shop", "Training & equipment"]
MONTHS = ["2026-07", "2026-08", "2026-09"]


def setup(sc):
    pc.sh(sc, "mkdir -p ~/Documents/Finance ~/Documents/Reports")
    up = f"{pc.WEBAPP_DIR}/{APP}/uploads"
    pc.upload(sc, [("x1/uploads/1_Management_Report_2026-07.pdf", f"{up}/1_Management_Report_2026-07.pdf"),
                   ("x1/uploads/2_Management_Report_2026-08.pdf", f"{up}/2_Management_Report_2026-08.pdf")])
    pc.start_webapp(sc, "x1/app_x1.py", APP, PORT, "x1/ledgerline_seed.db")
    pc.start_thunderbird(sc, "x1/thunderbird-profile.tar.gz", "pgx1")
    pc.open_chrome(sc, [f"http://127.0.0.1:{PORT}/login"])


# ---------------------------------------------------------------- A: bank reconciliation
def _status(s) -> str:
    f = str(s or "").lower()
    if "already" in f:
        return "Already posted"
    if "unmatch" in f or "not match" in f or "no match" in f:
        return "Unmatched"
    if "posted" in f:
        return "Posted"
    return ""


def _invs(s) -> set:
    return set(re.findall(r"(?<!\d)0[34]\d{2}(?!\d)", str(s or "")))


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    n = xc.numbers(str(v))
    return n[0] if len(n) == 1 else None


def _rec_parse(path) -> dict:
    """Header-driven read of the reconciliation workbook (the agent may move columns or add a sheet)."""
    from openpyxl import load_workbook
    out = {"rows": [], "summary": {}}
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for r, row in enumerate(grid):
            labels = [xc.flat(v) for v in row]
            for key in ("openingbalance", "totalpaidin", "totalpaidout", "closingbalance"):
                if labels and any(lb.startswith(key) for lb in labels if lb) and key not in out["summary"]:
                    start = next(i for i, lb in enumerate(labels) if lb and lb.startswith(key))
                    val = next((_num(v) for v in row[start + 1:] if _num(v) is not None), None)
                    if val is not None:
                        out["summary"][key] = val
            if any("invoice" in lb for lb in labels) and any("status" in lb for lb in labels):
                col = lambda word: next((i for i, lb in enumerate(labels) if word in lb), None)  # noqa: E731
                ca, ci, cs = col("amount") if col("amount") is not None else col("paidin"), col("invoice"), col("status")
                cc = col("comment")
                for below in grid[r + 1:]:
                    if ca is None or ca >= len(below):
                        continue
                    amt = _num(below[ca])
                    if amt is None:
                        continue
                    get = lambda c: below[c] if c is not None and c < len(below) else None  # noqa: E731
                    out["rows"].append({"amount": amt, "invoices": _invs(get(ci)), "status": _status(get(cs)),
                                        "raw_status": str(get(cs) or ""), "comment": str(get(cc) or "")})
    return out


def evaluate_recon(env, dest: str) -> dict:
    exp = pc.expected("x1")
    checks = {}
    q = pc.sql(env, pc.webapp_db(APP), {
        "pays": "SELECT p.id, i.number, p.date, p.amount, p.deleted FROM payments p JOIN invoices i ON i.id=p.invoice_id"})
    if q is None:
        return pc.result(0.0, {"app_db": "unreadable"})
    seed_max = exp["seed_payment_max_id"]
    new = [p for p in q["pays"] if p["id"] > seed_max and not p["deleted"]]
    want_invs = {e["invoice"] for e in exp["payments"]}
    hits, used = [], set()
    for e in exp["payments"]:
        m = next((p for p in new if p["id"] not in used and p["number"] == e["invoice"] and abs(p["amount"] - e["amount"]) < 0.006
                  and p["date"] == e["date"]), None)
        if m:
            used.add(m["id"])
        hits.append(bool(m))
    wrong = [f"{p['number']} {p['amount']} {p['date']}" for p in new if p["number"] not in want_invs]
    dupes = [inv for inv in want_invs if sum(p["number"] == inv for p in new) > 1]
    seed_removed = [p["number"] for p in q["pays"] if p["id"] <= seed_max and p["deleted"]]
    checks["payments_ok"] = f"{sum(hits)}/{len(hits)}"
    checks["payments_missing_or_wrong"] = [e["invoice"] for e, h in zip(exp["payments"], hits) if not h]
    checks["wrong_invoice_postings"] = wrong
    checks["duplicate_postings"] = dupes
    checks["seed_receipts_deleted"] = seed_removed
    no_wrong = not wrong and not dupes and not seed_removed and sum(hits) > 0

    # workbook
    path = pc.fetch(env, REC, dest + "_rec.xlsx")
    rec = None
    if path:
        try:
            rec = _rec_parse(path)
        except Exception as e:  # not a readable xlsx
            checks["rec_error"] = repr(e)[:200]
    rows_ok, summ_ok = 0.0, 0
    if rec:
        detail = []
        for c in exp["credits"]:
            row = next((r for r in rec["rows"] if abs(r["amount"] - c["amount"]) < 0.006), None)
            if not row:
                detail.append(f"{c['desc'][:28]}: no row")
                continue
            inv_ok = row["invoices"] == set(c["invoices"])
            st_ok = row["status"] == c["status"]
            rows_ok += 1.0 if inv_ok and st_ok else 0.5 if inv_ok or st_ok else 0.0
            if not (inv_ok and st_ok):
                detail.append(f"{c['desc'][:28]}: invoices {sorted(row['invoices'])} status '{row['raw_status']}'")
        st = exp["statement"]
        want = {"openingbalance": st["opening"], "totalpaidin": st["paid_in"], "totalpaidout": st["paid_out"],
                "closingbalance": st["closing"]}
        summ_ok = sum(abs(rec["summary"].get(k, -1e9) - v) < 0.006 for k, v in want.items())
        checks["rec_rows_ok"] = f"{rows_ok:g}/{len(exp['credits'])}"
        checks["rec_row_problems"] = detail[:12]
        checks["rec_summary"] = rec["summary"]
    checks["rec_file"] = bool(rec)

    # reply to Priya
    boxes = xc.mailboxes(env) or {}
    mine = [m for m in xc.outgoing(boxes, exp["msg"].values()) if xc.addressed(m, exp["priya"])]
    best, att_ok = None, False
    for k, m in enumerate(mine):
        for i, a in enumerate(m["attachments"]):
            if a["name"].lower().endswith((".xlsx", ".xls", ".ods")):
                p = xc.fetch_attachment(env, a, f"{dest}_att{k}_{i}")
                try:
                    vals = xc.cell_numbers(xc.xlsx_cells(p)) if p and p.endswith(".xlsx") else []
                except Exception:
                    vals = []
                n_cred = sum(xc.has_amount(vals, c["amount"]) for c in exp["credits"])
                if n_cred >= 10 and xc.has_amount(vals, exp["statement"]["closing"]):
                    best, att_ok = m, True
        best = best or m
    text = best["text"] if best else ""
    mentions = [xc.has_amount(text, 2145.30), xc.has_amount(text, 250.00), xc.has_amount(text, exp["outstanding_0421"])]
    is_reply = bool(best) and exp["msg"]["priya"] in best["irt"]
    checks.update({"reply_draft": bool(best), "reply_has_rec_attached": att_ok, "reply_mentions_vat_deposit_short": mentions,
                   "is_reply_to_priya": is_reply, "drafts_seen": [m["subject"][:50] for m in xc.outgoing(boxes, exp["msg"].values())][:6]})

    score = (0.40 * sum(hits) / len(hits) + 0.10 * no_wrong + 0.20 * rows_ok / len(exp["credits"]) + 0.10 * summ_ok / 4
             + 0.10 * att_ok + 0.075 * sum(mentions) / 3 + 0.025 * is_reply)
    score = xc.capped(score, [(sum(hits) == 0, 0.3), (not rec, 0.5)])
    return pc.result(score, checks)


# ---------------------------------------------------------------- B: management report
def _report_text_checks(text: str, exp: dict) -> dict:
    r = exp["report"]
    vals = xc.numbers(text)
    chan = [xc.has_amount(vals, r["rev"][m][c]) for m in MONTHS for c in CHANNELS]
    sep_total = xc.has_amount(vals, r["totals"]["2026-09"])
    pct = any(abs(abs(p) - abs(r["pct_sep_vs_aug"])) <= 0.051 for p in xc.percents(text))
    names = [n for n, _ in r["top3"]]
    top_order = xc.in_order(text, names)
    top_amounts = all(xc.has_amount(vals, v) for _, v in r["top3"])
    counts = [int(x) for x in re.findall(r"(\d+)\s+(?:sales\s+|new\s+)?invoices", text, re.I)]
    for m in re.finditer(r"invoices", text, re.I):  # "Invoices issued in September 2026: 23" / table cell on the next line
        counts += [int(x) for x in re.findall(r"(?<![\d.,£$])(\d{1,3})(?![\d.,%])", text[m.end(): m.end() + 70])]
    return {"channel_values": sum(chan), "sep_total": sep_total, "pct_change": pct, "top3_in_order": top_order,
            "top3_amounts": top_amounts, "sep_invoice_count": r["sep_count"] in counts}


def evaluate_report(env, dest: str) -> dict:
    exp = pc.expected("x1")
    r = exp["report"]
    checks = {}
    # workbook
    xp = pc.fetch(env, PACK, dest + "_pack.xlsx")
    chan_x = tot_x = 0
    chart_x = False
    if xp:
        try:
            vals = xc.cell_numbers(xc.xlsx_cells(xp))
            chan_x = sum(xc.has_amount(vals, r["rev"][m][c]) for m in MONTHS for c in CHANNELS)
            tot_x = sum(xc.has_amount(vals, v) for v in list(r["totals"].values()) + [r["q3_total"]])
            chart_x = xc.xlsx_charts(xp) > 0
        except Exception as e:
            checks["xlsx_error"] = repr(e)[:200]
    checks.update({"xlsx_found": bool(xp), "xlsx_channel_values": f"{chan_x}/9", "xlsx_totals": f"{tot_x}/4", "xlsx_chart": chart_x})
    # Writer document
    dp = pc.fetch(env, DOCX, dest + "_report.docx")
    vis = xc.docx_visuals(dp) if dp else {}
    docx_ok = bool(dp) and bool(xc.docx_text(dp).strip())
    docx_vis = docx_ok and sum(vis.values()) > 0
    checks.update({"docx_found": docx_ok, "docx_visuals": vis})
    # PDF
    pp = pc.fetch(env, PDF, dest + "_report.pdf")
    pdf_ok = xc.is_pdf(pp)
    t = _report_text_checks(xc.pdf_text(pp), exp) if pdf_ok else {}
    checks.update({"pdf_found": pdf_ok, "pdf_checks": t})
    # board pack in Ledgerline
    q = pc.sql(env, pc.webapp_db(APP), {"bp": "SELECT id, period, title, filename, path FROM boardpacks WHERE deleted=0 AND id>2"}) or {}
    up_ok = False
    for b in (q.get("bp") or []):
        if b["period"] != "2026-09":
            continue
        lp = pc.fetch(env, b["path"], f"{dest}_upload{b['id']}.pdf")
        if xc.is_pdf(lp) and xc.has_amount(xc.pdf_text(lp), r["totals"]["2026-09"]):
            up_ok = True
    checks["boardpack_uploads"] = [(b["period"], b["filename"]) for b in (q.get("bp") or [])]
    checks["boardpack_sep_ok"] = up_ok
    # reply to Tom with the PDF
    boxes = xc.mailboxes(env) or {}
    mine = [m for m in xc.outgoing(boxes, exp["msg"].values()) if xc.addressed(m, exp["tom"])]
    att_ok, is_reply = False, False
    for k, m in enumerate(mine):
        for i, a in enumerate(m["attachments"]):
            if a["name"].lower().endswith(".pdf"):
                lp = xc.fetch_attachment(env, a, f"{dest}_att{k}_{i}")
                if xc.is_pdf(lp) and xc.has_amount(xc.pdf_text(lp), r["totals"]["2026-09"]):
                    att_ok = True
                    is_reply = is_reply or exp["msg"]["pack"] in m["irt"]
    checks.update({"reply_to_tom": bool(mine), "reply_has_report_pdf": att_ok, "is_reply_to_tom": is_reply})

    score = (0.12 * chan_x / 9 + 0.04 * tot_x / 4 + 0.06 * chart_x + 0.04 * docx_ok + 0.06 * docx_vis + 0.04 * pdf_ok
             + (0.10 * t["channel_values"] / 9 + 0.04 * t["sep_total"] + 0.04 * t["pct_change"] + 0.06 * t["top3_in_order"]
                + 0.03 * t["top3_amounts"] + 0.03 * t["sep_invoice_count"] if t else 0)
             + 0.14 * up_ok + 0.15 * att_ok + 0.05 * is_reply)
    figures_wrong = bool(t) and (t["channel_values"] < 9 or not t["pct_change"] or not t["top3_in_order"])
    score = xc.capped(score, [(not (pdf_ok or up_ok or att_ok), 0.3), (figures_wrong, 0.75)])
    return pc.result(score, checks)
