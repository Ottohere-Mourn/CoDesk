"""a3 - Copper Lane Cycles: one parts-stock workbook open in LibreOffice Calc (~/Documents/Stock/parts_stock.xlsx).

A (pga3goodsin)   book three scanned delivery notes: add received qty to On hand, Last received, Goods-in log rows,
                  back-order list (~/Documents/Stock/backorders_2026-11.txt).
B (pga3stocktake) apply Friday's handwritten stock-take sheets: On hand += variance (counted - printed system qty),
                  Last counted, delete discontinued rows, insert new SKUs in order, sheet 'Variance 2026-11'.
Shared live state: one soffice process with the workbook open (one model, one file). Both change the same On hand
cells (read-modify-write) and B's row deletes/inserts move the rows A edits. Checkers read the FINAL workbook and
find SKUs by code; each tolerates the other task's correct delta being present or absent.
"""
from __future__ import annotations

import datetime as dt
import re

from evaluation_examples.task_class import pg_common as pc

DIR = f"{pc.HOME}/Documents/Stock"
XLSX = f"{DIR}/parts_stock.xlsx"
BACKORDERS = f"{DIR}/backorders_2026-11.txt"
EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def setup(sc):
    exp = pc.expected("a3")
    stage = f"{pc.HOME}/.cache/pga3"
    pairs = [("a3/parts_stock.xlsx", f"{stage}/parts_stock.xlsx")]
    pairs += [(f"a3/{rel}", f"{pc.HOME}/{dst}") for rel, dst in sorted(exp["files"].items())]
    pc.upload(sc, pairs)
    # the workbook is copied only once: a dual run calls setup twice and must not reset it
    pc.sh(sc, f"mkdir -p {DIR} && cp -n {stage}/parts_stock.xlsx {XLSX}")
    sc._open_setup(XLSX)


# ---------------------------------------------------------------- workbook parsing (header-driven)
def norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*", str(v))
    return float(m.group(1)) if m else None


def sku_of(v) -> str:
    m = re.search(r"\b([A-Z]{3})\s*-?\s*(\d{3})\b", str(v or "").upper())
    return f"{m.group(1)}-{m.group(2)}" if m else ""


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 40000 < v < 60000:
        return (EXCEL_EPOCH + dt.timedelta(days=float(v))).date()
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%d %B %Y", "%d %b %Y", "%m/%d/%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def date_ok(v, want: str) -> bool:
    """Right date; the day/month swap an en-US Calc makes of a typed dd/mm/yyyy is tolerated (display ambiguity)."""
    d, w = to_date(v), dt.date.fromisoformat(want)
    return d is not None and (d == w or (w.day <= 12 and d == dt.date(w.year, w.day, w.month)))


def _header(ws, want: dict):
    for r in range(1, 8):
        heads = {}
        for c in ws[r]:
            if c.value is not None:
                heads[c.column - 1] = norm(c.value)
        hits = {}
        for name, pat in want.items():
            col = next((k for k, h in sorted(heads.items()) if re.search(pat, h)), None)
            if col is not None:
                hits[name] = col
        if len(hits) >= min(len(want), 3) and "sku" in hits:
            return r, hits
    return None, {}


STOCK_WANT = {"sku": r"^sku|^code|^item", "desc": r"desc", "cat": r"categ", "bin": r"^bin|location",
              "onhand": r"onhand|^qty|quantity|stock", "last_rec": r"lastrec|received", "last_cnt": r"lastcount|counted"}


def read_book(path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {"sheets": wb.sheetnames, "stock": None, "goods": [], "variance": None, "stock_sheet": None}
    cands = sorted(wb.worksheets, key=lambda ws: (ws.title.strip().lower() != "stock",))
    for ws in cands:
        hr, cols = _header(ws, STOCK_WANT)
        if hr and "onhand" in cols and "last_rec" in cols:
            rows = []
            for r, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), hr + 1):
                g = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731
                sku = sku_of(g("sku"))
                if not sku:
                    continue
                rows.append({"row": r, "sku": sku, "desc": g("desc"), "cat": g("cat"), "bin": g("bin"),
                             "onhand": num(g("onhand")), "last_rec": g("last_rec"), "last_cnt": g("last_cnt")})
            out["stock"], out["stock_sheet"] = rows, ws.title
            break
    for ws in wb.worksheets:
        t = norm(ws.title)
        if "goods" in t:
            hr, cols = _header(ws, {"date": r"date", "supplier": r"supplier", "note": r"note|deliv|dn",
                                     "sku": r"^sku|ref|item|code", "qty": r"qty|quant"})
            if hr:
                for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                    g = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731
                    if sku_of(g("sku")):
                        out["goods"].append({"date": g("date"), "supplier": str(g("supplier") or ""),
                                             "note": str(g("note") or ""), "sku": sku_of(g("sku")), "qty": num(g("qty"))})
        if "variance" in t and out["variance"] is None:
            hr, cols = _header(ws, {"sku": r"^sku|code|item", "system": r"system|book|expected",
                                     "counted": r"count", "variance": r"varian|diff"})
            rows = []
            if hr:
                for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                    g = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731
                    if sku_of(g("sku")):
                        rows.append({"sku": sku_of(g("sku")), "system": num(g("system")), "counted": num(g("counted")),
                                     "variance": num(g("variance"))})
            out["variance"] = {"title": ws.title, "rows": rows}
    return out


def _load(env, dest):
    path = pc.fetch(env, XLSX, dest + ".xlsx")
    if not path:
        return None, "workbook missing"
    try:
        return read_book(path), ""
    except Exception as e:  # noqa: BLE001
        return None, f"unreadable: {e!r}"[:200]


def _by_sku(book):
    out = {}
    for r in book["stock"] or []:
        out.setdefault(r["sku"], []).append(r)
    return out


def _onhand(rows):
    return rows[0]["onhand"] if rows else None


def _capped(score, caps):
    for cond, cap in caps:
        if cond:
            score = min(score, cap)
    return score


# ---------------------------------------------------------------- A: goods in
def evaluate_goodsin(env, dest: str) -> dict:
    exp = pc.expected("a3")
    book, err = _load(env, dest)
    if not book:
        return pc.result(0.0, {"workbook": err})
    if not book["stock"]:
        return pc.result(0.0, {"stock_sheet": "not found", "sheets": book["sheets"]})
    by = _by_sku(book)
    seed, rec, var = exp["seed"], exp["receipts"], exp["variances"]
    checks = {}
    # 1. On hand of every received SKU = seed + received (+ the stock-take variance, if it has been applied)
    ok_oh, bad = 0, {}
    for sku, q in rec.items():
        v = _onhand(by.get(sku))
        allowed = {seed[sku]["onhand"] + q, seed[sku]["onhand"] + q + var.get(sku, 0)}
        if v is not None and any(abs(v - a) < 1e-6 for a in allowed):
            ok_oh += 1
        else:
            bad[sku] = f"{v} (want {sorted(allowed)})"
    checks["onhand_ok"] = f"{ok_oh}/{len(rec)}"
    checks["onhand_wrong"] = bad
    # 2. Last received
    lr = sum(bool(by.get(s)) and date_ok(by[s][0]["last_rec"], d) for s, d in exp["last_received"].items())
    checks["last_received_ok"] = f"{lr}/{len(exp['last_received'])}"
    # 3. Goods-in log rows (one per booked line; matched on SKU + qty + date + note number)
    old_notes = {"MCP554871", "BW76902", "VT20514"}  # October deliveries already in the seeded log
    new_goods = [g for g in book["goods"] if not any(n in norm(g["note"]).upper() for n in old_notes)]
    used, gl = set(), 0.0
    for want in exp["booked_lines"]:
        hit = None
        for i, g in enumerate(new_goods):
            if i in used or g["sku"] != want["sku"] or g["qty"] is None or abs(g["qty"] - want["qty"]) > 1e-6:
                continue
            hit = i
            break
        if hit is not None:
            used.add(hit)
            g = new_goods[hit]
            gl += 0.6 + 0.2 * date_ok(g["date"], want["date"]) + 0.2 * (norm(want["note"]) in norm(g["note"]))
    extra_goods = [f"{g['sku']} {g['qty']}" for i, g in enumerate(new_goods) if i not in used and (g["qty"] or 0) > 0]
    checks["goods_in_rows"] = f"{gl:.1f}/{len(exp['booked_lines'])}"
    checks["goods_in_extra"] = extra_goods
    gl_score = max(0.0, gl - 0.5 * len(extra_goods)) / len(exp["booked_lines"])
    # 4. traps (only meaningful once something was booked)
    started = ok_oh >= len(rec) // 2
    dmg_sku = next(iter(exp["damaged"]))
    dmg = _onhand(by.get(dmg_sku))
    trap_damaged = started and dmg is not None and all(
        abs(dmg - (seed[dmg_sku]["onhand"] + rec[dmg_sku] + exp["damaged"][dmg_sku] + v)) > 1e-6 for v in (0, var.get(dmg_sku, 0)))
    bo_skus = [b["sku"] for b in exp["backorders"] if b["sku"] not in rec]
    trap_bo = started and all(_onhand(by.get(s)) == seed[s]["onhand"] for s in bo_skus)
    typo_bad, typo_good = next(iter(exp["ref_typo"].items()))
    trap_typo = started and typo_bad not in by and ok_oh and typo_good not in bad
    checks["trap_returned_item_not_booked"] = trap_damaged
    checks["trap_backorder_not_booked"] = trap_bo
    checks["trap_wrong_ref_matched"] = bool(trap_typo)
    # 5. untouched SKUs unchanged (writes landing on the wrong row show up here)
    touched = set(rec) | set(var) | set(exp["discontinued"]) | {n["sku"] for n in exp["new"]}
    others = [s for s in seed if s not in touched]
    collateral = [s for s in others if _onhand(by.get(s)) != seed[s]["onhand"]]
    checks["other_skus_changed_or_missing"] = collateral[:10]
    untouched_ok = started and not collateral
    # 6. back-order list
    txt = pc.guest(env, f"cat {BACKORDERS} 2>/dev/null")
    bo_hits = 0
    for b in exp["backorders"]:
        m = [mm.end() for mm in re.finditer(re.escape(b["sku"]), txt.upper())]
        desc_hit = any(w in txt.lower() for w in b["desc"].lower().split()[1:3])
        near = any(re.search(rf"(?<![\d.]){b['qty']}(?![\d.])", txt[max(0, p - 120): p + 120]) for p in m)
        bo_hits += 0.5 * bool(m or desc_hit) + 0.5 * near
    wrong_bo = [s for s in rec if s not in {b["sku"] for b in exp["backorders"]} and s in txt.upper()]
    checks["backorder_file"] = bool(txt.strip()) and "No such file" not in txt
    checks["backorders_listed"] = f"{bo_hits}/{len(exp['backorders'])}"
    checks["backorder_file_extra_skus"] = wrong_bo
    bo_score = max(0.0, bo_hits - 0.25 * len(wrong_bo)) / len(exp["backorders"])

    score = (0.40 * ok_oh / len(rec) + 0.10 * lr / len(exp["last_received"]) + 0.20 * gl_score
             + 0.04 * trap_damaged + 0.03 * trap_bo + 0.03 * bool(trap_typo) + 0.05 * untouched_ok + 0.15 * bo_score)
    score = _capped(score, [(ok_oh < len(rec) / 2, 0.4), (ok_oh < len(rec) - 2, 0.75)])
    return pc.result(score, checks)


# ---------------------------------------------------------------- B: stock-take
def evaluate_stocktake(env, dest: str) -> dict:
    exp = pc.expected("a3")
    book, err = _load(env, dest)
    if not book:
        return pc.result(0.0, {"workbook": err})
    if not book["stock"]:
        return pc.result(0.0, {"stock_sheet": "not found", "sheets": book["sheets"]})
    by = _by_sku(book)
    seed, rec, var = exp["seed"], exp["receipts"], exp["variances"]
    checks = {}
    # 1. On hand of the counted SKUs = seed + variance (+ received goods, if booked)
    nz = [s for s in exp["counted"] if var[s] != 0]
    ok_oh, bad, overwritten = 0, {}, []
    for sku in exp["counted"]:
        v = _onhand(by.get(sku))
        allowed = {seed[sku]["onhand"] + var[sku], seed[sku]["onhand"] + var[sku] + rec.get(sku, 0)}
        if v is not None and any(abs(v - a) < 1e-6 for a in allowed):
            ok_oh += 1
        else:
            bad[sku] = f"{v} (want {sorted(allowed)})"
            if v is not None and abs(v - exp["counted_qty"][sku]) < 1e-6:
                overwritten.append(sku)
    nz_ok = sum(1 for s in nz if s not in bad)
    checks["onhand_ok"] = f"{ok_oh}/{len(exp['counted'])}"
    checks["onhand_variance_skus_ok"] = f"{nz_ok}/{len(nz)}"
    checks["onhand_wrong"] = bad
    checks["overwritten_with_counted"] = overwritten
    # 2. Last counted on every counted SKU (incl. the new ones)
    want_lc = list(exp["counted"]) + [n["sku"] for n in exp["new"]]
    lc = sum(bool(by.get(s)) and date_ok(by[s][0]["last_cnt"], exp["count_date"]) for s in want_lc)
    checks["last_counted_ok"] = f"{lc}/{len(want_lc)}"
    # 3. discontinued rows deleted, nothing else deleted, no duplicate rows
    disc_gone = sum(s not in by for s in exp["discontinued"])
    missing = [s for s in seed if s not in exp["discontinued"] and s not in by]
    dups = [s for s, rows in by.items() if len(rows) > 1]
    checks["discontinued_deleted"] = f"{disc_gone}/{len(exp['discontinued'])}"
    checks["other_rows_missing"] = missing[:10]
    checks["duplicate_rows"] = dups
    # 4. new SKUs: row present, on hand = counted (+ nothing else), sensible category/description, in SKU order
    order = [r["sku"] for r in book["stock"]]
    new_score, new_detail = 0.0, {}
    for n in exp["new"]:
        rows = by.get(n["sku"])
        if not rows:
            new_detail[n["sku"]] = "missing"
            continue
        r = rows[0]
        i = order.index(n["sku"])
        prev_ok = i == 0 or order[i - 1] < n["sku"]
        next_ok = i == len(order) - 1 or order[i + 1] > n["sku"]
        parts = {"onhand": r["onhand"] is not None and abs(r["onhand"] - n["counted"]) < 1e-6,
                 "in_order": prev_ok and next_ok,
                 "desc": any(w.lower() in str(r["desc"] or "").lower() for w in n["desc"].split()[:2]),
                 "category": norm(n["cat"])[:5] in norm(r["cat"]) or norm(r["cat"])[:5] in norm(n["cat"])}
        new_score += 0.3 + 0.3 * parts["onhand"] + 0.25 * parts["in_order"] + 0.075 * parts["desc"] + 0.075 * parts["category"]
        new_detail[n["sku"]] = {k: v for k, v in parts.items() if not v} or "ok"
    checks["new_skus"] = new_detail
    # 5. Variance sheet
    want_var = {s: (exp["printed_system"][s], exp["counted_qty"][s], var[s]) for s in nz}
    want_var.update({n["sku"]: (0, n["counted"], n["counted"]) for n in exp["new"]})
    vs = book["variance"]
    var_ok = 0.0
    if vs:
        rows = {r["sku"]: r for r in vs["rows"]}
        for s, (sy, ct, vv) in want_var.items():
            r = rows.get(s)
            if not r:
                continue
            var_ok += 0.6 * (r["variance"] is not None and abs(r["variance"] - vv) < 1e-6)
            var_ok += 0.2 * (r["system"] is not None and abs(r["system"] - sy) < 1e-6)
            var_ok += 0.2 * (r["counted"] is not None and abs(r["counted"] - ct) < 1e-6)
        wrong_rows = [s for s, r in rows.items() if s not in want_var and (r["variance"] or 0) != 0]
        checks["variance_sheet"] = vs["title"]
        checks["variance_extra_nonzero_rows"] = wrong_rows
        var_ok = max(0.0, var_ok - 0.5 * len(wrong_rows))
    else:
        checks["variance_sheet"] = None
    checks["variance_rows_ok"] = f"{var_ok:.1f}/{len(want_var)}"
    title_ok = bool(vs) and norm(vs["title"]) == norm("Variance 2026-11")

    started = nz_ok >= len(nz) / 2
    score = (0.35 * ok_oh / len(exp["counted"]) + 0.08 * lc / len(want_lc)
             + 0.12 * disc_gone / len(exp["discontinued"]) + 0.15 * new_score / len(exp["new"])
             + 0.17 * var_ok / len(want_var) + 0.03 * title_ok
             + 0.10 * (started and not missing and not dups))
    score = _capped(score, [(nz_ok < len(nz) / 2, 0.4), (len(overwritten) >= 3, 0.5), (nz_ok < len(nz) - 2, 0.75)])
    return pc.result(score, checks)
