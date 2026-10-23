"""Evaluation helpers shared by the F-family pairs (pg_f1..pg_f4): spreadsheets, JPEG/XCF inspection, image matching.

Runs on the host in the OSWorld venv (PIL, numpy, cv2, openpyxl, odfpy/pandas, PyMuPDF, python-pptx/docx).
"""
from __future__ import annotations

import io
import os
import re
import struct

import numpy as np
from PIL import Image


# ---------------------------------------------------------------- spreadsheets
def sheet_rows(path: str, sheet: str | None = None) -> list[list] | None:
    """All rows (lists of cell values) of a sheet in an .xlsx/.ods file; sheet matched case-insensitively.
    Returns None if the file or sheet cannot be read."""
    if not path or not os.path.exists(path):
        return None
    try:
        if path.lower().endswith(".ods"):
            import pandas as pd
            book = pd.read_excel(path, engine="odf", sheet_name=None, header=None)
            for name, df in book.items():
                if sheet is None or name.strip().lower() == sheet.lower():
                    return [[None if (isinstance(v, float) and v != v) else v for v in row] for row in df.values.tolist()]
            return None
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb.worksheets[0] if sheet is None else next(
            (w for w in wb.worksheets if w.title.strip().lower() == sheet.lower()), None)
        if ws is None:
            return None
        return [list(r) for r in ws.iter_rows(values_only=True)]
    except Exception:
        return None


def sheet_names(path: str) -> list[str]:
    try:
        if path.lower().endswith(".ods"):
            import pandas as pd
            return list(pd.read_excel(path, engine="odf", sheet_name=None, header=None).keys())
        from openpyxl import load_workbook
        return [w.title for w in load_workbook(path, read_only=True).worksheets]
    except Exception:
        return []


def table(rows: list[list] | None, key_header: str) -> tuple[dict, list[dict]]:
    """Find the header row containing key_header; return ({normalised header: column index}, [row dicts])."""
    if not rows:
        return {}, []
    for i, r in enumerate(rows):
        heads = [norm_head(v) for v in r]
        if norm_head(key_header) in heads:
            cols = {h: j for j, h in enumerate(heads) if h}
            out = []
            for r2 in rows[i + 1:]:
                if r2 is None or all(v in (None, "") for v in r2):
                    continue
                out.append({h: (r2[j] if j < len(r2) else None) for h, j in cols.items()})
            return cols, out
    return {}, []


def norm_head(v) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).strip()


def cell_text(v) -> str:
    return "" if v is None else str(v).strip()


def num(v):
    """Cell value -> float (accepts '1,234.50', '1.234,50', '€ 12', 12); None if not numeric."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s or not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if re.search(r",\d{1,2}$", s) else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def iso_date(v) -> str | None:
    """Cell value -> 'YYYY-MM-DD' (date/datetime cells, ISO text, day-first text like 03/10/2026 or 3 Oct 2026)."""
    import datetime as dt
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 40000 < v < 60000:  # spreadsheet serial number
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = cell_text(v)
    if not s:
        return None
    s = re.sub(r"\s+", " ", s.split("T")[0].strip())
    for f in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%d %B %Y", "%d %b %Y", "%d-%b-%Y",
              "%d/%m/%y", "%d.%m.%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, f).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def alnum(v) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())


# ---------------------------------------------------------------- JPEG / XCF
def jpeg_quality(path: str) -> int | None:
    """Estimate the IJG quality setting from the luminance quantisation table (exact for libjpeg/GIMP/PIL output)."""
    try:
        im = Image.open(path)
        if im.format != "JPEG":
            return None
        q0 = list(im.quantization[0])
    except Exception:
        return None
    best, err = None, None
    grey = Image.new("L", (16, 16))
    for q in range(1, 101):
        buf = io.BytesIO()
        grey.save(buf, "JPEG", quality=q)
        t = list(Image.open(io.BytesIO(buf.getvalue())).quantization[0])
        e = sum(abs(a - b) for a, b in zip(sorted(t), sorted(q0)))
        if err is None or e < err:
            best, err = q, e
    return best


def xcf_info(path: str) -> dict | None:
    """Parse a GIMP .xcf: {'version', 'width', 'height', 'layers': [{'name', 'w', 'h', 'text': str|None}]}.
    Text layers are recognised by their 'gimp-text-layer' parasite (the text itself is extracted from it)."""
    try:
        data = open(path, "rb").read()
    except OSError:
        return None
    if not data.startswith(b"gimp xcf "):
        return None
    ver_s = data[9:13]
    version = 0 if ver_s == b"file" else int(ver_s[1:].decode()) if ver_s.startswith(b"v") else 0
    p = 14
    w, h, _base = struct.unpack(">III", data[p:p + 12])
    p += 12
    if version >= 4:
        p += 4
    psize = 8 if version >= 11 else 4

    def props(p):
        out = []
        while p + 8 <= len(data):
            t, ln = struct.unpack(">II", data[p:p + 8])
            out.append((t, data[p + 8:p + 8 + ln]))
            p += 8 + ln
            if t == 0:
                break
        return out, p

    _, p = props(p)
    ptrs = []
    while p + psize <= len(data):
        v = int.from_bytes(data[p:p + psize], "big")
        p += psize
        if v == 0:
            break
        ptrs.append(v)
    layers = []
    for lp in ptrs:
        try:
            lw, lh, _lt = struct.unpack(">III", data[lp:lp + 12])
            n = struct.unpack(">I", data[lp + 12:lp + 16])[0]
            name = data[lp + 16:lp + 16 + n].rstrip(b"\0").decode("utf-8", "replace")
            pl, _ = props(lp + 16 + n)
            text = None
            for t, payload in pl:
                if t == 21:  # PROP_PARASITES
                    q = 0
                    while q + 4 <= len(payload):
                        nl = struct.unpack(">I", payload[q:q + 4])[0]
                        pname = payload[q + 4:q + 4 + nl].rstrip(b"\0").decode("utf-8", "replace")
                        q += 4 + nl
                        _flags, size = struct.unpack(">II", payload[q:q + 8])
                        pdata = payload[q + 8:q + 8 + size]
                        q += 8 + size
                        if pname == "gimp-text-layer":
                            text = text_from_parasite(pdata.decode("utf-8", "replace"))
            layers.append({"name": name, "w": lw, "h": lh, "text": text})
        except Exception:
            continue
    return {"version": version, "width": w, "height": h, "layers": layers}


def text_from_parasite(s: str) -> str:
    m = re.search(r'\(text "((?:[^"\\]|\\.)*)"\)', s)
    if m:
        return m.group(1).encode().decode("unicode_escape", "ignore") if "\\u" in m.group(1) else m.group(1).replace('\\"', '"')
    m = re.search(r'\(markup "((?:[^"\\]|\\.)*)"\)', s)
    if m:
        return re.sub(r"<[^>]+>", "", m.group(1).replace('\\"', '"'))
    return s


# ---------------------------------------------------------------- image comparison
def gray(im: Image.Image) -> np.ndarray:
    return np.asarray(im.convert("L"), dtype=np.float32)


def thumb_corr(a: Image.Image, b: Image.Image, size=(96, 64)) -> float:
    """Pearson correlation of small grey thumbnails (orientation-sensitive: compare same aspect only)."""
    if (a.size[0] > a.size[1]) != (b.size[0] > b.size[1]):
        return -1.0
    s = size if a.size[0] >= a.size[1] else (size[1], size[0])
    x = gray(a.resize(s, Image.BILINEAR)).ravel()
    y = gray(b.resize(s, Image.BILINEAR)).ravel()
    x, y = x - x.mean(), y - y.mean()
    d = float(np.sqrt((x * x).sum() * (y * y).sum()))
    return float((x * y).sum() / d) if d else 0.0


def find_square_crop(out: Image.Image, src: Image.Image, smin: int = 300):
    """Locate the square region of src that out (a resized square crop of src, possibly with a small overlay) was
    made from. Returns (score, x, y, side) in src pixels; score = normalised cross-correlation (1 = identical)."""
    import cv2
    o = gray(out)
    s = gray(src)
    H, W = s.shape
    f = 4
    s4 = cv2.resize(s, (W // f, H // f), interpolation=cv2.INTER_AREA)
    best = (-1.0, 0, 0, 0)
    for side in range(smin, min(W, H) + 1, 12):
        t = max(8, round(side / f))
        if t > s4.shape[0] or t > s4.shape[1]:
            continue
        tm = cv2.resize(o, (t, t), interpolation=cv2.INTER_AREA)
        r = cv2.matchTemplate(s4, tm, cv2.TM_CCOEFF_NORMED)
        _, mv, _, ml = cv2.minMaxLoc(r)
        if mv > best[0]:
            best = (float(mv), ml[0] * f, ml[1] * f, t * f)
    # refine at 1/2 then full resolution in a local window
    for f, span, step in ((2, 16, 1), (1, 4, 1)):
        _, bx, by, bs = best
        level = (-1.0, bx, by, bs)
        sf = cv2.resize(s, (W // f, H // f), interpolation=cv2.INTER_AREA) if f > 1 else s
        for side in range(bs - span * f, bs + span * f + 1, step * f):
            t = round(side / f)
            if t < 8 or t > sf.shape[0] or t > sf.shape[1]:
                continue
            x0, y0 = max(0, bx // f - span), max(0, by // f - span)
            x1, y1 = min(sf.shape[1], bx // f + t + span), min(sf.shape[0], by // f + t + span)
            win = sf[y0:y1, x0:x1]
            if win.shape[0] < t or win.shape[1] < t:
                continue
            tm = cv2.resize(o, (t, t), interpolation=cv2.INTER_AREA)
            r = cv2.matchTemplate(win, tm, cv2.TM_CCOEFF_NORMED)
            _, mv, _, ml = cv2.minMaxLoc(r)
            if mv > level[0]:
                level = (float(mv), (x0 + ml[0]) * f, (y0 + ml[1]) * f, t * f)
        if level[0] > 0:
            best = level
    return best


def overlay_mask(out: Image.Image, src: Image.Image, box, thr: float = 40.0) -> np.ndarray:
    """Boolean mask (out-sized) of pixels where out differs clearly from the matching crop of src (e.g. added text)."""
    import cv2
    x, y, side = box
    rec = src.crop((x, y, x + side, y + side)).resize(out.size, Image.LANCZOS)
    a = cv2.GaussianBlur(gray(out), (5, 5), 0)
    b = cv2.GaussianBlur(gray(rec), (5, 5), 0)
    return np.abs(a - b) > thr


def text_overlay_in_corner(mask: np.ndarray, corner: str = "br") -> dict:
    """Is there an overlay (text) concentrated in the given corner region of the image?"""
    h, w = mask.shape
    reg = mask[int(h * 0.72):, int(w * 0.40):] if corner == "br" else mask[int(h * 0.72):, :int(w * 0.6)]
    inside = int(reg.sum())
    total = int(mask.sum())
    return {"corner_px": inside, "total_px": total, "ok": inside >= 250 and inside >= 0.6 * total}


def detail_corr(a: Image.Image, b: Image.Image, size=(640, 360), ignore=None, blur: float = 1.0) -> float:
    """Correlation of the fine detail (text strokes) of two images: both resized to `size`, band-pass filtered
    (blur - heavy blur) so that shared layout (backgrounds, bars) does not dominate; an optional box (x0, y0, x1, y1,
    fractions of the image, e.g. a webcam inset) is ignored."""
    import cv2

    def bp(im):
        g = gray(im.convert("RGB").resize(size, Image.BILINEAR))
        return cv2.GaussianBlur(g, (0, 0), blur) - cv2.GaussianBlur(g, (0, 0), 8)
    x, y = bp(a), bp(b)
    m = np.ones(x.shape, bool)
    if ignore:
        w, h = size
        m[int(ignore[1] * h):int(ignore[3] * h), int(ignore[0] * w):int(ignore[2] * w)] = False
    x, y = x[m], y[m]
    x, y = x - x.mean(), y - y.mean()
    d = float(np.sqrt((x * x).sum() * (y * y).sum()))
    return float((x * y).sum() / d) if d else 0.0
