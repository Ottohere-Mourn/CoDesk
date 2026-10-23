"""F3 - a community fair's photos: recap slideshow + photo album (slides) vs next year's poster (poster).

Shared live state: one LibreOffice process (Impress with the recap deck and Writer with the poster - two windows of the
same soffice instance), one Nautilus window on the photo folder; task A deletes and re-files the photos task B picks from.
"""
from __future__ import annotations

import io
import os
import re
import zipfile

from evaluation_examples.task_class.pg_common import ASSETS, HOME, expected, fetch, result, sh, tree, upload
from evaluation_examples.task_class import pg_fimg as fi

PHOTOS = f"{HOME}/Pictures/MakersFair_2026"
DOCS = f"{HOME}/Documents/MakersFair"
DECK = f"{DOCS}/MakersFair_2026_recap.pptx"
POSTER = f"{DOCS}/poster_2027.docx"
PDF = f"{DOCS}/poster_2027.pdf"
KEYS = {"opening": "opening", "pottery": "pottery", "choir": "choir", "ukulele": "ukulele", "bread": "bread",
        "repair": "repair", "raffle": "raffle", "forge": "blacksmith"}
WHO = {"pottery": r"penhallow", "choir": r"st\.?\s*mary", "ukulele": r"falmouth\s+ukulele"}


def setup(sc):
    upload(sc, [("f3/f3.tar.gz", "/tmp/pgf3.tar.gz")])
    sh(sc, "mkdir -p ~/.cache; [ -f ~/.cache/.pgf3_ready ] || { tar xzf /tmp/pgf3.tar.gz -C ~ && "
           "touch ~/.cache/.pgf3_ready; }; rm -f /tmp/pgf3.tar.gz")
    sc.launch(f"[ -e /tmp/.pgf3_nautilus ] || {{ touch /tmp/.pgf3_nautilus; exec nautilus '{PHOTOS}'; }}", shell=True)
    sc._open_setup(DECK)
    sc._open_setup(POSTER)


# ---------------------------------------------------------------- photo identification
_refs = {}


def _photo_refs():
    from PIL import Image
    if not _refs:
        for p in expected("f3")["photos"]:
            im = Image.open(os.path.join(ASSETS, "f3/photos", p["name"]))
            im.load()
            _refs[p["name"]] = (p, im.convert("RGB").resize((160, 120)))
    return _refs


def identify(blob_or_img):
    """Which original photo is this embedded image? -> (photo dict, corr) or (None, best corr)."""
    from PIL import Image
    try:
        im = blob_or_img if isinstance(blob_or_img, Image.Image) else Image.open(io.BytesIO(blob_or_img))
        im = im.convert("RGB")
    except Exception:
        return None, 0.0
    if im.size[0] < 80 or im.size[1] < 60:
        return None, 0.0
    best, bc = None, -1.0
    for name, (p, ref) in _photo_refs().items():
        c = fi.thumb_corr(im, ref, size=(80, 60))
        if c > bc:
            best, bc = p, c
    return (best, bc) if bc >= 0.85 else (None, bc)


# ---------------------------------------------------------------- presentation parsing (pptx, or odp if re-saved)
def _slides_pptx(path):
    from pptx import Presentation
    out = []
    for s in Presentation(path).slides:
        texts, imgs = [], []
        stack = list(s.shapes)
        while stack:
            sh_ = stack.pop(0)
            if getattr(sh_, "shape_type", None) == 6:  # group
                stack.extend(sh_.shapes)
            if getattr(sh_, "has_text_frame", False) and sh_.has_text_frame:
                texts.append(sh_.text_frame.text)
            try:
                imgs.append(sh_.image.blob)
            except Exception:
                pass
        out.append({"text": " | ".join(texts), "images": imgs})
    return out


def _slides_odp(path):
    import xml.etree.ElementTree as ET
    ns = {"draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
          "xlink": "http://www.w3.org/1999/xlink", "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0"}
    z = zipfile.ZipFile(path)
    root = ET.fromstring(z.read("content.xml"))
    out = []
    for page in root.iter(f"{{{ns['draw']}}}page"):
        texts = [" ".join(t.itertext()) for t in page.iter("{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p")]
        imgs = []
        for im in page.iter(f"{{{ns['draw']}}}image"):
            href = im.get(f"{{{ns['xlink']}}}href", "")
            if href in z.namelist():
                imgs.append(z.read(href))
        out.append({"text": " | ".join(texts), "images": imgs})
    return out


def read_deck(env, dest):
    p = fetch(env, DECK, dest + ".pptx")
    if p:
        try:
            return _slides_pptx(p), p
        except Exception:
            pass
    p = fetch(env, DECK[:-5] + ".odp", dest + ".odp")
    if p:
        try:
            return _slides_odp(p), p
        except Exception:
            pass
    return None, None


# ---------------------------------------------------------------- task A: album + slideshow
def evaluate_slides(env, dest: str) -> dict:
    exp = expected("f3")
    photos = exp["photos"]
    prog = {it["key"]: it for it in exp["programme"]}
    t = tree(env, ["~/Pictures"])
    by_md5 = {}
    for p, m in t.items():
        by_md5.setdefault(m, []).append(p)
    filed, wrong = 0, []
    keep = [p for p in photos if not p["umbrella"]]
    for p in keep:
        locs = [x for x in by_md5.get(p["md5"], []) if x.startswith("Pictures/MakersFair_2026/")]
        it = prog.get(p["item"])
        ok = False
        for x in locs:
            parts = x.split("/")
            if len(parts) != 4:
                continue
            folder = parts[2].strip().lower()
            if it and it["time"]:
                hh, mm = it["time"].split(":")
                ok |= bool(re.match(rf"^{hh}[-._: ]?{mm}\b", folder)) and KEYS[p["item"]][:4] in folder.replace("'", "")
            else:
                ok |= folder.startswith("general")
        filed += ok
        if not ok:
            wrong.append([p["name"], p["item"], locs[:2]])
    umb = [p for p in photos if p["umbrella"]]
    umb_gone = sum(not by_md5.get(p["md5"]) for p in umb)
    # the deck
    slides, path = read_deck(env, dest)
    deck = {"found": bool(slides), "slides": len(slides or [])}
    title_ok = photo_ok = 0
    order_ok = cancelled_ok = template_ok = changed = False
    used_umbrella = []
    happened = [it for it in exp["programme"] if it["time"]]
    if slides:
        changed = len(slides) > 2
        idx = []
        per = {}
        for it in happened:
            kw = KEYS[it["key"]]
            hits = [i for i, s in enumerate(slides) if kw in s["text"].lower() and it["time"] in s["text"].replace(".", ":")]
            per[it["key"]] = {"slide": hits[0] + 1 if hits else None}
            if not hits:
                continue
            i = hits[0]
            idx.append(i)
            title_ok += 1
            ids = [identify(b) for b in slides[i]["images"]]
            got = [p["item"] for p, _ in ids if p]
            per[it["key"]]["photos"] = got
            photo_ok += any(p and p["item"] == it["key"] and not p["umbrella"] for p, _ in ids)
        for s in slides:
            for b in s["images"]:
                p, _ = identify(b)
                if p and p["umbrella"]:
                    used_umbrella.append(p["name"])
        order_ok = len(idx) == len(happened) and idx == sorted(idx)
        cancelled_ok = not any(re.search(r"blacksmith|forge", s["text"], re.I) for s in slides)
        template_ok = "riverside makers fair 2026" in slides[0]["text"].lower() and "see you next year" in slides[-1]["text"].lower()
        deck["per_item"] = per
    n = len(happened)
    score = (0.25 * filed / len(keep) + 0.10 * umb_gone / len(umb) + 0.20 * title_ok / n + 0.25 * photo_ok / n
             + 0.05 * order_ok + 0.05 * (changed and cancelled_ok) + 0.05 * (changed and template_ok) + 0.05 * changed)
    if used_umbrella:
        score = min(score, 0.5)
    if not changed:
        score = min(score, 0.4)
    checks = {"photos_filed": f"{filed}/{len(keep)}", "misfiled": wrong[:8], "umbrella_photos_deleted": f"{umb_gone}/{len(umb)}",
              "deck": deck, "slides_with_time_and_item": f"{title_ok}/{n}", "slides_with_right_photo": f"{photo_ok}/{n}",
              "programme_order": order_ok, "no_cancelled_item": cancelled_ok, "title_and_closing_kept": template_ok,
              "umbrella_photos_in_deck": used_umbrella}
    return result(score, checks)


# ---------------------------------------------------------------- task B: poster
DATE_RE = r"(16(th)?\s*(of\s*)?oct(ober)?\.?,?\s*2027|oct(ober)?\.?\s*16(th)?,?\s*2027|16[/.]10[/.](20)?27|2027-10-16)"
TIME_RE = r"(10(:00|\.00)?\s*(am|a\.m\.)?\s*(-|–|—|to|until)\s*(4(:00|\.00)?\s*(pm|p\.m\.)|16[:.]00))"


def evaluate_poster(env, dest: str) -> dict:
    exp = expected("f3")
    c = {}
    pdf = fetch(env, PDF, dest + ".pdf")
    text, pdf_imgs, page_ok = "", [], False
    if pdf:
        try:
            import fitz
            doc = fitz.open(pdf)
            page = doc[0]
            w, h = page.rect.width, page.rect.height
            page_ok = len(doc) == 1 and abs(w - 595.3) < 6 and abs(h - 841.9) < 6
            text = " ".join(pg.get_text() for pg in doc)
            for pg in doc:
                for img in pg.get_images(full=True):
                    try:
                        pdf_imgs.append(doc.extract_image(img[0])["image"])
                    except Exception:
                        pass
            c["pages"], c["page_size"] = len(doc), [round(w), round(h)]
        except Exception as e:  # noqa: BLE001
            c["pdf_error"] = repr(e)
    flat = re.sub(r"\s+", " ", text.lower())
    date_ok = bool(re.search(DATE_RE, flat))
    venue_ok = "princess pavilion" in flat
    time_ok = bool(re.search(TIME_RE, flat))
    no_ph = bool(text) and not re.search(r"\[(date|time|venue|photo|caption)\]", flat)
    ids = [identify(b) for b in pdf_imgs]
    found = {p["item"]: p for p, _ in ids if p and not p["umbrella"]}
    umbrella = [p["name"] for p, _ in ids if p and p["umbrella"]]
    acts_ok = {a: a in found for a in exp["poster_acts"]}
    who_ok = {a: bool(re.search(WHO[a], flat)) for a in exp["poster_acts"]}
    # the saved document (docx, or odt if re-saved)
    doc_ok, doc_info = False, {}
    dp = fetch(env, POSTER, dest + ".docx") or fetch(env, POSTER[:-5] + ".odt", dest + ".odt")
    if dp:
        try:
            z = zipfile.ZipFile(dp)
            names = z.namelist()
            body = z.read("word/document.xml" if "word/document.xml" in names else "content.xml").decode("utf-8", "replace")
            media = [n for n in names if n.startswith(("word/media/", "Pictures/"))]
            plain = re.sub(r"<[^>]+>", "", body)
            doc_info = {"images": len(media), "placeholders_left": re.findall(r"\[(?:DATE|TIME|VENUE|photo|caption)\]", plain)}
            doc_ok = len(media) >= 3 and not doc_info["placeholders_left"] and "princess pavilion" in plain.lower()
        except Exception as e:  # noqa: BLE001
            doc_info = {"error": repr(e)}
    score = (0.10 * page_ok + 0.15 * date_ok + 0.05 * venue_ok + 0.05 * time_ok + 0.05 * no_ph
             + 0.10 * sum(acts_ok.values()) + 0.05 * sum(who_ok.values()) + 0.15 * doc_ok)
    if not pdf:
        score = min(score, 0.3)
    if umbrella:
        score = min(score, 0.5)
    c.update(pdf=bool(pdf), one_a4_page=page_ok, date=date_ok, venue=venue_ok, time=time_ok, placeholders_gone=no_ph,
             photos_in_pdf=[(p["name"], p["item"]) if p else round(v, 2) for p, v in ids], acts=acts_ok, captions=who_ok,
             umbrella_photos=umbrella, document_saved=doc_ok, document=doc_info)
    return result(score, c)
