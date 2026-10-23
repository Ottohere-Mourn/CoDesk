"""W4 - one applicant-tracking web app (Quillon Talent on :8104) in one Chrome session: entering scanned interview
feedback, moving candidates and booking onsite loops vs scheduling hiring-manager interviews and building the weekly
recruiting summary. Both book into the same interviewer calendars (Priya Raman). See spec_w4.md."""
from __future__ import annotations

import datetime as dt
import re

from evaluation_examples.task_class.pg_common import (HOME, expected, fetch, open_chrome, result, sh, sql,
                                                      start_webapp, upload, webapp_db)

NAME, PORT = "w4", 8104
URL = f"http://127.0.0.1:{PORT}/"
DOCS = ["Interview_Decision_Rubric.pdf"]
XLSX = f"{HOME}/Documents/recruiting_weekly_wk43.xlsx"


def setup(sc):
    start_webapp(sc, f"{NAME}/app_w4.py", NAME, PORT, f"{NAME}/seed.db")
    upload(sc, [(f"{NAME}/{d}", f"{HOME}/.cache/pg{NAME}/{d}") for d in DOCS])
    sh(sc, f"mkdir -p ~/Documents && for f in {' '.join(DOCS)}; do cp -n ~/.cache/pg{NAME}/$f ~/Documents/$f; done")
    open_chrome(sc, [URL])


def _state(env, exp):
    m = exp["seed_max"]
    return sql(env, webapp_db(NAME), {
        "cands": "SELECT id, name, req_id, stage, reason FROM cands",
        "scorecards": f"SELECT id, cand_id, person_id, s1, s2, s3, s4, rec FROM scorecards WHERE id > {m['scorecards']} ORDER BY id",
        "interviews": "SELECT id, cand_id, kind, date, start, minutes, updated FROM interviews ORDER BY id",
        "panel": "SELECT interview_id, person_id FROM panel",
        "avail": "SELECT person_id, weekday, start, end FROM avail",
        "ooo": "SELECT person_id, date FROM ooo",
        "notes": f"SELECT id, kind, ref_id, user_id, body FROM notes WHERE id > {m['notes']} ORDER BY id"})


def _m(t):
    try:
        h, mm = str(t).split(":")[:2]
        return int(h) * 60 + int(mm)
    except (ValueError, AttributeError):
        return None


class Calendar:
    """Bookings + availability as stored in the app; checks one booking against them."""

    def __init__(self, st, seed_max):
        self.ivs = {i["id"]: i for i in st["interviews"]}
        # save order: seed bookings first, then by last save time (the Guest clock), then id
        self.order = {i["id"]: (0, "", i["id"]) if i["id"] <= seed_max else (1, i.get("updated") or "", i["id"])
                      for i in st["interviews"]}
        self.panel = {}
        for p in st["panel"]:
            self.panel.setdefault(p["interview_id"], set()).add(p["person_id"])
        self.avail = st["avail"]
        self.ooo = {(o["person_id"], o["date"]) for o in st["ooo"]}

    def span(self, iv):
        s = _m(iv["start"])
        return (s, s + (iv["minutes"] or 0)) if s is not None else (None, None)

    def in_availability(self, iv):
        s, e = self.span(iv)
        try:
            wd = dt.date.fromisoformat(iv["date"]).weekday()
        except (TypeError, ValueError):
            return False
        if s is None:
            return False
        for pid in self.panel.get(iv["id"], ()):
            if (pid, iv["date"]) in self.ooo:
                return False
            if not any(a["person_id"] == pid and a["weekday"] == wd and _m(a["start"]) <= s and e <= _m(a["end"]) for a in self.avail):
                return False
        return True

    def overlaps(self, iv):
        """Other interviews sharing a panelist that overlap this one (ids)."""
        s, e = self.span(iv)
        out = []
        for j in self.ivs.values():
            if j["id"] == iv["id"] or j["date"] != iv["date"] or not (self.panel.get(j["id"], set()) & self.panel.get(iv["id"], set())):
                continue
            js, je = self.span(j)
            if s is not None and js is not None and js < e and s < je:
                out.append(j["id"])
        return out

    def blame(self, iv):
        """Overlapping bookings that were saved BEFORE this one: this booking created the double-booking (the app
        warned when it was saved), so it - not the earlier one - is at fault."""
        return [j for j in self.overlaps(iv) if self.order[j] < self.order[iv["id"]]]


def _cand_ok(iv, cons, cal):
    s, e = cal.span(iv)
    if s is None:
        return False
    if "dates" in cons and iv["date"] not in cons["dates"]:
        return False
    if iv["date"] in cons.get("not_dates", []):
        return False
    if "min_start" in cons and s < cons["min_start"]:
        return False
    if "max_end" in cons and e > cons["max_end"]:
        return False
    return True


def _booking(cal, ivs, cons, minutes, panel_ok):
    """Score the latest of a candidate's new bookings of the right type: format/panel 0.4 (0.3 for B + stage),
    availability, overlaps, candidate constraint."""
    if not ivs:
        return {"booked": False}, {}
    iv = ivs[-1]
    fmt = iv["date"] in cal.week and iv["minutes"] == minutes and panel_ok(cal.panel.get(iv["id"], set()))
    parts = {"format_and_panel": fmt, "in_availability": cal.in_availability(iv), "no_overlap": not cal.blame(iv),
             "candidate_availability": _cand_ok(iv, cons, cal), "single_booking": len(ivs) == 1}
    info = {"booked": f"{iv['date']} {iv['start']} {iv['minutes']}min panel={sorted(cal.panel.get(iv['id'], set()))}",
            "overlaps_with": cal.overlaps(iv), "created_double_booking": cal.blame(iv)}
    return info, parts


def evaluate_feedback(env, dest: str) -> dict:
    exp = expected(NAME)
    st = _state(env, exp)
    if not st:
        return result(0.0, {"app_db": "unreadable"})
    cal = Calendar(st, exp["seed_max"]["interviews"])
    cal.week = exp["week"]
    cands = {str(c["id"]): c for c in st["cands"]}
    new_ivs = [i for i in st["interviews"] if i["id"] > exp["seed_max"]["interviews"]]
    checks, sc_ok, st_ok, book, n_onsite, wrong = {}, 0, 0, 0.0, 0, []
    for cid, want in exp["a_cands"].items():
        c = cands.get(cid, {})
        cards = [s for s in st["scorecards"] if str(s["cand_id"]) == cid]
        card_ok = bool(cards) and (cards[-1]["person_id"] == want["interviewer"] and [cards[-1][k] for k in ("s1", "s2", "s3", "s4")]
                                    == want["scores"] and cards[-1]["rec"] == want["rec"])
        stage_ok = c.get("stage") == want["stage"] and (want["stage"] != "Rejected" or c.get("reason") == "Did not pass interview")
        sc_ok += card_ok
        st_ok += stage_ok
        onsite = [i for i in new_ivs if str(i["cand_id"]) == cid and i["kind"] == "Onsite loop"]
        ch = {"scorecard": card_ok, "stage": f"{c.get('stage')}/{c.get('reason')} (want {want['stage']})"}
        if cards and not card_ok:
            last = cards[-1]
            ch["scorecard_got"] = [last["person_id"], last["s1"], last["s2"], last["s3"], last["s4"], last["rec"]]
        if want["stage"] == "Onsite":
            n_onsite += 1
            spec = exp["onsite"][str(want["req"])]
            req, pool = set(spec["required"]), set(spec["pool"])
            panel_ok = lambda p: req <= p and (len(p - req) == 1 and (p - req) <= pool if pool else p == req)  # noqa: E731
            info, parts = _booking(cal, onsite, want["constraint"], spec["minutes"], panel_ok)
            if parts:
                b = 0.4 * parts["format_and_panel"] + 0.2 * parts["in_availability"] + 0.25 * parts["no_overlap"] + 0.15 * parts["candidate_availability"]
                book += b * (1 if parts["single_booking"] else 0.5)
            ch.update(info, **parts)
        elif onsite:
            wrong.append(want["name"])
        checks[want["name"]] = ch
    n = len(exp["a_cands"])
    # notes on the two requisitions
    notes_ok = 0
    for req in sorted({str(w["req"]) for w in exp["a_cands"].values()}):
        text = "\n".join(x["body"] for x in st["notes"] if x["kind"] == "req" and str(x["ref_id"]) == req and x["user_id"] == exp["jordan_id"])
        names = [w["name"].split()[-1] for w in exp["a_cands"].values() if str(w["req"]) == req]
        ok = bool(text) and all(re.search(nm, text, re.I) for nm in names) and re.search(r"onsite", text, re.I) is not None \
            and re.search(r"review", text, re.I) is not None and re.search(r"reject", text, re.I) is not None
        notes_ok += ok
        checks[f"note_on_{exp['reqs'][req]}"] = ok
    done = st_ok / n
    checks.update({"scorecards_ok": f"{sc_ok}/{n}", "stages_ok": f"{st_ok}/{n}", "onsite_booking_points": round(book, 3),
                   "onsite_for_wrong_candidates": wrong})
    score = 0.28 * sc_ok / n + 0.30 * st_ok / n + 0.30 * book / n_onsite + 0.04 * done * (not wrong) + 0.08 * notes_ok / 2
    clash = [k for k, v in checks.items() if isinstance(v, dict) and v.get("created_double_booking")]
    if clash:
        score = min(score, 0.7)
        checks["cap_double_booking"] = f"{clash} double-booked an interviewer -> score <= 0.7"
    if st_ok < n / 2:
        score = min(score, 0.4)
        checks["cap"] = "fewer than half of the candidates in the right stage -> score <= 0.4"
    return result(score, checks)


# ---------------------------------------------------------------- task B
def _txt(v):
    if isinstance(v, dt.datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and 0 < v < 1:  # an Excel time stored as a day fraction
        m = round(v * 1440)
        return f"{m // 60:02d}:{m % 60:02d}"
    return "" if v is None else str(v)


def _date_in(text, iso):
    d = dt.date.fromisoformat(iso)
    pats = [iso, f"{d.month}/{d.day}/{d.year}", f"{d.month:02d}/{d.day:02d}/{d.year}", f"{d.month}/{d.day}(?!\\d)",
            f"{d.day}\\s*{d.strftime('%b')}", f"{d.strftime('%b')}\\w*\\.?\\s*{d.day}(?!\\d)", f"{d.day}/{d.month}/{d.year}",
            f"{d.day:02d}\\.{d.month:02d}\\.{d.year}"]
    return any(re.search(p, text, re.I) for p in pats)


def _time_in(text, hhmm):
    h, m = int(hhmm[:2]), int(hhmm[3:])
    h12, ap = (h - 1) % 12 + 1, "a" if h < 12 else "p"
    pats = [rf"(?<!\d)0?{h}[:.]{m:02d}(?!\s*[ap]\.?m)(?!\d)", rf"(?<!\d){h12}(:{m:02d})?\s*{ap}\.?m"]
    if m:
        pats = [pats[0], rf"(?<!\d){h12}:{m:02d}\s*{ap}\.?m"]
    return any(re.search(p, text, re.I) for p in pats)


def _sheets(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    return {ws.title: [[c for c in r] for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def evaluate_schedule(env, dest: str) -> dict:
    exp = expected(NAME)
    st = _state(env, exp)
    if not st:
        return result(0.0, {"app_db": "unreadable"})
    cal = Calendar(st, exp["seed_max"]["interviews"])
    cal.week = exp["week"]
    cands = {str(c["id"]): c for c in st["cands"]}
    new_ivs = [i for i in st["interviews"] if i["id"] > exp["seed_max"]["interviews"]]
    pri = exp["priya_id"]
    checks, part, booked_any, latest = {}, 0.0, False, {}
    for cid, want in exp["b_cands"].items():
        ivs = [i for i in new_ivs if str(i["cand_id"]) == cid and i["kind"] == "Hiring manager interview"]
        info, parts = _booking(cal, ivs, want["constraint"], exp["b_minutes"], lambda p: p == {pri})
        stage_ok = cands.get(cid, {}).get("stage") == "HM interview"
        if parts:
            booked_any = True
            latest[cid] = ivs[-1]
            b = 0.3 * parts["format_and_panel"] + 0.2 * parts["in_availability"] + 0.25 * parts["no_overlap"] + 0.15 * parts["candidate_availability"]
            part += b * (1 if parts["single_booking"] else 0.5)
        part += 0.1 * stage_ok
        checks[want["name"]] = dict(info, **parts, stage_hm_interview=stage_ok)
    nb = len(exp["b_cands"])
    # the summary workbook
    apps_score, iv_score = 0.0, 0.0
    path = fetch(env, XLSX, dest + ".xlsx")
    if not path:
        checks["xlsx"] = "missing"
    else:
        try:
            sheets = _sheets(path)
        except Exception as e:  # noqa: BLE001
            sheets, checks["xlsx_error"] = {}, str(e)[:200]
        # applications table
        cand_sheets = sorted(sheets, key=lambda t: 0 if "applic" in t.lower() else 1)
        best, best_detail = 0.0, None
        for title in cand_sheets:
            rows = sheets[title]
            hdr_i = next((i for i, r in enumerate(rows) if sum(any(_txt(c).strip().lower() == s.lower() for c in r) for s in exp["sources"]) >= 3), None)
            if hdr_i is None:
                continue
            col = {s: next((j for j, c in enumerate(rows[hdr_i]) if _txt(c).strip().lower() == s.lower()), None) for s in exp["sources"]}
            good, total, bad = 0, 0, []
            for code, want in exp["applications"].items():
                row = next((r for r in rows[hdr_i + 1:] if any(code.lower() in _txt(c).lower() for c in r)), None)
                for s in exp["sources"]:
                    total += 1
                    v = row[col[s]] if row is not None and col[s] is not None and col[s] < len(row) else None
                    got = 0 if v in (None, "", "-", "—") else (int(v) if isinstance(v, (int, float)) and float(v).is_integer()
                                                              else (int(v) if isinstance(v, str) and v.strip().isdigit() else "?"))
                    if row is not None and got == want.get(s, 0):
                        good += 1
                    else:
                        bad.append(f"{code}/{s}: {v!r} want {want.get(s, 0)}")
            if good / total > best:
                best, best_detail = good / total, {"sheet": title, "cells_ok": f"{good}/{total}", "wrong": bad[:8]}
        apps_score = best
        checks["applications_table"] = best_detail or "no sheet with a requisition x source table"
        # interviews list
        ok_rows = 0
        for cid, iv in latest.items():
            name = exp["b_cands"][cid]["name"]
            hit = False
            for rows in sheets.values():
                for r in rows:
                    text = " | ".join(_txt(c) for c in r)
                    if re.search(re.escape(name), text, re.I) and _date_in(text, iv["date"]) and _time_in(text, iv["start"]):
                        hit = True
            ok_rows += hit
        iv_score = ok_rows / nb
        checks["interview_rows_matching_app"] = f"{ok_rows}/{nb}"
    score = 0.55 * part / nb + 0.25 * apps_score + 0.20 * iv_score
    clash = [k for k, v in checks.items() if isinstance(v, dict) and v.get("created_double_booking")]
    if clash:
        score = min(score, 0.7)
        checks["cap_double_booking"] = f"{clash} double-booked Priya -> score <= 0.7"
    if not booked_any:
        score = min(score, 0.3)
        checks["cap"] = "no hiring-manager interview booked -> score <= 0.3"
    return result(score, checks)
