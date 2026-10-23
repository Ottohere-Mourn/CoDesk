"""X4 - research-group admin at the University of Northbridge: ECCV travel claims vs the publication lists.

build(out_dir) writes:
  app_x4.py                       LabDesk web app (copied)
  labdesk_seed.db                 members, grants, FX table, claim history, publication records of both groups
  thunderbird-profile.tar.gz      Jamie Ross's mailbox (receipts as photos / scans, camera-ready PDFs, decisions, policy)
  ECCV2026_travel_summary.xlsx    summary template Elena wants filled in
  marsh-lab-website.tar.gz        the Marsh Lab website git repo (_bibliography/papers.bib + README conventions)
  expected.json                   ground truth (host only)
"""
from __future__ import annotations

import datetime as dt
import json
import os
import random
import shutil
import subprocess
import sys
import tarfile
import tempfile
from email.utils import parsedate_to_datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lib_x as X  # noqa: E402
from lib_thunderbird import make_profile  # noqa: E402

ME = ("Jamie Ross", "jamie.ross@northbridge.ac.uk")
ELENA = ("Elena Marsh", "e.marsh@northbridge.ac.uk")
CHIDI = ("Chidi Okafor", "c.okafor@northbridge.ac.uk")
FIN = ("School Finance - Claims", "finance-claims@inf.northbridge.ac.uk")
LIST = "ml-groups@lists.northbridge.ac.uk"
BUSINESS_DATE = "2026-10-05"
TZ = dt.timezone(dt.timedelta(hours=1))

GROUPS = [("marsh", "Marsh Lab (Vision & Learning)", "Prof. Elena Marsh"), ("okafor", "Okafor Lab (Language Agents)", "Dr Chidi Okafor")]
GRANTS = [("EP/X031452/1", "Robust Visual Learning from Imperfect Data", "EPSRC", "marsh", "2028-03-31", 12000),
          ("ERC-101115274", "SCENEGRAPH: Structured Scene Understanding", "ERC Starting Grant", "marsh", "2029-08-31", 20000),
          ("SOI-MARSH-DISC", "Marsh Lab discretionary fund", "School of Informatics", "marsh", "2027-07-31", 3000),
          ("EP/Y004117/1", "Trustworthy Language Agents", "EPSRC", "okafor", "2028-09-30", 10000),
          ("EP/W524311/1", "EPSRC DTP studentship (M. Tanaka)", "EPSRC DTP", "okafor", "2029-09-30", 1500)]
# id: (name, email, group, role, default grant)
MEMBERS = {1: ("Elena Marsh", "e.marsh@northbridge.ac.uk", "marsh", "Professor (PI)", "EP/X031452/1"),
           2: ("Tomás Ruiz", "t.ruiz@northbridge.ac.uk", "marsh", "Research Fellow", "EP/X031452/1"),
           3: ("Aisha Rahman", "a.rahman@northbridge.ac.uk", "marsh", "PhD student", "EP/X031452/1"),
           4: ("Ben Carter", "b.carter@northbridge.ac.uk", "marsh", "PhD student", "ERC-101115274"),
           5: ("Lena Vogel", "l.vogel@northbridge.ac.uk", "marsh", "PhD student", "ERC-101115274"),
           6: ("Wei Zhang", "w.zhang@northbridge.ac.uk", "marsh", "PhD student", "SOI-MARSH-DISC"),
           7: ("Chidi Okafor", "c.okafor@northbridge.ac.uk", "okafor", "Reader (PI)", "EP/Y004117/1"),
           8: ("Priya Natarajan", "p.natarajan@northbridge.ac.uk", "okafor", "PhD student", "EP/Y004117/1"),
           9: ("Sam O'Neill", "s.oneill@northbridge.ac.uk", "okafor", "Research Fellow", "EP/Y004117/1"),
           10: ("Mei Tanaka", "m.tanaka@northbridge.ac.uk", "okafor", "PhD student", "EP/W524311/1")}

# What each traveller may claim (policy applied): member id -> grant, lines (date, currency, amount, what)
CLAIMS = {
    2: ("EP/X031452/1", [("2026-08-04", "GBP", 262.40, "BA e-ticket LHR-CPH return, fare + taxes (seat selection £24.00 excluded)"),
                         ("2026-09-12", "SEK", 9000.00, "Hotel Kvarnhuset 5 nights, capped at SEK 1,800/night (minibar excluded)"),
                         ("2026-09-09", "SEK", 570.00, "Dinner Saltkråkan, food only (wine SEK 210 excluded)")]),
    3: ("ERC-101115274", [("2026-08-11", "GBP", 241.90, "BA e-ticket LHR-CPH return"),
                          ("2026-09-11", "SEK", 6600.00, "Hotel Kvarnhuset 4 nights x SEK 1,650"),
                          ("2026-09-10", "SEK", 650.00, "Meals 10 Sep (SEK 185 + 690) capped at SEK 650/day"),
                          ("2026-09-07", "SEK", 125.00, "Öresundståg CPH airport - Malmö C")]),
    4: ("ERC-101115274", [("2026-09-12", "SEK", 380.00, "Taxi Skåne hotel - Malmö C")]),
    8: ("EP/Y004117/1", [("2026-08-20", "GBP", 312.75, "SAS e-ticket LHR-CPH return"),
                         ("2026-09-11", "SEK", 5400.00, "Hotel Live Öresund 3 nights, capped at SEK 1,800/night"),
                         ("2026-09-09", "SEK", 445.00, "Dinner Smakverket, food only (beer SEK 95 excluded)")]),
}

P_NEW = {  # B: the papers that must end up in LabDesk (and, for Marsh Lab, in papers.bib)
    "P1": {"group": "marsh", "title": "Sparse Prompt Tuning for Open-Vocabulary Semantic Segmentation",
           "old_title": "Sparse Prompt Tuning for Open-Vocabulary Segmentation",
           "authors": ["Tomás Ruiz", "Aisha Rahman", "Wei Zhang", "Elena Marsh"], "venue_keys": ["neurips", "neuralinformationprocessing"],
           "year": 2026, "arxiv": "2605.08812", "bibkey": "ruiz2026sparse", "type": "conf"},
    "P2": {"group": "okafor", "title": "Calibrated Refusals for Tool-Using Language Agents",
           "authors": ["Priya Natarajan", "Sam O'Neill", "Chidi Okafor"], "venue_keys": ["emnlp"], "year": 2026, "arxiv": "2609.11873",
           "bibkey": None, "type": "conf"},
    "P3": {"group": "marsh", "title": "Few-Shot Scene Graph Generation with Relation Priors",
           "authors": ["Ben Carter", "Lena Vogel", "Elena Marsh"], "venue_keys": ["wacv", "winterconferenceonapplications"], "year": 2027,
           "arxiv": "", "bibkey": "carter2027fewshot", "type": "conf"},
    "P4": {"group": "marsh", "title": "Benchmarking Monocular Depth Foundation Models under Adverse Weather",
           "authors": ["Lena Vogel", "Tomás Ruiz", "Elena Marsh"], "venue_keys": ["arxiv"], "year": 2026, "arxiv": "2609.15502",
           "bibkey": "vogel2026benchmarking", "type": "preprint"},
}
REJECTED = {"title": "Token Merging for Efficient Video Transformers", "authors": ["Wei Zhang", "Elena Marsh"]}

# existing publication records: (group, title, authors, venue, year, type, arxiv, bibtype, bibvenue)
OLD_PUBS = [
    ("marsh", "Sparse Prompt Tuning for Open-Vocabulary Segmentation", "Tomás Ruiz, Aisha Rahman, Elena Marsh", "arXiv preprint", 2026,
     "Preprint", "2605.08812", "article", "arXiv preprint arXiv:2605.08812"),
    ("marsh", "Relation-Aware Distillation for Panoptic Scene Graphs", "Ben Carter, Elena Marsh", "CVPR 2026", 2026, "Conference paper", "2603.02214",
     "inproceedings", "IEEE/CVF Conference on Computer Vision and Pattern Recognition (CVPR)"),
    ("marsh", "Learning with Noisy Correspondence for Video-Text Retrieval", "Wei Zhang, Tomás Ruiz, Elena Marsh", "ICCV 2025", 2025,
     "Conference paper", "2507.01133", "inproceedings", "IEEE/CVF International Conference on Computer Vision (ICCV)"),
    ("marsh", "Self-Supervised Depth Completion in Fog", "Lena Vogel, Elena Marsh", "BMVC 2025", 2025, "Conference paper", "2508.09920",
     "inproceedings", "British Machine Vision Conference (BMVC)"),
    ("marsh", "Prompt Ensembles Are Robust Zero-Shot Segmenters", "Aisha Rahman, Tomás Ruiz, Elena Marsh", "WACV 2026", 2026,
     "Conference paper", "2510.04471", "inproceedings", "IEEE/CVF Winter Conference on Applications of Computer Vision (WACV)"),
    ("marsh", "A Survey of Robust Visual Learning under Label Noise", "Elena Marsh, Tomás Ruiz", "International Journal of Computer Vision", 2025,
     "Journal article", "2411.07762", "article", "International Journal of Computer Vision"),
    ("marsh", "Graph Priors for Long-Tail Predicate Classification", "Ben Carter, Lena Vogel, Elena Marsh", "ECCV 2024 Workshops", 2024,
     "Workshop paper", "", "inproceedings", "European Conference on Computer Vision Workshops (ECCVW)"),
    ("marsh", "Weakly Supervised Instance Segmentation from Point Labels", "Tomás Ruiz, Elena Marsh", "NeurIPS 2024", 2024,
     "Conference paper", "2410.05518", "inproceedings", "Advances in Neural Information Processing Systems (NeurIPS)"),
    ("okafor", "Tool-Use Hallucinations in LLM Agents: A Taxonomy", "Sam O'Neill, Priya Natarajan, Chidi Okafor", "ACL 2026", 2026,
     "Conference paper", "2602.08841", "", ""),
    ("okafor", "When Should an Agent Ask? Clarification Policies for Web Tasks", "Mei Tanaka, Chidi Okafor", "arXiv preprint", 2026,
     "Preprint", "2607.03318", "", ""),
    ("okafor", "Measuring Over-Refusal in Instruction-Tuned Models", "Priya Natarajan, Chidi Okafor", "EMNLP 2025", 2025,
     "Conference paper", "2509.11206", "", ""),
    ("okafor", "Grounded Planning with Retrieval for Household Agents", "Sam O'Neill, Chidi Okafor", "TACL", 2025, "Journal article",
     "2502.06617", "", ""),
    ("okafor", "Interactive Evaluation of Language Agents", "Chidi Okafor", "NAACL 2024", 2024, "Conference paper", "2404.01123", "", ""),
]


def bibkey(title, authors, year):
    import re
    import unicodedata
    last = authors.split(",")[0].strip().split()[-1]
    words = [w for w in re.split(r"\s+", title) if w.lower() not in ("a", "an", "the")]
    norm = lambda s: re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower())  # noqa: E731
    return f"{norm(last)}{year}{norm(words[0])}"


def bib_authors(authors):
    out = []
    for a in authors.split(","):
        parts = a.strip().split()
        name = f"{parts[-1]}, {' '.join(parts[:-1])}"
        out.append(name)
    s = " and ".join(out)
    return s.replace("á", "{\\'a}")


def bib_entry(p):
    g, title, authors, venue, year, typ, arxiv, bt, bv = p
    key = bibkey(title, authors, year)
    field = "booktitle" if bt == "inproceedings" else "journal"
    s = f"@{bt}{{{key},\n  title     = {{{title}}},\n  author    = {{{bib_authors(authors)}}},\n  {field:<9} = {{{bv}}},\n  year      = {{{year}}},\n"
    if arxiv:
        s += f"  arxiv     = {{{arxiv}}},\n"
    return s + "}\n"


# ---------------------------------------------------------------- documents
def slip(lines, width=560, size=22, seed=0, mono=True):
    """A till receipt: lines are str or (left, right) tuples; '---' draws a rule; ('#', text) is bold centred."""
    h = 80 + sum(int(size * 1.6) if l != "---" else 18 for l in lines) + 40
    p = X.Page(width, h, bg=(252, 251, 246))
    y = 50
    for l in lines:
        if l == "---":
            p.d.line((30, y + 6, width - 30, y + 6), fill=(90, 90, 90), width=2)
            y += 18
            continue
        if isinstance(l, tuple) and l[0] == "#":
            p.text(width // 2, y, l[1], size=size + 4, bold=True, anchor="ma")
        elif isinstance(l, tuple):
            p.text(30, y, l[0], size=size, mono=mono)
            p.text(width - 30, y, l[1], size=size, mono=mono, anchor="ra")
        else:
            p.text(width // 2, y, l, size=size - 2, mono=mono, anchor="ma")
        y += int(size * 1.6)
    return p


def on_a4(sl, note=None):
    p = X.Page()
    p.im.paste(sl.im, (150, 120))
    if note:
        p.handwriting(760, 200, note, size=30, seed=3)
    return p


def eticket(airline, color, ref, pax, segs, fare, total, issued, card):
    p = X.Page()
    p.d.rectangle((0, 0, X.W, 150), fill=color)
    p.text(80, 50, airline, size=46, bold=True, fill=(255, 255, 255))
    p.text(1160, 62, "e-ticket receipt", size=28, fill=(255, 255, 255), anchor="ra")
    y = p.lines(80, 210, [f"Booking reference: {ref}", f"Passenger: {pax}", f"Date of issue: {issued}"], size=26)
    y = p.table(80, y + 30, [("Flight", 170, "l"), ("Date", 190, "l"), ("From", 250, "l"), ("To", 250, "l"), ("Class", 220, "l")],
                segs, size=21, row_h=40)
    p.text(80, y + 50, "Payment summary", size=28, bold=True)
    y = p.table(80, y + 100, [("Item", 760, "l"), ("Amount", 320, "r")], fare + [("Total paid", total)], size=22, row_h=42)
    p.lines(80, y + 40, [f"Paid by {card}", "Seat selection and insurance are optional extras chosen at booking.",
                         "This receipt is valid as proof of payment."], size=20, fill=(80, 80, 80))
    return p


def hotel_invoice(name, addr, inv, guest, arrive, depart, rows, total, date, vat):
    p = X.Page()
    p.text(620, 90, name, size=50, bold=True, serif=True, anchor="ma", fill=(60, 40, 30))
    p.text(620, 160, addr, size=20, anchor="ma", fill=(90, 90, 90))
    p.hline(80, 1160, 205)
    p.lines(80, 240, [f"FAKTURA / INVOICE  {inv}", f"Guest: {guest}", f"Arrival: {arrive}    Departure: {depart}", f"Date: {date}"], size=25)
    y = p.table(80, 470, [("Description", 560, "l"), ("Qty", 110, "r"), ("Price SEK", 200, "r"), ("Amount SEK", 210, "r")], rows,
                size=21, row_h=42)
    p.text(1160, y + 30, f"Total SEK {total}", size=30, bold=True, anchor="ra")
    p.lines(80, y + 110, [vat, "Paid with card - thank you for staying with us!"], size=20, fill=(80, 80, 80))
    p.stamp(780, y + 200, "BETALD / PAID", angle=8)
    return p


def registration_receipt(name):
    p = X.Page()
    p.d.rectangle((0, 0, X.W, 170), fill=(26, 74, 140))
    p.text(80, 55, "ECCV 2026  ·  Malmö", size=48, bold=True, fill=(255, 255, 255))
    p.lines(80, 230, ["Registration payment receipt", "", f"Attendee: {name}", "Order: ECCV26-R-18842", "Date: 15 Jul 2026",
                      "Item: Full conference registration (early, non-student)", "Amount: EUR 750.00", "",
                      "Paid by card: VISA ****8812  (UNIV OF NORTHBRIDGE INF PCARD / J ROSS)"], size=26)
    return p


def paper_pdf(path, title, authors, affil, footer, abstract):
    import fitz
    doc = fitz.open()
    pg = doc.new_page(width=612, height=792)
    pg.insert_textbox(fitz.Rect(72, 80, 540, 150), title, fontsize=16, fontname="hebo", align=1)
    pg.insert_textbox(fitz.Rect(72, 160, 540, 200), "   ".join(authors), fontsize=11, fontname="helv", align=1)
    pg.insert_textbox(fitz.Rect(72, 200, 540, 230), affil, fontsize=9, fontname="helv", align=1)
    pg.insert_textbox(fitz.Rect(72, 250, 540, 270), "Abstract", fontsize=11, fontname="hebo", align=1)
    pg.insert_textbox(fitz.Rect(100, 272, 512, 520), abstract, fontsize=9.5, fontname="helv", align=3)
    pg.insert_text((72, 560), "1  Introduction", fontsize=11, fontname="hebo")
    pg.insert_textbox(fitz.Rect(72, 570, 540, 720), "Recent progress " + abstract[:600], fontsize=9.5, fontname="helv", align=3)
    pg.insert_textbox(fitz.Rect(72, 740, 540, 760), footer, fontsize=8, fontname="helv", align=1)
    pg2 = doc.new_page(width=612, height=792)
    pg2.insert_textbox(fitz.Rect(72, 72, 540, 720), (abstract + " ") * 4, fontsize=9.5, fontname="helv", align=3)
    doc.set_metadata({"title": title, "author": ", ".join(authors), "creator": "LaTeX with hyperref", "producer": "pdfTeX-1.40.26"})
    doc.save(path, garbage=3, deflate=True)


def summary_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "ECCV 2026"
    ws["A1"] = "ECCV 2026 (Malmö, 7-12 Sep) - travel claims summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "For Prof. Marsh's records - one row per traveller, totals as submitted in LabDesk."
    for j, h in enumerate(["Traveller", "Group", "Grant", "LabDesk claim no.", "Total claimed (GBP)"], 1):
        c = ws.cell(4, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6EE")
    ws["A11"] = "Totals by grant"
    ws["A11"].font = Font(bold=True)
    for j, h in enumerate(["Grant", "Total (GBP)"], 1):
        c = ws.cell(12, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6EE")
    for col, w in zip("ABCDE", (22, 30, 20, 20, 20)):
        ws.column_dimensions[col].width = w
    wb.save(path)


# ---------------------------------------------------------------- website repo
README = """# Marsh Lab website

Jekyll site (al-folio theme) for the Vision & Learning group of Prof. Elena Marsh. The publications page is generated
from `_bibliography/papers.bib`.

## Adding publications

* Only **Marsh Lab** papers go on this site (the Okafor Lab has its own site).
* BibTeX key: surname of the first author + year + first word of the title, ignoring a leading "A", "An" or "The";
  all lowercase, letters and digits only (accents and hyphens dropped). Example: `marsh2025survey`.
* Conference and workshop papers: `@inproceedings` with `booktitle` = the full venue name followed by the acronym in
  parentheses, e.g. `IEEE/CVF Winter Conference on Applications of Computer Vision (WACV)`.
* Preprints: `@article` with `journal = {arXiv preprint arXiv:<id>}`.
* Always give `title`, `author` (`Surname, Given and Surname, Given`, in paper order), `year`, and `arxiv = {<id>}` when
  the paper is on arXiv.
* When a preprint is accepted, update its existing entry (keep the key) instead of adding a second one.
* Keep the file sorted newest first. Commit bibliography changes on their own, e.g. "Add new publications (Oct 2026)".
"""


def build_repo(d):
    tmp = tempfile.mkdtemp()
    repo = f"{tmp}/marsh-lab-website"
    env = {**os.environ, "GIT_AUTHOR_NAME": "Jamie Ross", "GIT_AUTHOR_EMAIL": ME[1], "GIT_COMMITTER_NAME": "Jamie Ross",
           "GIT_COMMITTER_EMAIL": ME[1]}

    def git(*a, when="2026-09-01T10:00:00"):
        subprocess.run(["git", "-C", repo, *a], check=True, capture_output=True, env={**env, "GIT_AUTHOR_DATE": when, "GIT_COMMITTER_DATE": when})

    os.makedirs(f"{repo}/_bibliography")
    os.makedirs(f"{repo}/_pages")
    subprocess.run(["git", "init", "-q", "-b", "main", repo], check=True)
    files = {"README.md": README,
             "_config.yml": "title: Marsh Lab - Vision & Learning\nurl: https://marshlab.inf.northbridge.ac.uk\nscholar:\n  bibliography: papers.bib\n  sort_by: year\n  order: descending\n",
             "_pages/publications.md": "---\nlayout: page\ntitle: Publications\npermalink: /publications/\n---\n\n{% bibliography %}\n",
             "index.md": "---\nlayout: about\ntitle: Home\n---\n\nWe study robust visual learning and structured scene understanding.\n"}
    for p, s in files.items():
        open(f"{repo}/{p}", "w").write(s)
    git("add", ".", when="2025-11-03T09:00:00")
    git("commit", "-q", "-m", "Initial site", when="2025-11-03T09:00:00")
    marsh = [p for p in OLD_PUBS if p[0] == "marsh"]
    bib = "% Marsh Lab publications - newest first\n\n" + "\n".join(bib_entry(p) for p in marsh)
    open(f"{repo}/_bibliography/papers.bib", "w").write(bib)
    git("add", ".", when="2026-06-15T14:00:00")
    git("commit", "-q", "-m", "Add publications (June 2026)", when="2026-06-15T14:00:00")
    with open(f"{repo}/.git/config", "a") as f:
        f.write(f"[user]\n\tname = Jamie Ross\n\temail = {ME[1]}\n")
    base = subprocess.run(["git", "-C", repo, "rev-parse", "HEAD"], capture_output=True, text=True).stdout.strip()
    with tarfile.open(f"{d}/marsh-lab-website.tar.gz", "w:gz") as t:
        t.add(repo, arcname="marsh-lab-website")
    shutil.rmtree(tmp)
    return base, [bibkey(p[1], p[2], p[4]) for p in marsh]


# ---------------------------------------------------------------- mailbox
def mailbox(d, docs):
    when = lambda s: dt.datetime.fromisoformat(s).replace(tzinfo=TZ)  # noqa: E731
    ids = {}

    def mid(key, dom="northbridge.ac.uk"):
        ids[key] = f"<x4.{key}.{len(ids) + 1:02d}@{dom}>"
        return ids[key]

    rd = lambda n: open(f"{d}/_docs/{n}", "rb").read()  # noqa: E731
    pdf, jpg = "application/pdf", "image/jpeg"
    M = X.mail
    inbox = [
        (M(("Informatics Seminars", "seminars@inf.northbridge.ac.uk"), [ME], "Seminar Thu 17 Sep: Causal representation learning",
           "Speaker: Dr Hana Kim (Edinburgh). 14:00, Forum G.07. All welcome.", when("2026-09-11T09:00:00"), mid("seminar", "inf.northbridge.ac.uk")), False),
        (M(("Tomás Ruiz", MEMBERS[2][1]), [ME], "ECCV receipts + good news",
           "Hi Jamie,\n\nAll my receipts from Malmö are in the attached PDF (scanned in the office). I put the registration in "
           "too for completeness, although I think that went on your card.\n\nPS: our NeurIPS paper got accepted!! The "
           "camera-ready is attached - note the final title is slightly different from the arXiv version and Wei is now a "
           "co-author.\n\nThanks!\nTomás", when("2026-09-14T10:41:00"), mid("tomas"),
           atts=[("Tomas_Ruiz_ECCV2026_receipts.pdf", rd("Tomas_Ruiz_ECCV2026_receipts.pdf"), pdf),
                 ("ruiz_neurips2026_camera_ready.pdf", rd("ruiz_neurips2026_camera_ready.pdf"), pdf)]), False),
        (M(("Aisha Rahman", MEMBERS[3][1]), [ME], "ECCV claim - receipts",
           "Hi Jamie,\n\nHere are my ECCV receipts: the BA confirmation and photos of the hotel bill and my lunch and dinner on the "
           "10th. Elena said my trip should go on the ERC grant this time, not the EPSRC one.\n\nBest,\nAisha",
           when("2026-09-14T16:05:00"), mid("aisha1"),
           atts=[("BA_confirmation_RAHMAN.pdf", rd("BA_confirmation_RAHMAN.pdf"), pdf), ("IMG_20260911_0812.jpg", rd("aisha_hotel.jpg"), jpg),
                 ("IMG_20260910_1302.jpg", rd("aisha_lunch.jpg"), jpg), ("IMG_20260910_2117.jpg", rd("aisha_dinner.jpg"), jpg)]), False),
        (M(("Ben Carter", MEMBERS[4][1]), [ME], "Receipts from Malmö",
           "Hi Jamie,\n\nFlights and hotel were booked by the department's travel agent, so I only have two receipts: the taxi to the "
           "station on the last day and some sun cream and toothpaste I had to buy (not sure that counts?).\n\nCheers,\nBen",
           when("2026-09-15T09:12:00"), mid("ben"),
           atts=[("taxi.jpg", rd("ben_taxi.jpg"), jpg), ("apotek.jpg", rd("ben_pharmacy.jpg"), jpg)]), True),
        (M(("Aisha Rahman", MEMBERS[3][1]), [ME], "Re: ECCV claim - receipts",
           "Sorry, forgot the train ticket from Copenhagen airport to Malmö - photo attached.\nAisha",
           when("2026-09-16T08:30:00"), mid("aisha2"), irt=ids["aisha1"], atts=[("IMG_20260907_1544.jpg", rd("aisha_train.jpg"), jpg)]), True),
        (M(("Priya Natarajan", MEMBERS[8][1]), [ME], "My ECCV receipts (Okafor lab)",
           "Hi Jamie,\n\nChidi said you'd kindly do my ECCV claim too - receipts scanned into one PDF. My usual grant.\n\n"
           "Also, our EMNLP paper got into Findings! Camera-ready attached; the arXiv version is 2609.11873.\n\nThank you!\nPriya",
           when("2026-09-17T11:20:00"), mid("priya"), cc=[CHIDI],
           atts=[("Natarajan_ECCV_receipts.pdf", rd("Natarajan_ECCV_receipts.pdf"), pdf),
                 ("natarajan_emnlp2026_findings_camera_ready.pdf", rd("natarajan_emnlp2026_findings_camera_ready.pdf"), pdf)]), True),
        (M(("Wei Zhang", MEMBERS[6][1]), [ME], "AAAI decision :(",
           "Hi Jamie,\n\nJust so you know - our AAAI-27 submission \"Token Merging for Efficient Video Transformers\" (me and Elena) "
           "was rejected. We'll revise it for CVPR, so please don't put it on the website.\n\nWei", when("2026-09-22T18:02:00"), mid("wei")), True),
        (M(ELENA, [ME], "Fwd: WACV 2027 Paper Decision - Paper #412",
           "Great news for Ben and Lena - please add it to the lists.\nE.\n\n---------- Forwarded message ----------\n"
           "From: WACV 2027 Program Chairs <wacv2027-pcs@openreview.net>\nSubject: WACV 2027 Paper Decision - Paper #412\n\n"
           "Dear Ben Carter,\n\nWe are pleased to inform you that your paper\n\n  #412 Few-Shot Scene Graph Generation with Relation Priors\n"
           "  Authors: Ben Carter (University of Northbridge), Lena Vogel (University of Northbridge), Elena Marsh (University of Northbridge)\n\n"
           "has been ACCEPTED to the IEEE/CVF Winter Conference on Applications of Computer Vision (WACV 2027), Tucson, Arizona, "
           "March 2027.\n\nThe camera-ready deadline is 30 October 2026.\n\nWACV 2027 Program Chairs", when("2026-09-24T07:45:00"), mid("wacv")), True),
        (M(FIN, [ME], "ECCV 2026 travel claims - deadline 16 October",
           "Dear colleagues,\n\nPlease submit ECCV 2026 travel claims in LabDesk by Friday 16 October. The relevant extract of the "
           "University Travel & Expenses Policy is attached - in particular the overseas accommodation and subsistence caps for "
           "Sweden. Once submitted, email us the list of LabDesk claim numbers with their GBP totals so we can batch them.\n\n"
           "School Finance Office", when("2026-09-28T09:00:00"), mid("finance", "inf.northbridge.ac.uk"),
           atts=[("Travel_Expenses_Policy_extract_2026.pdf", rd("Travel_Expenses_Policy_extract_2026.pdf"), pdf)]), True),
        (M(("Lena Vogel", MEMBERS[5][1]), [ME], "Our weather benchmark is on arXiv",
           "Hi Jamie, the depth benchmark preprint went out last night - announcement below. Could it go on the website?\nLena\n\n"
           "------------------------------------------------------------------------------\narXiv:2609.15502\nDate: Wed, 30 Sep 2026 "
           "17:59:02 GMT   (8412kb)\n\nTitle: Benchmarking Monocular Depth Foundation Models under Adverse Weather\nAuthors: Lena Vogel, "
           "Tomás Ruiz, Elena Marsh\nCategories: cs.CV\nComments: 14 pages, 9 figures\n\n  We evaluate eleven monocular depth "
           "foundation models on rain, fog and snow ...\n( https://arxiv.org/abs/2609.15502 )", when("2026-10-01T08:10:00"), mid("arxiv")), True),
        (M(("IT Services", "it-help@northbridge.ac.uk"), [ME], "Planned outage: VPN Saturday 10 Oct",
           "The VPN will be unavailable 07:00-09:00 on Saturday 10 October.", when("2026-10-02T12:00:00"), mid("it")), True),
        (M(ELENA, [ME], "ECCV claims",
           "Hi Jamie,\n\nCould you do the ECCV travel claims in LabDesk for Tomás, Aisha and Ben? Chidi asked whether you could do "
           "Priya's as well, which is fine by me. Their receipts are in your inbox; Finance's policy extract has the caps.\n\n"
           "The registrations were all paid on the School card back in July, so they must not be claimed again. Aisha's trip goes "
           "on the ERC grant (SCENEGRAPH); everyone else on their usual grant.\n\nPlease fill in the summary sheet in "
           "Documents/Lab for my records and send Finance the claim numbers.\n\nThanks!\nElena", when("2026-10-02T15:30:00"), mid("elena_claims")), True),
        (M(ELENA, [ME], "Website publications update",
           "Hi Jamie,\n\nBefore the open day, could you bring our publication lists up to date? All the news from the last few weeks "
           "is in your inbox (acceptances, Lena's arXiv preprint, camera-readies).\n\n- LabDesk Publications for both groups (Chidi "
           "asked for his group's paper too): exact title and author order as in the camera-ready or arXiv listing, venue, year "
           "and arXiv id where there is one. If a paper is already listed as a preprint, update that record rather than adding "
           "a second one.\n- Our website repo: add/update the BibTeX entries in _bibliography/papers.bib following the README, and "
           "commit.\n- Then a short announcement to the groups' list (ml-groups@lists.northbridge.ac.uk) listing the new papers.\n\n"
           "Nothing that was rejected, obviously.\n\nElena", when("2026-10-05T08:55:00"), mid("elena_pubs")), True),
    ]
    old = [(M(ME, [FIN], "BMVC 2025 claims", "Claim numbers CLM-2025-0412, CLM-2025-0413 submitted.", when("2025-11-29T10:00:00"),
              mid("bmvc_sent")), False)]
    folders = {"Inbox": [(m, parsedate_to_datetime(m["Date"]), u) for m, u in inbox],
               "Sent": [(m, parsedate_to_datetime(m["Date"]), False) for m, u in old],
               "Claims": [],
               "Claims/BMVC 2025": [(M(("Lena Vogel", MEMBERS[5][1]), [ME], "BMVC receipts", "Receipts attached (already claimed).",
                                       when("2025-11-27T09:00:00"), mid("bmvc")), parsedate_to_datetime("Fri, 27 Nov 2025 09:00:00 +0100"), False)]}
    make_profile(f"{d}/thunderbird-profile.tar.gz", ME, folders)
    return ids


def make_docs(d):
    o = f"{d}/_docs"
    os.makedirs(o)
    # Tomás: one scanned 4-page PDF
    pages = [registration_receipt("Tomás Ruiz"),
             eticket("British Airways", (7, 52, 107), "XK7P2Q", "RUIZ/TOMAS MR",
                     [("BA816", "07 Sep 2026", "London LHR", "Copenhagen CPH", "Economy"), ("BA817", "12 Sep 2026", "Copenhagen CPH", "London LHR", "Economy")],
                     [("Air fare", "GBP 201.00"), ("Taxes, fees and carrier charges", "GBP 61.40"), ("Seat selection (2 seats, extra legroom)", "GBP 24.00")],
                     "GBP 286.40", "04 Aug 2026", "VISA ****4410 (personal card)"),
             hotel_invoice("Hotel Kvarnhuset", "Kvarngatan 14, 211 47 Malmö · org.nr 556712-3301", "55102", "Tomás Ruiz", "07.09.2026", "12.09.2026",
                           [("Superior double room", "5", "1 950,00", "9 750,00"), ("Minibar 11.09", "1", "120,00", "120,00")],
                           "9 870,00", "12.09.2026", "incl. moms 12% (room) / 25% (minibar)"),
             on_a4(slip([("#", "RESTAURANG SALTKRÅKAN"), "Lilla Torg 7, Malmö", "Kvitto 20481   09.09.2026 21:14", "---",
                         ("1 Toast Skagen", "165,00"), ("1 Grillad torsk", "260,00"), ("1 Ostbricka", "145,00"), ("2 Glas Riesling", "210,00"), "---",
                         ("TOTALT SEK", "780,00"), ("Kort VISA ****4410", "780,00"), "Varav moms 12%/25%", "Tack för besöket!"], seed=1))]
    X.scanned_pdf(f"{o}/Tomas_Ruiz_ECCV2026_receipts.pdf", pages, seed=41)
    paper_pdf(f"{o}/ruiz_neurips2026_camera_ready.pdf", P_NEW["P1"]["title"], P_NEW["P1"]["authors"],
              "School of Informatics, University of Northbridge, UK", "40th Conference on Neural Information Processing Systems (NeurIPS 2026).",
              "Open-vocabulary semantic segmentation models inherit rich visual-language alignment from large pretrained encoders, but "
              "adapting them to new domains usually requires tuning millions of parameters. We introduce sparse prompt tuning, which "
              "learns a small set of prompts attached to a sparse subset of attention heads selected by a gradient-based saliency "
              "criterion. With 0.1% of the trainable parameters of full fine-tuning, our method improves mIoU on ADE-847, PC-459 and "
              "Cityscapes and transfers across backbones.")
    # Aisha: scanned confirmation + phone photos
    X.scanned_pdf(f"{o}/BA_confirmation_RAHMAN.pdf", [eticket(
        "British Airways", (7, 52, 107), "LM3R8D", "RAHMAN/AISHA MS",
        [("BA816", "07 Sep 2026", "London LHR", "Copenhagen CPH", "Economy"), ("BA815", "11 Sep 2026", "Copenhagen CPH", "London LHR", "Economy")],
        [("Air fare", "GBP 180.50"), ("Taxes, fees and carrier charges", "GBP 61.40")], "GBP 241.90", "11 Aug 2026", "MASTERCARD ****0077")], seed=42)
    hotel_a = hotel_invoice("Hotel Kvarnhuset", "Kvarngatan 14, 211 47 Malmö · org.nr 556712-3301", "55087", "Aisha Rahman", "07.09.2026",
                            "11.09.2026", [("Standard single room", "4", "1 650,00", "6 600,00")], "6 600,00", "11.09.2026", "incl. moms 12%")
    X.photo_jpg(f"{o}/aisha_hotel.jpg", hotel_a, seed=5, size=(1200, 1600))
    X.photo_jpg(f"{o}/aisha_lunch.jpg", slip([("#", "MALMÖ SALUHALL"), "Gibraltargatan 6", "Kvitto 7713   10.09.2026 13:02", "---",
                                               ("Fiskgryta", "165,00"), ("Mineralvatten", "20,00"), "---", ("TOTAL SEK", "185,00"),
                                               ("Mastercard", "185,00")], seed=2), seed=6, size=(900, 1300))
    X.photo_jpg(f"{o}/aisha_dinner.jpg", slip([("#", "BORD 21 · VÄSTRA HAMNEN"), "Dockplatsen 1, Malmö", "Nota 3302   10.09.2026 21:17", "---",
                                                ("Löjromstoast", "195,00"), ("Halstrad röding", "285,00"), ("2 Alkoholfri öl", "90,00"),
                                                ("Kaffe & tryffel", "120,00"), "---", ("ATT BETALA SEK", "690,00"), ("Mastercard", "690,00")], seed=3),
                  seed=7, size=(900, 1400))
    X.photo_jpg(f"{o}/aisha_train.jpg", slip([("#", "Öresundståg"), "Skånetrafiken", "Köpenhamn Lufthavn -> Malmö C", "07.09.2026 15:44", "---",
                                               ("1 Vuxen enkel", "125,00"), "---", ("Totalt SEK", "125,00"), "Giltig 07.09.2026"], seed=4),
                  seed=8, size=(900, 1100))
    X.photo_jpg(f"{o}/ben_taxi.jpg", slip([("#", "TAXI SKÅNE"), "Bil 2214   Org.nr 556010-1144", "12.09.2026  07:52", "Kvarngatan 14 -> Malmö C",
                                            "---", ("Taxameter", "320,00"), ("Bagage/väntetid", "60,00"), "---", ("TOTALT SEK", "380,00"),
                                            ("Kort", "380,00")], seed=5), seed=9, size=(900, 1200))
    X.photo_jpg(f"{o}/ben_pharmacy.jpg", slip([("#", "APOTEK HJÄRTAT"), "Södra Förstadsgatan 2", "09.09.2026  18:21", "---",
                                                ("Solkräm SPF30", "50,00"), ("Tandkräm", "39,00"), "---", ("SUMMA SEK", "89,00"), ("Kort", "89,00")],
                                               seed=6), seed=10, size=(900, 1100))
    # Priya: one scanned 3-page PDF
    X.scanned_pdf(f"{o}/Natarajan_ECCV_receipts.pdf", [
        eticket("SAS Scandinavian Airlines", (0, 33, 82), "8PQ4ZT", "NATARAJAN/PRIYA MS",
                [("SK502", "08 Sep 2026", "London LHR", "Copenhagen CPH", "SAS Go"), ("SK505", "11 Sep 2026", "Copenhagen CPH", "London LHR", "SAS Go")],
                [("Fare", "GBP 249.00"), ("Taxes and charges", "GBP 63.75")], "GBP 312.75", "20 Aug 2026", "VISA ****9031"),
        hotel_invoice("Hotel Live Öresund", "Dag Hammarskjölds Torg 2, 211 18 Malmö", "LO-88215", "Priya Natarajan", "08.09.2026", "11.09.2026",
                      [("Moderate double room", "3", "2 100,00", "6 300,00")], "6 300,00", "11.09.2026", "incl. moms 12%"),
        on_a4(slip([("#", "SMAKVERKET"), "Stortorget 9, Malmö", "Kvitto 1187   09.09.2026 20:40", "---", ("Svamprisotto", "245,00"),
                    ("Pannacotta", "110,00"), ("Kaffe", "45,00"), ("Mineralvatten", "45,00"), ("Stor stark (öl)", "95,00"), "---",
                    ("TOTAL SEK", "540,00"), ("Visa", "540,00")], seed=7))], seed=43)
    paper_pdf(f"{o}/natarajan_emnlp2026_findings_camera_ready.pdf", P_NEW["P2"]["title"], P_NEW["P2"]["authors"],
              "School of Informatics, University of Northbridge, UK", "Findings of the Association for Computational Linguistics: EMNLP 2026",
              "Tool-using language agents must decide when to refuse a request whose execution would be unsafe or impossible. We show "
              "that current agents are badly calibrated: they over-refuse benign tool calls and under-refuse harmful ones. We propose "
              "a calibration objective over refusal decisions and a benchmark of 2,400 tool-use episodes, reducing expected "
              "calibration error by 41% without hurting task success.")
    X.text_pdf(f"{o}/Travel_Expenses_Policy_extract_2026.pdf", X.flow([
        ("h1", "University of Northbridge - Travel & Expenses Policy"), ("small", "Extract for School of Informatics staff and students, version 2026.2"),
        ("h2", "3. Before you travel"),
        ("p", "3.1 Book economy travel. 3.2 Conference registration fees paid on a School purchasing card or by invoice must not be "
              "claimed again by the traveller."),
        ("h2", "4. What can be reimbursed"),
        ("li", "4.1 Air and rail fares in economy/standard class. Seat selection, upgrades, priority boarding and travel insurance "
               "are not reimbursable (University insurance covers business travel)."),
        ("li", "4.2 Local transport (taxi, metro, airport trains) for business journeys."),
        ("li", "4.3 Accommodation: the room rate up to the overseas nightly cap for the country, times the number of nights. Sweden: "
               "SEK 1,800 per night. Any amount above the cap, and extras such as minibar, laundry or room service, are not reimbursable."),
        ("li", "4.4 Subsistence (meals and non-alcoholic drinks) at actual cost up to the daily cap per person. Sweden: SEK 650 per "
               "day. Add up all meal receipts of the same day and claim at most the cap. Alcoholic drinks are never reimbursable - "
               "deduct them from the receipt before applying the cap."),
        ("li", "4.5 Personal items (toiletries, medicines, sun cream, souvenirs, clothing) are not reimbursable."),
        ("h2", "5. Making a claim"),
        ("p", "5.1 Use LabDesk. Create one claim per traveller per trip and add one line per receipt (or per day for capped meals) "
              "with the receipt date, the currency shown on the receipt and the reimbursable amount in that currency; LabDesk "
              "converts to GBP at the Finance rate for the receipt date. Do not convert amounts yourself."),
        ("p", "5.2 Charge the claim to the traveller's default grant shown in LabDesk unless the grant holder has asked for another "
              "grant. 5.3 Submit the claim in LabDesk; Finance needs the claim number."),
        ("p", "5.4 Keep the original receipts for six years."),
    ]), title="Travel & Expenses Policy extract")


def make_fx(rng):
    rows, eur, sek, usd = [], 0.8580, 0.07810, 0.7460
    day = dt.date(2026, 5, 1)
    while day <= dt.date.fromisoformat(BUSINESS_DATE):
        eur += rng.uniform(-0.0018, 0.0018)
        sek += rng.uniform(-0.00022, 0.00022)
        usd += rng.uniform(-0.0020, 0.0020)
        if day.weekday() < 5:  # Finance publishes rates on working days; weekends use Friday's rate
            rows += [{"date": day.isoformat(), "currency": "EUR", "rate": round(eur, 4)},
                     {"date": day.isoformat(), "currency": "SEK", "rate": round(sek, 5)},
                     {"date": day.isoformat(), "currency": "USD", "rate": round(usd, 4)}]
        day += dt.timedelta(days=1)
    return rows


def rate_for(fx, cur, date):
    if cur == "GBP":
        return 1.0
    return max((r for r in fx if r["currency"] == cur and r["date"] <= date), key=lambda r: r["date"])["rate"]


def build(d):
    os.makedirs(d, exist_ok=True)
    X.copy_app("app_x4.py", d)
    rng = random.Random(40425)
    make_docs(d)
    ids = mailbox(d, None)
    fx = make_fx(rng)
    summary_template(f"{d}/ECCV2026_travel_summary.xlsx")
    base, seeded_keys = build_repo(d)
    shutil.rmtree(f"{d}/_docs")
    # ---- db
    rows = {"settings": [{"key": "business_date", "value": BUSINESS_DATE}, {"key": "next_claim", "value": "139"}],
            "users": [{"id": 1, "username": "jross", "name": "Jamie Ross", "salt": "nB4w", "pw": X.pw_hash("Northbridge#26", "nB4w")},
                      {"id": 2, "username": "emarsh", "name": "Elena Marsh", "salt": "eM1q", "pw": X.pw_hash("Vision!Lab2026", "eM1q")}],
            "groups": [{"id": g, "name": n, "pi": pi} for g, n, pi in GROUPS],
            "grants": [{"code": c, "title": t, "funder": f, "group_id": g, "end_date": e, "travel_budget": b} for c, t, f, g, e, b in GRANTS],
            "members": [{"id": k, "name": v[0], "email": v[1], "group_id": v[2], "role": v[3], "default_grant": v[4]} for k, v in MEMBERS.items()],
            "fx": fx, "claims": [], "claim_lines": [], "pubs": [], "audit": []}
    history = [  # (member, title, status, submitted, lines)
        (1, "CVPR 2026, Nashville", "Paid", "2026-06-24", [("2026-06-10", "Registration", "USD", 925.00), ("2026-06-16", "Accommodation", "USD", 1140.00)]),
        (2, "CVPR 2026, Nashville", "Paid", "2026-06-24", [("2026-05-20", "Travel (air/rail)", "GBP", 684.30), ("2026-06-16", "Accommodation", "USD", 980.00)]),
        (9, "ACL 2026, Vienna", "Paid", "2026-07-21", [("2026-07-04", "Travel (air/rail)", "GBP", 212.80), ("2026-07-09", "Accommodation", "EUR", 540.00)]),
        (5, "BMVA summer school, Edinburgh", "Approved", "2026-07-30", [("2026-07-27", "Travel (air/rail)", "GBP", 96.40), ("2026-07-29", "Accommodation", "GBP", 255.00)]),
        (8, "ACL 2026, Vienna", "Approved", "2026-07-22", [("2026-07-04", "Travel (air/rail)", "GBP", 198.10), ("2026-07-09", "Subsistence", "EUR", 61.50)]),
        (4, "Data collection, Lake District", "Paid", "2026-08-12", [("2026-08-06", "Local transport", "GBP", 74.20)]),
        (10, "ICML 2026 workshop poster printing", "Submitted", "2026-09-02", [("2026-08-28", "Other", "GBP", 38.00)]),
        (6, "Visit to Oxford VGG", "Submitted", "2026-09-18", [("2026-09-15", "Travel (air/rail)", "GBP", 64.90)]),
    ]
    lid = 0
    for cid, (mem, title, status, sub, lines) in enumerate(history, 1):
        m = MEMBERS[mem]
        rows["claims"].append({"id": cid, "number": f"CLM-2026-{130 + cid:04d}", "member_id": mem, "group_id": m[2], "grant_code": m[4],
                               "title": title, "status": status, "created_at": f"{sub} 10:0{cid}:00", "created_by": "jross",
                               "submitted_at": f"{sub} 11:1{cid}:00"})
        for date, cat, cur, amt in lines:
            lid += 1
            r = rate_for(fx, cur, date)
            rows["claim_lines"].append({"id": lid, "claim_id": cid, "date": date, "category": cat, "description": cat, "currency": cur,
                                        "amount": amt, "rate": r, "gbp": round(amt * r + 1e-9, 2), "deleted": 0})
    for pid, p in enumerate(OLD_PUBS, 1):
        g, title, authors, venue, year, typ, arxiv = p[:7]
        rows["pubs"].append({"id": pid, "group_id": g, "title": title, "authors": authors, "venue": venue, "year": year, "type": typ,
                             "arxiv": arxiv, "doi": "", "deleted": 0, "updated_at": "2026-06-15 14:02:00", "updated_by": "jross"})
    rows["audit"] = [{"ts": "2026-09-18 11:16:00", "user": "jross", "action": "claim submitted", "detail": "CLM-2026-0138 (draft #8)"}]
    X.seed_db(f"{d}/labdesk_seed.db", SCHEMA, rows)
    # ---- expected
    trav = {}
    for mem, (grant, lines) in CLAIMS.items():
        sums = {}
        total = 0.0
        for date, cur, amt, _ in lines:
            sums[cur] = round(sums.get(cur, 0) + amt, 2)
            total += round(amt * rate_for(fx, cur, date) + 1e-9, 2)
        trav[str(mem)] = {"name": MEMBERS[mem][0], "group": MEMBERS[mem][2], "grant": grant, "sums": sums, "gbp_total": round(total, 2),
                          "lines": lines}
    receipt_mails = [ids[k] for k in ("tomas", "aisha1", "aisha2", "ben", "priya")]
    exp = {"me": ME[1], "finance": FIN[1], "list": LIST, "msg": ids, "receipt_mails": receipt_mails, "travellers": trav,
           "seed_claim_max_id": len(history), "seed_pub_max_id": len(OLD_PUBS), "papers": P_NEW, "rejected": REJECTED,
           "repo_base": base, "seeded_bib_keys": seeded_keys}
    json.dump(exp, open(f"{d}/expected.json", "w"), indent=1, ensure_ascii=False)


SCHEMA = """
CREATE TABLE settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE users(id INTEGER PRIMARY KEY, username TEXT UNIQUE, name TEXT, salt TEXT, pw TEXT);
CREATE TABLE groups(id TEXT PRIMARY KEY, name TEXT, pi TEXT);
CREATE TABLE grants(code TEXT PRIMARY KEY, title TEXT, funder TEXT, group_id TEXT, end_date TEXT, travel_budget REAL);
CREATE TABLE members(id INTEGER PRIMARY KEY, name TEXT, email TEXT, group_id TEXT, role TEXT, default_grant TEXT);
CREATE TABLE fx(date TEXT, currency TEXT, rate REAL, PRIMARY KEY(date, currency));
CREATE TABLE claims(id INTEGER PRIMARY KEY, number TEXT, member_id INTEGER, group_id TEXT, grant_code TEXT, title TEXT, status TEXT,
    created_at TEXT, created_by TEXT, submitted_at TEXT);
CREATE TABLE claim_lines(id INTEGER PRIMARY KEY, claim_id INTEGER, date TEXT, category TEXT, description TEXT, currency TEXT,
    amount REAL, rate REAL, gbp REAL, deleted INTEGER DEFAULT 0);
CREATE TABLE pubs(id INTEGER PRIMARY KEY, group_id TEXT, title TEXT, authors TEXT, venue TEXT, year INTEGER, type TEXT, arxiv TEXT,
    doi TEXT, deleted INTEGER DEFAULT 0, updated_at TEXT, updated_by TEXT);
CREATE TABLE sessions(sid TEXT PRIMARY KEY, user_id INTEGER, created TEXT, state TEXT);
CREATE TABLE audit(id INTEGER PRIMARY KEY, ts TEXT, user TEXT, action TEXT, detail TEXT);
"""

if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp())
