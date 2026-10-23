"""o2 - one Writer document: the Brightline master services agreement (negotiation copy with tracked changes from
three reviewers and their comments). A resolves the tracked changes per the head of legal's rules and her scanned
call notes; B inserts the new Data Protection clause from the client's scanned rider, renumbers clauses and
cross-references, updates the table of contents and exports a PDF."""
from __future__ import annotations

import copy
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_oscan import Page, save_scan  # noqa: E402

PRIYA, DANIEL, MARTA = "Priya Raman", "Daniel Okafor", "Marta Lindqvist"
DATES = {PRIYA: "2026-10-08T09:14:00Z", DANIEL: "2026-10-10T15:02:00Z", MARTA: "2026-10-12T11:40:00Z"}
Q = "“", "”"


def d(term):
    return f"{Q[0]}{term}{Q[1]}"


# Clause text. Each paragraph is a list of segments: str (plain), ("ins", text, author), ("del", text, author),
# ("com", text, author, comment) = plain text carrying a reviewer comment. "P:" prefix marks a whole inserted paragraph.
CLAUSES = [
    ("Definitions and Interpretation", [
        ["1.1 In this Agreement the following definitions apply:"],
        ["DEF", d("Affiliate") + " means, in relation to a party, any entity that controls, is controlled by, or is under "
                "common control with that party from time to time."],
        ["DEF", ("ins", d("Business Day") + " means a day other than a Saturday, Sunday or public holiday in England.", PRIYA),
         "INS_PARA", PRIYA],
        ["DEF", d("Charges") + " means the charges for the Services set out in Schedule 1."],
        ["DEF", ("com", d("Confidential Information"), PRIYA, "Business Day added above - it is used in clause 5.1."),
         " means all information of a confidential nature disclosed by one party to the other in connection with this "
         "Agreement, in any form."],
        ["DEF", d("Contract Year") + " means each period of twelve (12) months starting on the Commencement Date or an "
                "anniversary of it."],
        ["DEF", d("Deliverables") + " means all documents, reports, models and other materials produced by the Supplier "
                "for the Customer as part of the Services."],
        ["DEF", d("Services") + " means the data analytics and reporting services described in Schedule 2, as varied "
                "from time to time in accordance with Clause 12.2."],
        ["1.2 Clause and schedule headings do not affect the interpretation of this Agreement. Words in the singular "
         "include the plural and vice versa."],
    ]),
    ("Services", [
        ["2.1 The Supplier shall provide the Services to the Customer from the Commencement Date in accordance with "
         "this Agreement and each Statement of Work."],
        ["2.2 The Supplier shall use ", ("ins", "all ", PRIYA), "reasonable endeavours to meet any performance dates "
         "specified in a Statement of Work, but any such dates are estimates only."],
        ["2.3 Any liability of the Supplier for failure to perform the Services is subject to Clause 9 (Limitation of "
         "Liability)."],
    ]),
    ("Term", [
        ["3.1 This Agreement commences on the Commencement Date and continues for an initial term of thirty-six (36) "
         "months, unless terminated earlier in accordance with Clause 10 (Termination)."],
        ["3.2 This Agreement shall then renew automatically for successive periods of twelve (12) months unless either "
         "party gives at least three (3) months' written notice of non-renewal."],
    ]),
    ("Charges and Payment", [
        ["4.1 In consideration of the provision of the Services, the Customer shall pay the Charges."],
        ["4.2 The Supplier shall ", ("com", "invoice the Customer", MARTA, "Billing in advance helps our cash flow - "
                                                                          "please keep."),
         " monthly in ", ("del", "arrears", MARTA), ("ins", "advance", MARTA), " for the Charges."],
        ["4.3 The Customer shall pay each valid invoice within ", ("del", "thirty (30)", DANIEL),
         ("ins", "sixty (60)", DANIEL), " ", ("com", "days of receipt", DANIEL,
                                              "Brightline's standard payment term for all suppliers is 60 days."), "."],
        ["4.4 If the Customer fails to pay any amount by its due date, the Supplier may charge interest on the overdue "
         "amount", ("ins", " at the rate of 4% per annum above the Bank of England base rate", MARTA),
         " and may suspend the Services in accordance with Clause 10.3."],
        [("ins", "4.5 The Charges shall be increased on each anniversary of the Commencement Date by the percentage "
                 "increase in the Consumer Prices Index over the preceding twelve (12) months.", MARTA), "INS_PARA", MARTA],
    ]),
    ("Service Levels", [
        ["5.1 The Supplier shall meet or exceed the Service Levels set out in Schedule 3.",
         ("ins", " The Supplier shall provide the Customer with a written service report within ten (10) Business Days "
                 "after the end of each month.", DANIEL)],
        ["5.2 If the Supplier fails to meet a Service Level, the Customer shall be entitled to the service credits set "
         "out in Schedule 3, which are the Customer's sole remedy for that failure save under Clause 10 (Termination)."],
    ]),
    ("Intellectual Property", [
        ["6.1 All Intellectual Property Rights in the Deliverables shall vest in the Customer on payment in full of the "
         "Charges for them."],
        ["6.2 ", ("del", "The Supplier retains ownership of all Intellectual Property Rights in its pre-existing tools, "
                         "methodologies and software (" + d("Supplier Materials") + ") and grants the Customer a "
                         "non-exclusive, royalty-free licence to use the Supplier Materials to the extent necessary to use "
                         "the Deliverables.", DANIEL)],
    ]),
    ("Confidentiality", [
        ["7.1 Each party shall keep the other party's Confidential Information confidential and shall not use it except "
         "to perform its obligations or exercise its rights under this Agreement."],
        ["7.2 A party may disclose Confidential Information to its employees, officers and professional advisers who need "
         "to know it, provided that they are bound by obligations of confidentiality.",
         ("del", " A party may disclose Confidential Information to its employees who need to know it.", PRIYA)],
        ["7.3 The obligations in this Clause 7 shall survive termination of this Agreement for five (5) years."],
    ]),
    ("Warranties", [
        ["8.1 The Supplier warrants that the Services will be performed with reasonable skill and care and in accordance "
         "with good industry practice."],
        ["8.2 Except as expressly set out in this Clause 8, all warranties, conditions and other terms implied by statute "
         "or common law are excluded to the fullest extent permitted by law."],
    ]),
    ("Limitation of Liability", [
        ["9.1 Nothing in this Agreement limits or excludes liability for death or personal injury caused by negligence, "
         "for fraud or fraudulent misrepresentation, or for any other liability that cannot be limited by law."],
        ["9.2 Subject to Clause 9.1, each party's total aggregate liability arising under or in connection with this "
         "Agreement in each Contract Year shall not exceed ", ("del", "100%", DANIEL), ("ins", "200%", DANIEL),
         " of the Charges paid or payable in that Contract Year."],
        ["9.3 Subject to Clause 9.1, neither party shall be liable for any loss of profits, loss of business, or any "
         "indirect or consequential loss."],
    ]),
    ("Termination", [
        ["10.1 Either party may terminate this Agreement with immediate effect by written notice if the other party "
         "commits a material breach which is not remedied within thirty (30) days of ", ("del", "recieving", PRIYA),
         ("ins", "receiving", PRIYA), " notice of it."],
        ["10.2 Either party may terminate this Agreement with immediate effect if the other party becomes insolvent or "
         "ceases to carry on business.", ("ins", " The Customer may also terminate this Agreement for convenience on not "
                                                 "less than ninety (90) days' written notice to the Supplier.", DANIEL)],
        ["10.3 The Supplier may suspend the Services on seven (7) days' written notice if any undisputed invoice remains "
         "unpaid for more than thirty (30) days after its due date."],
        ["10.4 Termination shall not affect any rights or remedies accrued before termination. Clauses 1, 6, 7, 9 and 11 "
         "shall survive termination of this Agreement."],
    ]),
    ("Governing Law and Jurisdiction", [
        ["11.1 This Agreement and any dispute or claim arising out of or in connection with it shall be governed by the "
         "laws of ", ("del", "England and Wales", DANIEL), ("ins", "Scotland", DANIEL), "."],
        ["11.2 Subject to Clause 11.3, the courts of England and Wales shall have exclusive jurisdiction to settle any "
         "dispute or claim arising out of or in connection with this Agreement."],
        ["11.3 Nothing in Clause 11.2 prevents either party from seeking urgent injunctive relief in any court of "
         "competent jurisdiction."],
    ]),
    ("General", [
        ["12.1 This Agreement, its schedules and each Statement of Work constitute the entire agreement between the "
         "parties and supersede all previous discussions between them."],
        ["12.2 No variation of this Agreement shall be effective unless it is in writing and signed by both parties."],
        ["12.3 Notices under this Agreement shall be in writing and delivered by hand, by recorded delivery post or by "
         "email to the addresses set out at the start of this Agreement."],
        ["12.4 A notice given under Clause 12.3 is deemed received at the time of delivery if delivered by hand or "
         "email, and two (2) Business Days after posting if sent by recorded delivery."],
        ["12.5 This Agreement may be executed in any number of counterparts, each of which is an original."],
    ]),
]

# The expected outcome of every tracked change (see the call notes and Hannah's rules)
CHANGES = [
    dict(id="P1", frags=["means a day other than"], author=PRIYA, accept=True, topic="Business Day definition",
         yes='"business day" means a day other than a saturday', no=None),
    dict(id="P2", frags=["=all"], author=PRIYA, accept=True, topic="all reasonable endeavours",
         yes="shall use all reasonable endeavours to meet", no="shall use reasonable endeavours to meet"),
    dict(id="P3", frags=["employees who need to know it"], author=PRIYA, accept=True, topic="duplicated sentence in 7.2",
         yes="bound by obligations of confidentiality.", no="to its employees who need to know it"),
    dict(id="P4", frags=["receiving", "recieving"], author=PRIYA, accept=True, topic="typo receiving", yes="days of receiving notice", no="recieving"),
    dict(id="D1", frags=["thirty (30)", "sixty (60)"], author=DANIEL, accept=False, topic="payment term 60 days",
         yes="invoice within thirty (30) days of receipt", no="sixty (60)"),
    dict(id="D2", frags=["100%", "200%"], author=DANIEL, accept=False, topic="liability cap 200%",
         yes="shall not exceed 100% of the charges", no="200%"),
    dict(id="D3", frags=["written service report"], author=DANIEL, accept=True, topic="monthly service reports",
         yes="written service report within ten (10) business days", no=None),
    dict(id="D4", frags=["for convenience"], author=DANIEL, accept=True, topic="termination for convenience",
         yes="terminate this agreement for convenience on not less than ninety (90) days", no=None),
    dict(id="D5", frags=["pre-existing tools"], author=DANIEL, accept=False, topic="deletion of the Supplier Materials clause",
         yes="the supplier retains ownership of all intellectual property rights in its pre-existing tools", no=None),
    dict(id="D6", frags=["england and wales", "scotland"], author=DANIEL, accept=False, topic="governing law Scotland (not covered -> reject)",
         yes="governed by the laws of england and wales", no="laws of scotland"),
    dict(id="M1", frags=["consumer prices index"], author=MARTA, accept=False, topic="CPI indexation", yes=None, no="consumer prices index"),
    dict(id="M2", frags=["bank of england base rate"], author=MARTA, accept=True, topic="late payment interest 4%",
         yes="at the rate of 4% per annum above the bank of england base rate", no=None),
    dict(id="M3", frags=["arrears", "advance"], author=MARTA, accept=True, topic="invoicing in advance (not covered -> accept)",
         yes="invoice the customer monthly in advance", no="monthly in arrears"),
]

RIDER = [
    "8.1 Each party shall comply with all applicable requirements of the Data Protection Legislation. This Clause 8 is "
    "in addition to, and does not relieve, remove or replace, a party's obligations under the Data Protection "
    "Legislation.",
    "8.2 The parties acknowledge that, for the purposes of the Data Protection Legislation, the Customer is the "
    "controller and the Supplier is the processor of any Customer Personal Data processed in connection with the "
    "Services.",
    "8.3 The Supplier shall process Customer Personal Data only on the documented written instructions of the Customer, "
    "shall ensure that all personnel who have access to it are bound by a duty of confidentiality, and shall not transfer "
    "it outside the United Kingdom without the Customer's prior written consent.",
    "8.4 The Supplier shall notify the Customer without undue delay, and in any event within forty-eight (48) hours, "
    "on becoming aware of a personal data breach affecting Customer Personal Data.",
    "8.5 On termination of this Agreement the Supplier shall, at the Customer's option, delete or return all Customer "
    "Personal Data, unless retention is required by law.",
]
RIDER_DEF = (d("Data Protection Legislation") + " means the UK GDPR, the Data Protection Act 2018 and any other "
             "applicable law relating to the processing of personal data.")

# Cross references B must update: (regex on the final visible text, expected group values)
XREFS = [
    (r"is subject to clause (\d+) \(limitation of liability\)", ["10"]),
    (r"in accordance with clause (\d+) \(termination\)", ["11"]),
    (r"suspend the services in accordance with clause (\d+\.\d+)", ["11.3"]),
    (r"save under clause (\d+) \(termination\)", ["11"]),
    (r"the obligations in this clause (\d+) shall survive", ["7"]),
    (r"except as expressly set out in this clause (\d+),", ["9"]),
    (r"subject to clause (\d+\.\d+), each party's total", ["10.1"]),
    (r"subject to clause (\d+\.\d+), neither party", ["10.1"]),
    (r"subject to clause (\d+\.\d+), the courts", ["12.3"]),
    (r"nothing in clause (\d+\.\d+) prevents", ["12.2"]),
    (r"a notice given under clause (\d+\.\d+) is deemed", ["13.3"]),
    (r"from time to time in accordance with clause (\d+\.\d+)", ["13.2"]),
]
SURVIVE_OLD = [1, 6, 7, 9, 11]
SURVIVE_NEW = [1, 6, 7, 8, 10, 12]

HANNAH = """From: Hannah Cole <hannah.cole@corvidanalytics.co.uk>
To: legal-ops@corvidanalytics.co.uk
Date: Wed, 14 Oct 2026 18:22
Subject: Brightline MSA v4 - please clean up the tracked changes

Hi,

The Brightline MSA (Contracts/Brightline_MSA_v4.docx) came back with tracked changes from three people:
Priya Raman (our counsel), Daniel Okafor (Brightline procurement) and Marta Lindqvist (our finance team).

Please resolve every one of their tracked changes before it goes back to Brightline:
 - Priya's edits are housekeeping: accept all of them.
 - Daniel's and Marta's edits: follow the outcome of today's negotiation call. My handwritten call notes
   are scanned in Contracts/call_notes_2026-10-14.pdf.
 - Anything of Daniel's that the call notes do not cover: reject it (we only concede what we agreed).
 - Anything of Marta's that the call notes do not cover: accept it.
Please also delete all the review comments, and leave track changes recording off.
Save it in place as Word (.docx), same file name.

Finally, put a short list of the points where we rejected Brightline's (Daniel's) edits in
Contracts/Brightline_rejected_points.txt - one line per point - so the account manager can explain
them to the client.

Thanks!
Hannah

--
Hannah Cole | Head of Legal | Corvid Analytics Ltd
"""


def _call_notes(path):
    p = Page(color=True)
    ink = (25, 40, 120)
    p.text(90, 80, "Corvid Analytics Ltd", 30, "bold")
    p.text(90, 125, "Meeting notes", 26)
    p.hline(90, 1150, 170, 2)
    y = p.hand(90, 200, "Brightline MSA - negotiation call 14/10/2026", 40, ink, angle=1, shear=0.3)
    y = p.hand(90, y + 5, "H. Cole, D. Okafor (Brightline), P. Raman", 30, ink, angle=0, shear=0.3)
    lines = [
        "1. Payment terms: they want 60 days. NOT agreed -",
        "     stays at 30 days.",
        "2. Late payment interest: OK for them, 4% over base agreed.",
        "3. Liability cap: asked 200%. NOT agreed - stays 100%",
        "     of the annual charges.",
        "4. Monthly service reports: agreed.",
        "5. Termination for convenience: agreed (90 days' notice).",
        "6. Our pre-existing IP / tools: not negotiable -",
        "     our clause stays exactly as drafted.",
        "7. CPI indexation (Marta's idea): Brightline will not",
        "     accept it -> we dropped it.",
    ]
    y += 30
    for ln in lines:
        y = p.hand(110, y, ln, 33, ink, angle=0.6, shear=0.28) + 12
    y += 30
    y = p.hand(110, y, "Next: clean version back to Daniel by Fri.", 33, ink, angle=1.2, shear=0.28)
    p.ink_line([(100, 470), (1120, 462)], color=ink, w=2)
    save_scan([p], path, seed=7201, color=True)


def _rider(path):
    pages = [Page(), Page()]
    p = pages[0]
    p.text(90, 80, "BRIGHTLINE RETAIL PLC", 34, "bold")
    p.text(90, 128, "Legal & Data Protection Office, 1 Carrow Way, Norwich NR1 2DH", 22)
    p.hline(90, 1150, 168, 3)
    p.text(90, 195, "RIDER DP-1 - DATA PROTECTION CLAUSE", 34, "bold")
    p.text(90, 245, "Master Services Agreement with Corvid Analytics Ltd - agreed wording, 12 October 2026", 23)
    y = 310
    y = p.para(90, y, "Instructions for the drafting party: insert the clause below as a new Clause 8 headed "
                      "\"Data Protection\", immediately after Clause 7 (Confidentiality), using the same heading style "
                      "as the other clauses. Renumber the subsequent clauses and their sub-clauses and update every "
                      "cross-reference accordingly. Add the new Clause 8 to the list of clauses that survive termination. "
                      "Add the definition below to Clause 1.1 in alphabetical order.", 24, width=1060)
    y += 30
    p.text(90, y, "Definition for Clause 1.1:", 25, "bold")
    y += 45
    y = p.para(120, y, RIDER_DEF, 25, "serif", width=1020)
    y += 30
    p.text(90, y, "8. Data Protection", 28, "serifbold")
    y += 55
    for para in RIDER[:3]:
        y = p.para(120, y, para, 25, "serif", width=1020) + 18
    p.text(1100, 1640, "1/2", 20)
    p = pages[1]
    y = 120
    for para in RIDER[3:]:
        y = p.para(120, y, para, 25, "serif", width=1020) + 18
    y += 50
    p.text(90, y, "Approved: S. Whitaker, Data Protection Officer, Brightline Retail plc", 23)
    p.hand(150, y + 50, "S. Whitaker", 44, (20, 20, 90), angle=3, shear=0.35)
    p.text(1100, 1640, "2/2", 20)
    save_scan(pages, path, seed=7202)


# ---------------------------------------------------------------- docx construction
def _w(tag):
    from docx.oxml.ns import qn
    return qn(tag)


class Builder:
    def __init__(self):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Cm, Pt, RGBColor
        self.doc = Document()
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.5))
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
        st = doc.styles
        normal = st["Normal"]
        normal.font.name, normal.font.size = "Liberation Serif", Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        title = st["Title"]
        title.font.name, title.font.size, title.font.color.rgb = "Liberation Sans", Pt(22), RGBColor(0, 0, 0)
        h1 = st["Heading 1"]
        h1.font.name, h1.font.size, h1.font.bold = "Liberation Sans", Pt(12.5), True
        h1.font.color.rgb = RGBColor(0, 0, 0)
        h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(14), Pt(6)
        self.body = st.add_style("Clause Text", WD_STYLE_TYPE.PARAGRAPH)
        self.body.base_style = normal
        self.body.paragraph_format.left_indent = Cm(1.0)
        self.body.paragraph_format.first_line_indent = Cm(-1.0)
        self.body.paragraph_format.alignment = 3  # justify
        self.defs = st.add_style("Definition", WD_STYLE_TYPE.PARAGRAPH)
        self.defs.base_style = normal
        self.defs.paragraph_format.left_indent = Cm(2.0)
        self.toc = st.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
        self.toc.base_style = normal
        self.toc.paragraph_format.space_after = Pt(3)
        self.toc.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), alignment=2, leader=1)  # right, dots
        self.rid = 100

    def _rev(self, kind, author):
        from docx.oxml import OxmlElement
        el = OxmlElement(f"w:{kind}")
        el.set(_w("w:id"), str(self.rid))
        el.set(_w("w:author"), author)
        el.set(_w("w:date"), DATES[author])
        self.rid += 1
        return el

    def _run(self, text, deleted=False):
        from docx.oxml import OxmlElement
        r = OxmlElement("w:r")
        t = OxmlElement("w:delText" if deleted else "w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = text
        r.append(t)
        return r

    def para(self, segs, style):
        doc = self.doc
        p = doc.add_paragraph(style=style)
        ins_para = "INS_PARA" in segs
        segs = [s for s in segs if s not in ("INS_PARA", "DEF", PRIYA, DANIEL, MARTA)]
        if ins_para:
            author = next(s[2] for s in segs if isinstance(s, tuple))
            ppr = p._p.get_or_add_pPr()
            from docx.oxml import OxmlElement
            rpr = OxmlElement("w:rPr")
            rpr.append(self._rev("ins", author))
            ppr.append(rpr)
        for s in segs:
            if isinstance(s, str):
                p.add_run(s)
            elif s[0] == "com":
                run = p.add_run(s[1])
                doc.add_comment([run], text=s[3], author=s[2], initials="".join(w[0] for w in s[2].split()))
            else:
                kind, text, author = s
                el = self._rev(kind, author)
                el.append(self._run(text, deleted=(kind == "del")))
                p._p.append(el)
        return p

    def toc_block(self, entries):
        """Contents heading + a TOC field (inside a docPartObj sdt) with cached entries."""
        from docx.oxml import OxmlElement
        doc = self.doc
        head = doc.add_paragraph("Contents", style="TOC Heading")
        entries_p = []
        for i, (text, page) in enumerate(entries):
            p = doc.add_paragraph(style="toc 1")
            if i == 0:
                for kind, instr in (("begin", None), (None, ' TOC \\o "1-1" \\h \\z \\u '), ("separate", None)):
                    r = OxmlElement("w:r")
                    if kind:
                        fc = OxmlElement("w:fldChar")
                        fc.set(_w("w:fldCharType"), kind)
                        r.append(fc)
                    else:
                        it = OxmlElement("w:instrText")
                        it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                        it.text = instr
                        r.append(it)
                    p._p.append(r)
            p.add_run(text)
            r = p.add_run()
            r._r.append(OxmlElement("w:tab"))
            p.add_run(str(page))
            entries_p.append(p)
        r = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(_w("w:fldCharType"), "end")
        r.append(fc)
        entries_p[-1]._p.append(r)
        # wrap heading + entries in an sdt (Word/LibreOffice table-of-contents container)
        sdt = OxmlElement("w:sdt")
        pr = OxmlElement("w:sdtPr")
        dpo = OxmlElement("w:docPartObj")
        g = OxmlElement("w:docPartGallery")
        g.set(_w("w:val"), "Table of Contents")
        u = OxmlElement("w:docPartUnique")
        dpo.append(g)
        dpo.append(u)
        pr.append(dpo)
        sdt.append(pr)
        content = OxmlElement("w:sdtContent")
        sdt.append(content)
        head._p.addprevious(sdt)
        for p in [head] + entries_p:
            content.append(p._p)


def build(dst):
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    os.makedirs(f"{dst}/Contracts", exist_ok=True)
    b = Builder()
    doc = b.doc
    t = doc.add_paragraph(style="Title")
    t.add_run("Master Services Agreement")
    doc.add_paragraph("between Corvid Analytics Ltd and Brightline Retail plc").runs[0].font.size = Pt(13)
    doc.add_paragraph("Negotiation copy v4 - 14 October 2026 - subject to contract")
    b.toc_block([(f"{i}. {title}", 2 + (i - 1) // 3) for i, (title, _) in enumerate(CLAUSES, 1)])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("THIS AGREEMENT is dated 2026 and made between:")
    doc.add_paragraph("(1) CORVID ANALYTICS LTD, a company incorporated in England and Wales with company number "
                      "09381127 whose registered office is at 40 Silk Street, Manchester M4 6BX (the " + d("Supplier") +
                      "); and", style="Definition")
    doc.add_paragraph("(2) BRIGHTLINE RETAIL PLC, a company incorporated in England and Wales with company number "
                      "02214870 whose registered office is at 1 Carrow Way, Norwich NR1 2DH (the " + d("Customer") + ").",
                      style="Definition")
    doc.add_paragraph("BACKGROUND: The Customer wishes to obtain, and the Supplier has agreed to provide, data analytics "
                      "and reporting services on the terms of this Agreement.")
    for i, (title, paras) in enumerate(CLAUSES, 1):
        doc.add_paragraph(f"{i}. {title}", style="Heading 1")
        for segs in paras:
            style = "Definition" if "DEF" in segs else "Clause Text"
            b.para(list(segs), style)
    doc.add_paragraph("Signed for and on behalf of Corvid Analytics Ltd: ............................  Date: ..........")
    doc.add_paragraph("Signed for and on behalf of Brightline Retail plc: ...........................  Date: ..........")
    doc.core_properties.author = "Hannah Cole"
    doc.core_properties.title = "Master Services Agreement - Corvid / Brightline"
    doc.save(f"{dst}/Contracts/Brightline_MSA_v4.docx")
    open(f"{dst}/Contracts/Email_from_Hannah.txt", "w").write(HANNAH)
    _call_notes(f"{dst}/Contracts/call_notes_2026-10-14.pdf")
    _rider(f"{dst}/Contracts/Rider_DP-1_Data_Protection.pdf")
    titles = [t for t, _ in CLAUSES]
    new_titles = titles[:7] + ["Data Protection"] + titles[7:]
    json.dump({"changes": CHANGES, "authors": [PRIYA, DANIEL, MARTA], "rider": RIDER, "rider_def": RIDER_DEF,
               "titles_old": titles, "titles_new": new_titles, "xrefs": XREFS, "survive_old": SURVIVE_OLD,
               "survive_new": SURVIVE_NEW,
               "rejected_points": {"payment": r"\b60\b|sixty|payment term", "liability": r"200|liabilit|\bcap\b",
                                   "ip": r"pre-?existing|intellectual|\bip\b|supplier materials|tools",
                                   "law": r"scotland|governing law|scots"}},
              open(f"{dst}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/o2")
