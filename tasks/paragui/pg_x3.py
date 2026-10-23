"""X3 - Brightwater Studio's anniversary client evening (Thunderbird + Brightwater Hub in Chrome + LibreOffice Calc).

A (pgx3rsvp): RSVP mails (one a photo of a paper reply card) -> guest list in Hub (status, plus-one, a substitute) ->
   mails filed into a new folder -> door list xlsx -> draft to the venue with the headcount, door list, Martin in copy.
B (pgx3catering): the same RSVP mails (meal choices, dietary needs) + the hosts' choices on the Hub event page ->
   finish and submit the caterer's purchase order in Hub -> budget workbook -> draft to the caterer (PO number, dietary list).
Shared live state: one Thunderbird (A moves the RSVP mails B is reading; both compose), one Chrome with one Hub login
session whose event selector and guest-list filters are stored per session.
"""
from __future__ import annotations

import re
from decimal import ROUND_HALF_UP, Decimal

from evaluation_examples.task_class import pg_common as pc
from evaluation_examples.task_class import pg_xcommon as xc

APP, PORT = "hub", 8123
DOOR = f"{pc.HOME}/Documents/Events/Anniversary_door_list.xlsx"
BUDGET = f"{pc.HOME}/Documents/Events/Anniversary_Budget_2026.xlsx"
FF, ANNIV, WINTER_PO = 1, 1, 3


def setup(sc):
    pc.sh(sc, "mkdir -p ~/Documents/Events")
    pc.upload(sc, [("x3/Anniversary_Budget_2026.xlsx", BUDGET)])
    pc.start_webapp(sc, "x3/app_x3.py", APP, PORT, "x3/hub_seed.db")
    pc.start_thunderbird(sc, "x3/thunderbird-profile.tar.gz", "pgx3")
    pc.open_chrome(sc, [f"http://127.0.0.1:{PORT}/login"])


def _name_parts(name):
    return [xc.flat(p) for p in name.split()]


def _cell_is(cell, name) -> bool:
    """A cell holds this person's name ("David Park", "Mr David Park", "Park, David", "David Park (+1)")."""
    f = xc.flat(cell)
    parts = _name_parts(name)
    return bool(f) and all(p in f for p in parts) and len(f) <= len("".join(parts)) + 6


def _row_has(row, name) -> bool:
    cells = [str(v) for v in row if v is not None and str(v).strip()]
    if any(_cell_is(c, name) for c in cells):
        return True
    first, last = _name_parts(name)[0], _name_parts(name)[-1]
    flats = [xc.flat(c) for c in cells]
    return first in flats and last in flats  # first and last name in separate columns


def _names_in_xlsx(path, names) -> tuple[list, list]:
    """(names found on some row, rows) for an xlsx; rows are lists of cell values."""
    try:
        rows = [r for _, r in xc.xlsx_rows(path)]
    except Exception:
        return [], []
    return [n for n in names if any(_row_has(r, n) for r in rows)], rows


def _lead_name(row) -> str:
    return next((str(v) for v in row if v is not None and str(v).strip()), "")


def _same_person(a, b) -> bool:
    return bool(a) and all(p in xc.flat(a) for p in _name_parts(b))


# ---------------------------------------------------------------- A: RSVPs
def evaluate_rsvp(env, dest: str) -> dict:
    exp = pc.expected("x3")
    checks = {}
    q = pc.sql(env, pc.webapp_db(APP), {"g": "SELECT id, event_id, kind, name, status, plus_one FROM guests"})
    if q is None:
        return pc.result(0.0, {"app_db": "unreadable"})
    by_id = {r["id"]: r for r in q["g"]}
    anniv = [r for r in q["g"] if r["event_id"] == ANNIV]
    ok, bad = 0, []
    for name, want in exp["outcome"].items():
        r = next((x for x in anniv if x["kind"] == "invitee" and x["name"] == name), None)
        got_p = (r["plus_one"] or "").strip() if r else ""
        p_ok = (not got_p) if not want["plus_one"] else _same_person(got_p, want["plus_one"])
        if r and r["status"] == want["status"] and p_ok:
            ok += 1
        else:
            bad.append(f"{name}: {r['status'] if r else 'missing'} +1 '{got_p}'")
    responders = set(exp["outcome"])
    collateral = [f"{s['name']} ({'anniv' if s['event_id'] == ANNIV else 'winter'})" for s in exp["seed_guests"]
                  if s["name"] not in responders or s["event_id"] != ANNIV
                  if by_id.get(s["id"]) and (by_id[s["id"]]["status"] != s["status"] or (by_id[s["id"]]["plus_one"] or "").strip() != s["plus_one"])]
    added = [r for r in q["g"] if r["id"] > exp["seed_guest_max_id"]]
    sub_ok = any(r["event_id"] == ANNIV and _same_person(r["name"], exp["substitute"]) and r["status"] == "Attending" for r in added)
    fwd_bad = any(r["event_id"] == ANNIV and _same_person(r["name"], exp["forwarded"]) and r["status"] == "Attending" for r in q["g"])
    checks.update({"responses_ok": f"{ok}/{len(exp['outcome'])}", "response_problems": bad, "collateral_changes": collateral[:8],
                   "substitute_added": sub_ok, "forwarded_guest_confirmed": fwd_bad,
                   "added_rows": [(r["name"], r["event_id"], r["status"]) for r in added][:6]})

    # mails filed into "Anniversary RSVPs"
    boxes = xc.mailboxes(env) or {}
    where = xc.where(boxes)
    is_target = lambda f: xc.flat(xc.leaf(f)) == "anniversaryrsvps"  # noqa: E731
    filed = sum(any(is_target(f) for f in where.get(i, ())) and not any(xc.leaf(f) == "inbox" for f in where.get(i, ()))
                for i in exp["rsvp_ids"])
    winter_misfiled = any(is_target(f) for f in where.get(exp["winter_id"], ()))
    mail_score = max(0.0, filed / len(exp["rsvp_ids"]) - 0.25 * winter_misfiled)
    checks.update({"rsvp_mails_filed": f"{filed}/{len(exp['rsvp_ids'])}", "winter_mail_misfiled": winter_misfiled,
                   "folders": sorted(boxes)[:12]})

    # door list
    path = pc.fetch(env, DOOR, dest + "_door.xlsx")
    found, rows = _names_in_xlsx(path, exp["door_list"]) if path else ([], [])
    forbidden = [n for n in exp["not_on_door_list"] if any(_cell_is(_lead_name(r), n) for r in rows)]
    door_score = max(0.0, len(found) / len(exp["door_list"]) - 0.25 * len(forbidden)) if path else 0.0
    checks.update({"door_list_file": bool(path), "door_names": f"{len(found)}/{len(exp['door_list'])}",
                   "door_missing": [n for n in exp["door_list"] if n not in found][:8], "door_wrong_people": forbidden})

    # draft to the venue
    mine = [m for m in xc.outgoing(boxes, exp["msg"].values()) if xc.addressed(m, exp["clara"])]
    best, att_ok = None, False
    for k, m in enumerate(mine):
        for i, a in enumerate(m["attachments"]):
            if a["name"].lower().endswith((".xlsx", ".csv", ".ods", ".xls")):
                lp = xc.fetch_attachment(env, a, f"{dest}_att{k}_{i}")
                if lp and lp.endswith(".csv"):
                    names = [n for n in exp["door_list"] if any(_cell_is(c, n) for c in re.split(r"[,\n;]", open(lp, errors="replace").read()))]
                elif lp and lp.endswith(".xlsx"):
                    names = _names_in_xlsx(lp, exp["door_list"])[0]
                else:
                    names = []
                if len(names) >= len(exp["door_list"]) - 2:
                    best, att_ok = m, True
        best = best or m
    text = best["text"] if best else ""
    head_ok = bool(re.search(rf"(?<![\d.,£]){exp['headcount']}(?![\d.,%])", text))
    martin_ok = bool(best) and xc.addressed(best, exp["martin"]) and "marcuswebb" in xc.flat(text)
    checks.update({"venue_draft": bool(best), "door_list_attached": att_ok, "headcount_stated": head_ok,
                   "martin_copied_and_told_about_forward": martin_ok})

    score = (0.35 * ok / len(exp["outcome"]) + 0.05 * (not collateral and ok > 0) + 0.08 * sub_ok + 0.04 * (not fwd_bad and ok > 0)
             + 0.13 * mail_score + 0.15 * door_score + 0.10 * att_ok + 0.05 * head_ok + 0.05 * martin_ok)
    score = xc.capped(score, [(ok == 0, 0.3)])
    return pc.result(score, checks)


# ---------------------------------------------------------------- B: catering
def _total(lines, pct) -> Decimal:
    q = lambda x: x.quantize(Decimal("0.01"), ROUND_HALF_UP)  # noqa: E731
    sub = sum((q(Decimal(str(p))) * n for _, n, p in lines), Decimal("0"))
    svc = q(sub * Decimal(str(pct)) / 100)
    return sub + svc + q((sub + svc) * Decimal("0.2"))


def _budget(path) -> dict:
    """Header-driven read of the budget sheet: catering amount/ref, total committed, remaining."""
    from openpyxl import load_workbook
    out = {}
    for ws in load_workbook(path, data_only=True).worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        ca = cr = None
        for row in grid:
            labels = [xc.flat(v) for v in row]
            if ca is None and any(lb.startswith("item") for lb in labels) and any("amount" in lb for lb in labels):
                ca = next(i for i, lb in enumerate(labels) if "amount" in lb)
                cr = next((i for i, lb in enumerate(labels) if lb.startswith("po") or "ref" in lb), None)
                continue
            if ca is None or not row or not labels[0]:
                continue
            val = row[ca] if ca < len(row) else None
            num = float(val) if isinstance(val, (int, float)) and not isinstance(val, bool) else (
                (xc.numbers(str(val)) or [None])[0] if val is not None else None)
            if labels[0].startswith("catering") and "catering" not in out:
                out["catering"] = num
                out["catering_ref"] = str(row[cr] if cr is not None and cr < len(row) and row[cr] is not None else "")
            elif labels[0].startswith("total") and "total" not in out:
                out["total"] = num
            elif labels[0].startswith("remaining") and "remaining" not in out:
                out["remaining"] = num
    return out


def evaluate_catering(env, dest: str) -> dict:
    exp = pc.expected("x3")
    checks = {}
    q = pc.sql(env, pc.webapp_db(APP), {
        "pos": "SELECT p.id, p.number, p.supplier_id, p.event_id, p.status, s.service_pct FROM pos p JOIN suppliers s ON s.id=p.supplier_id",
        "lines": "SELECT l.po_id, c.code, l.qty, c.price FROM po_lines l JOIN catalogue c ON c.id=l.item_id"})
    if q is None:
        return pc.result(0.0, {"app_db": "unreadable"})
    want = exp["lines"]
    cands = []
    for p in q["pos"]:
        if p["supplier_id"] != FF or p["status"] == "cancelled" or p["id"] == WINTER_PO:
            continue
        lines = [(l["code"], l["qty"], l["price"]) for l in q["lines"] if l["po_id"] == p["id"] and l["qty"] > 0]
        have = {c: n for c, n, _ in lines}
        codes = set(want) | set(have)
        frac = sum(have.get(c) == want.get(c) for c in codes) / len(codes)
        cands.append((p["status"] == "submitted" and p["event_id"] == ANNIV, frac, p, have, lines))
    cands.sort(key=lambda c: (c[0], c[1]), reverse=True)
    best = cands[0] if cands else None
    active_anniv = [c for c in cands if c[2]["event_id"] == ANNIV]
    winter = next((p for p in q["pos"] if p["id"] == WINTER_PO), {})
    no_dupes = len(active_anniv) == 1 and winter.get("status") == "submitted"
    if best:
        _, frac, p, have, lines = best
        total = _total(lines, p["service_pct"])
        submitted, charged = p["status"] == "submitted", p["event_id"] == ANNIV
        checks.update({"po": p["number"], "po_status": p["status"], "charged_to_anniversary": charged, "line_match": round(frac, 3),
                       "line_diff": {c: [have.get(c, 0), want.get(c, 0)] for c in set(want) | set(have) if have.get(c) != want.get(c)},
                       "po_total": float(total)})
    else:
        frac, submitted, charged, total, p = 0.0, False, False, None, None
        checks["po"] = None
    checks["single_active_order"] = no_dupes

    # budget sheet
    path = pc.fetch(env, BUDGET, dest + "_budget.xlsx")
    b = {}
    if path:
        try:
            b = _budget(path)
        except Exception as e:
            checks["budget_error"] = repr(e)[:200]
    bud = [False] * 4
    if total is not None and b:
        t = float(total)
        bud = [b.get("catering") is not None and abs(b["catering"] - t) < 0.011,
               p["number"].lower() in str(b.get("catering_ref", "")).lower(),
               b.get("total") is not None and abs(b["total"] - (exp["budget_others"] + t)) < 0.011,
               b.get("remaining") is not None and abs(b["remaining"] - (exp["approved"] - exp["budget_others"] - t)) < 0.011]
    checks.update({"budget_read": b, "budget_ok": bud})

    # draft to the caterer
    boxes = xc.mailboxes(env) or {}
    mine = [m for m in xc.outgoing(boxes, exp["msg"].values()) if xc.addressed(m, exp["dev"])]
    text = " ".join(m["text"] for m in mine)
    po_in_mail = bool(p) and p["number"].lower() in text.lower()
    names = [n for n in exp["dietary_names"] if all(x in xc.flat(text) for x in _name_parts(n))]
    checks.update({"caterer_draft": bool(mine), "po_number_in_draft": po_in_mail, "dietary_names": f"{len(names)}/{len(exp['dietary_names'])}"})

    score = (0.45 * frac + 0.08 * submitted + 0.07 * (charged and submitted) + 0.05 * (no_dupes and submitted)
             + 0.07 * bud[0] + 0.03 * bud[1] + 0.02 * bud[2] + 0.03 * bud[3]
             + 0.05 * bool(mine) + 0.07 * po_in_mail + 0.08 * len(names) / len(exp["dietary_names"]))
    score = xc.capped(score, [(best is None or not charged, 0.3), (not submitted, 0.6), (frac < 0.7, 0.6)])
    return pc.result(score, checks)
