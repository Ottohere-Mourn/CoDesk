"""D3 - kestrel-site: static website of a physiotherapy practice + live-rebuilding preview server (port 8113).

A (pgd3restyle): restyle the landing page to the 2026 design: tokens from a scanned style guide (PDF, image only),
                 copy/layout from a mockup PNG, phone button in the shared header.
B (pgd3pages):   three new pages from the practice manager's Writer document (the price list is a scanned image
                 inside it), navigation order, and fix every broken link the site's link checker reports.
Both edit templates/base.html (header vs nav/footer) and look at the same live preview in the same Chrome.
"""
import io
import json
import os
import random
import sys

from PIL import Image, ImageDraw, ImageFilter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_devrepo import dev_sh, font, make_repo, scan_pdf  # noqa: E402

PORT = 8113
MAREIKE = ("Mareike Holst", "mareike@kestrel-physio.example")
NIKLAS = ("Niklas Wendt", "hello@wendt-web.example")

OLD_TOKENS = {"color-primary": "#2a6f97", "color-accent": "#e07a5f", "color-ink": "#1b1b1b", "color-surface": "#ffffff",
              "font-size-body": "16px", "font-size-h1": "36px", "font-size-h2": "24px", "line-height-body": "1.5",
              "space-section": "48px", "radius-card": "4px", "radius-button": "4px"}
NEW_TOKENS = {"color-primary": "#1f6f78", "color-accent": "#f2a541", "color-ink": "#22303c", "color-surface": "#faf7f2",
              "font-size-body": "17px", "font-size-h1": "44px", "font-size-h2": "28px", "line-height-body": "1.6",
              "space-section": "72px", "radius-card": "12px", "radius-button": "999px"}
HERO = {"headline": "Move without pain again.",
        "subline": "Physiotherapy, manual therapy and sports rehab in Bremen-Neustadt – appointments within five working days.",
        "cta": "Book an appointment", "cta_href": "/contact.html", "phone": "0421 555 0193", "tel": "+494215550193"}
CARDS = [("Manual therapy", "Hands-on treatment for joints and spine."),
         ("Sports rehabilitation", "Back to your sport – stronger than before."),
         ("Back school", "Small groups, big effect: six sessions for a healthy back.")]
TREATMENTS = [
    ("Physiotherapy (KG)", "Individual exercise therapy after injuries and operations and for chronic complaints."),
    ("Manual therapy", "Gentle mobilisation of joints and spine to relieve pain and restore movement."),
    ("Sports rehabilitation", "Return-to-sport programmes with strength, stability and running analysis."),
    ("Lymphatic drainage", "Manual lymphatic drainage to reduce swelling after surgery or with lymphoedema."),
    ("Back school", "Small-group course of six sessions for a stronger, pain-free back at work and at home."),
    ("Pelvic floor therapy", "Discreet one-to-one therapy for women and men, before and after childbirth or surgery."),
]
TEAM = [("Mareike Holst", "Practice owner, physiotherapist", "Manual therapy (OMT), sports physiotherapy (DOSB)"),
        ("Jonas Feddersen", "Physiotherapist", "Lymphatic drainage (MLD/CDT), pelvic floor therapy"),
        ("Aylin Demir", "Sports therapist", "M.Sc. Sports Science, running analysis"),
        ("Piet Ohlsen", "Physiotherapist", "Bobath, back school instructor")]
PRICES = [("Physiotherapy (KG)", "25 min", "38.00"), ("Manual therapy", "30 min", "52.00"),
          ("Sports rehabilitation", "45 min", "68.00"), ("Lymphatic drainage", "45 min", "55.00"),
          ("Lymphatic drainage", "60 min", "72.00"), ("Pelvic floor therapy", "45 min", "64.00"),
          ("Back school", "6 sessions", "129.00")]
NAV = [("/", "Home"), ("/treatments.html", "Treatments"), ("/team.html", "Team"), ("/prices.html", "Prices"),
       ("/news.html", "News"), ("/contact.html", "Contact")]

# ------------------------------------------------------------------ site generator (stdlib)
MARKDOWN_PY = r'''"""A small Markdown subset: # headings, paragraphs, - lists, 1. lists, | tables |, **bold**, *italic*,
[links](url), ![images](src), and raw HTML blocks (lines starting with '<' are copied as they are)."""
import html
import re


def inline(text):
    text = html.escape(text, quote=False)
    text = re.sub(r'!\[([^\]]*)\]\(([^)\s]+)\)', r'<img src="\2" alt="\1">', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)\s]+)\)', r'<a href="\2">\1</a>', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', text)
    text = re.sub(r'(?<![*\w])\*([^*]+)\*(?!\*)', r'<em>\1</em>', text)
    return text


def to_html(md):
    out, para, lst, table = [], [], None, []

    def flush():
        nonlocal para, lst, table
        if para:
            out.append('<p>%s</p>' % inline(' '.join(para)))
            para = []
        if lst:
            tag, items = lst
            out.append('<%s>%s</%s>' % (tag, ''.join('<li>%s</li>' % inline(i) for i in items), tag))
            lst = None
        if table:
            rows = [[c.strip() for c in r.strip().strip('|').split('|')] for r in table]
            body = [r for r in rows[1:] if not all(re.fullmatch(r':?-+:?', c) for c in r)]
            out.append('<table><thead><tr>%s</tr></thead><tbody>%s</tbody></table>' % (
                ''.join('<th>%s</th>' % inline(c) for c in rows[0]),
                ''.join('<tr>%s</tr>' % ''.join('<td>%s</td>' % inline(c) for c in r) for r in body)))
            table = []

    for line in md.splitlines():
        s = line.strip()
        if not s:
            flush()
        elif s.startswith('<'):
            flush()
            out.append(line)
        elif s.startswith('|'):
            if para or lst:
                flush()
            table.append(s)
        elif re.match(r'#{1,4} ', s):
            flush()
            level = len(s.split(' ')[0])
            out.append('<h%d>%s</h%d>' % (level, inline(s[level + 1:]), level))
        elif re.match(r'[-*] ', s) or re.match(r'\d+\. ', s):
            tag = 'ul' if s[0] in '-*' else 'ol'
            if para or table or (lst and lst[0] != tag):
                flush()
            if not lst:
                lst = (tag, [])
            lst[1].append(re.sub(r'^([-*]|\d+\.) ', '', s))
        else:
            if lst or table:
                flush()
            para.append(s)
    flush()
    return '\n'.join(out)
'''

BUILD_PY = r'''"""Build the site: content/*.md (front matter + Markdown) -> _site/*.html through templates/, plus static/.

    python3 -m sitegen.build            # one-off build into _site/
"""
import os
import re
import shutil

from .markdown import to_html

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, '_site')
SOURCES = ('content', 'templates', 'static')


def front_matter(text):
    meta = {}
    if text.startswith('---\n'):
        head, _, text = text[4:].partition('\n---\n')
        for line in head.splitlines():
            if ':' in line:
                k, v = line.split(':', 1)
                meta[k.strip()] = v.strip()
    return meta, text


def render(template, values):
    def sub(m):
        key = m.group(1)
        if key not in values:
            raise KeyError('unknown placeholder {{ %s }}' % key)
        return values[key]
    return re.sub(r'\{\{\s*(\w+)\s*\}\}', sub, template)


def load(name):
    with open(os.path.join(ROOT, 'templates', name), encoding='utf-8') as f:
        return f.read()


def build(out=OUT):
    tmp = out + '.tmp'
    shutil.rmtree(tmp, ignore_errors=True)
    shutil.copytree(os.path.join(ROOT, 'static'), os.path.join(tmp, 'static'))
    base = load('base.html')
    pages = 0
    for name in sorted(os.listdir(os.path.join(ROOT, 'content'))):
        if not name.endswith('.md'):
            continue
        with open(os.path.join(ROOT, 'content', name), encoding='utf-8') as f:
            meta, body = front_matter(f.read())
        content = to_html(body)
        tpl = meta.get('template', 'page')
        if tpl != 'page':
            content = render(load(tpl + '.html'), {'content': content, 'title': meta.get('title', '')})
        else:
            content = '<article class="page">%s</article>' % content
        html = render(base, {'title': meta.get('title', name[:-3]), 'content': content,
                             'description': meta.get('description', '')})
        with open(os.path.join(tmp, name[:-3] + '.html'), 'w', encoding='utf-8') as f:
            f.write(html)
        pages += 1
    shutil.rmtree(out, ignore_errors=True)
    os.replace(tmp, out)
    return pages


def source_stamp():
    stamp = 0
    for d in SOURCES:
        for root, _, files in os.walk(os.path.join(ROOT, d)):
            for f in files:
                stamp = max(stamp, os.stat(os.path.join(root, f)).st_mtime_ns)
            stamp = max(stamp, os.stat(root).st_mtime_ns)
    return stamp


if __name__ == '__main__':
    print('built %d pages into %s' % (build(), OUT))
'''

SERVE_PY = r'''"""Live preview: serves _site/ and rebuilds it whenever a file under content/, templates/ or static/ changed.

    python3 -m sitegen.serve --port 8113        (or ./dev.sh start|restart)

Changes to the generator itself (sitegen/*.py) need a restart.
"""
import argparse
import html
import mimetypes
import os
import sys
import threading
import traceback
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from .build import OUT, build, source_stamp

LOCK = threading.Lock()
STATE = {'stamp': None, 'error': None}


def ensure_built():
    with LOCK:
        stamp = source_stamp()
        if stamp != STATE['stamp']:
            try:
                n = build()
                STATE['error'] = None
                sys.stderr.write('rebuilt %d pages\n' % n)
            except Exception:
                STATE['error'] = traceback.format_exc()
                sys.stderr.write(STATE['error'])
            STATE['stamp'] = stamp
        return STATE['error']


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        sys.stderr.write('%s %s\n' % (self.log_date_time_string(), fmt % args))

    def send(self, status, data, ctype):
        self.send_response(status)
        self.send_header('Content-Type', ctype)
        self.send_header('Content-Length', str(len(data)))
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        err = ensure_built()
        if err:
            body = '<h1>Build failed</h1><pre>%s</pre>' % html.escape(err)
            return self.send(500, body.encode(), 'text/html; charset=utf-8')
        path = urllib.parse.unquote(urllib.parse.urlsplit(self.path).path)
        if path.endswith('/'):
            path += 'index.html'
        full = os.path.normpath(os.path.join(OUT, path.lstrip('/')))
        if not full.startswith(OUT) or not os.path.isfile(full):
            return self.send(404, b'<h1>404 Not Found</h1>', 'text/html; charset=utf-8')
        with open(full, 'rb') as f:
            data = f.read()
        ctype = mimetypes.guess_type(full)[0] or 'application/octet-stream'
        if ctype.startswith('text/'):
            ctype += '; charset=utf-8'
        return self.send(200, data, ctype)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--port', type=int, default=8113)
    args = ap.parse_args()
    ensure_built()
    print('preview on http://127.0.0.1:%d/' % args.port, flush=True)
    ThreadingHTTPServer(('127.0.0.1', args.port), Handler).serve_forever()


if __name__ == '__main__':
    main()
'''

LINKCHECK_PY = r'''#!/usr/bin/env python3
"""Crawl the site from a start URL and report broken internal links and images.

    python3 tools/linkcheck.py http://127.0.0.1:8113/
"""
import sys
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        for k, v in attrs:
            if k in ('href', 'src') and v:
                self.links.append(v)


def fetch(url):
    try:
        with urllib.request.urlopen(url, timeout=10) as r:
            return r.status, r.headers.get('Content-Type', ''), r.read()
    except urllib.error.HTTPError as e:
        return e.code, '', b''
    except Exception:
        return 0, '', b''


def main(start):
    host = urllib.parse.urlsplit(start).netloc
    todo, seen, broken = [start], set(), []
    status = {}
    while todo:
        url = todo.pop()
        if url in seen:
            continue
        seen.add(url)
        code, ctype, body = fetch(url)
        status[url] = code
        if code != 200 or 'html' not in ctype:
            continue
        p = Links()
        p.feed(body.decode('utf-8', 'replace'))
        for link in p.links:
            target = urllib.parse.urljoin(url, link).split('#')[0]
            parts = urllib.parse.urlsplit(target)
            if parts.scheme not in ('http', 'https') or parts.netloc != host:
                continue
            if target not in status:
                status[target] = fetch(target)[0]
            if status[target] != 200:
                broken.append((url, link, status[target]))
            elif target.endswith('.html') or target.endswith('/'):
                todo.append(target)
    for page, link, code in broken:
        print('BROKEN %s  on %s  (HTTP %s)' % (link, page, code))
    print('%d pages checked, %d broken links' % (len([u for u in seen]), len(broken)))
    return 1 if broken else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else 'http://127.0.0.1:8113/'))
'''

TOKENS_CSS = """/* Design tokens - the only place where colours, type sizes and spacing are defined. */
:root {
  /* Brand colours */
  --color-primary: %(color-primary)s;     /* header background, headings */
  --color-accent: %(color-accent)s;      /* buttons, highlights */
  --color-ink: %(color-ink)s;         /* body text */
  --color-surface: %(color-surface)s;     /* page background */

  /* Typography */
  --font-family: "Lato", "DejaVu Sans", Arial, sans-serif;
  --font-size-body: %(font-size-body)s;
  --font-size-h1: %(font-size-h1)s;
  --font-size-h2: %(font-size-h2)s;
  --line-height-body: %(line-height-body)s;

  /* Spacing and shapes */
  --space-section: %(space-section)s;    /* vertical padding of page sections */
  --radius-card: %(radius-card)s;
  --radius-button: %(radius-button)s;
}
""" % OLD_TOKENS

SITE_CSS = """@import url("tokens.css");
* { box-sizing: border-box; }
body { margin: 0; font-family: var(--font-family); font-size: var(--font-size-body);
  line-height: var(--line-height-body); color: var(--color-ink); background: var(--color-surface); }
.wrap { max-width: 1120px; margin: 0 auto; padding: 0 24px; }
.site-header { background: var(--color-primary); color: #fff; }
.site-header .wrap { display: flex; align-items: center; gap: 32px; height: 72px; }
.logo { color: #fff; font-weight: bold; font-size: 22px; text-decoration: none; margin-right: auto; }
.site-nav a { color: #fff; text-decoration: none; margin-left: 20px; opacity: .9; }
.site-nav a:hover { opacity: 1; text-decoration: underline; }
h1 { font-size: var(--font-size-h1); line-height: 1.15; margin: 0 0 16px; color: var(--color-primary); }
h2 { font-size: var(--font-size-h2); color: var(--color-primary); }
.button { display: inline-block; background: var(--color-accent); color: var(--color-ink); padding: 12px 26px;
  border-radius: var(--radius-button); text-decoration: none; font-weight: bold; }
.button-outline { background: transparent; color: #fff; border: 2px solid #fff; padding: 6px 16px; }
.hero { padding: var(--space-section) 0; }
.hero .lead { font-size: 1.25em; max-width: 640px; }
.cards { display: grid; grid-template-columns: repeat(auto-fit, minmax(240px, 1fr)); gap: 24px; padding-bottom: var(--space-section); }
.card { background: #fff; border-radius: var(--radius-card); padding: 24px; box-shadow: 0 1px 3px rgba(0,0,0,.08); }
.news-teaser { padding-bottom: var(--space-section); }
article.page { padding: var(--space-section) 0; max-width: 820px; }
article.page img { max-width: 100%; }
table { border-collapse: collapse; } th, td { padding: 8px 14px; border-bottom: 1px solid #ddd; text-align: left; }
.site-footer { border-top: 1px solid #ddd; padding: 24px 0 48px; font-size: 14px; color: #666; }
.site-footer a { color: inherit; }
"""

BASE_HTML = """<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{{ title }} | Kestrel Physio Bremen</title>
  <meta name="description" content="{{ description }}">
  <link rel="stylesheet" href="/static/css/site.css">
</head>
<body>
  <header class="site-header">
    <div class="wrap">
      <a class="logo" href="/">Kestrel Physio</a>
      <nav class="site-nav">
        <a href="/">Home</a>
        <a href="/news.html">News</a>
        <a href="/contact.html">Contact</a>
      </nav>
    </div>
  </header>
  <main class="wrap">
{{ content }}
  </main>
  <footer class="site-footer">
    <div class="wrap">
      &copy; 2026 Kestrel Physio, Westerstraße 41, 28199 Bremen &middot;
      <a href="/imprint.htm">Imprint</a> &middot; <a href="/privacy.html">Privacy</a>
    </div>
  </footer>
</body>
</html>
"""

HOME_HTML = """<section class="hero">
  <h1>Welcome to Kestrel Physio</h1>
  <p class="lead">Your physiotherapy practice in Bremen.</p>
  <a class="button" href="/contact.html">Contact us</a>
</section>
<section class="cards">
  <article class="card">
    <h2>Physiotherapy</h2>
    <p>Exercise therapy for every age.</p>
  </article>
  <article class="card">
    <h2>Massage</h2>
    <p>Classic massage and heat therapy.</p>
  </article>
</section>
<section class="news-teaser">
  <h2>Latest news</h2>
  {{ content }}
  <p><a href="/news.html">All news</a></p>
</section>
"""

CONTENT = {
    "content/index.md": "---\ntitle: Physiotherapy in Bremen-Neustadt\ntemplate: home\n"
                        "description: Physiotherapy, manual therapy and sports rehab in Bremen-Neustadt.\n---\n"
                        "From 1 November we are open on Saturday mornings too (9:00-13:00).\n",
    "content/contact.md": "---\ntitle: Contact\n---\n# Contact\n\nKestrel Physio, Westerstraße 41, 28199 Bremen\n\n"
                          "Phone: 0421 555 0193 (Mon-Fri 8:00-18:00)  \nE-mail: praxis@kestrel-physio.example\n\n"
                          "To book an appointment, call us or send an e-mail with your prescription. Please read our "
                          "[privacy notice](/privacy) before sending health information by e-mail.\n\n"
                          "## Getting here\n\nTram 1 and 8, stop *Westerstraße*. Two parking spaces in the courtyard.\n",
    "content/news.md": "---\ntitle: News\n---\n# News\n\n## Saturday opening hours (October 2026)\n\n"
                       "From 1 November we are open on Saturdays from 9:00 to 13:00.\n\n"
                       "## Team summer run (August 2026)\n\n![The team at the summer run](/static/img/team-2025.jpg)\n\n"
                       "Six of us ran the 10 km of the Weser summer run for the children's hospital.\n\n"
                       "## New website (March 2025)\n\nOur new website is online.\n",
    "content/imprint.md": "---\ntitle: Imprint\n---\n# Imprint\n\nKestrel Physio, Mareike Holst, Westerstraße 41, "
                          "28199 Bremen. Phone 0421 555 0193.\n\nResponsible for content: Mareike Holst.\n",
    "content/privacy.md": "---\ntitle: Privacy\n---\n# Privacy notice\n\nWe only process the data you send us to "
                          "arrange and document your treatment. See the [imprint](/imprint.html) for our contact details.\n",
}

README = """# kestrel-site

Website of Kestrel Physio, Bremen. Plain static HTML generated by a tiny Python generator (stdlib only).

| path | |
|---|---|
| `content/*.md` | pages: front matter (`title`, optional `template`, `description`) + Markdown; `foo.md` -> `/foo.html` |
| `templates/base.html` | page frame: header with the navigation, footer |
| `templates/home.html` | landing page sections (used by `content/index.md`) |
| `static/css/tokens.css` | design tokens (colours, type sizes, spacing) |
| `static/css/site.css` | layout and components, uses the tokens |
| `docs/design/` | design hand-over from the agency |

## Preview

    ./dev.sh start          # http://127.0.0.1:8113/  (stop|restart|status|log)

The preview rebuilds `_site/` automatically when anything under content/, templates/ or static/ changes - just
reload the page. Only changes to `sitegen/*.py` need `./dev.sh restart`.

## Checks

    python3 tools/linkcheck.py http://127.0.0.1:8113/     # broken internal links and images (run against the preview)
"""

GITIGNORE = "__pycache__/\n*.pyc\n_site/\n_site.tmp/\nlogs/\n.devserver.pid\n"


def svg(color, label):
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="640" height="360" viewBox="0 0 640 360">'
            f'<rect width="640" height="360" fill="{color}"/><text x="40" y="200" font-family="sans-serif" '
            f'font-size="36" fill="#fff">{label}</text></svg>\n')


# ------------------------------------------------------------------ design hand-over images
def hexrgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))


def mockup_png(path):
    t = NEW_TOKENS
    W, H = 1440, 1000
    im = Image.new("RGB", (W, H), hexrgb(t["color-surface"]))
    d = ImageDraw.Draw(im)
    prim, acc, ink = hexrgb(t["color-primary"]), hexrgb(t["color-accent"]), hexrgb(t["color-ink"])
    d.rectangle((0, 0, W, 72), fill=prim)
    x0 = 160
    d.text((x0, 22), "Kestrel Physio", fill="white", font=font(22, family="lato", bold=True))
    x = 760
    for label in ("Home", "News", "Contact"):
        d.text((x, 25), label, fill=(225, 235, 236), font=font(17, family="lato"))
        x += 90
    d.rounded_rectangle((1060, 16, 1280, 56), 20, outline="white", width=2)
    d.text((1088, 25), "Call " + HERO["phone"], fill="white", font=font(17, family="lato", bold=True))
    y = 72 + 72
    d.text((x0, y), HERO["headline"], fill=prim, font=font(44, family="lato", bold=True))
    y += 70
    sub = HERO["subline"]
    cut = sub.index("Neustadt") + len("Neustadt")
    for part in (sub[:cut], sub[cut + 1:]):
        d.text((x0, y), part.strip(), fill=ink, font=font(21, family="lato"))
        y += 34
    y += 26
    d.rounded_rectangle((x0, y, x0 + 270, y + 52), 26, fill=acc)
    d.text((x0 + 34, y + 14), HERO["cta"], fill=ink, font=font(18, family="lato", bold=True))
    cta_y = y
    y += 52 + 72
    cw, gap = 352, 24
    for i, (title, text) in enumerate(CARDS):
        cx = x0 + i * (cw + gap)
        d.rounded_rectangle((cx + 2, y + 3, cx + cw + 2, y + 190 + 3), 12, fill=(230, 226, 220))
        d.rounded_rectangle((cx, y, cx + cw, y + 190), 12, fill="white")
        d.text((cx + 24, y + 26), title, fill=prim, font=font(28, family="lato", bold=True))
        words, line, ly = text.split(), "", y + 82
        for w_ in words:
            if d.textlength(line + " " + w_, font=font(17, family="lato")) > cw - 48:
                d.text((cx + 24, ly), line.strip(), fill=ink, font=font(17, family="lato"))
                ly += 27
                line = ""
            line += " " + w_
        d.text((cx + 24, ly), line.strip(), fill=ink, font=font(17, family="lato"))
    y += 190 + 72
    d.text((x0, y), "Latest news", fill=prim, font=font(28, family="lato", bold=True))
    d.text((x0, y + 50), "(news teaser - unchanged)", fill=(150, 150, 150), font=font(17, family="lato"))
    d.text((x0, y + 84), "All news", fill=prim, font=font(17, family="lato"))
    # designer's red-line annotations
    red = (214, 40, 57)
    d.line((x0 + 270, cta_y + 26, x0 + 420, cta_y + 26), fill=red, width=2)
    d.text((x0 + 430, cta_y + 14), "links to the Contact page (/contact.html)", fill=red, font=font(16, bold=True))
    d.line((1170, 56, 1170, 110), fill=red, width=2)
    d.text((900, 114), "phone button on every page: tel:+49 421 555 0193", fill=red, font=font(16, bold=True))
    d.text((x0, H - 60), "Kestrel Physio - landing page 2026 - v3 (final) - Studio Wendt", fill=(150, 150, 150),
           font=font(15))
    d.text((x0, H - 36), "Colours, type sizes and spacing: see the style guide. Navigation items are managed by the "
                         "practice and are not part of this design.", fill=(150, 150, 150), font=font(15))
    im.save(path, optimize=True)


def style_guide_pdf(path):
    t = NEW_TOKENS
    W, H = 1240, 1754
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    d.text((100, 90), "Kestrel Physio", fill="black", font=font(40, bold=True))
    d.text((100, 145), "Brand & UI style guide 2026 - excerpt for the website (p. 7)", fill="black", font=font(24))
    d.line((100, 190, W - 100, 190), fill="black", width=2)
    d.text((100, 220), "1  Colours", fill="black", font=font(28, bold=True))
    sw = [("Primary", "Kestrel Teal", t["color-primary"], "header background, headings"),
          ("Accent", "Sunrise", t["color-accent"], "buttons and highlights"),
          ("Text", "Slate", t["color-ink"], "body text"),
          ("Background", "Paper", t["color-surface"], "page background")]
    y = 280
    for role, name, hx, use in sw:
        d.rectangle((110, y, 230, y + 90), fill=hexrgb(hx), outline="black")
        d.text((260, y + 4), f"{role} - \"{name}\"", fill="black", font=font(24, bold=True))
        d.text((260, y + 38), f"HEX {hx.upper()}     RGB {', '.join(str(v) for v in hexrgb(hx))}", fill="black",
               font=font(22, mono=True))
        d.text((260, y + 68), f"use: {use}", fill=(60, 60, 60), font=font(19))
        y += 125
    y += 20
    d.text((100, y), "2  Typography (web)", fill="black", font=font(28, bold=True))
    y += 56
    for ln in [f"Body text        {t['font-size-body']}, line height {t['line-height-body']}",
               f"Heading 1 (H1)   {t['font-size-h1']}",
               f"Heading 2 (H2)   {t['font-size-h2']}",
               "Typeface         Lato (fallback: system sans-serif)"]:
        d.text((130, y), ln, fill="black", font=font(23, mono=True))
        y += 40
    y += 30
    d.text((100, y), "3  Spacing and shapes", fill="black", font=font(28, bold=True))
    y += 56
    for ln in [f"Vertical padding of page sections   {t['space-section']}",
               f"Card corner radius                  {t['radius-card']}",
               f"Buttons                             pill shape, border-radius {t['radius-button']}"]:
        d.text((130, y), ln, fill="black", font=font(23, mono=True))
        y += 40
    d.text((100, H - 120), "Studio Wendt - kestrel_styleguide_2026_v2.pdf - page 7 of 16", fill=(90, 90, 90), font=font(18))
    # photocopy look, but keep colour
    r = random.Random(7)
    im = im.rotate(0.4, resample=Image.BICUBIC, fillcolor="white").filter(ImageFilter.GaussianBlur(0.6))
    px = im.load()
    for _ in range(W * H // 500):
        x, yy = r.randrange(W), r.randrange(H)
        v = r.randrange(120, 200)
        px[x, yy] = (v, v, v)
    scan_pdf([im], path)


def price_scan_png():
    W, H = 1100, 760
    im = Image.new("L", (W, H), 250)
    d = ImageDraw.Draw(im)
    d.text((60, 50), "Kestrel Physio - Price list for self-payers", fill=0, font=font(34, serif=True, bold=True))
    d.text((60, 100), "valid from 1 November 2026, all prices incl. VAT where applicable", fill=40, font=font(20, serif=True))
    y = 170
    d.line((60, y - 10, W - 60, y - 10), fill=0, width=2)
    for name, dur, price in PRICES:
        d.text((70, y), name, fill=0, font=font(26, serif=True))
        d.text((560, y), dur, fill=0, font=font(26, serif=True))
        p = "€ " + price
        d.text((W - 80 - d.textlength(p, font=font(26, serif=True)), y), p, fill=0, font=font(26, serif=True))
        y += 62
        d.line((60, y - 14, W - 60, y - 14), fill=150, width=1)
    d.text((60, y + 20), "Mareike Holst, practice owner", fill=60, font=font(20, serif=True))
    r = random.Random(11)
    im = im.rotate(-0.8, resample=Image.BICUBIC, fillcolor=250).filter(ImageFilter.GaussianBlur(0.7))
    px = im.load()
    for _ in range(W * H // 300):
        x, yy = r.randrange(W), r.randrange(H)
        px[x, yy] = max(0, px[x, yy] - r.randrange(30, 110))
    buf = io.BytesIO()
    im.convert("RGB").save(buf, "PNG")
    return buf.getvalue()


def writer_doc(path):
    from docx import Document
    from docx.shared import Cm, Pt
    doc = Document()
    doc.core_properties.author = "Ilka Brandes"
    doc.add_heading("Website update - October 2026", 0)
    doc.add_paragraph("From: Ilka Brandes (practice management), 12 October 2026")
    doc.add_paragraph("Hello! Here are the texts for the three new pages. Please put each on its own page at "
                      "/treatments.html, /team.html and /prices.html, with the page titles Treatments, Team and Prices. "
                      "The menu should then read, from left to right: Home, Treatments, Team, Prices, News, Contact.")
    doc.add_paragraph("Our web agency also says the site has a few broken links - please fix them while you are at it.")
    doc.add_heading("1. Treatments", 1)
    doc.add_paragraph("We treat patients of all ages with a prescription from their doctor or as self-payers.")
    for name, text in TREATMENTS:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(name + ": ").bold = True
        p.add_run(text)
    doc.add_heading("2. Team", 1)
    doc.add_paragraph("Four therapists, one team:")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for c, h in zip(table.rows[0].cells, ("Name", "Role", "Qualifications")):
        c.text = h
    for row in TEAM:
        cells = table.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    doc.add_heading("3. Prices", 1)
    doc.add_paragraph("Prices for private patients and self-payers. With statutory health insurance and a prescription "
                      "you only pay the legal co-payment.")
    doc.add_paragraph("Here is the new price list (scanned from the printed version on the counter - the "
                      "page should show it as a proper table, not as a picture):")
    doc.add_picture(io.BytesIO(price_scan_png()), width=Cm(15))
    doc.add_paragraph("Appointments cancelled less than 24 hours in advance are charged at 50 %.")
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size is None:
                r.font.size = Pt(11)
    doc.save(path)


def build(out):
    os.makedirs(out, exist_ok=True)
    tmp_m, tmp_s = os.path.join(out, "_m.png"), os.path.join(out, "_s.pdf")
    mockup_png(tmp_m)
    style_guide_pdf(tmp_s)
    writer_doc(os.path.join(out, "Website_update_Oct2026.docx"))
    files = {".gitignore": GITIGNORE, "README.md": README, "sitegen/__init__.py": "", "sitegen/markdown.py": MARKDOWN_PY,
             "sitegen/build.py": BUILD_PY, "sitegen/serve.py": SERVE_PY, "tools/linkcheck.py": LINKCHECK_PY,
             "templates/base.html": BASE_HTML, "templates/home.html": HOME_HTML, "static/css/tokens.css": TOKENS_CSS,
             "static/css/site.css": SITE_CSS, "static/img/team-summer-run-2026.svg": svg("#2a6f97", "Summer run 2026"),
             "static/img/favicon.svg": svg("#2a6f97", "K"), **CONTENT,
             "content/news.md": "---\ntitle: News\n---\n# News\n\n## New website (March 2025)\n\nOur new website is online.\n"}
    history = [
        ("2025-03-11T10:00:00", NIKLAS, files, "Kestrel Physio website: generator, templates, first pages"),
        ("2026-08-25T19:30:00", MAREIKE, {"content/news.md": CONTENT["content/news.md"]}, "News: summer run"),
        ("2026-10-06T09:15:00", NIKLAS, {"dev.sh": dev_sh("sitegen.serve", PORT, "kestrel-site preview")},
         "Add dev.sh for the live preview"),
        ("2026-10-15T17:40:00", NIKLAS, {"docs/design/landing-2026-mockup.png": open(tmp_m, "rb").read(),
                                         "docs/design/styleguide-2026-p7.pdf": open(tmp_s, "rb").read()},
         "Design hand-over: 2026 landing page mockup and style guide excerpt"),
    ]
    info = make_repo(os.path.join(out, "kestrel-site.tar.gz"), "kestrel-site", history, MAREIKE)
    os.remove(tmp_m)
    os.remove(tmp_s)
    json.dump({"base": info["base"], "port": PORT, "tokens": NEW_TOKENS, "old_tokens": OLD_TOKENS, "hero": HERO,
               "cards": [c[0] for c in CARDS], "treatments": [t[0] for t in TREATMENTS], "team": TEAM,
               "prices": PRICES, "nav": [h for h, _ in NAV]},
              open(os.path.join(out, "expected.json"), "w"), indent=1, ensure_ascii=False)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/d3_build")
