"""o4 - one Impress presentation (a PhD student's workshop talk): A applies the advisor's handwritten comments from a
scanned annotated handout and exports a review PDF; B prepares the deck for the conference AV team (16:9, speaker notes
for every slide from a scanned speaker script, funding acknowledgement slide, notes-pages PDF)."""
from __future__ import annotations

import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_oscan import Page, font, save_scan  # noqa: E402

TITLE = "StreamJoin: Low-Latency Joins over Bursty Data Streams"
SLIDES = [  # (title, bullets) of the deck as it is now (4:3)
    (TITLE, ["Lina Park - Institute for Data Systems, University of Lindenau",
             "DSSW 2026 - Data Stream Systems Workshop, 12 November 2026"]),
    ("Motivaton", ["Sensor, taxi and click streams arrive in bursts",
                   "Window joins must answer within milliseconds",
                   "Today's operators over-provision for the peak"]),
    ("Problem Setting", ["Two unbounded streams R and S, sliding windows of size W",
                         "Goal: minimise p99 join latency under bursts",
                         "We assume a single-node setting",
                         "Fixed memory budget per operator"]),
    ("Key Idea: Adaptive Window Partitioning", ["Split each window into micro-partitions by key range",
                                                "Re-balance partitions as soon as a burst is detected",
                                                "Reduces p99 latency by 82.1% vs. static partitioning"]),
    ("System Overview", ["Burst detector (EWMA of the arrival rate)", "Partition manager",
                         "Join workers with a shared state store"]),
    ("Experimental Setup", ["Data: NYC Taxi (1.2B events), Twitter sample (310M tweets)",
                            "Hardware: 4 nodes, Intel Xeon 6248, 192 GB RAM each",
                            "Baselines: Flink interval join, SplitJoin, BiStream"]),
    ("Results: Latency", ["p99 latency 18 ms vs. 112 ms for Flink", "Stable under 10x input bursts"]),
    ("Results: Throughput", ["1.9M tuples/s per node", "Within 6% of the fastest baseline"]),
    ("Backup: Parameter Sweep", ["Partition size 64-4096 keys", "Detection threshold 1.5x-4x mean rate"]),
    ("Related Work", ["SplitJoin [Najafi et al., 2016]", "BiStream [Lin et al., 2015]", "Flink interval joins"]),
    ("Limitations and Future Work", ["Skewed keys can still create hot partitions", "Extend to multi-way joins"]),
    ("Thank You", ["lina.park@ids.uni-lindenau.de"]),
]

# Advisor's comments on the handout -> expected deck after task A
FINAL_TITLES = [TITLE, "Motivation", "Problem Setting", "Key Idea: Adaptive Window Partitioning", "System Overview",
                "Experimental Setup", "Results: Throughput", "Results: Latency", "Related Work",
                "Limitations and Future Work", "Questions?"]
COMMENTS = {  # original slide index (1-based) -> annotation lines
    2: ["typo! -> Motivation"],
    3: ["remove this bullet -", "we run on 4 nodes now"],
    4: ["82.1% is from the old run;", "final number is 84.3%"],
    6: ["CPU is wrong: it's a", "Xeon Gold 6338"],
    7: ["swap 7 and 8: show", "throughput first, then latency"],
    9: ["drop this slide", "(keep it out of the talk)"],
    10: ["add: Chen et al., VLDB 2025", "(learned partitioning)"],
    11: ["add the code link:", "github.com/lpark/streamjoin"],
    12: ["make this 'Questions?'"],
}

NOTES = [  # (id, title aliases, notes text) - Lina's speaker script, in the order of the original deck
    ("title", [TITLE], "Good morning, I am Lina Park from the Institute for Data Systems. In the next fifteen minutes I "
                       "will show how StreamJoin keeps window joins fast when the input suddenly bursts."),
    ("motivation", ["Motivation", "Motivaton"], "Many real streams, like taxi trips or clicks, arrive in bursts. A join "
                   "operator has to answer within milliseconds, so today systems reserve resources for the worst peak, "
                   "which is expensive."),
    ("problem", ["Problem Setting"], "We join two unbounded streams over sliding windows. Our goal is to minimise the "
                "99th percentile latency while each operator keeps a fixed memory budget."),
    ("idea", ["Key Idea: Adaptive Window Partitioning"], "The key idea is to cut every window into small key-range "
             "partitions and to move them between workers as soon as a burst is detected, instead of partitioning "
             "once up front."),
    ("overview", ["System Overview"], "The system has three parts: a burst detector based on a moving average of the "
                 "arrival rate, a partition manager, and the join workers that share one state store."),
    ("setup", ["Experimental Setup"], "We ran all experiments on four nodes with two real data sets, NYC taxi and a "
              "Twitter sample, and compared against Flink, SplitJoin and BiStream."),
    ("latency", ["Results: Latency"], "Under bursts StreamJoin keeps the 99th percentile latency at about eighteen "
                "milliseconds, while Flink goes above one hundred milliseconds."),
    ("throughput", ["Results: Throughput"], "Throughput stays competitive: almost two million tuples per second per "
                   "node, within six percent of the fastest baseline."),
    ("backup", ["Backup: Parameter Sweep"], "This backup slide shows how the partition size and the detection "
               "threshold affect latency. I only show it if someone asks."),
    ("related", ["Related Work"], "SplitJoin and BiStream parallelise joins statically. Recent learned partitioning "
                "work is closest to ours but assumes stable input rates."),
    ("limits", ["Limitations and Future Work"], "Heavily skewed keys can still create hot partitions, and we want to "
               "extend the approach to multi-way joins. The code is open source."),
    ("ack", ["Acknowledgements", "Acknowledgments"], "I would like to thank the German Research Foundation for funding "
            "this work and my colleagues in the systems group for their feedback."),
    ("closing", ["Questions?", "Questions", "Thank You", "Thank you"], "Thank you for listening. I am happy to take "
                "your questions."),
]
ACK_TEXT = ["This work was funded by the German Research Foundation (DFG), grant PA 4471/2-1.",
            "Thanks to the IDS Systems Group and the anonymous DSSW reviewers."]

RULES = """Institute for Data Systems - rules for conference talks (AV team, v2026)

1. Format: 16:9 widescreen ("Screen 16:9" in LibreOffice Impress). Old 4:3 decks must be converted.
2. Every slide must carry speaker notes. The AV team turns them into live captions and the
   recording transcript, so put the full text of your speaker script into the notes of each slide.
3. Publicly funded work: add an "Acknowledgements" slide with the funding statement, placed
   immediately before your closing slide.
4. Send the AV team a PDF of the notes pages (slide + notes on each page) together with the .pptx.
"""


def _thumb(i, title, bullets, w=470, h=352):
    """A printed 4:3 slide as it appears on the handout."""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w - 1, h - 1), outline=(70, 70, 70), width=2)
    d.rectangle((0, 0, w - 1, 58), fill=(31, 56, 100))
    tf = font(16 if len(title) > 40 else 19 if len(title) > 32 else 22, "bold")
    d.text((14, 16), title if len(title) < 46 else title[:44] + "...", fill="white", font=tf)
    y = 80
    for b in bullets:
        words, line = b.split(), ""
        first = True
        for wd in words:
            cand = (line + " " + wd).strip()
            if d.textlength(cand, font=font(15)) > w - 60 and line:
                d.text((40 if not first else 22, y), ("- " if first else "  ") + line, fill=(25, 25, 25), font=font(15))
                y += 22
                line, first = wd, False
            else:
                line = cand
        d.text((40 if not first else 22, y), ("- " if first else "  ") + line, fill=(25, 25, 25), font=font(15))
        y += 30
    d.text((w - 30, h - 24), str(i), fill=(90, 90, 90), font=font(13))
    return im


def _handout(path):
    red = (205, 30, 30)
    pages = []
    for pg in range(3):
        p = Page(color=True)
        p.text(80, 40, f"StreamJoin - DSSW 2026 talk (draft)          page {pg + 1}/3", 22, "sans", fill=(80, 80, 80))
        if pg == 0:
            p.hand(640, 70, "Lina - my comments in red. MW", 30, red, angle=2, shear=0.3)
        for k in range(4):
            i = pg * 4 + k + 1
            title, bullets = SLIDES[i - 1]
            x, y = 70, 140 + k * 400
            p.paste(_thumb(i, title, bullets), (x, y))
            notes = COMMENTS.get(i)
            if not notes:
                p.dr.line((600, y + 170, 1150, y + 170), fill=(200, 200, 200), width=1)
                continue
            yy = y + 40
            for ln in notes:
                yy = p.hand(590, yy, ln, 30, red, angle=1.5, shear=0.3) + 4
            # pen marks on the thumbnail itself
            if i == 2:
                p.ink_ellipse((x + 8, y + 6, x + 150, y + 54), red, 3)
            elif i == 3:
                p.ink_line([(x + 18, y + 80 + 2 * 30 + 10 + 22), (x + 380, y + 80 + 2 * 30 + 10 + 22)], red, 4)
            elif i == 4:
                x0 = x + 22 + p.dr.textlength("- Reduces p99 latency by ", font=font(15))
                x1 = x0 + p.dr.textlength("82.1%", font=font(15))
                p.ink_ellipse((x0 - 10, y + 80 + 60 - 8, x1 + 10, y + 80 + 60 + 26), red, 3)
            elif i == 6:
                p.ink_ellipse((x + 200, y + 100, x + 330, y + 160), red, 3)
            elif i == 7:
                p.ink_line([(x + 475, y + 300), (x + 540, y + 360), (x + 475, y + 420)], red, 4)
            elif i == 9:
                p.ink_line([(x + 10, y + 10), (x + 460, y + 340)], red, 5)
                p.ink_line([(x + 460, y + 10), (x + 10, y + 340)], red, 5)
            elif i in (10, 11):
                p.ink_line([(x + 30, y + 300), (x + 60, y + 285), (x + 90, y + 300)], red, 4)
            elif i == 12:
                p.ink_ellipse((x + 8, y + 6, x + 170, y + 54), red, 3)
        pages.append(p)
    save_scan(pages, path, seed=7401, color=True)


def _script(path):
    pages, p, y = [], Page(), 0

    def newpage():
        nonlocal p, y
        p = Page()
        pages.append(p)
        y = 90
    newpage()
    p.text(90, y, "DSSW 2026 talk - speaker script (L. Park), final", 32, "bold")
    y += 60
    y = p.para(90, y, "One block per slide, in the order of the current deck (before Markus' comments). Copy each block "
                      "into the speaker notes of that slide. The acknowledgements slide is new: insert it just before "
                      "the closing slide with the title \"Acknowledgements\" and the text given below.", 23, width=1060)
    y += 25
    for nid, aliases, text in NOTES:
        name = {"motivation": "Motivation", "closing": "Closing slide (Thank You / Questions?)",
                "ack": "Acknowledgements (NEW slide)"}.get(nid, aliases[0])
        need = 70 + 34 * (len(text) // 70 + 2) + (70 if nid == "ack" else 0)
        if y + need > 1650:
            newpage()
        p.text(90, y, "Slide: " + name, 25, "bold")
        y += 40
        if nid == "ack":
            p.text(120, y, "Slide text:", 23, "bold")
            y += 34
            for ln in ACK_TEXT:
                y = p.para(150, y, ln, 23, "serif", width=980)
            y += 6
            p.text(120, y, "Notes:", 23, "bold")
            y += 34
        y = p.para(120, y, text, 24, "serif", width=1010) + 26
    for i, pg in enumerate(pages, 1):
        pg.text(1080, 1660, f"{i}/{len(pages)}", 20)
    save_scan(pages, path, seed=7402)


def build(dst):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    os.makedirs(f"{dst}/Talk", exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    for i, (title, bullets) in enumerate(SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        if i == 0:
            for para in body.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(20)
    prs.core_properties.author = "Lina Park"
    prs.core_properties.title = TITLE
    prs.save(f"{dst}/Talk/StreamJoin_DSSW.pptx")
    open(f"{dst}/Talk/IDS_talk_rules.txt", "w").write(RULES)
    _handout(f"{dst}/Talk/handout_comments_MW.pdf")
    _script(f"{dst}/Talk/speaker_script.pdf")
    json.dump({"slides": SLIDES, "final_titles": FINAL_TITLES, "notes": NOTES, "ack_text": ACK_TEXT,
               "grant": "PA 4471/2-1"}, open(f"{dst}/expected.json", "w"), indent=1)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/o4")
