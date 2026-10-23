"""F4 - lecture-capture recordings in one VLC: building the lecture library (library) vs revision notes (revision).

Shared live state: one VLC instance whose playlist holds all 7 recordings (both tasks play/seek in it; A renames,
moves and deletes the files behind the playlist entries and saves the playlist; B seeks to slides and takes
snapshots), one Writer window with the revision notes, one Nautilus window on the capture folder.
"""
from __future__ import annotations

import io
import os
import re
import urllib.parse

from evaluation_examples.task_class.pg_common import ASSETS, HOME, expected, fetch, result, sh, tree, upload
from evaluation_examples.task_class import pg_fimg as fi

CAPTURE = f"{HOME}/Videos/lecture_capture"
LIB = f"{HOME}/Videos/CS101"
PLAYLIST = f"{LIB}/CS101_lectures.xspf"
INDEX = f"{LIB}/index.txt"
NOTES = f"{HOME}/Documents/CS101/revision_notes.docx"
SNAPS = f"{HOME}/Documents/CS101/revision"
TOPIC_RE = {1: r"python|basic|welcome", 2: r"conditional", 3: r"loop|iteration", 4: r"function", 5: r"list|dictionar"}


def setup(sc):
    upload(sc, [("f4/f4.tar.gz", "/tmp/pgf4.tar.gz")])
    sh(sc, "mkdir -p ~/.cache ~/Pictures; [ -f ~/.cache/.pgf4_ready ] || { tar xzf /tmp/pgf4.tar.gz -C ~ && "
           "touch ~/.cache/.pgf4_ready; }; rm -f /tmp/pgf4.tar.gz")
    sc.launch(f"[ -e /tmp/.pgf4_nautilus ] || {{ touch /tmp/.pgf4_nautilus; exec nautilus '{CAPTURE}'; }}", shell=True)
    sc.launch(f"pgrep -x vlc >/dev/null || exec vlc --no-video-title-show {CAPTURE}/REC_*.mp4", shell=True)
    sc._open_setup(NOTES)


# ---------------------------------------------------------------- task A: lecture library
def _xspf_locations(path) -> list[str]:
    import xml.etree.ElementTree as ET
    try:
        root = ET.parse(path).getroot()
    except Exception:
        return []
    out = []
    for el in root.iter():
        if el.tag.endswith("}location") or el.tag == "location":
            loc = urllib.parse.unquote((el.text or "").strip())
            out.append(loc[7:] if loc.startswith("file://") else loc)
    return out


def evaluate_library(env, dest: str) -> dict:
    exp = expected("f4")
    recs = exp["recordings"]
    lects = sorted([r for r in recs if r["kind"] == "lecture"], key=lambda r: r["lecture"])
    junk = [r for r in recs if r["kind"] != "lecture"]
    t = tree(env, ["~/Videos"])
    by_md5 = {}
    for p, m in t.items():
        by_md5.setdefault(m, []).append(p)
    named = [r["lecture"] for r in lects if t.get(f"Videos/CS101/CS101_L{r['lecture']:02d}.mp4") == r["md5"]]
    junk_gone = sum(not by_md5.get(r["md5"]) for r in junk)
    left = [p for p in t if p.startswith("Videos/lecture_capture/") and p.lower().endswith(".mp4")]
    # playlist
    pl = fetch(env, PLAYLIST, dest + ".xspf")
    locs = _xspf_locations(pl) if pl else []
    pos_ok = 0
    for i, r in enumerate(lects):  # entry i must be the renamed file of lecture i+1, and that file must be lecture i+1
        if i >= len(locs):
            break
        loc = locs[i] if locs[i].startswith("/") else f"{LIB}/{locs[i]}"
        want = f"{LIB}/CS101_L{r['lecture']:02d}.mp4"
        pos_ok += os.path.normpath(loc) == want and t.get(want[len(HOME) + 1:]) == r["md5"]
    want = lects
    playlist_ok = pos_ok / len(want) * (1.0 if len(locs) == len(want) else 0.5)
    # index
    ip = fetch(env, INDEX, dest + "_index.txt")
    lines = open(ip, encoding="utf-8", errors="replace").read().splitlines() if ip else []
    topic_ok = dur_ok = 0
    idx_detail = {}
    for r in lects:
        n = r["lecture"]
        ln = next((x for x in lines if re.search(rf"\bL0?{n}\b|lecture\s*0?{n}\b", x, re.I)), None)
        if ln is None:
            idx_detail[n] = "missing"
            continue
        tp = bool(re.search(TOPIC_RE[n], ln, re.I))
        durs = [int(a) * 60 + int(b) for a, b in re.findall(r"\b(\d{1,2}):(\d{2})\b", ln)]
        du = any(abs(x - r["seconds"]) <= 2 for x in durs)
        topic_ok += tp
        dur_ok += du
        idx_detail[n] = {"line": ln[:80], "topic": tp, "duration": du}
    k = len(lects)
    score = (0.30 * len(named) / k + 0.10 * junk_gone / len(junk) + 0.05 * (not left) + 0.25 * playlist_ok
             + 0.15 * topic_ok / k + 0.15 * dur_ok / k)
    if not named:
        score = min(score, 0.3)
    checks = {"renamed_lectures": named, "test_and_blank_deleted": f"{junk_gone}/{len(junk)}", "left_in_capture": left,
              "playlist": locs, "playlist_positions_ok": f"{pos_ok}/{k}", "index": idx_detail}
    return result(score, checks)


# ---------------------------------------------------------------- task B: revision notes
_refs = {}


def _refs_for():
    from PIL import Image
    if not _refs:
        d = os.path.join(ASSETS, "f4")
        for f in sorted(os.listdir(d)):
            if f.endswith(".png") and (f.startswith("summary_") or f.startswith("slide_")):
                im = Image.open(os.path.join(d, f))
                im.load()
                _refs[f[:-4]] = im.convert("RGB")
    return _refs


def which_slide(img) -> tuple[str | None, float]:
    """Best matching reference slide ('summary_L03', 'slide_L03_1', ...) for a snapshot/embedded image: correlation at
    640x360 (band-passed text strokes) with the presenter inset ignored; >= 0.7 and >= 0.15 above the runner-up."""
    from PIL import Image
    try:
        im = img if isinstance(img, Image.Image) else Image.open(io.BytesIO(img))
        im = im.convert("RGB")
    except Exception:
        return None, 0.0
    if im.size[0] < 160:
        return None, 0.0
    scores = sorted(((fi.detail_corr(im, ref, ignore=(0.78, 0.68, 1.0, 1.0)), name) for name, ref in _refs_for().items()),
                    reverse=True)
    (bc, best), (c2, _) = scores[0], scores[1]
    return (best, bc) if bc >= 0.7 and bc - c2 >= 0.15 else (None, bc)


def _docx_sections(path):
    """[(heading text, text of the section, [image blobs])] in document order."""
    from docx import Document
    doc = Document(path)
    rels = {r.rId: r for r in doc.part.rels.values()}
    secs = [("", "", [])]
    for p in doc.paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        text = p.text
        if style.lower().startswith(("heading", "title")) or re.match(r"^\s*lecture\s+\d+\s*[-–—:]", text, re.I):
            secs.append((text, "", []))
            continue
        h, tx, ims = secs[-1]
        for m in re.finditer(r'r:(?:embed|id)="(rId\d+)"', p._p.xml):
            r = rels.get(m.group(1))
            if r is not None and "image" in r.reltype:
                ims.append(r.target_part.blob)
        secs[-1] = (h, tx + "\n" + text, ims)
    return secs


def evaluate_revision(env, dest: str) -> dict:
    exp = expected("f4")
    lecs = exp["revision"]
    c, total = {}, 0.0
    np_ = fetch(env, NOTES, dest + ".docx")
    secs = []
    if np_:
        try:
            secs = _docx_sections(np_)
        except Exception as e:  # noqa: BLE001
            c["docx_error"] = repr(e)
    all_text = "\n".join(h + "\n" + t for h, t, _ in secs)
    any_img = any(ims for _, _, ims in secs)
    for n in lecs:
        cc = {}
        sp = fetch(env, f"{SNAPS}/L{n:02d}_summary.png", f"{dest}_L{n:02d}.png")
        name, corr = which_slide(open(sp, "rb").read()) if sp else (None, 0.0)
        snap_ok = name == f"summary_L{n:02d}"
        cc.update(snapshot=bool(sp), snapshot_is=name, corr=round(corr, 3))
        sec = next(((h, t, ims) for h, t, ims in secs if re.search(rf"lecture\s+{n}\b", h, re.I)), None)
        img_ok = kp = 0
        if sec:
            ids = [which_slide(b)[0] for b in sec[2]]
            img_ok = f"summary_L{n:02d}" in ids
            body = re.sub(r"\s+", " ", sec[1].lower())
            hits = [bool(re.search(rx, body)) for rx in exp["keypoints"][str(n)]]
            kp = sum(hits) / len(hits)
            cc.update(section_images=ids, keypoints=hits)
        else:
            cc["section"] = "missing"
        s = 0.12 * snap_ok + 0.10 * img_ok + 0.10 * kp
        total += s
        cc["score"] = round(s, 3)
        c[f"L{n:02d}"] = cc
    no_ph = bool(secs) and not re.search(r"\[(insert snapshot|type the three)", all_text, re.I)
    total += 0.04 * no_ph
    if not any_img:
        total = min(total, 0.4)
    c.update(notes_found=bool(np_), placeholders_gone=no_ph)
    return result(total, c)
